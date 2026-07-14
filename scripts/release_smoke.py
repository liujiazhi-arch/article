from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
import socket
import subprocess
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from typing import Any, Sequence

from cli_json_output import build_failed_json_payload as _build_failed_payload
from cli_json_output import emit_json_payload as _emit_payload
from smoke_workdir_utils import prepare_work_dir

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORK_DIR = Path(".release_smoke")
DEFAULT_COMMAND_TIMEOUT_SECONDS = 300.0
SMOKE_BODY_ANCHORS = (
    "Thesis format smoke anchor one verifies document and PDF version matching",
    "Second distinct paragraph covers rendered page evidence in the local workflow",
    "Third paragraph confirms the installed bundle can extract searchable PDF text",
)


def _venv_python(venv_dir: Path) -> Path:
    if sys.platform == "win32":
        return venv_dir / "Scripts" / "python.exe"
    return venv_dir / "bin" / "python"


def _venv_script(venv_dir: Path, script_name: str) -> Path:
    if sys.platform == "win32":
        return venv_dir / "Scripts" / f"{script_name}.exe"
    return venv_dir / "bin" / script_name


def _run(
    command: Sequence[str | Path],
    *,
    cwd: Path | None = None,
    env: dict[str, str] | None = None,
    capture_output: bool = False,
    text: bool = True,
    check: bool = True,
    timeout: float | None = None,
) -> subprocess.CompletedProcess[str]:
    run_kwargs: dict[str, Any] = {}
    if text:
        run_kwargs["encoding"] = "utf-8"
        run_kwargs["errors"] = "replace"
    return subprocess.run(
        [str(item) for item in command],
        cwd=str(cwd) if cwd is not None else None,
        env=env,
        capture_output=capture_output,
        text=text,
        check=check,
        timeout=timeout,
        **run_kwargs,
    )


def _json_run(
    command: Sequence[str | Path],
    *,
    cwd: Path | None = None,
    timeout: float | None = None,
) -> dict[str, Any]:
    completed = _run(command, cwd=cwd, capture_output=True, timeout=timeout)
    return json.loads(completed.stdout)


def _available_port() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
        sock.bind(("127.0.0.1", 0))
        return int(sock.getsockname()[1])


def _request_json(
    url: str,
    *,
    method: str = "GET",
    payload: dict[str, Any] | None = None,
    timeout: float = 10.0,
) -> dict[str, Any]:
    data = None
    headers = {"Accept": "application/json"}
    if payload is not None:
        data = json.dumps(payload).encode("utf-8")
        headers["Content-Type"] = "application/json"
    request = urllib.request.Request(url, data=data, headers=headers, method=method)
    with urllib.request.urlopen(request, timeout=timeout) as response:
        return json.loads(response.read().decode("utf-8"))


def _request_bytes(url: str, *, timeout: float = 10.0) -> bytes:
    request = urllib.request.Request(url, method="GET")
    with urllib.request.urlopen(request, timeout=timeout) as response:
        return response.read()


def _wait_for_ready(base_url: str, *, timeout_seconds: float = 30.0) -> dict[str, Any]:
    deadline = time.monotonic() + timeout_seconds
    last_error: Exception | None = None
    while time.monotonic() < deadline:
        try:
            payload = _request_json(f"{base_url}/ready", timeout=2.0)
            if payload.get("status") == "ready":
                return payload
        except (OSError, urllib.error.URLError) as exc:
            last_error = exc
        time.sleep(0.25)
    raise RuntimeError(f"article-api did not become ready within {timeout_seconds}s: {last_error}")


def _wait_for_job(base_url: str, job_id: str, *, timeout_seconds: float = 60.0) -> dict[str, Any]:
    deadline = time.monotonic() + timeout_seconds
    while time.monotonic() < deadline:
        payload = _request_json(f"{base_url}/jobs/{urllib.parse.quote(job_id)}", timeout=5.0)
        if payload.get("status") in {"succeeded", "failed"}:
            return payload
        time.sleep(0.25)
    raise RuntimeError(f"job did not finish within {timeout_seconds}s: {job_id}")


def _post_multipart_file(
    url: str,
    *,
    field_name: str,
    file_path: Path,
    content_type: str,
    params: dict[str, str],
) -> dict[str, Any]:
    boundary = f"----release-smoke-{int(time.time() * 1000)}"
    query = urllib.parse.urlencode(params)
    target = f"{url}?{query}" if query else url
    content = file_path.read_bytes()
    head = (
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="{field_name}"; filename="{file_path.name}"\r\n'
        f"Content-Type: {content_type}\r\n\r\n"
    ).encode("utf-8")
    tail = f"\r\n--{boundary}--\r\n".encode("utf-8")
    request = urllib.request.Request(
        target,
        data=head + content + tail,
        headers={
            "Accept": "application/json",
            "Content-Type": f"multipart/form-data; boundary={boundary}",
        },
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=20.0) as response:
        return json.loads(response.read().decode("utf-8"))


def _build_smoke_docx(docx_path: Path, *, python_executable: Path) -> None:
    script = f"""
from docx import Document
doc = Document()
doc.add_paragraph("封面标题")
h1 = doc.add_paragraph("1 绪论")
h1.style = doc.styles["Heading 1"]
h2 = doc.add_paragraph("1.1 研究背景")
h2.style = doc.styles["Heading 2"]
for text in {SMOKE_BODY_ANCHORS!r}:
    doc.add_paragraph(text)
doc.save({str(docx_path)!r})
"""
    _run([str(python_executable), "-c", script])


def _build_smoke_pdf(
    pdf_path: Path,
    *,
    python_executable: Path,
    include_layout_issue: bool = False,
) -> None:
    pages = [list(SMOKE_BODY_ANCHORS)]
    if include_layout_issue:
        pages = [[*SMOKE_BODY_ANCHORS, "x = y + z"], ["(1)"]]
    script = f"""
import ctypes
import pypdfium2 as pdfium
from pypdfium2 import raw

document = pdfium.PdfDocument.new()
font = raw.FPDFText_LoadStandardFont(document.raw, b"Helvetica")
for lines in {pages!r}:
    page = document.new_page(595, 842)
    for index, line in enumerate(lines):
        text_object = raw.FPDFPageObj_CreateTextObj(document.raw, font, 12)
        text = (ctypes.c_ushort * (len(line) + 1))(*(ord(char) for char in line), 0)
        assert raw.FPDFText_SetText(text_object, text)
        raw.FPDFPageObj_Transform(text_object, 1, 0, 0, 1, 72, 760 - index * 24)
        raw.FPDFPage_InsertObject(page.raw, text_object)
    assert raw.FPDFPage_GenerateContent(page.raw)
    page.close()
document.save({str(pdf_path)!r})
raw.FPDFFont_Close(font)
document.close()
"""
    _run([str(python_executable), "-c", script])


def _start_http_server(
    *,
    venv_python: Path,
    state_root: Path,
    runtime_root: Path,
    smoke_dir: Path,
) -> tuple[str, subprocess.Popen[str]]:
    base_url = f"http://127.0.0.1:{_available_port()}"
    process = subprocess.Popen(
        [
            str(venv_python),
            "-m",
            "article_api.local_app",
            "serve",
            "--host",
            "127.0.0.1",
            "--port",
            base_url.rsplit(":", 1)[-1],
            "--state-root",
            str(state_root),
            "--runtime-root",
            str(runtime_root),
        ],
        cwd=str(smoke_dir),
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    return base_url, process


def _stop_http_server(process: subprocess.Popen[str]) -> None:
    process.terminate()
    try:
        process.wait(timeout=5)
    except subprocess.TimeoutExpired:
        process.kill()
        process.wait(timeout=5)


def _run_apply_http_smoke(base_url: str, docx_path: Path, runtime_root: Path) -> dict[str, Any]:
    upload = _post_multipart_file(
        f"{base_url}/uploads/docx",
        field_name="file",
        file_path=docx_path,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        params={"runtime_root": str(runtime_root)},
    )
    job = _request_json(
        f"{base_url}/uploads/{urllib.parse.quote(upload['upload_id'])}/jobs/apply",
        method="POST",
        payload={"scopes": ["headings"], "runtime_root": str(runtime_root), "stage_input": True},
        timeout=10.0,
    )
    status = _wait_for_job(base_url, job["job_id"])
    if status.get("status") != "succeeded":
        raise RuntimeError(f"release smoke apply job failed: {status}")
    result = _request_json(f"{base_url}/jobs/{urllib.parse.quote(job['job_id'])}/result", timeout=10.0)
    output = _request_bytes(
        f"{base_url}/jobs/{urllib.parse.quote(job['job_id'])}/artifacts/output/download",
        timeout=20.0,
    )
    if not output:
        raise RuntimeError("release smoke output download returned no bytes")
    output_path = docx_path.with_name(f"{docx_path.stem}_fixed.docx")
    output_path.write_bytes(output)
    output_upload = _post_multipart_file(
        f"{base_url}/uploads/docx",
        field_name="file",
        file_path=output_path,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        params={"runtime_root": str(runtime_root)},
    )
    return {
        "upload_id": upload["upload_id"],
        "output_upload_id": output_upload["upload_id"],
        "job_id": job["job_id"],
        "job_status": status["status"],
        "business_status": (result.get("summary") or {}).get("business_status"),
        "download_bytes": len(output),
    }


def _run_render_http_smoke(
    base_url: str,
    docx_upload_id: str,
    pdf_path: Path,
    runtime_root: Path,
) -> dict[str, Any]:
    pdf_upload = _post_multipart_file(
        f"{base_url}/uploads/pdf",
        field_name="file",
        file_path=pdf_path,
        content_type="application/pdf",
        params={"runtime_root": str(runtime_root)},
    )
    job = _request_json(
        f"{base_url}/uploads/{urllib.parse.quote(docx_upload_id)}/render-review-jobs",
        method="POST",
        payload={"pdf_upload_id": pdf_upload["upload_id"], "pdf_matches_docx_confirmed": True},
        timeout=10.0,
    )
    status = _wait_for_job(base_url, job["job_id"])
    if status.get("status") != "succeeded":
        raise RuntimeError(f"release smoke render-review job failed: {status}")
    result_payload = _request_json(f"{base_url}/jobs/{urllib.parse.quote(job['job_id'])}/result", timeout=10.0)
    result = result_payload.get("result") or {}
    page_count = int(result.get("page_count") or (result.get("summary") or {}).get("page_count") or 0)
    if page_count != 1:
        raise RuntimeError(f"release smoke render-review returned {page_count} pages")
    summary = result.get("summary") or {}
    if (
        result.get("evidence_trust") != "user-confirmed"
        or result.get("layout_decision_eligible") is not True
        or result.get("pdf_matches_docx_confirmed") is not True
        or summary.get("pdf_matches_docx_confirmed") is not True
        or summary.get("pdf_content_match_status") != "matched"
        or summary.get("pdf_content_matched") is not True
    ):
        raise RuntimeError("release smoke render-review did not preserve confirmed PDF evidence")
    return {
        "render_upload_id": pdf_upload["upload_id"],
        "render_job_id": job["job_id"],
        "render_job_status": status["status"],
        "render_page_count": page_count,
        "render_evidence_trust": result["evidence_trust"],
        "render_pdf_matches_docx_confirmed": result["pdf_matches_docx_confirmed"],
        "render_pdf_content_match_status": summary["pdf_content_match_status"],
        "render_layout_decision_eligible": result["layout_decision_eligible"],
    }


def _run_http_smoke(
    *,
    venv_python: Path,
    state_root: Path,
    runtime_root: Path,
    smoke_dir: Path,
) -> dict[str, Any]:
    docx_path = smoke_dir / "release_smoke.docx"
    pdf_path = smoke_dir / "release_smoke.pdf"
    _build_smoke_docx(docx_path, python_executable=venv_python)
    _build_smoke_pdf(pdf_path, python_executable=venv_python)
    base_url, process = _start_http_server(
        venv_python=venv_python,
        state_root=state_root,
        runtime_root=runtime_root,
        smoke_dir=smoke_dir,
    )
    try:
        ready = _wait_for_ready(base_url)
        apply_result = _run_apply_http_smoke(base_url, docx_path, runtime_root)
        render_result = _run_render_http_smoke(base_url, apply_result["output_upload_id"], pdf_path, runtime_root)
        return {
            "status": "ok",
            "ready": ready.get("status"),
            **apply_result,
            **render_result,
        }
    finally:
        _stop_http_server(process)


def _resolve_latest_wheel(wheel_dir: Path) -> Path:
    wheels = sorted(wheel_dir.glob("thesis_format_tool-*.whl"), key=lambda item: item.stat().st_mtime, reverse=True)
    if not wheels:
        raise RuntimeError(f"No wheel was built in {wheel_dir}")
    return wheels[0]


def run_release_smoke(
    *,
    work_dir: str | Path = DEFAULT_WORK_DIR,
    python_executable: str | Path = sys.executable,
    state_root: str | Path | None = None,
    runtime_root: str | Path | None = None,
    wheelhouse: str | Path | None = None,
    keep_work_dir: bool = False,
    command_timeout_seconds: float = DEFAULT_COMMAND_TIMEOUT_SECONDS,
) -> dict[str, Any]:
    smoke_dir = Path(work_dir).expanduser().resolve()
    wheelhouse_path = Path(wheelhouse).expanduser().resolve() if wheelhouse is not None else None
    if wheelhouse_path is not None and not keep_work_dir:
        try:
            wheelhouse_path.relative_to(smoke_dir)
        except ValueError:
            pass
        else:
            raise RuntimeError("Prebuilt wheelhouse must be outside work-dir unless --keep-work-dir is set")
    prepare_work_dir(smoke_dir, keep_existing=keep_work_dir)
    venv_dir = smoke_dir / "venv"
    wheel_dir = wheelhouse_path if wheelhouse_path is not None else smoke_dir / "wheelhouse"
    state = Path(state_root).expanduser().resolve() if state_root is not None else smoke_dir / "state"
    runtime = Path(runtime_root).expanduser().resolve() if runtime_root is not None else smoke_dir / "runtime"

    command_timeout = float(command_timeout_seconds)
    _run([str(python_executable), "-m", "venv", str(venv_dir)], timeout=command_timeout)
    venv_python = _venv_python(venv_dir)
    _run([str(venv_python), "-m", "pip", "install", "-U", "wheel"], timeout=command_timeout)
    if wheelhouse is None:
        _run(
            [str(venv_python), "-m", "pip", "wheel", "--wheel-dir", str(wheel_dir), ".[api]"],
            cwd=PROJECT_ROOT,
            timeout=command_timeout,
        )
    elif not wheel_dir.exists():
        raise RuntimeError(f"Prebuilt wheelhouse does not exist: {wheel_dir}")
    wheel_path = _resolve_latest_wheel(wheel_dir)
    _run(
        [
            str(venv_python),
            "-m",
            "pip",
            "install",
            "--force-reinstall",
            "--no-index",
            "--find-links",
            str(wheel_dir),
            f"{wheel_path}[api]",
        ],
        timeout=command_timeout,
    )

    article_local = _venv_script(venv_dir, "lnu-thesis-local")
    thesis_workbench = _venv_script(venv_dir, "thesis-workbench")
    doctor = _json_run(
        [str(article_local), "doctor", "--state-root", str(state), "--runtime-root", str(runtime)],
        cwd=smoke_dir,
        timeout=command_timeout,
    )
    profiles_stdout = _run(
        [str(thesis_workbench), "profiles"],
        cwd=smoke_dir,
        capture_output=True,
        timeout=command_timeout,
    ).stdout
    if "lnu-checker-2026" not in profiles_stdout:
        raise RuntimeError("thesis-workbench profiles did not expose lnu-checker-2026")
    http_smoke = _run_http_smoke(
        venv_python=venv_python,
        state_root=state,
        runtime_root=runtime,
        smoke_dir=smoke_dir,
    )
    return {
        "status": "ok",
        "service_version": str(doctor.get("version") or ""),
        "work_dir": str(smoke_dir),
        "install": {
            "mode": "wheel",
            "venv": str(venv_dir),
        },
        "distribution": {
            "wheel_path": str(wheel_path),
            "wheelhouse": str(wheel_dir),
        },
        "roots": {
            "state_root": str(state),
            "runtime_root": str(runtime),
        },
        "checks": {
            "doctor": {
                "status": doctor.get("status") or doctor.get("summary", {}).get("status") or "ok",
                "headline": doctor.get("summary", {}).get("headline"),
            },
            "profiles": {
                "status": "ok",
                "stdout": profiles_stdout,
            },
            "http_smoke": http_smoke,
        },
    }


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="release_smoke.py")
    parser.add_argument("--work-dir", default=str(DEFAULT_WORK_DIR))
    parser.add_argument("--python", default=sys.executable)
    parser.add_argument("--state-root")
    parser.add_argument("--runtime-root")
    parser.add_argument("--wheelhouse", help="Reuse a prebuilt wheelhouse and install the project wheel from it")
    parser.add_argument("--keep-work-dir", action="store_true")
    parser.add_argument("--command-timeout-seconds", type=float, default=DEFAULT_COMMAND_TIMEOUT_SECONDS)
    parser.add_argument("--json-output", type=Path, help="Write the final JSON payload to this path")
    return parser


_RELEASE_SMOKE_TIMEOUT_NEXT_STEPS = [
    "Retry with a larger --command-timeout-seconds value if the network or package index is slow.",
    "Check that the dependency wheelhouse can download python-docx, lxml, FastAPI, uvicorn, and python-multipart.",
    "Use CI pip caching so the release smoke gate is deterministic enough for repeated release checks.",
]

_RELEASE_SMOKE_ERROR_NEXT_STEPS = [
    "Inspect the failed release smoke step and rerun after fixing the packaging or local API issue.",
]


def _format_release_error(exc: Exception) -> dict[str, Any]:
    if isinstance(exc, subprocess.TimeoutExpired):
        return _build_failed_payload(
            exc,
            next_steps=_RELEASE_SMOKE_TIMEOUT_NEXT_STEPS,
            timeout_type="command_timeout",
            timeout_message="Release smoke command timed out before the package install flow completed.",
        )
    return _build_failed_payload(exc, next_steps=_RELEASE_SMOKE_ERROR_NEXT_STEPS)


def main(argv: list[str] | None = None) -> int:
    args = _build_parser().parse_args(argv)
    try:
        payload = run_release_smoke(
            work_dir=args.work_dir,
            python_executable=args.python,
            state_root=args.state_root,
            runtime_root=args.runtime_root,
            wheelhouse=args.wheelhouse,
            keep_work_dir=bool(args.keep_work_dir),
            command_timeout_seconds=float(args.command_timeout_seconds),
        )
    except Exception as exc:
        _emit_payload(_format_release_error(exc), json_output=args.json_output)
        return 1
    _emit_payload(payload, json_output=args.json_output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
