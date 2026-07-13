from __future__ import annotations

import argparse
import re
from collections import OrderedDict
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from backmatter_title_utils import (
    is_acknowledgement_title,
    is_appendix_title,
    is_reference_title,
)
from citation_text_utils import (
    CITATION_NUMBER_GROUP_RE as CITATION_RE,
    citation_numbers_from_group as _expand_citation_numbers,
    format_citation_numbers as _compress_citation_numbers,
)


REFERENCE_RE = re.compile(r"^\[(\d+)\]\s*")
_CHAPTER_HEADING_RE = re.compile(r"^第.+[章节篇]")


def _normalize_backmatter_text(text: str | None) -> str:
    return re.sub(r"[\s\u3000]+", "", str(text or "")).lower()


def _is_reference_heading_text(text: str | None) -> bool:
    return is_reference_title(text)


def _is_reference_stop_text(text: str | None) -> bool:
    stripped = str(text or "").strip()
    if not stripped:
        return False
    normalized = _normalize_backmatter_text(stripped)
    return (
        is_acknowledgement_title(stripped)
        or is_appendix_title(stripped)
        or normalized == "abstract"
        or bool(_CHAPTER_HEADING_RE.match(stripped))
    )


def _iter_body_paragraphs_until_references(doc: Document):
    for paragraph in doc.paragraphs:
        if _is_reference_heading_text(paragraph.text):
            break
        yield paragraph


def _build_reference_order_map(doc: Document) -> dict[int, int]:
    seen: OrderedDict[int, None] = OrderedDict()
    for paragraph in _iter_body_paragraphs_until_references(doc):
        for match in CITATION_RE.finditer(paragraph.text):
            for number in _expand_citation_numbers(match.group(1)):
                seen.setdefault(number, None)

    referenced_numbers = list(seen.keys())
    reference_entries = []
    in_references = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if _is_reference_heading_text(text):
            in_references = True
            continue
        if not in_references:
            continue
        if _is_reference_stop_text(text):
            break
        match = REFERENCE_RE.match(text)
        if match:
            reference_entries.append(int(match.group(1)))

    for number in reference_entries:
        seen.setdefault(number, None)

    return {old: new for new, old in enumerate(seen.keys(), start=1)}


def _replace_citation_text(text: str, order_map: dict[int, int]) -> str:
    def repl(match: re.Match[str]) -> str:
        old_numbers = _expand_citation_numbers(match.group(1))
        new_numbers = [order_map[number] for number in old_numbers if number in order_map]
        return _compress_citation_numbers(new_numbers)

    return CITATION_RE.sub(repl, text)


def _update_body_citations(doc: Document, order_map: dict[int, int]) -> None:
    for paragraph in _iter_body_paragraphs_until_references(doc):
        for run in paragraph.runs:
            if "[" not in run.text or "]" not in run.text:
                continue
            updated = _replace_citation_text(run.text, order_map)
            if updated != run.text:
                run.text = updated


def _collect_reference_entry_blocks(doc: Document):
    body = doc._body._element
    children = list(body.iterchildren())
    ref_heading_paragraph = next((p for p in doc.paragraphs if _is_reference_heading_text(p.text)), None)
    if ref_heading_paragraph is None:
        raise ValueError("文档中未找到“参考文献”标题。")
    ref_heading_idx = next((idx for idx, child in enumerate(children) if child is ref_heading_paragraph._p), None)
    if ref_heading_idx is None:
        raise ValueError("未能在正文节点序列中定位参考文献标题。")

    blocks = []
    in_references = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if _is_reference_heading_text(text):
            in_references = True
            continue
        if not in_references:
            continue
        if text and _is_reference_stop_text(text):
            break
        child = paragraph._p
        if not text:
            continue
        match = REFERENCE_RE.match(text)
        if match:
            blocks.append({"old_number": int(match.group(1)), "elem": child, "text": text})
        elif text == "——.":
            blocks.append({"old_number": None, "elem": child, "text": text})
    return body, blocks, ref_heading_idx


def _update_reference_entry_number(elem, new_number: int, *, trailing_space: bool = True) -> None:
    text_elems = elem.findall(".//" + qn("w:t"))
    if not text_elems:
        return
    first = text_elems[0]
    original = first.text or ""
    separator = " " if trailing_space else ""
    if REFERENCE_RE.match(original):
        first.text = REFERENCE_RE.sub(f"[{new_number}]{separator}", original, count=1)
        return

    full_text = "".join(t.text or "" for t in text_elems)
    match = REFERENCE_RE.match(full_text)
    if match is None:
        return

    remaining_prefix = match.end()
    for index, text_elem in enumerate(text_elems):
        text = text_elem.text or ""
        consumed = min(len(text), remaining_prefix)
        suffix = text[consumed:]
        text_elem.text = f"[{new_number}]{separator}{suffix}" if index == 0 else suffix
        remaining_prefix -= consumed
        if remaining_prefix <= 0:
            break


def _reorder_reference_entries(doc: Document, order_map: dict[int, int], *, trailing_space: bool = True) -> None:
    body, blocks, ref_heading_idx = _collect_reference_entry_blocks(doc)
    valid_blocks = [block for block in blocks if block["old_number"] is not None]
    stray_blocks = [block for block in blocks if block["old_number"] is None]
    child_index = {id(child): idx for idx, child in enumerate(list(body.iterchildren()))}
    if valid_blocks:
        insert_idx = min(child_index[id(block["elem"])] for block in valid_blocks)
    else:
        insert_idx = ref_heading_idx + 1

    for block in blocks:
        body.remove(block["elem"])

    for block in sorted(valid_blocks, key=lambda item: order_map[item["old_number"]]):
        _update_reference_entry_number(
            block["elem"],
            order_map[block["old_number"]],
            trailing_space=trailing_space,
        )
        body.insert(insert_idx, block["elem"])
        insert_idx += 1

    for block in stray_blocks:
        # 删除明显的参考文献占位残留，如“——.”
        continue


def reorder_references_in_document(doc: Document, *, trailing_space: bool = True) -> dict[int, int]:
    order_map = _build_reference_order_map(doc)
    _update_body_citations(doc, order_map)
    try:
        _reorder_reference_entries(doc, order_map, trailing_space=trailing_space)
    except ValueError:
        return order_map
    return order_map


def reorder_references_by_appearance(input_path: str, output_path: str, *, trailing_space: bool = True) -> dict[int, int]:
    doc = Document(input_path)
    order_map = reorder_references_in_document(doc, trailing_space=trailing_space)
    doc.save(output_path)
    return order_map


def main() -> int:
    parser = argparse.ArgumentParser(description="按正文首次出现顺序重排参考文献编号，并同步更新正文引文。")
    parser.add_argument("input_docx", help="输入 docx")
    parser.add_argument("--output", required=True, help="输出 docx")
    args = parser.parse_args()

    order_map = reorder_references_by_appearance(args.input_docx, args.output)
    print(f"已输出: {args.output}")
    print(f"共重排编号: {len(order_map)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
