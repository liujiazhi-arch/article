#!/usr/bin/env python3

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook


SRC_DIR_DEFAULT = "/Users/apple/Desktop/gelatin_practice_data"
DST_DIR_DEFAULT = "/Users/apple/Desktop/gelatin_practice_data_中文"


SHEET_TITLE_MAP = {
    "README": "说明",
    "Files": "文件清单",
    "Prism_Raw": "Prism原始数据",
    "Long_Format": "长表格式",
    "Summary": "汇总",
    "UV_XY_Raw": "UV_XY原始数据",
    "UV_XY_GroupMean": "UV_XY组均值",
    "PeakAbs_230": "230nm峰吸光度",
    "PeakAbs_280": "280nm峰吸光度",
    "Ratio_280_230": "A280_A230比值",
    "PeakSummary": "峰值汇总",
    "FTIR_XY_Raw": "FTIR_XY原始数据",
    "FTIR_XY_GroupMean": "FTIR_XY组均值",
    "Amide_A_Pos": "Amide_A峰位",
    "Amide_I_Pos": "Amide_I峰位",
    "Amide_II_Pos": "Amide_II峰位",
    "Amide_III_Pos": "Amide_III峰位",
}

STRING_MAP = {
    "MASTER_SUMMARY_SIMULATED_PRACTICE_ONLY": "总索引_仅供练习模拟",
    "Status": "状态",
    "SIMULATED_PRACTICE_ONLY": "仅供练习模拟（SIMULATED_PRACTICE_ONLY）",
    "Note 1": "说明 1",
    "Note 2": "说明 2",
    "Note 3": "说明 3",
    "Note 4": "说明 4",
    "Generated files": "已生成文件",
    "Experiment": "实验项目",
    "File name": "文件名",
    "Purpose": "用途",
    "Column data for Prism": "Prism 柱状图数据",
    "XY curves plus peak metrics": "XY 曲线及峰值指标",
    "XY curves plus peak positions": "XY 曲线及峰位数据",
    "01_WHC_SIMULATED_PRACTICE_ONLY": "01_WHC_仅供练习模拟",
    "02_OHC_SIMULATED_PRACTICE_ONLY": "02_OHC_仅供练习模拟",
    "03_EAI_SIMULATED_PRACTICE_ONLY": "03_EAI_仅供练习模拟",
    "04_ESI_SIMULATED_PRACTICE_ONLY": "04_ESI_仅供练习模拟",
    "05_UV_SIMULATED_PRACTICE_ONLY": "05_UV_仅供练习模拟",
    "06_FTIR_SIMULATED_PRACTICE_ONLY": "06_FTIR_仅供练习模拟",
    "This workbook indexes the generated synthetic practice files.": "本工作簿用于索引已生成的练习用模拟文件。",
    "Each experiment was exported as a separate Excel file to match a Prism-first workflow.": "每个实验都单独导出为一个 Excel 文件，便于按 Prism 的使用流程整理。",
    "Use these files for layout rehearsal, import testing, and sheet-structure planning only.": "这些文件仅用于版式演练、导入测试和工作表结构规划。",
    "This workbook contains synthetic practice data derived from group mean and SD values only.": "本工作簿包含依据各组均值和 SD 推导得到的练习用模拟数据。",
    "These values are for Prism workflow rehearsal, figure layout testing, and file-structure setup.": "这些数值用于 Prism 流程演练、图形排版测试和文件结构搭建。",
    "They are not measured laboratory observations and should not be represented as real experimental raw data.": "这些数值不是实验室实测结果，不应表述为真实实验原始数据。",
    "The three replicates in each group were generated to reproduce the supplied mean and sample SD exactly.": "每组 3 个平行样通过模拟生成，用于尽量复现给定的均值和样本 SD。",
    "Prism_Raw (Prism-ready raw table)": "Prism原始数据（可直接导入 Prism）",
    "Replicate (g/g)": "平行样（g/g）",
    "Replicate (m^2/g)": "平行样（m^2/g）",
    "Replicate (min)": "平行样（min）",
    "Mean": "均值",
    "SD": "SD",
    "Letter": "显著性字母",
    "Long_Format (long format)": "长表格式",
    "Group": "组别",
    "Replicate": "平行样",
    "Value (g/g)": "数值（g/g）",
    "Value (m^2/g)": "数值（m^2/g）",
    "Value (min)": "数值（min）",
    "Summary": "汇总",
    "Mean (g/g)": "均值（g/g）",
    "Mean (m^2/g)": "均值（m^2/g）",
    "Mean (min)": "均值（min）",
    "05_UV_SIMULATED_PRACTICE_ONLY": "05_UV_仅供练习模拟",
    "This workbook contains synthetic practice UV data derived from the supplied schematic spectral parameters.": "本工作簿包含依据给定示意光谱参数生成的 UV 练习用模拟数据。",
    "The XY sheet is suitable for Prism XY plotting practice.": "XY 工作表适合用于 Prism 的 XY 作图练习。",
    "The peak-metric sheets provide column-format tables for group comparison in Prism.": "峰值指标工作表提供可用于 Prism 组间比较的柱状表格。",
    "These are synthetic curves and synthetic derived metrics, not instrument-exported measurements.": "这些数据为模拟曲线及其派生指标，不是仪器直接导出的实测数据。",
    "UV_XY_Raw": "UV_XY原始数据",
    "Wavelength (nm)": "波长（nm）",
    "UV_XY_GroupMean": "UV_XY组均值",
    "PeakAbs_230": "230nm峰吸光度",
    "PeakAbs_280": "280nm峰吸光度",
    "Ratio_280_230": "A280/A230 比值",
    "PeakSummary": "峰值汇总",
    "Absorbance near 230 nm": "230 nm 附近吸光度",
    "Absorbance near 280 nm": "280 nm 附近吸光度",
    "A280/A230 ratio": "A280/A230 比值",
    "Target p1 (nm)": "目标 p1 峰位（nm）",
    "Target p2 (nm)": "目标 p2 峰位（nm）",
    "06_FTIR_SIMULATED_PRACTICE_ONLY": "06_FTIR_仅供练习模拟",
    "This workbook contains synthetic practice FTIR data based on the supplied representative peak positions.": "本工作簿包含依据给定代表性峰位生成的 FTIR 练习用模拟数据。",
    "The XY sheet is suitable for Prism XY plotting rehearsal.": "XY 工作表适合用于 Prism 的 XY 作图演练。",
    "Peak-position sheets provide groupwise tables that can be used for summary comparison figures.": "峰位工作表提供按组整理的数据，可用于绘制汇总比较图。",
    "All values are synthetic practice data rather than spectrometer-exported observations.": "所有数值均为练习用模拟数据，不是仪器直接导出的实测结果。",
    "FTIR_XY_Raw": "FTIR_XY原始数据",
    "Wavenumber (cm^-1)": "波数（cm^-1）",
    "FTIR_XY_GroupMean": "FTIR_XY组均值",
    "Amide A peak position (cm^-1)": "Amide A 峰位（cm^-1）",
    "Amide I peak position (cm^-1)": "Amide I 峰位（cm^-1）",
    "Amide II peak position (cm^-1)": "Amide II 峰位（cm^-1）",
    "Amide III peak position (cm^-1)": "Amide III 峰位（cm^-1）",
}

HEADER_REPLACEMENTS = {
    "G1_R1": "G1_平行1",
    "G1_R2": "G1_平行2",
    "G1_R3": "G1_平行3",
    "G2_R1": "G2_平行1",
    "G2_R2": "G2_平行2",
    "G2_R3": "G2_平行3",
    "G3_R1": "G3_平行1",
    "G3_R2": "G3_平行2",
    "G3_R3": "G3_平行3",
    "G4_R1": "G4_平行1",
    "G4_R2": "G4_平行2",
    "G4_R3": "G4_平行3",
    "G1_Mean": "G1_均值",
    "G2_Mean": "G2_均值",
    "G3_Mean": "G3_均值",
    "G4_Mean": "G4_均值",
    "R1": "平行1",
    "R2": "平行2",
    "R3": "平行3",
}

FILE_NAME_MAP = {
    "00_MASTER_SUMMARY_SIMULATED_PRACTICE_ONLY.xlsx": "00_总索引_仅供练习模拟.xlsx",
    "01_WHC_SIMULATED_PRACTICE_ONLY.xlsx": "01_WHC_仅供练习模拟.xlsx",
    "02_OHC_SIMULATED_PRACTICE_ONLY.xlsx": "02_OHC_仅供练习模拟.xlsx",
    "03_EAI_SIMULATED_PRACTICE_ONLY.xlsx": "03_EAI_仅供练习模拟.xlsx",
    "04_ESI_SIMULATED_PRACTICE_ONLY.xlsx": "04_ESI_仅供练习模拟.xlsx",
    "05_UV_SIMULATED_PRACTICE_ONLY.xlsx": "05_UV_仅供练习模拟.xlsx",
    "06_FTIR_SIMULATED_PRACTICE_ONLY.xlsx": "06_FTIR_仅供练习模拟.xlsx",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Translate gelatin practice Excel files to Chinese.")
    parser.add_argument("--src-dir", default=SRC_DIR_DEFAULT)
    parser.add_argument("--dst-dir", default=DST_DIR_DEFAULT)
    return parser.parse_args()


def replace_string(value: str) -> str:
    if value in HEADER_REPLACEMENTS:
        return HEADER_REPLACEMENTS[value]
    if value in STRING_MAP:
        return STRING_MAP[value]
    return value


def translate_formula(value: str) -> str:
    translated = value
    for old_name, new_name in SHEET_TITLE_MAP.items():
        translated = translated.replace(f"{old_name}!", f"{new_name}!")
        translated = translated.replace(f"'{old_name}'!", f"'{new_name}'!")
    return translated


def process_workbook(src_path: Path, dst_path: Path) -> None:
    wb = load_workbook(src_path)

    for ws in wb.worksheets:
        if ws.title in SHEET_TITLE_MAP:
            ws.title = SHEET_TITLE_MAP[ws.title]

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    if cell.value.startswith("="):
                        cell.value = translate_formula(cell.value)
                    else:
                        cell.value = replace_string(cell.value)

    wb.save(dst_path)


def main() -> None:
    args = parse_args()
    src_dir = Path(args.src_dir).expanduser().resolve()
    dst_dir = Path(args.dst_dir).expanduser().resolve()
    dst_dir.mkdir(parents=True, exist_ok=True)

    generated: list[Path] = []
    for src_path in sorted(src_dir.glob("*.xlsx")):
        if src_path.name.startswith("~") or src_path.name.startswith(".~"):
            continue
        dst_name = FILE_NAME_MAP.get(src_path.name, src_path.name)
        dst_path = dst_dir / dst_name
        process_workbook(src_path, dst_path)
        generated.append(dst_path)

    print("Generated Chinese workbooks:")
    for path in generated:
        print(path)


if __name__ == "__main__":
    main()
