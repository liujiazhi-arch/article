from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUTPUT_PATH = "/Users/apple/Desktop/毕业论文/第二部分实验方法_溶胀率与溶失率试验_原始.docx"


def set_run_fonts(run, east_asia="宋体", ascii_font="Times New Roman", size_pt=12, bold=False, italic=False, superscript=False):
    run.font.name = ascii_font
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.superscript = superscript
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), ascii_font)
    r_fonts.set(qn("w:hAnsi"), ascii_font)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    r_fonts.set(qn("w:cs"), ascii_font)


def add_text(paragraph, text, east_asia="宋体", ascii_font="Times New Roman", size_pt=12, bold=False, italic=False, superscript=False):
    run = paragraph.add_run(text)
    set_run_fonts(run, east_asia=east_asia, ascii_font=ascii_font, size_pt=size_pt, bold=bold, italic=italic, superscript=superscript)
    return run


def add_body_paragraph(doc, parts):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.5
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for part in parts:
        add_text(
            p,
            part["text"],
            east_asia=part.get("east_asia", "宋体"),
            ascii_font=part.get("ascii_font", "Times New Roman"),
            size_pt=part.get("size_pt", 12),
            bold=part.get("bold", False),
            italic=part.get("italic", False),
            superscript=part.get("superscript", False),
        )
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    if level == 1:
        set_run_fonts(run, east_asia="黑体", ascii_font="Times New Roman", size_pt=16, bold=True)
    elif level == 2:
        set_run_fonts(run, east_asia="黑体", ascii_font="Times New Roman", size_pt=14, bold=True)
    else:
        set_run_fonts(run, east_asia="黑体", ascii_font="Times New Roman", size_pt=12, bold=True)
    return p


def add_formula(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing = 1.5
    add_text(p, text, east_asia="宋体", ascii_font="Times New Roman", size_pt=12, italic=False)
    return p


def add_reference(doc, index, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.left_indent = Pt(0)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.hanging_indent = Pt(21)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_text(p, f"[{index}] ", east_asia="宋体", ascii_font="Times New Roman", size_pt=10.5)
    add_text(p, text, east_asia="宋体", ascii_font="Times New Roman", size_pt=10.5)
    return p


def configure_page(document):
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.gutter = Cm(0.5)


def configure_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    for style_name, size_pt in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size_pt)
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")


def main():
    doc = Document()
    configure_page(doc)
    configure_styles(doc)

    add_heading(doc, "第2章 溶胀率与溶失率试验", 1)

    add_heading(doc, "2.1 实验设计说明", 2)
    add_body_paragraph(
        doc,
        [
            {"text": "考虑到样品数量和材料形态，本研究将溶胀率试验设计为连续取样与重复称重法，即每组设置3个平行样，在不同时间点反复取出称重后再放回原液继续浸泡；将溶失率试验设计为24 h浸泡后测残余干质量法。上述两种方法均可用于明胶基材料的水稳定性评价"},
            {"text": "[1]", "superscript": True},
            {"text": "[5]", "superscript": True},
            {"text": "[7]", "superscript": True},
            {"text": "。"},
        ],
    )

    add_heading(doc, "2.2 溶胀率试验", 2)
    add_heading(doc, "2.2.1 实验目的", 3)
    add_body_paragraph(doc, [{"text": "测定不同改性处理后的鹿皮明胶冻干凝胶在水中的吸水膨胀能力，比较各组样品的吸水行为和网络稳定性。"}])

    add_heading(doc, "2.2.2 实验分组", 3)
    add_body_paragraph(doc, [{"text": "实验共设置6组，分别为G、G-HT、A1、A2、B1和B2。每组设置3个平行样，因此溶胀率试验共需18个样品。"}])
    add_formula(doc, "6组 × 3个平行 = 18个样品")

    add_heading(doc, "2.2.3 试剂与仪器", 3)
    add_body_paragraph(doc, [{"text": "试剂与仪器包括蒸馏水或去离子水、分析天平（精度0.1 mg）、镊子、滤纸、编号样品瓶或烧杯、25 ± 1 ℃恒温环境和干燥器。"}])

    add_heading(doc, "2.2.4 样品准备", 3)
    for sentence in [
        "从各组冻干凝胶中切取尺寸和厚度尽量一致的样品。",
        "每组取3个平行样并分别编号。",
        "将样品置于干燥器中平衡24 h。",
        "用分析天平称量每个样品的初始干质量，记为W0。",
    ]:
        add_body_paragraph(doc, [{"text": sentence}])

    add_heading(doc, "2.2.5 浸泡条件", 3)
    for sentence in [
        "每个样品单独放入一个编号容器中。",
        "每个容器加入20 mL蒸馏水。",
        "于25 ± 1 ℃条件下静置浸泡。",
        "设定测定时间点为0.5、1、2、4、8和24 h。",
    ]:
        add_body_paragraph(doc, [{"text": sentence}])

    add_heading(doc, "2.2.6 连续称重步骤", 3)
    steps = [
        "用镊子轻轻取出样品。",
        "用同型号滤纸轻轻吸去样品表面自由水，两面各轻触1次，总接触时间控制在5~10 s，不可按压或挤压样品。",
        "立即称量样品湿质量，记为该时间点的Wt。",
        "称量后立刻将样品放回原容器中，继续浸泡至下一时间点。",
        "记录样品外观变化，如是否软化、边缘碎裂或体积明显膨大等。",
    ]
    for sentence in steps:
        add_body_paragraph(doc, [{"text": sentence}])
    add_body_paragraph(
        doc,
        [
            {"text": "采用重量法测定明胶水凝胶溶胀并在称量前去除表面液体，是已有研究中常见的处理方式；mTG交联明胶海绵或支架的吸水测定也采用了类似思路"},
            {"text": "[1]", "superscript": True},
            {"text": "[7]", "superscript": True},
            {"text": "。"},
        ],
    )

    add_heading(doc, "2.2.7 计算公式", 3)
    add_formula(doc, "SR(%) = (Wt - W0) / W0 × 100")
    add_body_paragraph(doc, [{"text": "式中，W0为样品初始干质量，Wt为样品在某一时间点的湿质量。"}])

    add_heading(doc, "2.2.8 结果表示", 3)
    add_body_paragraph(doc, [{"text": "每组设置3个平行样，结果以平均值 ± 标准差表示。"}])

    add_heading(doc, "2.2.9 注意事项", 3)
    for sentence in [
        "每次吸去表面水的动作必须尽量一致，否则误差较大。",
        "每次从滤纸到天平的称量速度应尽可能快，避免样品在空气中继续失水或吸湿。",
        "每个样品始终放回原来的容器中，不能混淆。",
        "若发现某组样品在前30 min吸水变化较快，可在预实验中补加10 min和20 min两个时间点。",
    ]:
        add_body_paragraph(doc, [{"text": sentence}])

    add_heading(doc, "2.3 溶失率试验", 2)
    add_heading(doc, "2.3.1 实验目的", 3)
    add_body_paragraph(doc, [{"text": "测定不同改性处理后的鹿皮明胶冻干凝胶在水中浸泡24 h后的质量损失情况，用于评价样品的水中稳定性。"}])

    add_heading(doc, "2.3.2 实验分组", 3)
    add_body_paragraph(doc, [{"text": "溶失率试验同样设置6组，每组3个平行样，共需18个样品。"}])
    add_formula(doc, "6组 × 3个平行 = 18个样品")

    add_heading(doc, "2.3.3 试剂与仪器", 3)
    add_body_paragraph(doc, [{"text": "试剂与仪器包括蒸馏水或去离子水、分析天平、中速定量滤纸或Whatman 1号滤纸、镊子、样品瓶或烧杯、105 ℃鼓风干燥箱和干燥器。"}])

    add_heading(doc, "2.3.4 样品准备", 3)
    for sentence in [
        "从各组冻干凝胶中切取尺寸和厚度尽量一致的样品。",
        "每组取3个平行样并编号。",
        "样品置于干燥器中平衡24 h后，称量初始干质量，记为W0。",
    ]:
        add_body_paragraph(doc, [{"text": sentence}])

    add_heading(doc, "2.3.5 浸泡步骤", 3)
    for sentence in [
        "将每个样品单独置于盛有20 mL蒸馏水的编号容器中。",
        "于25 ± 1 ℃条件下静置浸泡24 h。",
    ]:
        add_body_paragraph(doc, [{"text": sentence}])

    add_heading(doc, "2.3.6 残余物回收步骤", 3)
    for sentence in [
        "24 h后，将样品连同浸泡液一起倒入预先称量好的滤纸中过滤。",
        "用少量蒸馏水冲洗容器内壁1~2次，以收集附着碎屑。",
        "将滤纸连同未溶解残余物一起转移至干燥皿中。",
        "若样品仍保持完整，也建议过滤回收，不要只夹取大块残余物，以免遗漏碎屑。",
    ]:
        add_body_paragraph(doc, [{"text": sentence}])

    add_heading(doc, "2.3.7 干燥至恒重", 3)
    for sentence in [
        "将滤纸和残余物置于105 ℃干燥箱中干燥。",
        "先干燥4 h，取出后放入干燥器冷却30 min，再称重。",
        "继续每次补干1 h，冷却后复称。",
        "当前后两次称量差值≤2 mg时，视为恒重。",
        "扣除滤纸质量后，得到样品残余干质量，记为Wr。",
    ]:
        add_body_paragraph(doc, [{"text": sentence}])

    add_heading(doc, "2.3.8 计算公式", 3)
    add_formula(doc, "WL(%) = (W0 - Wr) / W0 × 100")
    add_body_paragraph(doc, [{"text": "式中，W0为样品初始干质量，Wr为浸泡后残余物的恒重干质量。"}])

    add_heading(doc, "2.3.9 结果表示", 3)
    add_body_paragraph(doc, [{"text": "每组设置3个平行样，结果以平均值 ± 标准差表示。"}])

    add_heading(doc, "2.3.10 注意事项", 3)
    add_body_paragraph(
        doc,
        [
            {"text": "若样品在水中碎裂或局部胶化，必须通过过滤完整回收残余物。24 h浸泡后回收未溶物，再于105 ℃干燥并按干基计算质量损失或水溶性的做法，在明胶基膜和其他蛋白基材料中较为常见"},
            {"text": "[5]", "superscript": True},
            {"text": "[6]", "superscript": True},
            {"text": "。"},
        ],
    )

    add_heading(doc, "2.4 样品总数汇总", 2)
    add_heading(doc, "2.4.1 溶胀率试验", 3)
    add_formula(doc, "6组 × 3个平行 = 18个样品")
    add_heading(doc, "2.4.2 溶失率试验", 3)
    add_formula(doc, "6组 × 3个平行 = 18个样品")
    add_heading(doc, "2.4.3 总样品数", 3)
    add_formula(doc, "18 + 18 = 36个样品")

    add_heading(doc, "2.5 可直接用于论文的方法学表述", 2)
    add_heading(doc, "2.5.1 溶胀率测定", 3)
    add_body_paragraph(
        doc,
        [
            {"text": "采用重量法测定鹿皮明胶冻干凝胶的溶胀率。将冻干样品切割成尺寸基本一致的试样，于干燥器中平衡24 h后称取其初始干质量，记为W0。每组设置3个平行样，分别置于盛有20 mL蒸馏水的密闭容器中，于25 ± 1 ℃条件下静置浸泡。在0.5、1、2、4、8和24 h取出样品，用滤纸轻轻吸去表面自由水后立即称量湿质量，记为Wt，随后将样品放回原容器继续浸泡。溶胀率按下式计算："},
        ],
    )
    add_formula(doc, "SR(%) = (Wt - W0) / W0 × 100")
    add_body_paragraph(doc, [{"text": "每组测定3个平行样，结果以平均值 ± 标准差表示。"}])

    add_heading(doc, "2.5.2 溶失率测定", 3)
    add_body_paragraph(
        doc,
        [
            {"text": "采用浸泡后残余干质量法测定样品溶失率。准确称取冻干凝胶样品初始干质量，记为W0。将样品置于20 mL蒸馏水中，于25 ± 1 ℃条件下浸泡24 h。浸泡结束后，将样品体系经预称量滤纸过滤回收未溶解残余物，并用少量蒸馏水冲洗容器内壁以收集附着碎屑。随后将残余物连同滤纸置于105 ℃干燥箱中烘干至恒重，冷却至室温后称重，扣除滤纸质量后记为残余干质量Wr。溶失率按下式计算："},
        ],
    )
    add_formula(doc, "WL(%) = (W0 - Wr) / W0 × 100")
    add_body_paragraph(doc, [{"text": "每组设置3个平行样，结果以平均值 ± 标准差表示。"}])

    add_heading(doc, "2.6 相关参考文献", 2)
    add_body_paragraph(doc, [{"text": "下面列出与明胶水凝胶、DAS或TGase改性以及溶胀或水稳定性评价密切相关、并可直接支撑本节实验方法的文献。"}])

    add_heading(doc, "参考文献", 1)
    references = [
        "Skopinska-Wisniewska J, Tuszynska M, Olewnik-Kruszkowska E. Comparative Study of Gelatin Hydrogels Modified by Various Cross-Linking Agents[J]. Materials, 2021, 14(2): 396.",
        "Cui T, Sun Y, Wu Y, Wang J, Ding Y, Cheng J, Guo M. Mechanical, Microstructural, and Rheological Characterization of Gelatin-Dialdehyde Starch Hydrogels Constructed by Dual Dynamic Crosslinking[J]. LWT, 2022, 161: 113374.",
        "Liu F, Majeed H, Antoniou J, Li Y, Ma Y, Yokoyama W H, Ma J, Zhong F. Tailoring Physical Properties of Transglutaminase-Modified Gelatin Films by Varying Drying Temperature[J]. Food Hydrocolloids, 2016, 58: 20-28.",
        "Ma Y, Yang R, Zhao W. Innovative Water-Insoluble Edible Film Based on Biocatalytic Crosslink of Gelatin Rich in Glutamine[J]. Foods, 2020, 9(4): 503.",
        "Tessaro L, Carvalho S M, et al. Improving the Properties of Gelatin-Based Films by the Addition of Broth from Chicken and Bovine Bones[J]. Polymers, 2024, 16: 1707.",
        "Ciurzynska A, et al. Development and Characteristics of Protein Edible Film Derived from Pork Gelatin and Beef Broth[J]. Foods, 2024, 13: 1198.",
        "Long H, Zhao Y, et al. Preparation and Characteristics of Gelatin Sponges Crosslinked by Microbial Transglutaminase[J]. PeerJ, 2017, 5: e3665.",
    ]
    for idx, ref in enumerate(references, start=1):
        add_reference(doc, idx, ref)

    add_heading(doc, "2.7 文献使用建议", 2)
    add_body_paragraph(doc, [{"text": "在方法学部分，溶胀率测定建议以文献[1]为主要依据；在结果讨论中，DAS改性部分建议引用文献[2]，TGase改性部分建议引用文献[3]和文献[4]；溶失率方法学部分建议引用文献[5]或文献[6]；若需强调冻干凝胶或海绵形态材料与本研究样品形态相近，则可补充引用文献[7]。"}])

    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
