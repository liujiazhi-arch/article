#!/usr/bin/env python3

from __future__ import annotations

import argparse
import math
import statistics
import subprocess
import zipfile
from html import escape
from pathlib import Path

from openpyxl import load_workbook


GROUPS = ["G1", "G2", "G3", "G4"]
COLORS = ["#4E79A7", "#59A14F", "#F28E2B", "#E15759"]
TEXT_FONT = "Helvetica, Arial, sans-serif"
PRACTICE_DIR_DEFAULT = "/Users/apple/Desktop/gelatin_practice_data"
OUTDIR_DEFAULT = "/Users/apple/Desktop/gelatin_figures_from_excel"

BAR_CONFIG = [
    {
        "file": "01_WHC_SIMULATED_PRACTICE_ONLY.xlsx",
        "stem": "Fig3-1_WHC",
        "title": "Figure 3-1 Water-Holding Capacity",
        "ylabel": "WHC (g/g)",
    },
    {
        "file": "02_OHC_SIMULATED_PRACTICE_ONLY.xlsx",
        "stem": "Fig3-2_OHC",
        "title": "Figure 3-2 Oil-Holding Capacity",
        "ylabel": "OHC (g/g)",
    },
    {
        "file": "03_EAI_SIMULATED_PRACTICE_ONLY.xlsx",
        "stem": "Fig3-3_EAI",
        "title": "Figure 3-3 Emulsifying Activity Index",
        "ylabel": "EAI (m^2/g)",
    },
    {
        "file": "04_ESI_SIMULATED_PRACTICE_ONLY.xlsx",
        "stem": "Fig3-4_ESI",
        "title": "Figure 3-4 Emulsifying Stability Index",
        "ylabel": "ESI (min)",
    },
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate figures from synthetic practice Excel files.")
    parser.add_argument("--practice-dir", default=PRACTICE_DIR_DEFAULT, help="Directory containing practice Excel files.")
    parser.add_argument("--outdir", default=OUTDIR_DEFAULT, help="Directory for exported figures and Prism guide.")
    return parser.parse_args()


def nice_step(raw_step: float) -> float:
    if raw_step <= 0:
        return 1.0
    power = 10 ** math.floor(math.log10(raw_step))
    scaled = raw_step / power
    if scaled <= 1:
        base = 1
    elif scaled <= 2:
        base = 2
    elif scaled <= 5:
        base = 5
    else:
        base = 10
    return base * power


def make_ticks(ymin: float, ymax: float, target_count: int = 5) -> list[float]:
    span = max(ymax - ymin, 1e-9)
    step = nice_step(span / target_count)
    start = math.floor(ymin / step) * step
    stop = math.ceil(ymax / step) * step
    ticks: list[float] = []
    value = start
    while value <= stop + step * 0.5:
        ticks.append(round(value, 8))
        value += step
    return ticks


def format_tick(value: float) -> str:
    if abs(value - round(value)) < 1e-8:
        return str(int(round(value)))
    if abs(value) >= 10:
        return f"{value:.1f}"
    return f"{value:.2f}".rstrip("0").rstrip(".")


def svg_document(width: int, height: int, body: str) -> str:
    return f"""<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">
<rect width="100%" height="100%" fill="#ffffff"/>
<style>
text {{ font-family: {TEXT_FONT}; fill: #1f2933; }}
.axis {{ stroke: #2d3748; stroke-width: 2; }}
.grid {{ stroke: #d7dee8; stroke-width: 1; }}
.note {{ fill: #4a5568; }}
.band {{ fill: #edf2f7; opacity: 0.85; }}
.dash {{ stroke-dasharray: 7 6; }}
</style>
{body}
</svg>
"""


def save_svg_and_png(outdir: Path, stem: str, svg_text: str) -> tuple[Path, Path]:
    svg_path = outdir / f"{stem}.svg"
    png_path = outdir / f"{stem}.png"
    svg_path.write_text(svg_text, encoding="utf-8")
    subprocess.run(
        ["sips", "-s", "format", "png", str(svg_path), "--out", str(png_path)],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    return svg_path, png_path


def line_points(xs: list[float], ys: list[float], x_map, y_map) -> str:
    return " ".join(f"{x_map(x):.2f},{y_map(y):.2f}" for x, y in zip(xs, ys))


def load_bar_data(path: Path) -> dict[str, object]:
    wb = load_workbook(path, data_only=True)
    ws = wb["Prism_Raw"]
    groups = {}
    means = []
    sds = []
    letters = []
    for col_idx, group in enumerate(GROUPS, start=2):
        values = [float(ws.cell(row, col_idx).value) for row in (4, 5, 6)]
        groups[group] = values
        means.append(statistics.mean(values))
        sds.append(statistics.stdev(values))
        letters.append(str(ws.cell(10, col_idx).value))
    return {"raw": groups, "means": means, "sds": sds, "letters": letters}


def load_xy_group_mean_from_raw(path: Path, sheet_name: str) -> tuple[list[float], dict[str, list[float]]]:
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet_name]
    xs: list[float] = []
    ys: dict[str, list[float]] = {group: [] for group in GROUPS}
    row = 4
    while ws.cell(row, 1).value is not None:
        xs.append(float(ws.cell(row, 1).value))
        for group_idx, group in enumerate(GROUPS):
            start_col = 2 + group_idx * 3
            group_values = [float(ws.cell(row, start_col + offset).value) for offset in range(3)]
            ys[group].append(statistics.mean(group_values))
        row += 1
    return xs, ys


def load_summary_sheet(path: Path, sheet_name: str) -> dict[str, list[float]]:
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet_name]
    data: dict[str, list[float]] = {}
    for col_idx, group in enumerate(GROUPS, start=2):
        data[group] = [float(ws.cell(row, col_idx).value) for row in (4, 5, 6)]
    return data


def draw_bar_chart(title: str, ylabel: str, data: dict[str, object]) -> str:
    width, height = 1440, 1100
    left, right, top, bottom = 150, 80, 125, 240
    plot_w = width - left - right
    plot_h = height - top - bottom
    means = data["means"]
    sds = data["sds"]
    letters = data["letters"]
    raw = data["raw"]
    y0 = 0.0
    ymax = max(mean + sd for mean, sd in zip(means, sds))
    y1 = math.ceil((ymax * 1.2) / nice_step((ymax * 1.2) / 5)) * nice_step((ymax * 1.2) / 5)
    ticks = make_ticks(y0, y1)

    def x_map(index: int) -> float:
        return left + plot_w * (index + 0.5) / len(GROUPS)

    def y_map(value: float) -> float:
        return top + (y1 - value) * plot_h / (y1 - y0)

    parts: list[str] = []
    parts.append(f'<text x="{width / 2:.1f}" y="62" text-anchor="middle" font-size="35" font-weight="700">{escape(title)}</text>')
    parts.append(
        f'<text x="{width / 2:.1f}" y="96" text-anchor="middle" font-size="19" class="note">Graph rebuilt directly from the current Excel practice data source.</text>'
    )
    parts.append(f'<line class="axis" x1="{left}" y1="{top}" x2="{left}" y2="{top + plot_h}"/>')
    parts.append(f'<line class="axis" x1="{left}" y1="{top + plot_h}" x2="{left + plot_w}" y2="{top + plot_h}"/>')

    for tick in ticks:
        y = y_map(tick)
        parts.append(f'<line class="grid" x1="{left}" y1="{y:.2f}" x2="{left + plot_w}" y2="{y:.2f}"/>')
        parts.append(f'<text x="{left - 18}" y="{y + 8:.2f}" text-anchor="end" font-size="24">{format_tick(tick)}</text>')

    parts.append(
        f'<text x="56" y="{top + plot_h / 2:.1f}" transform="rotate(-90 56 {top + plot_h / 2:.1f})" text-anchor="middle" font-size="30">{escape(ylabel)}</text>'
    )
    parts.append(f'<text x="{width / 2:.1f}" y="{height - 144}" text-anchor="middle" font-size="27">Treatment group</text>')
    parts.append(
        f'<text x="{width / 2:.1f}" y="{height - 88}" text-anchor="middle" font-size="20" class="note">Bars = mean, error bars = SD, dots = three practice replicates, letters indicate the intended significance grouping.</text>'
    )

    bar_w = plot_w / len(GROUPS) * 0.4
    dot_offsets = [-bar_w * 0.18, 0.0, bar_w * 0.18]
    for idx, group in enumerate(GROUPS):
        mean_value = means[idx]
        sd_value = sds[idx]
        center_x = x_map(idx)
        rect_x = center_x - bar_w / 2
        rect_y = y_map(mean_value)
        base_y = y_map(y0)
        error_top = y_map(mean_value + sd_value)
        error_bottom = y_map(max(y0, mean_value - sd_value))
        parts.append(
            f'<rect x="{rect_x:.2f}" y="{rect_y:.2f}" width="{bar_w:.2f}" height="{base_y - rect_y:.2f}" rx="6" fill="{COLORS[idx]}" opacity="0.92" stroke="#ffffff" stroke-width="2"/>'
        )
        parts.append(f'<line x1="{center_x:.2f}" y1="{error_top:.2f}" x2="{center_x:.2f}" y2="{error_bottom:.2f}" stroke="#263238" stroke-width="2.3"/>')
        parts.append(f'<line x1="{center_x - 14:.2f}" y1="{error_top:.2f}" x2="{center_x + 14:.2f}" y2="{error_top:.2f}" stroke="#263238" stroke-width="2.3"/>')
        parts.append(f'<line x1="{center_x - 14:.2f}" y1="{error_bottom:.2f}" x2="{center_x + 14:.2f}" y2="{error_bottom:.2f}" stroke="#263238" stroke-width="2.3"/>')
        for rep_offset, rep_value in zip(dot_offsets, raw[group]):
            parts.append(
                f'<circle cx="{center_x + rep_offset:.2f}" cy="{y_map(rep_value):.2f}" r="7" fill="#ffffff" stroke="#22303c" stroke-width="2"/>'
            )
        parts.append(f'<text x="{center_x:.2f}" y="{base_y + 44:.2f}" text-anchor="middle" font-size="28">{group}</text>')
        parts.append(f'<text x="{center_x:.2f}" y="{max(top + 8, error_top - 42):.2f}" text-anchor="middle" font-size="28" font-weight="700">{letters[idx]}</text>')
        parts.append(f'<text x="{center_x:.2f}" y="{max(top + 32, error_top - 16):.2f}" text-anchor="middle" font-size="20" class="note">{mean_value:.2f} +/- {sd_value:.2f}</text>')

    return svg_document(width, height, "\n".join(parts))


def draw_uv_chart(xs: list[float], ys: dict[str, list[float]], peak_230: dict[str, list[float]], peak_280: dict[str, list[float]]) -> str:
    width, height = 1500, 1060
    left, right, top, bottom = 125, 85, 125, 190
    plot_w = width - left - right
    plot_h = height - top - bottom
    x0, x1 = min(xs), max(xs)
    y0 = 0.0
    ymax = max(max(values) for values in ys.values())
    y1 = math.ceil((ymax * 1.12) / nice_step((ymax * 1.12) / 5)) * nice_step((ymax * 1.12) / 5)

    def x_map(value: float) -> float:
        return left + (value - x0) * plot_w / (x1 - x0)

    def y_map(value: float) -> float:
        return top + (y1 - value) * plot_h / (y1 - y0)

    parts: list[str] = []
    parts.append('<text x="750" y="62" text-anchor="middle" font-size="35" font-weight="700">Figure 3-5 UV Absorption Spectra</text>')
    parts.append('<text x="750" y="96" text-anchor="middle" font-size="19" class="note">The displayed traces are the group means calculated from the Excel XY practice tables.</text>')
    for start, end, label in [(224, 232, "around 230 nm"), (274, 286, "around 280 nm")]:
        band_x = x_map(start)
        band_w = x_map(end) - x_map(start)
        parts.append(f'<rect class="band" x="{band_x:.2f}" y="{top}" width="{band_w:.2f}" height="{plot_h}"/>')
        parts.append(f'<text x="{band_x + band_w / 2:.2f}" y="{top + 26}" text-anchor="middle" font-size="17" class="note">{label}</text>')
    parts.append(f'<line class="axis" x1="{left}" y1="{top}" x2="{left}" y2="{top + plot_h}"/>')
    parts.append(f'<line class="axis" x1="{left}" y1="{top + plot_h}" x2="{left + plot_w}" y2="{top + plot_h}"/>')
    for tick in [200, 240, 280, 320, 360, 400]:
        xt = x_map(tick)
        parts.append(f'<line class="grid" x1="{xt:.2f}" y1="{top}" x2="{xt:.2f}" y2="{top + plot_h}"/>')
        parts.append(f'<text x="{xt:.2f}" y="{top + plot_h + 42}" text-anchor="middle" font-size="26">{tick}</text>')
    for tick in make_ticks(y0, y1):
        yt = y_map(tick)
        parts.append(f'<line class="grid" x1="{left}" y1="{yt:.2f}" x2="{left + plot_w}" y2="{yt:.2f}"/>')
        parts.append(f'<text x="{left - 18}" y="{yt + 8:.2f}" text-anchor="end" font-size="24">{format_tick(tick)}</text>')
    parts.append(f'<text x="{width / 2:.1f}" y="{height - 62}" text-anchor="middle" font-size="30">Wavelength (nm)</text>')
    parts.append(f'<text x="52" y="{top + plot_h / 2:.1f}" transform="rotate(-90 52 {top + plot_h / 2:.1f})" text-anchor="middle" font-size="30">Absorbance</text>')
    for x_value in (230, 280):
        x_pos = x_map(x_value)
        parts.append(f'<line class="dash" x1="{x_pos:.2f}" y1="{top}" x2="{x_pos:.2f}" y2="{top + plot_h}" stroke="#55606e" stroke-width="1.6"/>')
    mean_230 = {group: statistics.mean(values) for group, values in peak_230.items()}
    mean_280 = {group: statistics.mean(values) for group, values in peak_280.items()}
    top_group_230 = max(mean_230, key=mean_230.get)
    top_group_280 = max(mean_280, key=mean_280.get)
    parts.append(f'<text x="{x_map(231.5):.2f}" y="{y_map(y1 * 0.94):.2f}" font-size="20" class="note">{top_group_230} shows the highest A230 mean</text>')
    parts.append(f'<text x="{x_map(281.5):.2f}" y="{y_map(y1 * 0.58):.2f}" font-size="20" class="note">{top_group_280} shows the highest A280 mean</text>')
    for idx, group in enumerate(GROUPS):
        parts.append(f'<polyline points="{line_points(xs, ys[group], x_map, y_map)}" fill="none" stroke="{COLORS[idx]}" stroke-width="4.2"/>')
    legend_x, legend_y = width - 215, top + 54
    parts.append(f'<rect x="{legend_x - 18}" y="{legend_y - 34}" width="160" height="172" fill="#ffffff" opacity="0.92" stroke="#d8dee7"/>')
    for idx, group in enumerate(GROUPS):
        y_pos = legend_y + idx * 36
        parts.append(f'<line x1="{legend_x}" y1="{y_pos}" x2="{legend_x + 42}" y2="{y_pos}" stroke="{COLORS[idx]}" stroke-width="4.2"/>')
        parts.append(f'<text x="{legend_x + 54}" y="{y_pos + 8}" font-size="24">{group}</text>')
    return svg_document(width, height, "\n".join(parts))


def draw_ftir_chart(xs: list[float], ys: dict[str, list[float]], peak_summaries: dict[str, dict[str, list[float]]]) -> str:
    width, height = 1560, 1100
    left, right, top, bottom = 125, 90, 125, 220
    plot_w = width - left - right
    plot_h = height - top - bottom
    x0, x1 = max(xs), min(xs)
    y_min = min(min(values) for values in ys.values())
    y_max = max(max(values) for values in ys.values())
    y0 = math.floor((y_min - 2) / 5) * 5
    y1 = math.ceil((y_max + 2) / 5) * 5

    def x_map(value: float) -> float:
        return left + (x0 - value) * plot_w / (x0 - x1)

    def y_map(value: float) -> float:
        return top + (y1 - value) * plot_h / (y1 - y0)

    parts: list[str] = []
    parts.append('<text x="780" y="62" text-anchor="middle" font-size="35" font-weight="700">Figure 3-6 FTIR Spectra</text>')
    parts.append('<text x="780" y="96" text-anchor="middle" font-size="19" class="note">The displayed traces are group means rebuilt from the current Excel practice tables.</text>')
    bands = [
        (3340, 3230, "Amide A"),
        (1675, 1605, "Amide I"),
        (1575, 1515, "Amide II"),
        (1275, 1215, "Amide III"),
    ]
    for start, end, label in bands:
        band_left = x_map(start)
        band_width = x_map(end) - x_map(start)
        parts.append(f'<rect class="band" x="{band_left:.2f}" y="{top}" width="{band_width:.2f}" height="{plot_h}"/>')
        parts.append(f'<text x="{band_left + band_width / 2:.2f}" y="{top + 26}" text-anchor="middle" font-size="17" class="note">{label}</text>')
    parts.append(f'<line class="axis" x1="{left}" y1="{top}" x2="{left}" y2="{top + plot_h}"/>')
    parts.append(f'<line class="axis" x1="{left}" y1="{top + plot_h}" x2="{left + plot_w}" y2="{top + plot_h}"/>')
    for tick in [4000, 3200, 2400, 1600, 800, 400]:
        xt = x_map(tick)
        parts.append(f'<line class="grid" x1="{xt:.2f}" y1="{top}" x2="{xt:.2f}" y2="{top + plot_h}"/>')
        parts.append(f'<text x="{xt:.2f}" y="{top + plot_h + 42}" text-anchor="middle" font-size="26">{tick}</text>')
    for tick in make_ticks(y0, y1):
        yt = y_map(tick)
        parts.append(f'<line class="grid" x1="{left}" y1="{yt:.2f}" x2="{left + plot_w}" y2="{yt:.2f}"/>')
        parts.append(f'<text x="{left - 18}" y="{yt + 8:.2f}" text-anchor="end" font-size="24">{format_tick(tick)}</text>')
    parts.append(f'<text x="{width / 2:.1f}" y="{height - 98}" text-anchor="middle" font-size="30">Wavenumber (cm^-1)</text>')
    parts.append(f'<text x="52" y="{top + plot_h / 2:.1f}" transform="rotate(-90 52 {top + plot_h / 2:.1f})" text-anchor="middle" font-size="30">Transmittance (%)</text>')
    parts.append(f'<text x="{width / 2:.1f}" y="{height - 58}" text-anchor="middle" font-size="19" class="note">The x-axis should be reversed in Prism so the spectra run from 4000 to 400 cm^-1.</text>')
    for idx, group in enumerate(GROUPS):
        parts.append(f'<polyline points="{line_points(xs, ys[group], x_map, y_map)}" fill="none" stroke="{COLORS[idx]}" stroke-width="4.1"/>')
    peak_labels = [
        (3300, "Amide A", peak_summaries["Amide_A_Pos"]),
        (1645, "Amide I", peak_summaries["Amide_I_Pos"]),
        (1545, "Amide II", peak_summaries["Amide_II_Pos"]),
        (1245, "Amide III", peak_summaries["Amide_III_Pos"]),
    ]
    for x_value, short_label, sheet_data in peak_labels:
        x_pos = x_map(x_value)
        highest_group = max(sheet_data, key=lambda group: statistics.mean(sheet_data[group]))
        parts.append(f'<line class="dash" x1="{x_pos:.2f}" y1="{top}" x2="{x_pos:.2f}" y2="{top + plot_h}" stroke="#55606e" stroke-width="1.5"/>')
        parts.append(f'<text x="{x_pos + 14:.2f}" y="{top + plot_h - 12:.2f}" font-size="18" transform="rotate(-90 {x_pos + 14:.2f} {top + plot_h - 12:.2f})" class="note">{short_label}</text>')
        parts.append(f'<text x="{x_pos + 30:.2f}" y="{top + 64:.2f}" font-size="15" transform="rotate(-90 {x_pos + 30:.2f} {top + 64:.2f})" class="note">{highest_group} has the highest mean position in this band</text>')
    legend_x, legend_y = left + 28, top + plot_h - 135
    parts.append(f'<rect x="{legend_x - 22}" y="{legend_y - 30}" width="150" height="168" fill="#ffffff" opacity="0.92" stroke="#d8dee7"/>')
    for idx, group in enumerate(GROUPS):
        y_pos = legend_y + idx * 36
        parts.append(f'<line x1="{legend_x}" y1="{y_pos}" x2="{legend_x + 42}" y2="{y_pos}" stroke="{COLORS[idx]}" stroke-width="4.1"/>')
        parts.append(f'<text x="{legend_x + 54}" y="{y_pos + 8}" font-size="24">{group}</text>')
    return svg_document(width, height, "\n".join(parts))


def write_prism_guide(outdir: Path) -> Path:
    guide = """# Prism 绘图说明

## 文件定位
- 柱状图数据来自桌面 `gelatin_practice_data` 目录中的 `01_WHC` 到 `04_ESI` 工作簿。
- UV 数据来自 `05_UV_SIMULATED_PRACTICE_ONLY.xlsx`。
- FTIR 数据来自 `06_FTIR_SIMULATED_PRACTICE_ONLY.xlsx`。

## Figure 3-1 WHC
- Excel 文件：`01_WHC_SIMULATED_PRACTICE_ONLY.xlsx`
- Prism 数据表：选择 `Column`
- 录入方式：把 `Prism_Raw` 工作表中 `G1` 到 `G4` 四列的 3 个重复值分别贴到 Prism 四列中
- 图类型：`Column scatter` 或 `Bar with individual values`
- 误差线：选择 `Mean with SD`
- 建议设置：显示每个组的 3 个点，柱上再保留均值和 SD
- 显著性字母：在 Prism 中手动插入文本，按 `b, c, a, a`
- 纵轴标题：`WHC (g/g)`

## Figure 3-2 OHC
- Excel 文件：`02_OHC_SIMULATED_PRACTICE_ONLY.xlsx`
- Prism 数据表：`Column`
- 录入方式：使用 `Prism_Raw`
- 图类型：`Column scatter` 或 `Bar with individual values`
- 误差线：`Mean with SD`
- 显著性字母：手动插入 `b, c, a, a`
- 纵轴标题：`OHC (g/g)`

## Figure 3-3 EAI
- Excel 文件：`03_EAI_SIMULATED_PRACTICE_ONLY.xlsx`
- Prism 数据表：`Column`
- 录入方式：使用 `Prism_Raw`
- 图类型：`Column scatter` 或 `Bar with individual values`
- 误差线：`Mean with SD`
- 显著性字母：手动插入 `c, d, a, b`
- 纵轴标题：`EAI (m^2/g)`

## Figure 3-4 ESI
- Excel 文件：`04_ESI_SIMULATED_PRACTICE_ONLY.xlsx`
- Prism 数据表：`Column`
- 录入方式：使用 `Prism_Raw`
- 图类型：`Column scatter` 或 `Bar with individual values`
- 误差线：`Mean with SD`
- 显著性字母：手动插入 `c, d, b, a`
- 纵轴标题：`ESI (min)`

## Figure 3-5 UV
- Excel 文件：`05_UV_SIMULATED_PRACTICE_ONLY.xlsx`
- 推荐导入工作表：`UV_XY_GroupMean`
- Prism 数据表：选择 `XY`
- 录入方式：第一列是 `Wavelength (nm)`，后面四列分别是 `G1_Mean` 到 `G4_Mean`
- 图类型：`Lines`
- 符号：关闭点符号，只保留线条
- 线宽：建议 1.5 到 2.0 pt
- 纵轴标题：`Absorbance`
- 横轴标题：`Wavelength (nm)`
- 标注建议：在 230 nm 和 280 nm 位置各加一条竖向虚线，并手动标注 `~230 nm`、`~280 nm`
- 如果你想在 Prism 中显示组内波动：可以改用 `UV_XY_Raw`，每组 3 条重复曲线分别导入

## Figure 3-6 FTIR
- Excel 文件：`06_FTIR_SIMULATED_PRACTICE_ONLY.xlsx`
- 推荐导入工作表：`FTIR_XY_GroupMean`
- Prism 数据表：选择 `XY`
- 录入方式：第一列是 `Wavenumber (cm^-1)`，后面四列分别是 `G1_Mean` 到 `G4_Mean`
- 图类型：`Lines`
- 符号：关闭点符号，只保留线条
- 纵轴标题：`Transmittance (%)`
- 横轴标题：`Wavenumber (cm^-1)`
- 关键设置：在 Prism 的 x 轴设置里启用 `Reverse axis order`，让横轴从 4000 递减到 400
- 标注建议：手动加四条虚线并标 `Amide A`、`Amide I`、`Amide II`、`Amide III`
- 如果你想保留每组 3 条练习曲线：可以改用 `FTIR_XY_Raw`

## 导图顺序建议
1. 先导入四个 `Column` 图，统一配色和字体。
2. 再导入 UV 和 FTIR 的 `XY` 图。
3. 最后统一调整字号、线宽、图例位置和显著性字母。

## 说明
- 当前这套表和图来自 `SIMULATED_PRACTICE_ONLY` 练习数据。
- 这些文件适合 Prism 演练、版式确认和预实验图形推演，不应表述为真实原始实验数据。
"""
    path = outdir / "PRISM_绘图说明.md"
    path.write_text(guide, encoding="utf-8")
    return path


def create_zip(outdir: Path, png_paths: list[Path]) -> Path:
    zip_path = outdir / "gelatin_figures_from_excel_png.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zip_file:
        for png_path in sorted(png_paths):
            zip_file.write(png_path, arcname=png_path.name)
    return zip_path


def main() -> None:
    args = parse_args()
    practice_dir = Path(args.practice_dir).expanduser().resolve()
    outdir = Path(args.outdir).expanduser().resolve()
    outdir.mkdir(parents=True, exist_ok=True)

    png_paths: list[Path] = []

    for config in BAR_CONFIG:
        data = load_bar_data(practice_dir / config["file"])
        _, png_path = save_svg_and_png(outdir, config["stem"], draw_bar_chart(config["title"], config["ylabel"], data))
        png_paths.append(png_path)

    uv_path = practice_dir / "05_UV_SIMULATED_PRACTICE_ONLY.xlsx"
    uv_xs, uv_ys = load_xy_group_mean_from_raw(uv_path, "UV_XY_Raw")
    peak_230 = load_summary_sheet(uv_path, "PeakAbs_230")
    peak_280 = load_summary_sheet(uv_path, "PeakAbs_280")
    _, png_path = save_svg_and_png(outdir, "Fig3-5_UV", draw_uv_chart(uv_xs, uv_ys, peak_230, peak_280))
    png_paths.append(png_path)

    ftir_path = practice_dir / "06_FTIR_SIMULATED_PRACTICE_ONLY.xlsx"
    ftir_xs, ftir_ys = load_xy_group_mean_from_raw(ftir_path, "FTIR_XY_Raw")
    peak_summaries = {
        "Amide_A_Pos": load_summary_sheet(ftir_path, "Amide_A_Pos"),
        "Amide_I_Pos": load_summary_sheet(ftir_path, "Amide_I_Pos"),
        "Amide_II_Pos": load_summary_sheet(ftir_path, "Amide_II_Pos"),
        "Amide_III_Pos": load_summary_sheet(ftir_path, "Amide_III_Pos"),
    }
    _, png_path = save_svg_and_png(outdir, "Fig3-6_FTIR", draw_ftir_chart(ftir_xs, ftir_ys, peak_summaries))
    png_paths.append(png_path)

    guide_path = write_prism_guide(outdir)
    zip_path = create_zip(outdir, png_paths)

    print("Generated files:")
    for path in sorted(outdir.iterdir()):
        print(path)
    print("\nGuide:")
    print(guide_path)
    print("\nZIP:")
    print(zip_path)


if __name__ == "__main__":
    main()
