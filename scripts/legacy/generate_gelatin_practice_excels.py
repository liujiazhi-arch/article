#!/usr/bin/env python3

from __future__ import annotations

import argparse
import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


GROUPS = ["G1", "G2", "G3", "G4"]
REPLICATE_LABELS = ["R1", "R2", "R3"]
OUTDIR_DEFAULT = "/Users/apple/Desktop/gelatin_practice_data"

BAR_EXPERIMENTS = [
    {
        "id": "01",
        "stem": "WHC",
        "title": "Water-Holding Capacity",
        "unit": "g/g",
        "mean": [6.34, 5.61, 7.46, 7.29],
        "sd": [0.20, 0.17, 0.23, 0.18],
        "letters": ["b", "c", "a", "a"],
    },
    {
        "id": "02",
        "stem": "OHC",
        "title": "Oil-Holding Capacity",
        "unit": "g/g",
        "mean": [3.63, 3.17, 4.02, 3.95],
        "sd": [0.11, 0.10, 0.12, 0.10],
        "letters": ["b", "c", "a", "a"],
    },
    {
        "id": "03",
        "stem": "EAI",
        "title": "Emulsifying Activity Index",
        "unit": "m^2/g",
        "mean": [28.63, 24.76, 33.85, 32.14],
        "sd": [0.87, 0.81, 0.96, 0.91],
        "letters": ["c", "d", "a", "b"],
    },
    {
        "id": "04",
        "stem": "ESI",
        "title": "Emulsifying Stability Index",
        "unit": "min",
        "mean": [22.36, 18.42, 27.54, 29.83],
        "sd": [0.71, 0.63, 0.79, 0.84],
        "letters": ["c", "d", "b", "a"],
    },
]

UV_PARAMS = {
    "G1": {"p1": 225.2, "a1": 0.92, "p2": 279.2, "a2": 0.25},
    "G2": {"p1": 223.8, "a1": 0.83, "p2": 278.5, "a2": 0.21},
    "G3": {"p1": 228.1, "a1": 1.07, "p2": 280.3, "a2": 0.31},
    "G4": {"p1": 227.4, "a1": 1.01, "p2": 279.8, "a2": 0.29},
}

FTIR_PEAKS = {
    "G1": {"A": 3291.8, "I": 1637.5, "II": 1543.2, "III": 1242.0},
    "G2": {"A": 3297.4, "I": 1633.2, "II": 1538.6, "III": 1237.8},
    "G3": {"A": 3283.6, "I": 1644.1, "II": 1550.4, "III": 1248.3},
    "G4": {"A": 3285.1, "I": 1641.8, "II": 1548.1, "III": 1246.5},
}

ASYMMETRY_R = [0.35, 0.62, 1.28, 0.78, 0.46, 1.14, 0.57, 1.36]
UV_REP_VARIATIONS = [
    {"shift1": -0.7, "amp1": 0.97, "shift2": -0.5, "amp2": 0.93, "base": -0.003},
    {"shift1": 0.0, "amp1": 1.00, "shift2": 0.0, "amp2": 1.00, "base": 0.000},
    {"shift1": 0.8, "amp1": 1.03, "shift2": 0.6, "amp2": 1.06, "base": 0.003},
]
FTIR_REP_VARIATIONS = [
    {"peak_shift": -2.1, "depth": 0.97, "base": -0.18},
    {"peak_shift": 0.0, "depth": 1.00, "base": 0.00},
    {"peak_shift": 1.8, "depth": 1.04, "base": 0.16},
]

RED_FILL = PatternFill(fill_type="solid", fgColor="FFC7CE")
YELLOW_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")
TITLE_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate practice-only synthetic Excel files for Prism.")
    parser.add_argument("--outdir", default=OUTDIR_DEFAULT, help="Output folder for Excel files.")
    return parser.parse_args()


def gauss(x: float, mu: float, amp: float, sigma: float) -> float:
    return amp * math.exp(-0.5 * ((x - mu) / sigma) ** 2)


def make_exact_replicates(mean_value: float, sd_value: float, asymmetry: float) -> list[float]:
    factor = math.sqrt(1 + asymmetry * asymmetry - asymmetry)
    a = sd_value / factor
    deltas = [-a, asymmetry * a, (1 - asymmetry) * a]
    return [round(mean_value + delta, 4) for delta in deltas]


def simulate_metric_table(means: list[float], sds: list[float], offset_seed: int) -> dict[str, list[float]]:
    table: dict[str, list[float]] = {}
    for idx, group in enumerate(GROUPS):
        asymmetry = ASYMMETRY_R[(offset_seed + idx) % len(ASYMMETRY_R)]
        table[group] = make_exact_replicates(means[idx], sds[idx], asymmetry)
    return table


def style_sheet_header(ws, title: str) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True)
    ws["A1"].fill = TITLE_FILL
    ws["A1"].alignment = Alignment(horizontal="left")


def write_readme_sheet(wb: Workbook, title: str, note_lines: list[str]) -> None:
    ws = wb.active
    ws.title = "README"
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 95
    style_sheet_header(ws, title)

    ws["A3"] = "Status"
    ws["B3"] = "SIMULATED_PRACTICE_ONLY"
    ws["A3"].font = Font(bold=True)
    ws["B3"].font = Font(bold=True, color="9C0006")
    ws["B3"].fill = RED_FILL

    for row, line in enumerate(note_lines, start=5):
        ws[f"A{row}"] = f"Note {row - 4}"
        ws[f"B{row}"] = line
        ws[f"A{row}"].font = Font(bold=True)
        ws[f"B{row}"].alignment = Alignment(wrap_text=True, vertical="top")


def add_prism_raw_sheet(wb: Workbook, sheet_name: str, unit: str, data: dict[str, list[float]], letters: list[str]) -> None:
    ws = wb.create_sheet(sheet_name)
    style_sheet_header(ws, f"{sheet_name} (Prism-ready raw table)")
    ws.freeze_panes = "A3"
    ws.column_dimensions["A"].width = 16
    for col_letter in ["B", "C", "D", "E"]:
        ws.column_dimensions[col_letter].width = 16

    ws["A3"] = f"Replicate ({unit})"
    ws["A3"].font = Font(bold=True)
    for idx, group in enumerate(GROUPS, start=2):
        cell = ws.cell(row=3, column=idx)
        cell.value = group
        cell.font = Font(bold=True)
        cell.fill = YELLOW_FILL

    for rep_idx, rep in enumerate(REPLICATE_LABELS, start=4):
        ws.cell(row=rep_idx, column=1).value = rep
        for col_idx, group in enumerate(GROUPS, start=2):
            ws.cell(row=rep_idx, column=col_idx).value = data[group][rep_idx - 4]

    summary_rows = {"Mean": 8, "SD": 9, "Letter": 10}
    for label, row_num in summary_rows.items():
        ws.cell(row=row_num, column=1).value = label
        ws.cell(row=row_num, column=1).font = Font(bold=True)

    for col_idx, group in enumerate(GROUPS, start=2):
        col_letter = chr(64 + col_idx)
        ws.cell(row=8, column=col_idx).value = f"=AVERAGE({col_letter}4:{col_letter}6)"
        ws.cell(row=9, column=col_idx).value = f"=STDEV.S({col_letter}4:{col_letter}6)"
        ws.cell(row=10, column=col_idx).value = letters[col_idx - 2]


def add_long_format_sheet(wb: Workbook, sheet_name: str, unit: str, data: dict[str, list[float]]) -> None:
    ws = wb.create_sheet(sheet_name)
    style_sheet_header(ws, f"{sheet_name} (long format)")
    headers = ["Group", "Replicate", f"Value ({unit})"]
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=3, column=col_idx).value = header
        ws.cell(row=3, column=col_idx).font = Font(bold=True)
        ws.cell(row=3, column=col_idx).fill = YELLOW_FILL
        ws.column_dimensions[chr(64 + col_idx)].width = 18

    row_num = 4
    for group in GROUPS:
        for rep_idx, value in enumerate(data[group], start=1):
            ws.cell(row=row_num, column=1).value = group
            ws.cell(row=row_num, column=2).value = f"R{rep_idx}"
            ws.cell(row=row_num, column=3).value = value
            row_num += 1


def create_bar_workbook(outdir: Path, exp: dict[str, object]) -> Path:
    wb = Workbook()
    title = f"{exp['id']}_{exp['stem']}_SIMULATED_PRACTICE_ONLY"
    readme_lines = [
        "This workbook contains synthetic practice data derived from group mean and SD values only.",
        "These values are for Prism workflow rehearsal, figure layout testing, and file-structure setup.",
        "They are not measured laboratory observations and should not be represented as real experimental raw data.",
        "The three replicates in each group were generated to reproduce the supplied mean and sample SD exactly.",
    ]
    write_readme_sheet(wb, title, readme_lines)

    raw_table = simulate_metric_table(exp["mean"], exp["sd"], int(exp["id"]))
    add_prism_raw_sheet(wb, "Prism_Raw", exp["unit"], raw_table, exp["letters"])
    add_long_format_sheet(wb, "Long_Format", exp["unit"], raw_table)

    summary = wb.create_sheet("Summary")
    style_sheet_header(summary, "Summary")
    headers = ["Group", f"Mean ({exp['unit']})", f"SD ({exp['unit']})", "Letter"]
    for col_idx, header in enumerate(headers, start=1):
        summary.cell(row=3, column=col_idx).value = header
        summary.cell(row=3, column=col_idx).font = Font(bold=True)
        summary.cell(row=3, column=col_idx).fill = YELLOW_FILL
        summary.column_dimensions[chr(64 + col_idx)].width = 20
    for idx, group in enumerate(GROUPS, start=4):
        summary.cell(row=idx, column=1).value = group
        summary.cell(row=idx, column=2).value = exp["mean"][idx - 4]
        summary.cell(row=idx, column=3).value = exp["sd"][idx - 4]
        summary.cell(row=idx, column=4).value = exp["letters"][idx - 4]

    path = outdir / f"{title}.xlsx"
    wb.save(path)
    return path


def uv_curve(params: dict[str, float], variation: dict[str, float], x_value: float) -> float:
    baseline = 0.02 + 0.01 * math.exp(-(x_value - 200) / 70) + variation["base"]
    peak_1 = gauss(x_value, params["p1"] + variation["shift1"], params["a1"] * variation["amp1"], 10.5)
    peak_2 = gauss(x_value, params["p2"] + variation["shift2"], params["a2"] * variation["amp2"], 14.5)
    return baseline + peak_1 + peak_2


def add_uv_metric_sheet(wb: Workbook, sheet_name: str, label: str, data: dict[str, list[float]]) -> None:
    ws = wb.create_sheet(sheet_name)
    style_sheet_header(ws, label)
    ws["A3"] = "Replicate"
    ws["A3"].font = Font(bold=True)
    for idx, group in enumerate(GROUPS, start=2):
        ws.cell(row=3, column=idx).value = group
        ws.cell(row=3, column=idx).font = Font(bold=True)
        ws.cell(row=3, column=idx).fill = YELLOW_FILL
        ws.column_dimensions[chr(64 + idx)].width = 14
    for rep_idx, rep in enumerate(REPLICATE_LABELS, start=4):
        ws.cell(row=rep_idx, column=1).value = rep
        for col_idx, group in enumerate(GROUPS, start=2):
            ws.cell(row=rep_idx, column=col_idx).value = data[group][rep_idx - 4]


def create_uv_workbook(outdir: Path) -> Path:
    wb = Workbook()
    title = "05_UV_SIMULATED_PRACTICE_ONLY"
    readme_lines = [
        "This workbook contains synthetic practice UV data derived from the supplied schematic spectral parameters.",
        "The XY sheet is suitable for Prism XY plotting practice.",
        "The peak-metric sheets provide column-format tables for group comparison in Prism.",
        "These are synthetic curves and synthetic derived metrics, not instrument-exported measurements.",
    ]
    write_readme_sheet(wb, title, readme_lines)

    ws = wb.create_sheet("UV_XY_Raw")
    style_sheet_header(ws, "UV_XY_Raw")
    ws.freeze_panes = "A3"
    ws.column_dimensions["A"].width = 16
    wavelengths = [round(200 + i * 0.5, 1) for i in range(401)]
    ws["A3"] = "Wavelength (nm)"
    ws["A3"].font = Font(bold=True)

    header_col = 2
    for group in GROUPS:
        for rep_idx, rep_label in enumerate(REPLICATE_LABELS):
            header = f"{group}_{rep_label}"
            ws.cell(row=3, column=header_col).value = header
            ws.cell(row=3, column=header_col).font = Font(bold=True)
            ws.cell(row=3, column=header_col).fill = YELLOW_FILL
            ws.column_dimensions[chr(64 + header_col)].width = 14
            header_col += 1

    peak_230: dict[str, list[float]] = {group: [] for group in GROUPS}
    peak_280: dict[str, list[float]] = {group: [] for group in GROUPS}
    ratio_280_230: dict[str, list[float]] = {group: [] for group in GROUPS}

    for row_idx, wavelength in enumerate(wavelengths, start=4):
        ws.cell(row=row_idx, column=1).value = wavelength

    for group_idx, group in enumerate(GROUPS):
        params = UV_PARAMS[group]
        for rep_idx, variation in enumerate(UV_REP_VARIATIONS):
            col_idx = 2 + group_idx * 3 + rep_idx
            a230 = uv_curve(params, variation, 230.0)
            a280 = uv_curve(params, variation, 280.0)
            peak_230[group].append(round(a230, 4))
            peak_280[group].append(round(a280, 4))
            ratio_280_230[group].append(round(a280 / a230, 4))

            for row_idx, wavelength in enumerate(wavelengths, start=4):
                ws.cell(row=row_idx, column=col_idx).value = round(uv_curve(params, variation, wavelength), 5)

    mean_sheet = wb.create_sheet("UV_XY_GroupMean")
    style_sheet_header(mean_sheet, "UV_XY_GroupMean")
    mean_sheet["A3"] = "Wavelength (nm)"
    mean_sheet["A3"].font = Font(bold=True)
    for idx, group in enumerate(GROUPS, start=2):
        mean_sheet.cell(row=3, column=idx).value = f"{group}_Mean"
        mean_sheet.cell(row=3, column=idx).font = Font(bold=True)
        mean_sheet.cell(row=3, column=idx).fill = YELLOW_FILL
    for row_idx, wavelength in enumerate(wavelengths, start=4):
        mean_sheet.cell(row=row_idx, column=1).value = wavelength
        for group_idx in range(4):
            start_col = 2 + group_idx * 3
            end_col = start_col + 2
            start_letter = chr(64 + start_col)
            end_letter = chr(64 + end_col)
            mean_sheet.cell(row=row_idx, column=2 + group_idx).value = f"=AVERAGE(UV_XY_Raw!{start_letter}{row_idx}:UV_XY_Raw!{end_letter}{row_idx})"

    add_uv_metric_sheet(wb, "PeakAbs_230", "Absorbance near 230 nm", peak_230)
    add_uv_metric_sheet(wb, "PeakAbs_280", "Absorbance near 280 nm", peak_280)
    add_uv_metric_sheet(wb, "Ratio_280_230", "A280/A230 ratio", ratio_280_230)

    summary = wb.create_sheet("PeakSummary")
    style_sheet_header(summary, "PeakSummary")
    headers = ["Group", "Mean A230", "Mean A280", "Mean A280/A230", "Target p1 (nm)", "Target p2 (nm)"]
    for col_idx, header in enumerate(headers, start=1):
        summary.cell(row=3, column=col_idx).value = header
        summary.cell(row=3, column=col_idx).font = Font(bold=True)
        summary.cell(row=3, column=col_idx).fill = YELLOW_FILL
        summary.column_dimensions[chr(64 + col_idx)].width = 18
    for row_idx, group in enumerate(GROUPS, start=4):
        summary.cell(row=row_idx, column=1).value = group
        summary.cell(row=row_idx, column=2).value = round(sum(peak_230[group]) / 3, 4)
        summary.cell(row=row_idx, column=3).value = round(sum(peak_280[group]) / 3, 4)
        summary.cell(row=row_idx, column=4).value = round(sum(ratio_280_230[group]) / 3, 4)
        summary.cell(row=row_idx, column=5).value = UV_PARAMS[group]["p1"]
        summary.cell(row=row_idx, column=6).value = UV_PARAMS[group]["p2"]

    path = outdir / f"{title}.xlsx"
    wb.save(path)
    return path


def ftir_curve(peaks: dict[str, float], variation: dict[str, float], x_value: float) -> float:
    baseline = 97 - 0.8 * math.sin((x_value - 400) / 850) + variation["base"]
    return (
        baseline
        - gauss(x_value, peaks["A"] + variation["peak_shift"], 10.0 * variation["depth"], 95)
        - gauss(x_value, peaks["I"] + variation["peak_shift"] * 0.7, 16.0 * variation["depth"], 55)
        - gauss(x_value, peaks["II"] + variation["peak_shift"] * 0.6, 12.0 * variation["depth"], 48)
        - gauss(x_value, peaks["III"] + variation["peak_shift"] * 0.5, 8.0 * variation["depth"], 45)
        - gauss(x_value, 2935 + variation["peak_shift"] * 0.4, 4.0 * variation["depth"], 70)
    )


def add_ftir_metric_sheet(wb: Workbook, sheet_name: str, label: str, data: dict[str, list[float]]) -> None:
    ws = wb.create_sheet(sheet_name)
    style_sheet_header(ws, label)
    ws["A3"] = "Replicate"
    ws["A3"].font = Font(bold=True)
    for idx, group in enumerate(GROUPS, start=2):
        ws.cell(row=3, column=idx).value = group
        ws.cell(row=3, column=idx).font = Font(bold=True)
        ws.cell(row=3, column=idx).fill = YELLOW_FILL
        ws.column_dimensions[chr(64 + idx)].width = 16
    for rep_idx, rep in enumerate(REPLICATE_LABELS, start=4):
        ws.cell(row=rep_idx, column=1).value = rep
        for col_idx, group in enumerate(GROUPS, start=2):
            ws.cell(row=rep_idx, column=col_idx).value = data[group][rep_idx - 4]


def create_ftir_workbook(outdir: Path) -> Path:
    wb = Workbook()
    title = "06_FTIR_SIMULATED_PRACTICE_ONLY"
    readme_lines = [
        "This workbook contains synthetic practice FTIR data based on the supplied representative peak positions.",
        "The XY sheet is suitable for Prism XY plotting rehearsal.",
        "Peak-position sheets provide groupwise tables that can be used for summary comparison figures.",
        "All values are synthetic practice data rather than spectrometer-exported observations.",
    ]
    write_readme_sheet(wb, title, readme_lines)

    ws = wb.create_sheet("FTIR_XY_Raw")
    style_sheet_header(ws, "FTIR_XY_Raw")
    ws.freeze_panes = "A3"
    ws.column_dimensions["A"].width = 18
    wavenumbers = [4000 - i * 4 for i in range(901)]
    ws["A3"] = "Wavenumber (cm^-1)"
    ws["A3"].font = Font(bold=True)

    header_col = 2
    for group in GROUPS:
        for rep_label in REPLICATE_LABELS:
            ws.cell(row=3, column=header_col).value = f"{group}_{rep_label}"
            ws.cell(row=3, column=header_col).font = Font(bold=True)
            ws.cell(row=3, column=header_col).fill = YELLOW_FILL
            ws.column_dimensions[chr(64 + header_col)].width = 15
            header_col += 1

    peak_tables = {
        "Amide_A_Pos": {group: [] for group in GROUPS},
        "Amide_I_Pos": {group: [] for group in GROUPS},
        "Amide_II_Pos": {group: [] for group in GROUPS},
        "Amide_III_Pos": {group: [] for group in GROUPS},
    }

    for row_idx, wavenumber in enumerate(wavenumbers, start=4):
        ws.cell(row=row_idx, column=1).value = wavenumber

    for group_idx, group in enumerate(GROUPS):
        peaks = FTIR_PEAKS[group]
        for rep_idx, variation in enumerate(FTIR_REP_VARIATIONS):
            col_idx = 2 + group_idx * 3 + rep_idx
            peak_tables["Amide_A_Pos"][group].append(round(peaks["A"] + variation["peak_shift"], 2))
            peak_tables["Amide_I_Pos"][group].append(round(peaks["I"] + variation["peak_shift"] * 0.7, 2))
            peak_tables["Amide_II_Pos"][group].append(round(peaks["II"] + variation["peak_shift"] * 0.6, 2))
            peak_tables["Amide_III_Pos"][group].append(round(peaks["III"] + variation["peak_shift"] * 0.5, 2))

            for row_idx, wavenumber in enumerate(wavenumbers, start=4):
                ws.cell(row=row_idx, column=col_idx).value = round(ftir_curve(peaks, variation, wavenumber), 5)

    mean_sheet = wb.create_sheet("FTIR_XY_GroupMean")
    style_sheet_header(mean_sheet, "FTIR_XY_GroupMean")
    mean_sheet["A3"] = "Wavenumber (cm^-1)"
    mean_sheet["A3"].font = Font(bold=True)
    for idx, group in enumerate(GROUPS, start=2):
        mean_sheet.cell(row=3, column=idx).value = f"{group}_Mean"
        mean_sheet.cell(row=3, column=idx).font = Font(bold=True)
        mean_sheet.cell(row=3, column=idx).fill = YELLOW_FILL
    for row_idx, wavenumber in enumerate(wavenumbers, start=4):
        mean_sheet.cell(row=row_idx, column=1).value = wavenumber
        for group_idx in range(4):
            start_col = 2 + group_idx * 3
            end_col = start_col + 2
            start_letter = chr(64 + start_col)
            end_letter = chr(64 + end_col)
            mean_sheet.cell(row=row_idx, column=2 + group_idx).value = f"=AVERAGE(FTIR_XY_Raw!{start_letter}{row_idx}:FTIR_XY_Raw!{end_letter}{row_idx})"

    add_ftir_metric_sheet(wb, "Amide_A_Pos", "Amide A peak position (cm^-1)", peak_tables["Amide_A_Pos"])
    add_ftir_metric_sheet(wb, "Amide_I_Pos", "Amide I peak position (cm^-1)", peak_tables["Amide_I_Pos"])
    add_ftir_metric_sheet(wb, "Amide_II_Pos", "Amide II peak position (cm^-1)", peak_tables["Amide_II_Pos"])
    add_ftir_metric_sheet(wb, "Amide_III_Pos", "Amide III peak position (cm^-1)", peak_tables["Amide_III_Pos"])

    path = outdir / f"{title}.xlsx"
    wb.save(path)
    return path


def create_master_summary(outdir: Path, bar_paths: list[Path], uv_path: Path, ftir_path: Path) -> Path:
    wb = Workbook()
    write_readme_sheet(
        wb,
        "MASTER_SUMMARY_SIMULATED_PRACTICE_ONLY",
        [
            "This workbook indexes the generated synthetic practice files.",
            "Each experiment was exported as a separate Excel file to match a Prism-first workflow.",
            "Use these files for layout rehearsal, import testing, and sheet-structure planning only.",
        ],
    )

    ws = wb.create_sheet("Files")
    style_sheet_header(ws, "Generated files")
    headers = ["Experiment", "File name", "Purpose"]
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=3, column=col_idx).value = header
        ws.cell(row=3, column=col_idx).font = Font(bold=True)
        ws.cell(row=3, column=col_idx).fill = YELLOW_FILL
        ws.column_dimensions[chr(64 + col_idx)].width = 44

    file_rows = [
        ("WHC", bar_paths[0].name, "Column data for Prism"),
        ("OHC", bar_paths[1].name, "Column data for Prism"),
        ("EAI", bar_paths[2].name, "Column data for Prism"),
        ("ESI", bar_paths[3].name, "Column data for Prism"),
        ("UV", uv_path.name, "XY curves plus peak metrics"),
        ("FTIR", ftir_path.name, "XY curves plus peak positions"),
    ]
    for row_idx, row in enumerate(file_rows, start=4):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx).value = value

    path = outdir / "00_MASTER_SUMMARY_SIMULATED_PRACTICE_ONLY.xlsx"
    wb.save(path)
    return path


def main() -> None:
    args = parse_args()
    outdir = Path(args.outdir).expanduser().resolve()
    outdir.mkdir(parents=True, exist_ok=True)

    bar_paths = [create_bar_workbook(outdir, exp) for exp in BAR_EXPERIMENTS]
    uv_path = create_uv_workbook(outdir)
    ftir_path = create_ftir_workbook(outdir)
    master_path = create_master_summary(outdir, bar_paths, uv_path, ftir_path)

    print("Generated files:")
    for path in [master_path, *bar_paths, uv_path, ftir_path]:
        print(path)


if __name__ == "__main__":
    main()
