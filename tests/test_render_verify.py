from pathlib import Path
import subprocess
import sys
from types import SimpleNamespace

from docx import Document
import pytest

import thesis_tool.render_verify as render_verify_module


SCRIPTS_DIR = Path(__file__).resolve().parents[1] / "scripts"


def test_render_verify_module_script_can_show_help_from_repo_root():
    result = subprocess.run(
        [
            sys.executable,
            str(SCRIPTS_DIR / "thesis_tool" / "render_verify.py"),
            "--help",
        ],
        cwd=SCRIPTS_DIR.parent,
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    assert "--rendered-pdf" in result.stdout


def test_find_external_tool_prefers_existing_env_path(monkeypatch, tmp_path):
    tool_path = tmp_path / "custom-tool"
    tool_path.write_text("#!/bin/sh\n", encoding="utf-8")
    monkeypatch.setenv("ARTICLE_CUSTOM_TOOL", str(tool_path))
    monkeypatch.setattr(
        render_verify_module.shutil,
        "which",
        lambda name: pytest.fail("env path should be used before PATH lookup"),
    )

    assert render_verify_module._find_external_tool(
        env_name="ARTICLE_CUSTOM_TOOL",
        binary="custom-tool",
        missing_message="missing custom tool",
    ) == str(tool_path.resolve())


def test_find_external_tool_falls_back_to_path_lookup(monkeypatch):
    monkeypatch.delenv("ARTICLE_CUSTOM_TOOL", raising=False)
    monkeypatch.setattr(render_verify_module.shutil, "which", lambda name: f"/opt/bin/{name}")

    assert render_verify_module._find_external_tool(
        env_name="ARTICLE_CUSTOM_TOOL",
        binary="custom-tool",
        missing_message="missing custom tool",
    ) == "/opt/bin/custom-tool"


def test_find_external_tool_missing_raises_original_message(monkeypatch):
    monkeypatch.delenv("ARTICLE_CUSTOM_TOOL", raising=False)
    monkeypatch.setattr(render_verify_module.shutil, "which", lambda name: None)

    with pytest.raises(RuntimeError, match="missing custom tool.*ARTICLE_CUSTOM_TOOL"):
        render_verify_module._find_external_tool(
            env_name="ARTICLE_CUSTOM_TOOL",
            binary="custom-tool",
            missing_message="missing custom tool; set ARTICLE_CUSTOM_TOOL",
        )


def test_convert_pdf_to_page_images_removes_stale_pages(monkeypatch, tmp_path):
    pdf_path = tmp_path / "proof.pdf"
    pdf_path.write_bytes(b"%PDF")
    page_dir = tmp_path / "pages"
    page_dir.mkdir()
    stale_page = page_dir / "page-01.png"
    stale_page.write_bytes(b"old")

    def fake_run(command, **_kwargs):
        Path(f"{command[-1]}-1.png").write_bytes(b"new")
        return SimpleNamespace(stdout="", stderr="")

    monkeypatch.setattr(render_verify_module, "_find_pdftoppm", lambda: "/usr/bin/pdftoppm")
    monkeypatch.setattr(render_verify_module.subprocess, "run", fake_run)

    render_verify_module._convert_pdf_to_page_images(str(pdf_path), str(page_dir))

    assert not stale_page.exists()
    assert (page_dir / "page-1.png").read_bytes() == b"new"


def test_convert_pdf_to_page_images_uses_pdfium_when_poppler_is_missing(monkeypatch, tmp_path):
    from PIL import Image

    pdf_path = tmp_path / "proof.pdf"
    Image.new("RGB", (80, 60), "white").save(pdf_path, "PDF", resolution=72)
    page_dir = tmp_path / "pages"

    def missing_poppler():
        raise RuntimeError("pdftoppm missing")

    monkeypatch.setattr(render_verify_module, "_find_pdftoppm", missing_poppler)

    render_verify_module._convert_pdf_to_page_images(str(pdf_path), str(page_dir))

    page_images = render_verify_module._collect_page_images(page_dir)
    assert len(page_images) == 1
    with Image.open(page_images[0]) as rendered:
        assert rendered.format == "PNG"
        assert rendered.width > 80


def test_extract_pdf_page_texts_uses_pdfium_when_pdftotext_is_missing(monkeypatch, tmp_path):
    pdf_path = tmp_path / "rendered.pdf"
    pdf_path.write_bytes(b"%PDF")

    def missing_poppler():
        raise RuntimeError("pdftotext missing")

    monkeypatch.setattr(render_verify_module, "_find_pdftotext", missing_poppler)
    monkeypatch.setattr(
        render_verify_module,
        "_extract_pdf_text_pages_with_pdfium",
        lambda _path: ["第一页正文", "第 2 章 实验材料与方法"],
        raising=False,
    )

    page_texts, summary = render_verify_module._extract_pdf_page_texts(str(pdf_path), page_count=2)

    assert page_texts == {1: "第一页正文", 2: "第 2 章 实验材料与方法"}
    assert summary["available"] is True
    assert summary["page_text_extraction_warning_count"] == 0


def test_word_pdf_export_script_targets_opened_file_not_active_document(monkeypatch, tmp_path):
    source_path = tmp_path / "render_verify_word_source.docx"
    source_path.write_bytes(b"docx")
    output_pdf = tmp_path / "render_verify_word.pdf"
    captured = {}

    def fake_run(command, **kwargs):
        captured["script"] = command[2]
        output_pdf.write_bytes(b"%PDF")
        return SimpleNamespace(stdout="", stderr="")

    monkeypatch.setattr(render_verify_module.sys, "platform", "darwin")
    monkeypatch.setattr(render_verify_module.shutil, "which", lambda name: "/usr/bin/osascript" if name == "osascript" else None)
    monkeypatch.setattr(render_verify_module.subprocess, "run", fake_run)

    render_verify_module._export_docx_to_pdf_with_word(str(source_path), str(output_pdf))

    assert "active document" not in captured["script"]
    assert "document 1" not in captured["script"]
    assert "close every document" in captured["script"]
    assert "did not expose the opened DOCX" in captured["script"]
    assert "save as docRef" in captured["script"]


def test_word_pdf_export_timeout_explains_word_automation_block(monkeypatch, tmp_path):
    source_path = tmp_path / "render_verify_word_timeout.docx"
    source_path.write_bytes(b"docx")
    output_pdf = tmp_path / "render_verify_word.pdf"

    def fake_run(command, **kwargs):
        raise render_verify_module.subprocess.TimeoutExpired(command, timeout=kwargs["timeout"])

    monkeypatch.setattr(render_verify_module.sys, "platform", "darwin")
    monkeypatch.setattr(render_verify_module.shutil, "which", lambda name: "/usr/bin/osascript" if name == "osascript" else None)
    monkeypatch.setattr(render_verify_module.subprocess, "run", fake_run)

    with pytest.raises(RuntimeError, match="Word 自动化没有完成"):
        render_verify_module._export_docx_to_pdf_with_word(str(source_path), str(output_pdf))


def test_build_pdf_toc_page_number_report_detects_mismatch(monkeypatch, tmp_path):
    rendered_pdf = tmp_path / "toc_mismatch.pdf"
    rendered_pdf.write_bytes(b"%PDF")

    def fake_extract_pdf_page_texts(pdf_path: str | None, *, page_count: int | None = None):
        assert pdf_path == str(rendered_pdf)
        assert page_count is None
        return (
            {
                1: "目录\n第1章 绪论 ........ 1",
                2: "第1章 绪论\n这是正文\n2",
            },
            {
                "source": "pdf",
                "available": True,
                "page_text_available_count": 2,
                "page_text_extraction_warning_count": 0,
                "warnings": [],
            },
        )

    monkeypatch.setattr(render_verify_module, "_extract_pdf_page_texts", fake_extract_pdf_page_texts)

    report = render_verify_module.build_pdf_toc_page_number_report(str(rendered_pdf))

    assert report["summary"]["toc_page_number_mismatch_count"] == 1
    assert report["findings"][0]["id"] == "toc_page_number_mismatch"
    assert "第1章 绪论" in report["findings"][0]["message"]


def test_build_render_verify_report_writes_markdown_and_collects_pages(monkeypatch, tmp_path):
    source_path = tmp_path / "render_verify_source.docx"
    doc = Document()
    doc.add_heading("Render Verify", level=1)
    doc.save(source_path)

    def fake_run_render_engine(input_docx: str, output_dir: str, renderer: str) -> dict:
        page_dir = Path(output_dir) / "word_pdf_pages"
        page_dir.mkdir(parents=True, exist_ok=True)
        (page_dir / "page-1.png").write_bytes(b"png")
        pdf_path = Path(output_dir) / "render_verify_word.pdf"
        pdf_path.write_bytes(b"pdf")
        return {
            "engine": "word-pdf",
            "page_dir": str(page_dir),
            "pdf_path": str(pdf_path),
            "fallback_used": False,
            "warnings": [],
        }

    monkeypatch.setattr(render_verify_module, "_run_render_engine", fake_run_render_engine)
    monkeypatch.setattr(
        render_verify_module,
        "analyze_page_images",
        lambda page_images, **kwargs: {
            "findings": [
                {
                    "id": "large_blank_region",
                    "severity": "warning",
                    "page": 1,
                    "region": {"x": 0.08, "y": 0.72, "w": 0.84, "h": 0.28},
                    "message": "页底存在大块连续空白。",
                    "suggested_scope": "figures_tables",
                    "evidence_source": kwargs.get("evidence_source"),
                }
            ],
            "summary": {
                "finding_count": 1,
                "highest_severity": "warning",
                "blank_page_count": 0,
                "large_blank_count": 1,
                "render_suspect_count": 0,
            },
        },
    )
    monkeypatch.setattr(
        render_verify_module,
        "build_document_preflight",
        lambda *args, **kwargs: {
            "preflight_status": "warning",
            "toc_status": "field_only",
            "preface_status": "not_detected",
            "heading_renumber_guard": {"status": "warn", "reason": "table_risk"},
            "style_conflict_count": 0,
            "table_heading_risk_count": 2,
            "recommended_actions": ["先复核目录域刷新。"],
            "diagnostics": {
                "toc": {"status": "field_only"},
                "heading_renumber_guard": {"status": "warn", "reason": "table_risk"},
            },
        },
    )
    monkeypatch.setattr(
        render_verify_module,
        "build_scope_verify",
        lambda *args, **kwargs: {
            "profile_id": "lnu-checker-2026",
            "requested_profile": "lnu",
            "fallback_used": False,
            "profile_display": "lnu-checker-2026 (requested: lnu)",
            "selected_scopes": ["headings"],
            "overall_status": "verified",
            "readiness": "structure-ready",
            "manual_review_rule_ids": [],
            "unsupported_rule_ids": [],
        },
    )

    report = render_verify_module.build_render_verify_report(
        str(source_path),
        profile_path="lnu",
        scopes=["headings"],
    )

    assert report["document"]["name"] == "render_verify_source.docx"
    assert report["render_engine"] == "word-pdf"
    assert report["evidence_trust"] == "authoritative"
    assert report["evidence_authoritative"] is True
    assert report["layout_decision_eligible"] is True
    assert report["render_pdf_path"].endswith("render_verify_word.pdf")
    assert report["render_page_dir"].endswith("word_pdf_pages")
    assert report["page_count"] == 1
    assert report["page_images"][0].endswith("page-1.png")
    assert report["evidence_source"] == "word-pdf"
    assert report["render_findings"] == []
    assert report["render_summary"]["finding_count"] == 0
    assert report["structure_readiness"] == "structure-ready"
    assert report["readiness"] == "render-check-required"
    assert report["preflight_status"] == "warning"
    assert report["render_evidence_status"] == "render-review-required"
    assert report["wild_doc"]["detected"] is True
    assert report["wild_doc"]["signals"][0]["id"] == "toc_structure"
    assert report["summary"]["wild_doc_signal_count"] == 2
    assert report["summary"]["render_finding_count"] == 0
    assert report["summary"]["render_highest_severity"] is None
    report_path = Path(report["report_path"])
    assert report_path.name == "render_verify_report.md"
    assert report_path.exists()
    assert Path(report["render_conclusion_report_path"]).exists()
    assert Path(report["render_ai_review_context_path"]).exists()
    report_text = report_path.read_text(encoding="utf-8")
    render_conclusion_text = Path(report["render_conclusion_report_path"]).read_text(encoding="utf-8")
    render_ai_text = Path(report["render_ai_review_context_path"]).read_text(encoding="utf-8")
    assert "文件: render_verify_source.docx" in report_text
    assert "自动页图判读" not in report_text
    assert "large_blank_region" not in report_text
    assert "大块连续空白" not in report_text
    assert "large_blank_region" not in render_conclusion_text
    assert "大块空白" not in render_conclusion_text
    assert "Render Findings" in render_ai_text
    assert not report_text.lstrip().startswith("{")


def test_filter_user_visible_render_findings_hides_all_large_blank_variants():
    visible = render_verify_module._filter_user_visible_render_findings(
        [
            {"id": "large_blank_region", "message": "页底存在大块连续空白。"},
            {"id": "large_blank_bottom", "message": "页底空白过大。"},
            {"id": "large_blank_middle", "message": "页面中部空白过大。"},
            {"id": "isolated_punctuation", "message": "页面文本存在单独成行的标点。"},
        ]
    )

    assert visible == [{"id": "isolated_punctuation", "message": "页面文本存在单独成行的标点。"}]


def test_build_evidence_items_validates_bbox_and_keeps_unlocatable_findings(tmp_path):
    page1 = tmp_path / "page-1.png"
    page2 = tmp_path / "page-2.png"
    page1.write_bytes(b"png")
    page2.write_bytes(b"png")

    items = render_verify_module._build_evidence_items(
        [
            {
                "id": "isolated_punctuation",
                "rule_id": "render.isolated_punctuation",
                "severity": "warning",
                "page": 2,
                "bbox": {"x": 0.12, "y": 0.64, "w": 0.72, "h": 0.18},
                "message": "页面文本存在单独成行的标点，疑似换行或排版挤压导致。",
                "suggested_action": "回到 WPS/Word 检查该处换行、字距和段落排版，调整后重新导出 PDF。",
                "image_path": str(page2),
            },
            {
                "id": "pdf_text_missing",
                "severity": "info",
                "page": 1,
                "bbox": {"x": 1.2, "y": 0.2, "w": 0.2, "h": 0.2},
                "message": "PDF 文本不可读",
            },
            {
                "id": "manual_review",
                "severity": "warning",
                "page": 1,
                "message": "需要人工复核页面",
            },
        ],
        page_images=[str(page1), str(page2)],
    )

    assert items[0] == {
        "page": 2,
        "screenshot_path": str(page2.resolve()),
        "rule_id": "render.isolated_punctuation",
        "bbox": {"x": 0.12, "y": 0.64, "w": 0.72, "h": 0.18},
        "message": "页面文本存在单独成行的标点，疑似换行或排版挤压导致。",
        "severity": "warning",
        "next_action": "回到 WPS/Word 检查该处换行、字距和段落排版，调整后重新导出 PDF。",
    }
    assert items[1]["screenshot_path"] == str(page1.resolve())
    assert items[1]["bbox"] is None
    assert items[1]["message"] == "PDF 文本不可读"
    assert items[2]["bbox"] is None
    assert items[2]["next_action"] == "回到 DOCX 调整对应版式问题后重新导出 PDF。"


def test_build_render_verify_report_adds_evidence_items(monkeypatch, tmp_path):
    source_path = tmp_path / "render_verify_evidence_source.docx"
    doc = Document()
    doc.add_heading("Render Verify", level=1)
    doc.save(source_path)

    def fake_run_render_engine(input_docx: str, output_dir: str, renderer: str) -> dict:
        page_dir = Path(output_dir) / "word_pdf_pages"
        page_dir.mkdir(parents=True, exist_ok=True)
        (page_dir / "page-1.png").write_bytes(b"png")
        (page_dir / "page-2.png").write_bytes(b"png")
        pdf_path = Path(output_dir) / "render_verify_word.pdf"
        pdf_path.write_bytes(b"pdf")
        return {
            "engine": "word-pdf",
            "page_dir": str(page_dir),
            "pdf_path": str(pdf_path),
            "fallback_used": False,
            "warnings": [],
        }

    def fake_analyze_page_images(page_images, **kwargs):
        return {
            "findings": [
                {
                    "id": "isolated_punctuation",
                    "rule_id": "render.isolated_punctuation",
                    "severity": "warning",
                    "page": 2,
                    "bbox": {"x": 0.08, "y": 0.72, "w": 0.84, "h": 0.2},
                    "message": "页面文本存在单独成行的标点，疑似换行或排版挤压导致。",
                    "actionable": True,
                    "suggested_action": "回到 WPS/Word 检查该处换行、字距和段落排版，调整后重新导出 PDF。",
                    "image_path": page_images[1],
                },
                {
                    "id": "manual_review",
                    "severity": "info",
                    "page": 1,
                    "message": "需要人工复核页面。",
                },
            ],
            "summary": {
                "finding_count": 2,
                "highest_severity": "warning",
                "blank_page_count": 0,
                "render_suspect_count": 0,
                "actionable_finding_count": 1,
            },
            "layout_score": {"score": 100, "penalty": 0, "actionable_finding_count": 0},
        }

    monkeypatch.setattr(render_verify_module, "_run_render_engine", fake_run_render_engine)
    monkeypatch.setattr(render_verify_module, "analyze_page_images", fake_analyze_page_images)
    monkeypatch.setattr(
        render_verify_module,
        "build_document_preflight",
        lambda *args, **kwargs: {
            "preflight_status": "ready",
            "toc_status": "generated_toc",
            "preface_status": "not_detected",
            "heading_renumber_guard": {"status": "clear"},
            "style_conflict_count": 0,
            "table_heading_risk_count": 0,
            "recommended_actions": [],
            "diagnostics": {"toc": {"status": "generated_toc"}, "heading_renumber_guard": {"status": "clear"}},
        },
    )
    monkeypatch.setattr(
        render_verify_module,
        "build_scope_verify",
        lambda *args, **kwargs: {
            "profile_id": "lnu-checker-2026",
            "requested_profile": "lnu",
            "fallback_used": False,
            "profile_display": "lnu-checker-2026 (requested: lnu)",
            "selected_scopes": ["figures_tables"],
            "overall_status": "verified",
            "readiness": "structure-ready",
            "manual_review_rule_ids": [],
            "unsupported_rule_ids": [],
        },
    )

    report = render_verify_module.build_render_verify_report(str(source_path), profile_path="lnu")

    assert report["summary"]["evidence_item_count"] == 2
    assert report["evidence_items"][0]["rule_id"] == "render.isolated_punctuation"
    assert report["evidence_items"][0]["bbox"] == {"x": 0.08, "y": 0.72, "w": 0.84, "h": 0.2}
    assert report["evidence_items"][0]["screenshot_path"].endswith("page-2.png")
    assert report["evidence_items"][1]["bbox"] is None


def test_extract_pdf_page_texts_splits_pdftotext_form_feeds(monkeypatch, tmp_path):
    pdf_path = tmp_path / "rendered.pdf"
    pdf_path.write_bytes(b"%PDF")

    def fake_run(command, **kwargs):
        assert command[-2:] == [str(pdf_path.resolve()), "-"]
        return SimpleNamespace(stdout="第一页正文\f第 2 章 实验材料与方法\f")

    monkeypatch.setattr(render_verify_module, "_find_pdftotext", lambda: "pdftotext")
    monkeypatch.setattr(render_verify_module.subprocess, "run", fake_run)

    page_texts, summary = render_verify_module._extract_pdf_page_texts(str(pdf_path), page_count=2)

    assert page_texts == {1: "第一页正文", 2: "第 2 章 实验材料与方法"}
    assert summary == {
        "source": "pdf",
        "available": True,
        "page_text_available_count": 2,
        "page_text_extraction_warning_count": 0,
        "warnings": [],
    }


def test_build_render_verify_report_passes_pdf_texts_to_analyzer(monkeypatch, tmp_path):
    source_path = tmp_path / "render_verify_text_source.docx"
    doc = Document()
    doc.add_heading("Render Verify", level=1)
    doc.save(source_path)
    captured = {}

    def fake_run_render_engine(input_docx: str, output_dir: str, renderer: str) -> dict:
        page_dir = Path(output_dir) / "word_pdf_pages"
        page_dir.mkdir(parents=True, exist_ok=True)
        (page_dir / "page-1.png").write_bytes(b"png")
        (page_dir / "page-2.png").write_bytes(b"png")
        pdf_path = Path(output_dir) / "render_verify_word.pdf"
        pdf_path.write_bytes(b"pdf")
        return {
            "engine": "word-pdf",
            "page_dir": str(page_dir),
            "pdf_path": str(pdf_path),
            "fallback_used": False,
            "warnings": [],
        }

    def fake_analyze_page_images(page_images, **kwargs):
        captured["page_texts"] = kwargs.get("page_texts")
        return {
            "findings": [],
            "summary": {
                "finding_count": 0,
                "highest_severity": None,
                "blank_page_count": 0,
                "large_blank_count": 0,
                "render_suspect_count": 0,
            },
            "layout_score": {
                "score": 100,
                "penalty": 0,
                "expected_blank_count": 0,
                "actionable_finding_count": 0,
                "object_flow_issue_count": 0,
                "heading_break_issue_count": 0,
                "render_integrity_issue_count": 0,
            },
        }

    monkeypatch.setattr(render_verify_module, "_run_render_engine", fake_run_render_engine)
    monkeypatch.setattr(
        render_verify_module,
        "_extract_pdf_page_texts",
        lambda pdf_path, *, page_count: (
            {1: "第一页正文", 2: "第 2 章 实验材料与方法"},
            {
                "source": "pdf",
                "available": True,
                "page_text_available_count": 2,
                "page_text_extraction_warning_count": 0,
                "warnings": [],
            },
        ),
        raising=False,
    )
    monkeypatch.setattr(render_verify_module, "analyze_page_images", fake_analyze_page_images)
    monkeypatch.setattr(
        render_verify_module,
        "build_document_preflight",
        lambda *args, **kwargs: {
            "preflight_status": "ready",
            "toc_status": "generated_toc",
            "preface_status": "not_detected",
            "heading_renumber_guard": {"status": "clear"},
            "style_conflict_count": 0,
            "table_heading_risk_count": 0,
            "recommended_actions": [],
            "diagnostics": {"toc": {"status": "generated_toc"}, "heading_renumber_guard": {"status": "clear"}},
        },
    )
    monkeypatch.setattr(
        render_verify_module,
        "build_scope_verify",
        lambda *args, **kwargs: {
            "profile_id": "lnu-checker-2026",
            "requested_profile": "lnu",
            "fallback_used": False,
            "profile_display": "lnu-checker-2026 (requested: lnu)",
            "selected_scopes": None,
            "overall_status": "verified",
            "readiness": "structure-ready",
            "manual_review_rule_ids": [],
            "unsupported_rule_ids": [],
        },
    )

    report = render_verify_module.build_render_verify_report(str(source_path), profile_path="lnu", renderer="word-pdf")

    assert captured["page_texts"] == {1: "第一页正文", 2: "第 2 章 实验材料与方法"}
    assert report["render_text_summary"]["available"] is True
    assert report["summary"]["page_text_available_count"] == 2
    assert report["layout_score"]["score"] == 100
    assert report["summary"]["layout_score"] == 100


def test_build_render_verify_report_flags_toc_declared_page_mismatch(monkeypatch, tmp_path):
    source_path = tmp_path / "render_verify_toc_source.docx"
    doc = Document()
    doc.add_heading("Render Verify", level=1)
    doc.save(source_path)

    def fake_run_render_engine(input_docx: str, output_dir: str, renderer: str) -> dict:
        page_dir = Path(output_dir) / "word_pdf_pages"
        page_dir.mkdir(parents=True, exist_ok=True)
        for page_number in range(1, 5):
            (page_dir / f"page-{page_number}.png").write_bytes(b"png")
        pdf_path = Path(output_dir) / "render_verify_word.pdf"
        pdf_path.write_bytes(b"pdf")
        return {
            "engine": "word-pdf",
            "page_dir": str(page_dir),
            "pdf_path": str(pdf_path),
            "fallback_used": False,
            "warnings": [],
        }

    monkeypatch.setattr(render_verify_module, "_run_render_engine", fake_run_render_engine)
    monkeypatch.setattr(
        render_verify_module,
        "_extract_pdf_page_texts",
        lambda pdf_path, *, page_count: (
            {
                1: "目 录\n摘要 ...... 2\n1 绪论 ...... 2\n2 材料与方法",
                2: "摘要\n本文研究论文格式检查工具。",
                3: "1 绪论\n研究背景与意义。\n- 3 -",
                4: "2 材料与方法\n实验设置。\n- 4 -",
            },
            {
                "source": "pdf",
                "available": True,
                "page_text_available_count": 4,
                "page_text_extraction_warning_count": 0,
                "warnings": [],
            },
        ),
        raising=False,
    )
    monkeypatch.setattr(
        render_verify_module,
        "analyze_page_images",
        lambda page_images, **kwargs: {
            "findings": [],
            "summary": {
                "finding_count": 0,
                "highest_severity": None,
                "blank_page_count": 0,
                "large_blank_count": 0,
                "render_suspect_count": 0,
            },
            "layout_score": {"score": 100, "penalty": 0, "actionable_finding_count": 0},
        },
    )
    monkeypatch.setattr(
        render_verify_module,
        "build_document_preflight",
        lambda *args, **kwargs: {
            "preflight_status": "ready",
            "toc_status": "generated_toc",
            "preface_status": "not_detected",
            "heading_renumber_guard": {"status": "clear"},
            "style_conflict_count": 0,
            "table_heading_risk_count": 0,
            "recommended_actions": [],
            "diagnostics": {"toc": {"status": "generated_toc"}, "heading_renumber_guard": {"status": "clear"}},
        },
    )
    monkeypatch.setattr(
        render_verify_module,
        "build_scope_verify",
        lambda *args, **kwargs: {
            "profile_id": "lnu-checker-2026",
            "requested_profile": "lnu",
            "fallback_used": False,
            "profile_display": "lnu-checker-2026 (requested: lnu)",
            "selected_scopes": ["toc"],
            "overall_status": "verified",
            "readiness": "structure-ready",
            "manual_review_rule_ids": [],
            "unsupported_rule_ids": [],
        },
    )

    report = render_verify_module.build_render_verify_report(str(source_path), profile_path="lnu", scopes=["toc"])

    mismatch = next(item for item in report["render_findings"] if item["id"] == "toc_page_number_mismatch")
    assert mismatch["page"] == 1
    assert mismatch["declared_page"] == 2
    assert mismatch["actual_page"] == 3
    assert mismatch["toc_entry"] == "1 绪论"
    assert "目录页码不一致" in mismatch["message"]
    assert report["summary"]["render_finding_count"] == 2
    assert any(item["rule_id"] == "render.toc_page_number_mismatch" for item in report["evidence_items"])

    unconfirmed = next(item for item in report["render_findings"] if item["id"] == "toc_page_number_unconfirmed")
    assert unconfirmed["toc_entry"] == "2 材料与方法"
    assert unconfirmed["severity"] == "warning"
    assert "缺少页码" in unconfirmed["message"]


def test_toc_page_number_check_uses_printed_page_number_not_pdf_index():
    findings = render_verify_module._build_toc_page_number_findings(
        {
            1: "目 录\n1 绪论 ...... 1",
            2: "摘 要\n摘要正文\nI",
            3: "1 绪论\n研究背景与意义。\n- 1 -",
        }
    )

    assert findings == []


def test_toc_page_number_check_does_not_treat_heading_number_as_printed_page_number():
    findings = render_verify_module._build_toc_page_number_findings(
        {
            1: "目 录\n1 绪论 ...... 1",
            2: "1 绪论\n研究背景与意义。",
        }
    )

    assert findings == [
        {
            "id": "toc_page_number_unconfirmed",
            "rule_id": "render.toc_page_number_unconfirmed",
            "classification": "render_integrity",
            "severity": "warning",
            "page": 1,
            "toc_entry": "1 绪论",
            "declared_page": 1,
            "actual_page": None,
            "message": "目录条目“1 绪论”已定位到 PDF 第 2 页，但未读取到该页显示页码，需人工复核。",
            "actionable": False,
            "suggested_scope": "toc",
            "suggested_action": "打开 PDF 对照目录和正文标题页码。",
        }
    ]


def test_toc_page_number_check_warns_when_heading_not_located():
    findings = render_verify_module._build_toc_page_number_findings(
        {
            1: "目录\n1 绪论 ...... 2",
            2: "摘要\n本文研究论文格式检查工具。",
        }
    )

    assert findings == [
        {
            "id": "toc_page_number_unconfirmed",
            "rule_id": "render.toc_page_number_unconfirmed",
            "classification": "render_integrity",
            "severity": "warning",
            "page": 1,
            "toc_entry": "1 绪论",
            "declared_page": 2,
            "actual_page": None,
            "message": "目录条目“1 绪论”未能在后续 PDF 页面中确认对应标题，需人工复核。",
            "actionable": False,
            "suggested_scope": "toc",
            "suggested_action": "打开 PDF 对照目录和正文标题页码。",
        }
    ]


def test_external_rendered_pdf_uses_neutral_manual_pdf_evidence_source(monkeypatch, tmp_path):
    pdf_path = tmp_path / "word-export.pdf"
    pdf_path.write_bytes(b"%PDF")

    monkeypatch.setattr(render_verify_module, "_convert_pdf_to_page_images", lambda *args, **kwargs: None)

    metadata = render_verify_module._run_external_pdf_render(str(pdf_path), str(tmp_path / "proof"))

    assert metadata["engine"] == "manual-pdf"


def test_build_render_review_items_explains_likely_causes():
    diagnostics = {
        "toc": {"status": "generated_toc"},
        "heading_renumber_guard": {"status": "clear"},
    }
    verification = {
        "manual_review_rule_ids": [],
        "unsupported_rule_ids": [],
        "render_findings": [
            {"id": "large_blank_region", "page": 3, "message": "页底存在大块连续空白。"},
            {"id": "object_overflow", "page": 5, "message": "图表整体被挤到后页。"},
            {"id": "heading_isolated", "page": 7, "message": "标题落在页尾，正文被挤到下一页。"},
        ],
    }

    items = render_verify_module._build_render_review_items(diagnostics, verification)
    text = "\n".join(items)

    assert "大块空白" not in text
    assert "对象塞不进当前页" not in text
    assert "图表挤页" in text
    assert "首次引用位置过晚" in text
    assert "标题孤页" in text
    assert "标题后正文不足" in text


def test_run_render_engine_auto_does_not_fallback_to_artifact(monkeypatch, tmp_path):
    def fake_word_pdf_render(input_docx: str, output_dir: str) -> dict:
        raise RuntimeError("Word PDF 渲染不可用: mock")

    monkeypatch.setattr(render_verify_module, "_run_word_pdf_render", fake_word_pdf_render)

    with pytest.raises(RuntimeError, match="Word PDF 渲染不可用"):
        render_verify_module._run_render_engine("demo.docx", str(tmp_path), "auto")


def test_run_render_engine_rejects_artifact_tool_renderer(tmp_path):
    with pytest.raises(ValueError, match="artifact-tool"):
        render_verify_module._run_render_engine("demo.docx", str(tmp_path), "artifact-tool")


def test_build_render_verify_report_accepts_external_page_images(monkeypatch, tmp_path):
    source_path = tmp_path / "render_verify_external_source.docx"
    doc = Document()
    doc.add_heading("Render Verify", level=1)
    doc.save(source_path)
    page_dir = tmp_path / "wps-pages"
    page_dir.mkdir()
    (page_dir / "page-1.png").write_bytes(b"png")

    monkeypatch.setattr(
        render_verify_module,
        "analyze_page_images",
        lambda page_images, **kwargs: {
            "findings": [],
            "summary": {
                "finding_count": 0,
                "highest_severity": None,
                "blank_page_count": 0,
                "large_blank_count": 0,
                "render_suspect_count": 0,
            },
        },
    )
    monkeypatch.setattr(
        render_verify_module,
        "build_document_preflight",
        lambda *args, **kwargs: {
            "preflight_status": "ready",
            "toc_status": "generated_toc",
            "preface_status": "not_detected",
            "heading_renumber_guard": {"status": "clear"},
            "style_conflict_count": 0,
            "table_heading_risk_count": 0,
            "recommended_actions": [],
            "diagnostics": {"toc": {"status": "generated_toc"}, "heading_renumber_guard": {"status": "clear"}},
        },
    )
    monkeypatch.setattr(
        render_verify_module,
        "build_scope_verify",
        lambda *args, **kwargs: {
            "profile_id": "lnu-checker-2026",
            "requested_profile": "lnu",
            "fallback_used": False,
            "profile_display": "lnu-checker-2026 (requested: lnu)",
            "selected_scopes": ["toc"],
            "overall_status": "verified",
            "readiness": "structure-ready",
            "manual_review_rule_ids": [],
            "unsupported_rule_ids": [],
        },
    )

    report = render_verify_module.build_render_verify_report(
        str(source_path),
        profile_path="lnu",
        scopes=["toc"],
        page_images_dir=str(page_dir),
    )

    assert report["render_engine"] == "wps-manual-images"
    assert report["evidence_source"] == "wps-manual-images"
    assert report["evidence_trust"] == "authoritative"
    assert report["layout_decision_eligible"] is True
    assert report["render_page_dir"] == str(page_dir.resolve())
    assert report["page_images"] == [str((page_dir / "page-1.png").resolve())]


def test_render_render_verify_report_includes_review_summary():
    report = {
        "document": {"name": "demo.docx"},
        "profile": {"display": "lnu-checker-2026 (requested: lnu)"},
        "render_engine": "word-pdf",
        "evidence_source": "word-pdf",
        "evidence_trust": "authoritative",
        "layout_decision_eligible": True,
        "render_page_dir": "/tmp/render-proof/word_pdf_pages",
        "render_pdf_path": "/tmp/render-proof/render_verify_word.pdf",
        "render_fallback_used": False,
        "render_warnings": [],
        "layout_score": {"score": 82, "penalty": 18},
        "render_text_summary": {
            "source": "pdf",
            "available": True,
            "page_text_available_count": 2,
            "page_text_extraction_warning_count": 0,
            "warnings": [],
        },
        "overall_status": "verified",
        "readiness": "render-check-required",
        "preflight_status": "warning",
        "render_evidence_status": "render-review-required",
        "output_dir": "/tmp/render-proof",
        "page_count": 2,
        "report_path": "/tmp/render-proof/render_verify_report.md",
        "selected_scopes": ["headings", "figures_tables"],
        "wild_doc": {
            "detected": True,
            "signals": [
                {
                    "id": "table_heading_candidates",
                    "severity": "warning",
                    "message": "发现 2 个表格伪标题候选。",
                }
            ],
        },
        "render_findings": [
            {
                "id": "large_blank_region",
                "severity": "warning",
                "page": 2,
                "region": {"x": 0.08, "y": 0.7, "w": 0.84, "h": 0.3},
                "message": "页底存在大块连续空白。",
            }
        ],
        "review_items": ["逐页检查页底空白。", "逐页检查公式横线。"],
    }

    rendered = render_verify_module.render_render_verify_report(report)

    assert "文件: demo.docx" in rendered
    assert "渲染引擎: word-pdf" in rendered
    assert "渲染证据来源: word-pdf" in rendered
    assert "渲染证据可信度: authoritative" in rendered
    assert "版式决策可用: 是" in rendered
    assert "Layout Score: 82 (penalty 18)" in rendered
    assert "PDF文本页数: 2" in rendered
    assert "渲染 PDF: /tmp/render-proof/render_verify_word.pdf" in rendered
    assert "自动页图判读：" not in rendered
    assert "large_blank_region [warning]" not in rendered
    assert "可提交状态: render-check-required" in rendered
    assert "预检状态: warning" in rendered
    assert "渲染证据状态: render-review-required" in rendered
    assert "渲染证据目录: /tmp/render-proof" in rendered
    assert "复核范围: headings, figures_tables" in rendered
    assert "野生文档信号：" in rendered
    assert "table_heading_candidates [warning]" in rendered
    assert "逐页检查页底空白" not in rendered
    assert "- 逐页检查公式横线。" in rendered


def test_render_render_verify_report_preserves_cause_oriented_review_items():
    report = {
        "document": {"name": "demo.docx"},
        "profile": {"display": "lnu"},
        "render_engine": "word-pdf",
        "evidence_source": "word-pdf",
        "evidence_trust": "authoritative",
        "layout_decision_eligible": True,
        "overall_status": "verified",
        "readiness": "render-check-required",
        "preflight_status": "warning",
        "render_evidence_status": "render-review-required",
        "output_dir": "/tmp/render-proof",
        "page_count": 1,
        "report_path": "/tmp/render-proof/render_verify_report.md",
        "selected_scopes": ["figures_tables"],
        "wild_doc": {"detected": False, "signals": []},
        "layout_score": {"score": 80, "penalty": 20},
        "render_text_summary": {"source": "pdf", "available": True, "page_text_available_count": 1, "page_text_extraction_warning_count": 0, "warnings": []},
        "render_findings": [
            {"id": "large_blank_region", "severity": "warning", "page": 1, "message": "页底存在大块连续空白。"}
        ],
        "review_items": [
            "图表挤页常见于首次引用位置过晚、对象块过大，或图题注整体绑定到后页。",
            "如果不进入候选稿排障，建议回 Word/WPS 检查图片环绕、分页符和段落间距设置。",
        ],
    }

    rendered = render_verify_module.render_render_verify_report(report)

    assert "大块空白" not in rendered
    assert "对象塞不进当前页" not in rendered
    assert "检查图片环绕、分页符和段落间距设置" in rendered
