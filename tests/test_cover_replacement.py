from __future__ import annotations

import hashlib
import io
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

import pytest
from docx import Document
from docx.shared import Inches

import fix_thesis
from article_engine import apply_fix
from thesis_fix.cover_template import CoverReplacementError, validate_cover_package
from thesis_fix.runtime import build_fix_runtime, build_scope_flags
from thesis_tool.scopes import get_scope_definition
from thesis_tool.workflow import apply_scoped_fix


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"w": W_NS, "r": R_NS}
ASSET_DIR = Path("config/templates/lnu/cover-assets")
COVER_FIELDS = {
    "thesis_title": "基于免疫信息学的疫苗设计",
    "college": "生命科学院",
    "major": "生物技术",
    "student_name": "测试学生",
    "advisor": "测试教师",
    "completion_date": "2026年6月",
}


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _source_docx(path: Path, *, confident_cover: bool = True) -> Path:
    doc = Document()
    if confident_cover:
        doc.add_paragraph("辽宁大学")
        doc.add_paragraph("毕业论文（设计）")
        doc.add_paragraph("题目：旧题目")
        doc.add_paragraph("学院：旧学院")
    else:
        doc.add_paragraph("这是一段可能属于正文的普通文字")
    doc.add_paragraph("摘  要")
    doc.add_paragraph("这是摘要正文。")
    doc.add_paragraph("第一章 绪论", style="Heading 1")
    doc.add_paragraph("这是正文，必须完整保留。")
    doc.sections[0].top_margin = Inches(1.1)
    doc.sections[0].footer.paragraphs[0].text = "正文页脚"
    doc.save(path)
    return path


def _source_without_cover(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("摘  要")
    doc.add_paragraph("这是摘要正文。")
    doc.add_paragraph("第一章 绪论", style="Heading 1")
    doc.add_paragraph("这是正文。")
    doc.save(path)
    return path


def _document_root(path: Path) -> ET.Element:
    with zipfile.ZipFile(path) as archive:
        return ET.fromstring(archive.read("word/document.xml"))


def _relationships(path: Path) -> set[tuple[str | None, ...]]:
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
    return {
        (rel.get("Id"), rel.get("Type"), rel.get("Target"), rel.get("TargetMode"))
        for rel in root.findall(f"{{{PKG_REL_NS}}}Relationship")
    }


def _body_section_xml(path: Path) -> bytes:
    section = _document_root(path).find("w:body/w:sectPr", NS)
    assert section is not None
    return ET.tostring(section)


def _package_part(path: Path, name: str) -> bytes:
    with zipfile.ZipFile(path) as archive:
        return archive.read(name)


def _extract_package(docx_path: Path, target: Path) -> Path:
    with zipfile.ZipFile(docx_path) as archive:
        archive.extractall(target)
    return target


def test_cover_scope_and_runtime_require_explicit_complete_fields() -> None:
    scope = get_scope_definition("cover")
    assert scope.rule_ids == ()
    assert build_scope_flags(build_fix_runtime(scopes=["page"]).requested_scopes).cover is False

    with pytest.raises(ValueError, match="cover_fields"):
        build_fix_runtime(scopes=["cover"])
    with pytest.raises(ValueError, match="completion_date"):
        build_fix_runtime(scopes=["cover"], cover_fields={key: value for key, value in COVER_FIELDS.items() if key != "completion_date"})
    with pytest.raises(ValueError, match="advisor"):
        build_fix_runtime(scopes=["cover"], cover_fields={**COVER_FIELDS, "advisor": "   "})
    with pytest.raises(ValueError, match="控制字符"):
        build_fix_runtime(scopes=["cover"], cover_fields={**COVER_FIELDS, "advisor": "测试\n教师"})
    with pytest.raises(ValueError, match="cover scope"):
        build_fix_runtime(scopes=["page"], cover_fields=COVER_FIELDS)

    runtime = build_fix_runtime(scopes=["cover"], cover_fields=COVER_FIELDS)
    assert runtime.cover_fields == COVER_FIELDS
    assert build_scope_flags(runtime.requested_scopes).cover is True


def test_official_cover_assets_are_clean_extractions() -> None:
    assert _sha256(ASSET_DIR / "emblem.png") == "8857ba2f44c9a3043941465ff4c5d9d2ca1617d94ce670b8043181c2b3ce1be8"
    assert _sha256(ASSET_DIR / "wordmark.jpeg") == "bef9a3dad0b1c8e2abcaa1283e804a99fe7d806442b0c1f059fa299dd0fd3e27"


def test_cover_replacement_builds_fixed_a4_page_and_preserves_body_package(tmp_path: Path) -> None:
    source = _source_docx(tmp_path / "source.docx")
    output = tmp_path / "fixed.docx"
    source_hash = _sha256(source)
    source_styles = _package_part(source, "word/styles.xml")
    source_relationships = _relationships(source)
    source_body_section = _body_section_xml(source)

    result = apply_scoped_fix(
        str(source),
        str(output),
        profile_path="lnu",
        scopes=["cover"],
        cover_fields=COVER_FIELDS,
    )

    assert result == str(output)
    assert _sha256(source) == source_hash
    Document(output)
    root = _document_root(output)
    xml_text = "".join(root.itertext())
    assert all(value in xml_text for value in COVER_FIELDS.values())
    assert "旧题目" not in xml_text
    assert "所有文本框阅后删除" not in xml_text
    assert "样本提供格式示范" not in xml_text
    assert len(root.findall(".//w:drawing", NS)) == 2

    cover_section = root.find("w:body/w:p/w:pPr/w:sectPr", NS)
    assert cover_section is not None
    assert cover_section.find("w:type", NS).get(f"{{{W_NS}}}val") == "nextPage"
    page_size = cover_section.find("w:pgSz", NS)
    assert page_size is not None
    assert page_size.get(f"{{{W_NS}}}w") == "11906"
    assert page_size.get(f"{{{W_NS}}}h") == "16838"
    assert cover_section.find("w:footerReference", NS) is None
    assert cover_section.find("w:pgNumType", NS) is None
    assert _body_section_xml(output) == source_body_section
    assert _package_part(output, "word/styles.xml") == source_styles
    assert source_relationships <= _relationships(output)

    with zipfile.ZipFile(output) as archive:
        assert archive.read("word/footer1.xml")
        document_xml = archive.read("word/document.xml")
        relationships_xml = archive.read("word/_rels/document.xml.rels")
        content_types_xml = archive.read("[Content_Types].xml")
        media = {name: archive.read(name) for name in archive.namelist() if name.startswith("word/media/")}
    assert f'<Relationships xmlns="{PKG_REL_NS}"'.encode() in relationships_xml
    assert b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"' in content_types_xml
    declared_prefixes = {
        prefix
        for _, (prefix, _) in ET.iterparse(io.BytesIO(document_xml), events=("start-ns",))
        if prefix
    }
    ignorable = set(_document_root(output).get("{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable", "").split())
    assert ignorable <= declared_prefixes
    assert hashlib.sha256(next(value for name, value in media.items() if name.endswith("emblem.png"))).hexdigest() == _sha256(ASSET_DIR / "emblem.png")
    assert hashlib.sha256(next(value for name, value in media.items() if name.endswith("wordmark.jpeg"))).hexdigest() == _sha256(ASSET_DIR / "wordmark.jpeg")


def test_cover_package_validator_accepts_complete_replacement(tmp_path: Path) -> None:
    source = _source_docx(tmp_path / "source.docx")
    output = tmp_path / "fixed.docx"
    apply_scoped_fix(str(source), str(output), scopes=["cover"], cover_fields=COVER_FIELDS)

    package_root = _extract_package(output, tmp_path / "package")

    validate_cover_package(str(package_root), {}, COVER_FIELDS)


def test_cover_package_validator_rejects_incomplete_structure(tmp_path: Path) -> None:
    source = _source_docx(tmp_path / "source.docx")
    output = tmp_path / "fixed.docx"
    apply_scoped_fix(str(source), str(output), scopes=["cover"], cover_fields=COVER_FIELDS)
    package_root = _extract_package(output, tmp_path / "package")
    document_root = ET.parse(package_root / "word/document.xml").getroot()
    advisor_text = next(item for item in document_root.findall(".//w:t", NS) if item.text == COVER_FIELDS["advisor"])
    advisor_text.text = ""
    broken_document = ET.tostring(document_root, encoding="utf-8", xml_declaration=True)

    with pytest.raises(CoverReplacementError, match="六项字段"):
        validate_cover_package(str(package_root), {"word/document.xml": broken_document}, COVER_FIELDS)

    cover_section = document_root.find("w:body/w:p/w:pPr/w:sectPr", NS)
    assert cover_section is not None
    advisor_text.text = COVER_FIELDS["advisor"]
    ET.SubElement(cover_section, f"{{{W_NS}}}footerReference", {f"{{{R_NS}}}id": "rIdFooter"})
    broken_document = ET.tostring(document_root, encoding="utf-8", xml_declaration=True)
    with pytest.raises(CoverReplacementError, match="封面分节"):
        validate_cover_package(str(package_root), {"word/document.xml": broken_document}, COVER_FIELDS)


def test_cover_package_validator_rejects_missing_duplicate_field_value(tmp_path: Path) -> None:
    source = _source_docx(tmp_path / "source.docx")
    output = tmp_path / "fixed.docx"
    duplicate_fields = {name: "相同字段" for name in COVER_FIELDS}
    apply_scoped_fix(str(source), str(output), scopes=["cover"], cover_fields=duplicate_fields)
    package_root = _extract_package(output, tmp_path / "package")
    document_root = ET.parse(package_root / "word/document.xml").getroot()
    value = next(item for item in document_root.findall(".//w:t", NS) if item.text == "相同字段")
    value.text = ""

    with pytest.raises(CoverReplacementError, match="六项字段"):
        validate_cover_package(
            str(package_root),
            {"word/document.xml": ET.tostring(document_root, encoding="utf-8", xml_declaration=True)},
            duplicate_fields,
        )


def test_cover_package_validator_rejects_broken_opc_images(tmp_path: Path) -> None:
    source = _source_docx(tmp_path / "source.docx")
    output = tmp_path / "fixed.docx"
    apply_scoped_fix(str(source), str(output), scopes=["cover"], cover_fields=COVER_FIELDS)
    package_root = _extract_package(output, tmp_path / "package")
    rels_path = package_root / "word/_rels/document.xml.rels"
    rels_root = ET.parse(rels_path).getroot()
    cover_relation = next(
        item
        for item in rels_root.findall(f"{{{PKG_REL_NS}}}Relationship")
        if str(item.get("Id", "")).startswith("rIdCover")
    )
    rels_root.remove(cover_relation)
    broken_relationships = ET.tostring(rels_root, encoding="utf-8", xml_declaration=True)

    with pytest.raises(CoverReplacementError, match="图片关系"):
        validate_cover_package(
            str(package_root),
            {"word/_rels/document.xml.rels": broken_relationships},
            COVER_FIELDS,
        )


def test_cover_postcondition_gate_runs_before_output_write(tmp_path: Path, monkeypatch) -> None:
    source = _source_docx(tmp_path / "source.docx")
    output = tmp_path / "must-not-exist.docx"

    def reject_package(*_args, **_kwargs):
        raise CoverReplacementError("封面结构后置条件失败")

    monkeypatch.setattr(fix_thesis, "validate_cover_package", reject_package)
    with pytest.raises(CoverReplacementError, match="结构后置条件"):
        apply_scoped_fix(str(source), str(output), scopes=["cover"], cover_fields=COVER_FIELDS)

    assert not output.exists()


def test_cover_and_page_scopes_keep_footer_relationships_resolvable(tmp_path: Path) -> None:
    source = _source_docx(tmp_path / "source.docx")
    output = tmp_path / "fixed.docx"

    apply_scoped_fix(
        str(source),
        str(output),
        profile_path="lnu",
        scopes=["cover", "page"],
        cover_fields=COVER_FIELDS,
    )

    root = _document_root(output)
    body = root.find("w:body", NS)
    assert body is not None
    children = list(body)
    field_table_index = next(index for index, child in enumerate(children) if child.tag == f"{{{W_NS}}}tbl")
    inline_section_paragraphs = [
        child
        for child in children
        if child.find("w:pPr/w:sectPr", NS) is not None
    ]
    assert len(inline_section_paragraphs) == 2
    assert children.index(inline_section_paragraphs[0]) == field_table_index + 1
    cover_section = inline_section_paragraphs[0].find("w:pPr/w:sectPr", NS)
    assert cover_section is not None
    assert cover_section.find("w:footerReference", NS) is None
    assert cover_section.find("w:pgNumType", NS) is None
    referenced_ids = {
        item.get(f"{{{R_NS}}}id")
        for item in root.findall(".//w:footerReference", NS)
    }
    relationship_ids = {item[0] for item in _relationships(output)}
    assert referenced_ids
    assert referenced_ids <= relationship_ids


def test_cover_and_page_scopes_preserve_cover_break_without_frontmatter(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    output = tmp_path / "fixed.docx"
    doc = Document()
    for text in ("辽宁大学", "毕业论文（设计）", "题目：旧题目", "学院：旧学院"):
        doc.add_paragraph(text)
    doc.add_paragraph("第一章 绪论", style="Heading 1")
    doc.add_paragraph("这是正文。")
    doc.save(source)

    apply_scoped_fix(
        str(source),
        str(output),
        profile_path="lnu",
        scopes=["cover", "page"],
        cover_fields=COVER_FIELDS,
    )

    root = _document_root(output)
    body = root.find("w:body", NS)
    assert body is not None
    children = list(body)
    table_index = next(index for index, child in enumerate(children) if child.tag == f"{{{W_NS}}}tbl")
    inline_sections = [child for child in children if child.find("w:pPr/w:sectPr", NS) is not None]
    assert len(inline_sections) == 1
    assert children.index(inline_sections[0]) == table_index + 1
    cover_section = inline_sections[0].find("w:pPr/w:sectPr", NS)
    assert cover_section is not None
    assert cover_section.find("w:footerReference", NS) is None
    assert cover_section.find("w:pgNumType", NS) is None


def test_cover_replacement_fails_closed_on_ambiguous_prefix(tmp_path: Path) -> None:
    source = _source_docx(tmp_path / "ambiguous.docx", confident_cover=False)
    output = tmp_path / "must-not-exist.docx"
    original = source.read_bytes()

    with pytest.raises(CoverReplacementError, match="封面边界"):
        apply_scoped_fix(
            str(source),
            str(output),
            scopes=["cover"],
            cover_fields=COVER_FIELDS,
        )

    assert source.read_bytes() == original
    assert not output.exists()


def test_cover_replacement_never_overwrites_input(tmp_path: Path) -> None:
    source = _source_docx(tmp_path / "same-path.docx")
    original = source.read_bytes()

    with pytest.raises(ValueError, match="输入.*输出"):
        apply_scoped_fix(
            str(source),
            str(source),
            scopes=["cover"],
            cover_fields=COVER_FIELDS,
        )

    assert source.read_bytes() == original


def test_non_cover_scope_does_not_replace_cover(tmp_path: Path) -> None:
    source = _source_docx(tmp_path / "not-selected.docx")
    output = tmp_path / "not-selected-fixed.docx"

    apply_scoped_fix(str(source), str(output), scopes=["headings"])

    text = "".join(_document_root(output).itertext())
    assert "题目：旧题目" in text
    assert COVER_FIELDS["thesis_title"] not in text


def test_article_engine_reports_cover_replacement_status(tmp_path: Path) -> None:
    source = _source_docx(tmp_path / "service-source.docx")
    output = tmp_path / "service-fixed.docx"

    payload = apply_fix(
        str(source),
        output_path=str(output),
        profile_path="lnu",
        scopes=["cover"],
        cover_fields=COVER_FIELDS,
        force=True,
    )

    assert payload["cover_replacement"] == {"status": "replaced", "field_count": 6}
    assert payload["output"]["path"] == str(output)
    assert output.exists()


def test_article_engine_reports_inserted_cover_status(tmp_path: Path) -> None:
    source = _source_without_cover(tmp_path / "service-source.docx")
    output = tmp_path / "service-fixed.docx"

    payload = apply_fix(
        str(source),
        output_path=str(output),
        profile_path="lnu",
        scopes=["cover"],
        cover_fields=COVER_FIELDS,
        force=True,
    )

    assert payload["cover_replacement"] == {"status": "inserted", "field_count": 6}
    assert output.exists()
