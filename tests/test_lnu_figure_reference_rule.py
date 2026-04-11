from __future__ import annotations

from pathlib import Path

import audit_thesis
from docx import Document


def _build_doc_with_caption_reference(path: str | Path, *, mention_before_caption: bool) -> Path:
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    heading = doc.add_paragraph("第一章 绪论")
    heading.style = "Heading 1"

    if mention_before_caption:
        doc.add_paragraph("如图1.1所示，实验装置由反应器和控制器组成。")
    else:
        doc.add_paragraph("本节介绍实验装置。")

    doc.add_paragraph("图1.1 实验装置示意图")
    doc.save(output_path)
    return output_path


def _rule_status(docx_path: str | Path, rule_id: str) -> dict:
    results, _score, _report = audit_thesis.audit_docx(str(docx_path), profile_path="lnu")
    return next(result for result in results if result["id"] == rule_id)


def test_lnu_f05_passes_when_caption_is_referenced_in_body(tmp_path):
    docx_path = _build_doc_with_caption_reference(
        tmp_path / "lnu_f05_pass.docx",
        mention_before_caption=True,
    )

    result = _rule_status(docx_path, "LNU_F05")
    assert result["passed"], result["issues"]


def test_lnu_f05_fails_when_caption_has_no_prior_body_reference(tmp_path):
    docx_path = _build_doc_with_caption_reference(
        tmp_path / "lnu_f05_fail.docx",
        mention_before_caption=False,
    )

    result = _rule_status(docx_path, "LNU_F05")
    assert not result["passed"]
    assert result["issues"]
