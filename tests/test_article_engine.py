from pathlib import Path

import pytest
from docx import Document

from article_engine import apply_fix, audit_document, normalize_document, plan_document, preflight_document, verify_document

from .conftest import RULE_MUTATORS, make_compliant_doc


def _build_mutated_doc(tmp_docx, *, filename: str, rule_ids: tuple[str, ...]) -> Path:
    path = tmp_docx(make_compliant_doc, filename=filename)
    doc = Document(path)
    for rule_id in rule_ids:
        RULE_MUTATORS[rule_id](doc)
    doc.save(path)
    return path


def _make_high_risk_lnu_doc(source_path: Path) -> Path:
    doc = Document()
    heading = doc.add_paragraph("1 绪论")
    heading.style = doc.styles["Heading 1"]
    doc.add_paragraph("这是正文示例。")
    table = doc.add_table(rows=5, cols=1)
    values = [
        "1,4-Diaminobutane",
        "2-Amino-2-Deoxy-D-Glucopyranose",
        "2.80E+08",
        "3.20E+01",
        "2.79E+08",
    ]
    for row, value in zip(table.rows, values):
        row.cells[0].text = value
    doc.save(source_path)
    return source_path


def _make_style_conflict_lnu_doc(source_path: Path) -> Path:
    doc = Document()
    heading = doc.add_paragraph("1 绪论")
    heading.style = doc.styles["Heading 1"]
    conflict = doc.add_paragraph("3.6 分子对接验证结果")
    conflict.style = doc.styles["Heading 1"]
    doc.add_paragraph("这是正文示例。")
    doc.save(source_path)
    return source_path


def _make_field_only_toc_lnu_doc(source_path: Path) -> Path:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    title = doc.add_paragraph("目  录")
    title_p_pr = title._p.get_or_add_pPr()
    title_style = OxmlElement("w:pStyle")
    title_style.set(qn("w:val"), "TOCHeading")
    title_p_pr.append(title_style)
    title_jc = OxmlElement("w:jc")
    title_jc.set(qn("w:val"), "center")
    title_p_pr.append(title_jc)
    title_run = title.runs[0]
    title_run.font.name = "Times New Roman"
    title_run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "黑体")
    title_run.font.size = None
    title_sz = OxmlElement("w:sz")
    title_sz.set(qn("w:val"), "32")
    title_run._element.get_or_add_rPr().append(title_sz)

    field = doc.add_paragraph("")
    field_p_pr = field._p.get_or_add_pPr()
    field_style = OxmlElement("w:pStyle")
    field_style.set(qn("w:val"), "TOCField")
    field_p_pr.append(field_style)
    run = field.add_run("")
    instr = OxmlElement("w:instrText")
    instr.text = ' TOC \\\\o "1-3" \\\\h \\\\z \\\\u '
    run._r.append(instr)

    heading = doc.add_paragraph("1 绪论")
    heading.style = doc.styles["Heading 1"]
    doc.save(source_path)
    return source_path


def test_audit_document_returns_structured_payload(tmp_docx):
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_engine_audit.docx",
        rule_ids=("H02", "KW01"),
    )

    payload = audit_document(str(source_path))

    assert payload["document"]["name"] == "article_engine_audit.docx"
    assert payload["summary"]["failed_rules"] >= 2
    assert payload["summary"]["total_rules"] >= payload["summary"]["failed_rules"]
    assert payload["profile"]["id"] == "cn-common"
    assert "headings" in payload["recommended_scope_order"]
    assert "abstract" in payload["recommended_scope_order"]

    failed_by_id = {item["id"]: item for item in payload["failed_results"]}
    assert failed_by_id["H02"]["action"] == "autofix"
    assert failed_by_id["KW01"]["action"] == "manual_review"


def test_audit_document_supports_lnu_profile_alias(tmp_docx):
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_engine_lnu_audit.docx",
        rule_ids=("H02",),
    )

    payload = audit_document(str(source_path), profile_path="lnu")

    assert payload["profile"]["id"] == "lnu-checker-2026"


def test_audit_document_defaults_to_strict_profile_for_explicit_profile(tmp_docx):
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_engine_missing_profile_audit.docx",
        rule_ids=("H02",),
    )

    with pytest.raises(ValueError, match="Profile 加载失败"):
        audit_document(str(source_path), profile_path="missing-profile.yaml")


def test_audit_document_can_explicitly_allow_profile_fallback(tmp_docx):
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_engine_missing_profile_fallback_audit.docx",
        rule_ids=("H02",),
    )

    payload = audit_document(
        str(source_path),
        profile_path="missing-profile.yaml",
        strict_profile=False,
    )

    assert payload["profile"]["id"] == "cn-common"
    assert payload["profile"]["fallback_used"] is True


def test_plan_document_can_focus_on_selected_scope(tmp_docx):
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_engine_plan.docx",
        rule_ids=("H02", "T01"),
    )

    payload = plan_document(str(source_path), scopes=["headings"])

    assert payload["document"]["name"] == "article_engine_plan.docx"
    assert payload["selected_scopes"] == ["headings"]
    assert payload["recommended_scope_order"] == ["headings"]
    assert [scope["id"] for scope in payload["scopes"]] == ["headings"]
    assert payload["scopes"][0]["failed_count"] >= 1


def test_preflight_document_surfaces_blocking_structure_risk(tmp_path):
    source_path = _make_style_conflict_lnu_doc(Path(tmp_path) / "article_engine_preflight_blocked.docx")

    payload = preflight_document(str(source_path), profile_path="lnu")

    assert payload["document"]["name"] == "article_engine_preflight_blocked.docx"
    assert payload["preflight_status"] == "blocked"
    assert payload["summary"]["preflight_status"] == "blocked"
    assert payload["heading_renumber_guard"]["status"] == "block"
    assert payload["summary"]["style_conflict_count"] >= 1
    assert payload["recommended_actions"]


def test_normalize_document_reduces_wild_doc_structure_risk(tmp_path):
    source_path = _make_style_conflict_lnu_doc(Path(tmp_path) / "article_engine_normalize.docx")
    output_path = Path(tmp_path) / "article_engine_normalized.docx"

    payload = normalize_document(
        str(source_path),
        output_path=str(output_path),
        profile_path="lnu",
    )

    assert output_path.exists()
    assert payload["document"]["name"] == "article_engine_normalize.docx"
    assert payload["output"]["name"] == "article_engine_normalized.docx"
    assert payload["changed"] is True
    assert any(item["id"] == "heading_styles" for item in payload["operations"])
    assert payload["before"]["preflight_status"] == "blocked"
    assert payload["after"]["preflight_status"] == "warning"
    assert payload["before"]["style_conflict_count"] >= 1
    assert payload["after"]["style_conflict_count"] == 0
    assert payload["next_steps"]


def test_verify_document_returns_scope_status(tmp_docx):
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_engine_verify.docx",
        rule_ids=("H02",),
    )

    payload = verify_document(str(source_path), scopes=["headings"])

    assert payload["document"]["name"] == "article_engine_verify.docx"
    assert payload["overall_status"] == "needs_fix"
    assert payload["readiness"] == "needs-fix"
    assert payload["selected_scopes"] == ["headings"]
    assert payload["summary"]["failed_rules"] >= 1
    assert payload["summary"]["readiness"] == "needs-fix"
    assert any(scope["id"] == "headings" for scope in payload["scopes"])


def test_verify_document_marks_field_only_toc_as_render_check_required(tmp_path):
    source_path = _make_field_only_toc_lnu_doc(Path(tmp_path) / "article_engine_verify_field_only_toc.docx")

    payload = verify_document(str(source_path), profile_path="lnu", scopes=["toc"])

    assert payload["overall_status"] == "verified"
    assert payload["readiness"] == "render-check-required"
    assert payload["summary"]["render_check_rules"] == 1
    assert payload["render_check_rule_ids"] == ["TOC_REFRESH_REQUIRED"]
    assert payload["render_check_rules"][0]["scope_id"] == "toc"


def test_apply_fix_returns_output_and_verification(tmp_docx, tmp_path):
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_engine_apply.docx",
        rule_ids=("H02", "T01"),
    )
    output_path = Path(tmp_path) / "article_engine_apply_fixed.docx"

    payload = apply_fix(
        str(source_path),
        output_path=str(output_path),
        scopes=["headings"],
    )

    assert payload["mode"] == "apply"
    assert payload["output"]["path"] == str(output_path)
    assert output_path.exists()
    assert payload["verification"]["selected_scopes"] == ["headings"]
    assert payload["verification"]["overall_status"] == "verified"
    assert payload["verification"]["readiness"] == "structure-ready"
    assert payload["readiness"] == "render-check-required"
    assert payload["guard"]["checked"] is True
    assert payload["guard"]["warnings"] == []
    assert payload["post_verify_notices"] == []


def test_apply_fix_dry_run_returns_preview_without_writing(tmp_docx, tmp_path):
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_engine_preview.docx",
        rule_ids=("H02",),
    )
    output_path = Path(tmp_path) / "article_engine_preview_fixed.docx"

    payload = apply_fix(
        str(source_path),
        output_path=str(output_path),
        scopes=["headings"],
        dry_run=True,
    )

    assert payload["mode"] == "preview"
    assert payload["output"]["path"] == str(output_path)
    assert payload["selected_scopes"] == ["headings"]
    assert not output_path.exists()


def test_apply_fix_warns_but_does_not_block_table_risk_only_heading_renumber(tmp_path):
    source_path = _make_high_risk_lnu_doc(Path(tmp_path) / "article_engine_guard_warn.docx")
    output_path = Path(tmp_path) / "article_engine_guard_warn_fixed.docx"

    payload = apply_fix(
        str(source_path),
        output_path=str(output_path),
        profile_path="lnu",
        scopes=["headings"],
        renumber_headings=True,
    )

    assert output_path.exists()
    assert payload["guard"]["checked"] is True
    assert payload["guard"]["blocked"] is False
    assert payload["guard"]["force_used"] is False
    assert any("表格伪标题候选" in line for line in payload["guard"]["warnings"])
    assert payload["verification"]["overall_status"] == "manual_review"
    assert payload["verification"]["readiness"] == "manual-review-required"
    assert payload["readiness"] == "manual-review-required"
    assert payload["post_verify_notices"]


def test_apply_fix_blocks_style_conflict_heading_renumber_without_force(tmp_path):
    source_path = _make_style_conflict_lnu_doc(Path(tmp_path) / "article_engine_guard_block.docx")
    output_path = Path(tmp_path) / "article_engine_guard_block_fixed.docx"

    with pytest.raises(RuntimeError, match="Apply blocked by structural risk"):
        apply_fix(
            str(source_path),
            output_path=str(output_path),
            profile_path="lnu",
            scopes=["headings"],
            renumber_headings=True,
        )

    assert not output_path.exists()


def test_apply_fix_allows_force_on_style_conflict_heading_renumber(tmp_path):
    source_path = _make_style_conflict_lnu_doc(Path(tmp_path) / "article_engine_guard_force.docx")
    output_path = Path(tmp_path) / "article_engine_guard_force_fixed.docx"

    payload = apply_fix(
        str(source_path),
        output_path=str(output_path),
        profile_path="lnu",
        scopes=["headings"],
        renumber_headings=True,
        force=True,
    )

    assert output_path.exists()
    assert payload["guard"]["checked"] is False
    assert payload["guard"]["force_used"] is True
    assert payload["guard"]["warnings"] == []
