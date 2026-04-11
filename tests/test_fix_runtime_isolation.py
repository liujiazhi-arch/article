from pathlib import Path

import audit_thesis
from docx import Document

from thesis_tool.workflow import apply_scoped_fix

from .conftest import (
    RULE_MUTATORS,
    _add_heading,
    _set_paragraph_properties,
    _set_reference_format,
    _set_run_format,
    make_compliant_doc,
    make_violating_doc,
)


def _audit_rule(docx_path: str | Path, rule_id: str, profile_path: str | None = None) -> dict:
    results, _score, _report = audit_thesis.audit_docx(str(docx_path), profile_path=profile_path)
    return next(item for item in results if item["id"] == rule_id)


def test_fix_docx_profile_runtime_isolation_between_calls(tmp_docx, tmp_path):
    # First call under LNU profile.
    lnu_source = tmp_docx(make_violating_doc, filename="h02_lnu_source.docx", rule_id="H02")
    lnu_fixed = Path(tmp_path) / "h02_lnu_fixed.docx"
    apply_scoped_fix(str(lnu_source), str(lnu_fixed), profile_path="lnu", scopes=["headings"])

    assert _audit_rule(lnu_fixed, "H02", profile_path="lnu")["passed"]

    # Second call without profile must not inherit LNU runtime state.
    default_source = tmp_docx(make_violating_doc, filename="h02_default_source.docx", rule_id="H02")
    default_fixed = Path(tmp_path) / "h02_default_fixed.docx"
    apply_scoped_fix(str(default_source), str(default_fixed), profile_path=None, scopes=["headings"])

    default_h02 = _audit_rule(default_fixed, "H02", profile_path=None)
    lnu_h02 = _audit_rule(default_fixed, "H02", profile_path="lnu")
    assert default_h02["passed"], default_h02["issues"]
    assert not lnu_h02["passed"]


def test_fix_docx_profile_runtime_isolation_under_alternating_profiles(tmp_docx, tmp_path):
    first_source = tmp_docx(make_violating_doc, filename="h02_default_first_source.docx", rule_id="H02")
    first_fixed = Path(tmp_path) / "h02_default_first_fixed.docx"
    apply_scoped_fix(str(first_source), str(first_fixed), profile_path=None, scopes=["headings"])

    first_default_h02 = _audit_rule(first_fixed, "H02", profile_path=None)
    first_lnu_h02 = _audit_rule(first_fixed, "H02", profile_path="lnu")
    assert first_default_h02["passed"], first_default_h02["issues"]
    assert not first_lnu_h02["passed"]

    second_source = tmp_docx(make_violating_doc, filename="h02_lnu_second_source.docx", rule_id="H02")
    second_fixed = Path(tmp_path) / "h02_lnu_second_fixed.docx"
    apply_scoped_fix(str(second_source), str(second_fixed), profile_path="lnu", scopes=["headings"])

    second_lnu_h02 = _audit_rule(second_fixed, "H02", profile_path="lnu")
    assert second_lnu_h02["passed"], second_lnu_h02["issues"]

    third_source = tmp_docx(make_violating_doc, filename="h02_default_third_source.docx", rule_id="H02")
    third_fixed = Path(tmp_path) / "h02_default_third_fixed.docx"
    apply_scoped_fix(str(third_source), str(third_fixed), profile_path=None, scopes=["headings"])

    third_default_h02 = _audit_rule(third_fixed, "H02", profile_path=None)
    third_lnu_h02 = _audit_rule(third_fixed, "H02", profile_path="lnu")
    assert third_default_h02["passed"], third_default_h02["issues"]
    assert not third_lnu_h02["passed"]


def test_scope_and_profile_combination_applies_only_selected_scope(tmp_docx, tmp_path):
    source_path = tmp_docx(make_compliant_doc, filename="scope_profile_source.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    RULE_MUTATORS["T01"](doc)
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "scope_profile_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["headings"],
    )

    lnu_h02 = _audit_rule(fixed_path, "H02", profile_path="lnu")
    lnu_t01 = _audit_rule(fixed_path, "T01", profile_path="lnu")
    assert lnu_h02["passed"], lnu_h02["issues"]
    assert not lnu_t01["passed"]


def test_references_scope_fixes_reference_paragraphs_without_reference_heading(tmp_path):
    source_path = Path(tmp_path) / "references_without_title_source.docx"
    fixed_path = Path(tmp_path) / "references_without_title_fixed.docx"

    doc = Document()
    _add_heading(doc, "第一章 绪论", level=1, size=30)
    body = doc.add_paragraph()
    _set_paragraph_properties(body)
    _set_run_format(body.add_run("这是正文示例。"), size=24)

    reference = doc.add_paragraph()
    _set_run_format(reference.add_run("[1]\tExample reference."), size=24)
    _set_reference_format(reference, left=0, hanging=0, tab_pos=200)
    doc.save(source_path)

    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["references"],
    )

    r01 = _audit_rule(fixed_path, "R01", profile_path="lnu")
    r03 = _audit_rule(fixed_path, "R03", profile_path="lnu")
    assert r01["passed"], r01["issues"]
    assert r03["passed"], r03["issues"]
