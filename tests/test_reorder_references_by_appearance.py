from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS_DIR = PROJECT_ROOT / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from reorder_references_by_appearance import reorder_references_in_document


def test_reorder_references_stops_before_acknowledgement_and_keeps_entries_under_reference_heading():
    doc = Document()
    doc.add_paragraph("正文中先引用[2]，后引用[1]。")
    doc.add_paragraph("参考文献")
    doc.add_paragraph("[1] First reference.")
    doc.add_paragraph("[2] Second reference.")
    doc.add_paragraph("致  谢")
    doc.add_paragraph("感谢老师指导。")

    reorder_references_in_document(doc, trailing_space=False)

    texts = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
    assert texts == [
        "正文中先引用[1]，后引用[2]。",
        "参考文献",
        "[1]Second reference.",
        "[2]First reference.",
        "致  谢",
        "感谢老师指导。",
    ]


def test_reorder_references_stops_before_appendix_variant():
    doc = Document()
    doc.add_paragraph("正文中先引用[2]，后引用[1]。")
    doc.add_paragraph("参考文献")
    doc.add_paragraph("[1] First reference.")
    doc.add_paragraph("[2] Second reference.")
    doc.add_paragraph("附录A")
    doc.add_paragraph("[3] Appendix note.")

    reorder_references_in_document(doc, trailing_space=False)

    texts = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
    assert texts == [
        "正文中先引用[1]，后引用[2]。",
        "参考文献",
        "[1]Second reference.",
        "[2]First reference.",
        "附录A",
        "[3] Appendix note.",
    ]
