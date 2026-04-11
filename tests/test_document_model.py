import xml.etree.ElementTree as ET

from docx import Document

from _thesis_utils import build_document_model, build_style_map

from .conftest import _add_heading, _set_reference_format, audit_thesis


def _build_model(doc: Document):
    document_root = ET.fromstring(doc._element.xml)
    styles_root = ET.fromstring(doc.styles.element.xml)
    style_map = build_style_map(styles_root)
    return document_root, style_map, build_document_model(document_root, style_map)


def test_build_document_model_classifies_abstract_modules():
    doc = Document()
    doc.add_paragraph("摘 要")
    doc.add_paragraph("这是中文摘要正文。")
    doc.add_paragraph("关键词：测试；模型；分类")
    doc.add_paragraph("Abstract")
    doc.add_paragraph("This is the English abstract body.")
    doc.add_paragraph("Key words: model; workflow; test")
    _add_heading(doc, "1 绪论", level=1, size=30)

    _, _, model = _build_model(doc)

    cn_modules = [node.module for node in model.section_nodes("abstract_cn")]
    en_modules = [node.module for node in model.section_nodes("abstract_en")]

    assert cn_modules == ["abstract_cn_title", "abstract_cn_body", "abstract_cn_keywords"]
    assert en_modules == ["abstract_en_title", "abstract_en_body", "abstract_en_keywords"]


def test_build_document_model_classifies_keywords_without_space_as_abstract_keywords():
    doc = Document()
    doc.add_paragraph("Abstract")
    doc.add_paragraph("This is the English abstract body.")
    doc.add_paragraph("Keywords: model; workflow; test")
    _add_heading(doc, "1 绪论", level=1, size=30)

    _, _, model = _build_model(doc)

    en_modules = [node.module for node in model.section_nodes("abstract_en")]
    assert en_modules == ["abstract_en_title", "abstract_en_body", "abstract_en_keywords"]


def test_build_document_model_resolves_backmatter_bucket_modules():
    doc = Document()
    _add_heading(doc, "1 绪论", level=1, size=30)
    doc.add_paragraph("这是正文段落。")
    _add_heading(doc, "参考文献", level=1, size=30)
    ref = doc.add_paragraph("[1]\tSome reference.")
    _set_reference_format(ref)

    _, _, model = _build_model(doc)

    reference_nodes = model.section_nodes("references")
    assert [node.module for node in reference_nodes] == ["references_title", "references_entry"]
    assert reference_nodes[1].container_section == "backmatter"
    assert reference_nodes[1].protected is True


def test_build_paragraph_contexts_keeps_raw_section_and_exposes_effective_section():
    doc = Document()
    _add_heading(doc, "1 绪论", level=1, size=30)
    doc.add_paragraph("这是正文段落。")
    _add_heading(doc, "参考文献", level=1, size=30)
    ref = doc.add_paragraph("[1]\tSome reference.")
    _set_reference_format(ref)

    document_root, style_map, _ = _build_model(doc)
    contexts = audit_thesis.build_paragraph_contexts(document_root, style_map)

    ref_ctx = next(ctx for ctx in contexts if ctx["kind"] == "reference")
    assert ref_ctx["section"] == "backmatter"
    assert ref_ctx["effective_section"] == "references"
    assert ref_ctx["module"] == "references_entry"


def test_build_document_model_exits_toc_on_first_non_toc_paragraph():
    doc = Document()
    doc.add_paragraph("目 录")
    doc.add_paragraph("1 绪论\t1")
    doc.add_paragraph("这是正文第一段。")
    doc.add_paragraph("这是正文第二段。")

    _, _, model = _build_model(doc)

    sections = [node.container_section for node in model.paragraphs]
    assert sections == ["toc", "toc", "body", "body"]
