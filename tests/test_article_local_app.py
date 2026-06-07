from __future__ import annotations

import json
import io
import os
import shlex
import sys
from pathlib import Path

import article_api.local_app as local_app
import article_api.local_env as local_env
import article_api.storage as storage
from article_api.uploads import resolve_runtime_root
from docx import Document
import pytest

def _make_style_conflict_lnu_doc(source_path: Path) -> Path:
    doc = Document()
    heading = doc.add_paragraph("1 绪论")
    heading.style = doc.styles["Heading 1"]
    conflict = doc.add_paragraph("3.6 分子对接验证结果")
    conflict.style = doc.styles["Heading 1"]
    doc.add_paragraph("这是正文示例。")
    doc.save(source_path)
    return source_path


def test_build_doctor_report_uses_requested_roots(monkeypatch, tmp_path):
    state_root = tmp_path / "state"
    runtime_root = tmp_path / "runtime"
    monkeypatch.delenv(storage.STATE_ROOT_ENV_VAR, raising=False)
    monkeypatch.delenv("ARTICLE_API_RUNTIME_ROOT", raising=False)

    report = local_app.build_doctor_report(
        state_root=str(state_root),
        runtime_root=str(runtime_root),
    )

    assert report["status"] == "ok"
    assert report["summary"]["headline"]
    assert report["summary"]["issue_count"] == 0
    assert report["summary"]["recommended_actions"]
    assert report["roots"]["shared_root"] is False
    assert report["roots"]["state_root"] == str(state_root.resolve())
    assert report["roots"]["runtime_root"] == str(runtime_root.resolve())
    assert report["checks"]["storage"]["status"] == "ok"
    assert report["storage"]["schema_version"] == storage.SCHEMA_VERSION
    assert report["storage"]["index_count"] >= 1
    assert report["runtime"]["worker_model"] == "single"
    assert report["runtime"]["status"] == "ok"
    assert report["inventory"]["runtime_root"]["managed_directory_count"] == 0
    assert report["roots"]["runtime_root_writable"] is True
    assert "article-local serve" in report["workflow"]["serve"]


def test_build_doctor_report_does_not_leak_root_envs(monkeypatch, tmp_path):
    import os as local_os

    monkeypatch.delenv(storage.STATE_ROOT_ENV_VAR, raising=False)
    monkeypatch.delenv("ARTICLE_API_RUNTIME_ROOT", raising=False)

    local_app.build_doctor_report(
        state_root=str(tmp_path / "state"),
        runtime_root=str(tmp_path / "runtime"),
    )

    assert storage.STATE_ROOT_ENV_VAR not in local_os.environ
    assert "ARTICLE_API_RUNTIME_ROOT" not in local_os.environ


def test_article_local_help_lists_expected_subcommands(capsys):
    with pytest.raises(SystemExit) as exc_info:
        local_app.main(["--help"])

    assert exc_info.value.code == 0
    captured = capsys.readouterr()
    assert "usage: article-local" in captured.out
    for subcommand in (
        "init",
        "serve",
        "doctor",
        "preflight",
        "normalize",
        "render-verify",
        "render-workflow-modes",
        "profiles",
        "maintain",
        "backup",
        "feedback",
        "restore",
    ):
        assert subcommand in captured.out
    assert "batch" not in captured.out


def test_module_entrypoint_exits_with_main_result(monkeypatch):
    received = []

    def fake_main(argv=None):
        received.append(argv)
        return 7

    monkeypatch.setattr(local_app, "main", fake_main)

    with pytest.raises(SystemExit) as exc_info:
        local_app.module_main()

    assert exc_info.value.code == 7
    assert received == [None]


def test_emit_json_writes_utf8_when_stdout_encoding_rejects_chinese(monkeypatch):
    output = io.BytesIO()
    cp1252_stdout = io.TextIOWrapper(output, encoding="cp1252", errors="strict")
    monkeypatch.setattr(sys, "stdout", cp1252_stdout)

    local_app._emit_json({"message": "论文格式检查"})

    cp1252_stdout.flush()
    payload = json.loads(output.getvalue().decode("utf-8"))
    assert payload["message"] == "论文格式检查"


@pytest.mark.parametrize(
    ("entrypoint", "argv0", "expected_usage", "expected_fragment"),
    [
        (local_app.serve_main, "article-api", "usage: article-local serve", "--host"),
        (local_app.doctor_main, "article-doctor", "usage: article-local doctor", "--state-root"),
        (local_app.maintain_main, "article-maintain", "usage: article-local maintain", "--vacuum"),
        (local_app.backup_main, "article-backup", "usage: article-local backup", "output"),
        (local_app.feedback_main, "article-feedback", "usage: article-local feedback", "output"),
        (local_app.restore_main, "article-restore", "usage: article-local restore", "--force"),
    ],
)
def test_console_script_alias_help_routes_to_expected_subcommand(
    monkeypatch,
    capsys,
    entrypoint,
    argv0,
    expected_usage,
    expected_fragment,
):
    monkeypatch.setattr(sys, "argv", [argv0, "--help"])

    with pytest.raises(SystemExit) as exc_info:
        entrypoint()

    assert exc_info.value.code == 0
    captured = capsys.readouterr()
    assert expected_usage in captured.out
    assert expected_fragment in captured.out


def test_initialize_local_workspace_creates_roots_and_env_file(tmp_path):
    env_path = tmp_path / "article-local.env"

    payload = local_app.initialize_local_workspace(
        state_root=str(tmp_path / "state"),
        runtime_root=str(tmp_path / "runtime"),
        write_env=str(env_path),
    )

    assert payload["status"] == "ok"
    assert payload["summary"]["headline"] == "本地运行壳初始化完成。"
    assert Path(payload["storage"]["db_path"]).exists()
    assert payload["runtime"]["directories"]["jobs"]["exists"] is True
    assert payload["runtime"]["directories"]["uploads"]["exists"] is True
    assert payload["runtime"]["directories"]["staging"]["exists"] is True
    assert payload["env"]["file_path"] == str(env_path.resolve())
    env_text = env_path.read_text(encoding="utf-8")
    assert "ARTICLE_API_STATE_ROOT" in env_text
    assert "ARTICLE_API_RUNTIME_ROOT" in env_text
    assert payload["doctor"]["status"] == "ok"
    assert len(payload["summary"]["next_steps"]) == 2


def test_initialize_local_workspace_reuses_identical_env_file(monkeypatch, tmp_path):
    env_path = tmp_path / "article-local.env"
    monkeypatch.setattr(local_env, "current_command_bin_dir", lambda: None)

    first_payload = local_app.initialize_local_workspace(
        state_root=str(tmp_path / "state"),
        runtime_root=str(tmp_path / "runtime"),
        write_env=str(env_path),
    )
    second_payload = local_app.initialize_local_workspace(
        state_root=str(tmp_path / "state"),
        runtime_root=str(tmp_path / "runtime"),
        write_env=str(env_path),
    )

    assert first_payload["env"]["file_path"] == str(env_path.resolve())
    assert second_payload["env"]["file_path"] == str(env_path.resolve())


def test_initialize_local_workspace_requires_overwrite_for_existing_env_file(tmp_path):
    env_path = tmp_path / "article-local.env"
    env_path.write_text("existing", encoding="utf-8")

    with pytest.raises(RuntimeError, match="Env file already exists"):
        local_app.initialize_local_workspace(
            state_root=str(tmp_path / "state"),
            runtime_root=str(tmp_path / "runtime"),
            write_env=str(env_path),
        )


def test_initialize_local_workspace_quotes_env_file_and_adds_command_bin(monkeypatch, tmp_path):
    env_path = tmp_path / "dir with spaces" / "article local.env"
    command_bin = tmp_path / "venv path" / "bin"
    monkeypatch.setattr(local_env, "current_command_bin_dir", lambda: command_bin)

    payload = local_app.initialize_local_workspace(
        state_root=str(tmp_path / "state root"),
        runtime_root=str(tmp_path / "runtime root"),
        write_env=str(env_path),
    )

    env_text = env_path.read_text(encoding="utf-8")
    assert f"export ARTICLE_API_STATE_ROOT={shlex.quote(str((tmp_path / 'state root').resolve()))}" in env_text
    assert f"export ARTICLE_API_RUNTIME_ROOT={shlex.quote(str((tmp_path / 'runtime root').resolve()))}" in env_text
    assert f"export PATH={shlex.quote(str(command_bin))}:$PATH" in env_text
    assert payload["env"]["file_path"] == str(env_path.resolve())


def test_main_init_writes_env_file_and_reports_next_steps(capsys, tmp_path):
    env_path = tmp_path / "article-local.env"

    exit_code = local_app.main(
        [
            "init",
            "--state-root",
            str(tmp_path / "state"),
            "--runtime-root",
            str(tmp_path / "runtime"),
            "--write-env",
            str(env_path),
        ]
    )

    assert exit_code == 0
    captured = capsys.readouterr()
    payload = json.loads(captured.out)
    assert payload["summary"]["headline"] == "本地运行壳初始化完成。"
    assert payload["env"]["file_path"] == str(env_path.resolve())
    assert payload["summary"]["next_steps"][0].startswith("article-local doctor")


def test_build_profile_catalog_lists_lnu_only_public_profile():
    payload = local_app.build_profile_catalog()

    profile_ids = [item["id"] for item in payload["profiles"]]
    assert payload["summary"]["default_profile_id"] == "lnu-checker-2026"
    assert payload["summary"]["profile_count"] == 1
    assert payload["summary"]["support_scenario_count"] == 1
    assert profile_ids == ["lnu-checker-2026"]
    assert "cn-common" not in str(payload).lower()
    assert "课程作业" not in str(payload)
    assert "综述" not in str(payload)
    profile = payload["profiles"][0]
    assert profile["support_level_label"] == "一等支持"
    assert any(item["label"] == "学校学位论文" for item in profile["support_scenarios"])


def test_run_preflight_detects_wild_doc_signals(tmp_path):
    source_path = _make_style_conflict_lnu_doc(tmp_path / "article_local_preflight.docx")

    payload = local_app.run_preflight(str(source_path), profile="lnu")

    assert payload["operation"] == "preflight"
    assert payload["status"] == "ok"
    assert payload["preflight_status"] == "blocked"
    assert payload["summary"]["wild_doc_detected"] is True
    assert any(item["id"] == "style_text_conflicts" for item in payload["wild_doc"]["signals"])


def test_run_normalize_returns_payload(monkeypatch):
    monkeypatch.setattr(
        local_app.app_module,
        "build_normalize_payload",
        lambda **kwargs: {
            "status": "ok",
            "operation": "normalize",
            "document": {"name": "demo.docx"},
            "output": {"name": "demo_normalized.docx"},
            "changed": True,
        },
    )

    payload = local_app.run_normalize(
        "/tmp/demo.docx",
        output_path="/tmp/demo_normalized.docx",
        profile="lnu",
    )

    assert payload["operation"] == "normalize"
    assert payload["changed"] is True
    assert payload["output"]["name"] == "demo_normalized.docx"


def test_run_render_verify_returns_payload(monkeypatch):
    monkeypatch.setattr(
        local_app.app_module,
        "build_render_verify_payload",
        lambda **kwargs: {
            "status": "ok",
            "operation": "render-verify",
            "document": {"name": "demo.docx"},
            "page_count": 1,
            "selected_scopes": kwargs.get("scopes"),
            "rendered_pdf": kwargs.get("rendered_pdf"),
            "page_images_dir": kwargs.get("page_images_dir"),
            "workflow_mode": kwargs.get("workflow_mode"),
        },
    )

    payload = local_app.run_render_verify(
        "/tmp/demo.docx",
        profile="lnu",
        scopes=["toc"],
        renderer="word-pdf",
        rendered_pdf="/tmp/export.pdf",
        page_images_dir=None,
        workflow_mode="default_user",
    )

    assert payload["operation"] == "render-verify"
    assert payload["page_count"] == 1
    assert payload["selected_scopes"] == ["toc"]
    assert payload["rendered_pdf"] == "/tmp/export.pdf"
    assert payload["workflow_mode"] == "default_user"


def test_main_profiles_outputs_json(capsys):
    exit_code = local_app.main(["profiles"])

    assert exit_code == 0
    payload = json.loads(capsys.readouterr().out)
    profile_ids = [item["id"] for item in payload["profiles"]]
    assert profile_ids == ["lnu-checker-2026"]
    assert [item["id"] for item in payload["summary"]["support_scenarios"]] == ["school_degree_thesis"]
    assert "cn-common" not in str(payload).lower()


def test_main_render_workflow_modes_outputs_json(capsys):
    exit_code = local_app.main(["render-workflow-modes"])

    assert exit_code == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["recommended_mode"] == "default_user"
    assert [item["id"] for item in payload["modes"]] == ["default_user", "agent_candidate"]


def test_main_preflight_outputs_json(capsys, tmp_path):
    source_path = _make_style_conflict_lnu_doc(tmp_path / "article_local_main_preflight.docx")

    exit_code = local_app.main(["preflight", str(source_path), "--profile", "lnu"])

    assert exit_code == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["operation"] == "preflight"
    assert payload["document"]["name"] == "article_local_main_preflight.docx"
    assert payload["preflight_status"] == "blocked"
    assert payload["summary"]["wild_doc_detected"] is True


def test_main_normalize_outputs_json(monkeypatch, capsys):
    monkeypatch.setattr(
        local_app.app_module,
        "build_normalize_payload",
        lambda **kwargs: {
            "status": "ok",
            "operation": "normalize",
            "document": {"name": "demo.docx"},
            "output": {"name": "demo_normalized.docx"},
            "changed": True,
        },
    )

    exit_code = local_app.main(["normalize", "/tmp/demo.docx", "--profile", "lnu"])

    assert exit_code == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["operation"] == "normalize"
    assert payload["changed"] is True
    assert payload["output"]["name"] == "demo_normalized.docx"


def test_main_render_verify_outputs_json(monkeypatch, capsys):
    monkeypatch.setattr(
        local_app.app_module,
        "build_render_verify_payload",
        lambda **kwargs: {
            "status": "ok",
            "operation": "render-verify",
            "document": {"name": "demo.docx"},
            "page_count": 2,
            "selected_scopes": kwargs.get("scopes"),
            "rendered_pdf": kwargs.get("rendered_pdf"),
            "page_images_dir": kwargs.get("page_images_dir"),
            "workflow_mode": kwargs.get("workflow_mode"),
        },
    )

    exit_code = local_app.main(
        [
            "render-verify",
            "/tmp/demo.docx",
            "--profile",
            "lnu",
            "--scope",
            "toc",
            "--rendered-pdf",
            "/tmp/export.pdf",
            "--workflow-mode",
            "default_user",
        ]
    )

    assert exit_code == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["operation"] == "render-verify"
    assert payload["rendered_pdf"] == "/tmp/export.pdf"
    assert payload["workflow_mode"] == "default_user"
    assert payload["page_count"] == 2
    assert payload["selected_scopes"] == ["toc"]


def test_build_doctor_report_quotes_workflow_paths(monkeypatch, tmp_path):
    state_root = tmp_path / "state root"
    runtime_root = tmp_path / "runtime root"
    monkeypatch.delenv(storage.STATE_ROOT_ENV_VAR, raising=False)
    monkeypatch.delenv("ARTICLE_API_RUNTIME_ROOT", raising=False)

    report = local_app.build_doctor_report(
        state_root=str(state_root),
        runtime_root=str(runtime_root),
    )

    assert shlex.quote(str(state_root.resolve())) in report["workflow"]["doctor"]
    assert shlex.quote(str(runtime_root.resolve())) in report["workflow"]["serve"]


def test_create_backup_and_restore_roundtrip(monkeypatch, tmp_path):
    state_root = tmp_path / "state"
    runtime_root = tmp_path / "runtime"
    restore_state_root = tmp_path / "restored-state"
    restore_runtime_root = tmp_path / "restored-runtime"
    monkeypatch.setenv(storage.STATE_ROOT_ENV_VAR, str(state_root))
    monkeypatch.setenv("ARTICLE_API_RUNTIME_ROOT", str(runtime_root))

    storage.init_storage()
    state_note = state_root / "notes.json"
    state_note.write_text('{"ok": true}', encoding="utf-8")
    runtime_file = runtime_root / "jobs" / "job-1" / "outputs" / "result.docx"
    runtime_file.parent.mkdir(parents=True, exist_ok=True)
    runtime_file.write_bytes(b"docx-bytes")
    upload_file = runtime_root / "uploads" / "upload.docx"
    upload_file.parent.mkdir(parents=True, exist_ok=True)
    upload_file.write_bytes(b"upload-bytes")

    archive = tmp_path / "article-backup.zip"
    backup_payload = local_app.create_backup_archive(str(archive))

    assert backup_payload["state_file_count"] >= 2
    assert backup_payload["runtime_file_count"] == 2
    assert archive.exists()
    assert backup_payload["size_bytes"] > 0

    restore_payload = local_app.restore_backup_archive(
        str(archive),
        state_root=str(restore_state_root),
        runtime_root=str(restore_runtime_root),
    )

    assert restore_payload["restored_state_files"] >= 2
    assert restore_payload["restored_runtime_files"] == 2
    assert (restore_state_root / storage.DB_FILENAME).exists()
    assert json.loads((restore_state_root / "notes.json").read_text(encoding="utf-8"))["ok"] is True
    assert (restore_runtime_root / "jobs" / "job-1" / "outputs" / "result.docx").read_bytes() == b"docx-bytes"
    assert (restore_runtime_root / "uploads" / "upload.docx").read_bytes() == b"upload-bytes"
    assert restore_payload["manifest"]["schema_version"] == storage.SCHEMA_VERSION


def test_restore_requires_empty_target_without_force(monkeypatch, tmp_path):
    state_root = tmp_path / "state"
    runtime_root = tmp_path / "runtime"
    monkeypatch.setenv(storage.STATE_ROOT_ENV_VAR, str(state_root))
    monkeypatch.setenv("ARTICLE_API_RUNTIME_ROOT", str(runtime_root))
    storage.init_storage()

    archive = tmp_path / "article-backup.zip"
    local_app.create_backup_archive(str(archive))

    occupied_state = tmp_path / "occupied-state"
    occupied_state.mkdir(parents=True, exist_ok=True)
    (occupied_state / "existing.txt").write_text("occupied", encoding="utf-8")

    try:
        local_app.restore_backup_archive(str(archive), state_root=str(occupied_state), runtime_root=str(tmp_path / "restored-runtime"))
    except RuntimeError as exc:
        assert "Restore target is not empty" in str(exc)
    else:  # pragma: no cover
        raise AssertionError("restore should require --force on non-empty targets")


def test_restore_rejects_backup_with_newer_schema_version(monkeypatch, tmp_path):
    state_root = tmp_path / "state"
    runtime_root = tmp_path / "runtime"
    monkeypatch.setenv(storage.STATE_ROOT_ENV_VAR, str(state_root))
    monkeypatch.setenv("ARTICLE_API_RUNTIME_ROOT", str(runtime_root))
    storage.init_storage()

    archive = tmp_path / "article-backup.zip"
    local_app.create_backup_archive(str(archive))
    rewritten_archive = tmp_path / "article-backup-rewritten.zip"
    with local_app.zipfile.ZipFile(archive) as src, local_app.zipfile.ZipFile(
        rewritten_archive, "w", compression=local_app.zipfile.ZIP_DEFLATED
    ) as dst:
        for member in src.namelist():
            payload = src.read(member)
            if member == "manifest.json":
                manifest = json.loads(payload.decode("utf-8"))
                manifest["schema_version"] = storage.SCHEMA_VERSION + 1
                payload = json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True).encode("utf-8")
            dst.writestr(member, payload)

    with pytest.raises(RuntimeError, match="newer than supported version"):
        local_app.restore_backup_archive(
            str(rewritten_archive),
            state_root=str(tmp_path / "restored-state"),
            runtime_root=str(tmp_path / "restored-runtime"),
        )


def test_restore_rejects_archive_member_that_escapes_target_root(tmp_path):
    archive = tmp_path / "article-backup-malicious.zip"
    with local_app.zipfile.ZipFile(archive, "w", compression=local_app.zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(
            "manifest.json",
            json.dumps(
                {
                    "created_at": "2026-04-22T00:00:00Z",
                    "service": "article-api",
                    "version": "0.1.0",
                    "schema_version": storage.SCHEMA_VERSION,
                    "supported_schema_versions": [1, storage.SCHEMA_VERSION],
                },
                ensure_ascii=False,
                indent=2,
                sort_keys=True,
            ).encode("utf-8"),
        )
        zf.writestr("state_root/../../escaped.txt", b"nope")

    escaped_path = tmp_path / "escaped.txt"
    with pytest.raises(RuntimeError, match="escapes restore root"):
        local_app.restore_backup_archive(
            str(archive),
            state_root=str(tmp_path / "restore-state"),
            runtime_root=str(tmp_path / "restore-runtime"),
        )

    assert escaped_path.exists() is False


def test_main_serve_runs_uvicorn_with_root_overrides(monkeypatch, tmp_path):
    calls: list[dict] = []
    observed_roots: dict[str, str] = {}

    class FakeUvicorn:
        @staticmethod
        def run(app_target: str, **kwargs):
            calls.append({"app_target": app_target, **kwargs})
            observed_roots["state_root"] = str(storage.resolve_state_root())
            observed_roots["runtime_root"] = str(resolve_runtime_root())

    monkeypatch.delenv(storage.STATE_ROOT_ENV_VAR, raising=False)
    monkeypatch.delenv("ARTICLE_API_RUNTIME_ROOT", raising=False)
    monkeypatch.setattr(local_app, "_load_uvicorn", lambda: FakeUvicorn)

    exit_code = local_app.main(
        [
            "serve",
            "--host",
            "127.0.0.1",
            "--port",
            "9123",
            "--state-root",
            str(tmp_path / "serve-state"),
            "--runtime-root",
            str(tmp_path / "serve-runtime"),
        ]
    )

    assert exit_code == 0
    assert calls[0]["app_target"] == "article_api.app:create_app"
    assert calls[0]["factory"] is True
    assert calls[0]["host"] == "127.0.0.1"
    assert calls[0]["port"] == 9123
    assert Path(observed_roots["state_root"]).name == "serve-state"
    assert Path(observed_roots["runtime_root"]).name == "serve-runtime"
    assert storage.STATE_ROOT_ENV_VAR not in os.environ
    assert "ARTICLE_API_RUNTIME_ROOT" not in os.environ


def test_run_storage_maintenance_reports_actions(monkeypatch, tmp_path):
    state_root = tmp_path / "state"
    runtime_root = tmp_path / "runtime"
    monkeypatch.delenv(storage.STATE_ROOT_ENV_VAR, raising=False)
    monkeypatch.delenv("ARTICLE_API_RUNTIME_ROOT", raising=False)

    payload = local_app.run_storage_maintenance(
        state_root=str(state_root),
        runtime_root=str(runtime_root),
        run_vacuum=False,
        run_analyze=True,
    )

    assert payload["roots"]["state_root"] == str(state_root.resolve())
    assert payload["roots"]["runtime_root"] == str(runtime_root.resolve())
    assert payload["summary"]["headline"] == "本地状态库维护完成。"
    assert payload["maintenance"]["actions"] == ["analyze"]
    assert payload["maintenance"]["storage"]["schema_version"] == storage.SCHEMA_VERSION
    assert payload["storage"]["schema_version"] == storage.SCHEMA_VERSION


def test_main_maintain_runs_storage_maintenance_with_requested_flags(monkeypatch, tmp_path):
    calls: list[dict] = []

    def fake_run_storage_maintenance(**kwargs):
        calls.append(kwargs)
        return {
            "service": "article-api",
            "stage": "prototype",
            "version": "0.1.0",
            "api_version": "v0",
            "observed_at": "2026-04-22T00:00:00Z",
            "roots": {
                "state_root": str((tmp_path / "state").resolve()),
                "runtime_root": str((tmp_path / "runtime").resolve()),
                "shared_root": False,
            },
            "maintenance": {
                "actions": ["vacuum"],
                "storage": {"schema_version": storage.SCHEMA_VERSION},
            },
        }

    monkeypatch.setattr(local_app, "run_storage_maintenance", fake_run_storage_maintenance)

    exit_code = local_app.main(
        [
            "maintain",
            "--state-root",
            str(tmp_path / "state"),
            "--runtime-root",
            str(tmp_path / "runtime"),
            "--vacuum",
            "--no-analyze",
        ]
    )

    assert exit_code == 0
    assert calls == [
        {
            "state_root": str(tmp_path / "state"),
            "runtime_root": str(tmp_path / "runtime"),
            "run_vacuum": True,
            "run_analyze": False,
        }
    ]
