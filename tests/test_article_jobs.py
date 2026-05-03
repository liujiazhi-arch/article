from concurrent.futures import Future
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from docx import Document

import article_api.jobs as jobs_module
import article_api.output_naming as output_naming
import article_api.storage as storage_module
from article_api.inspection import build_job_inspection, get_job_inspection, get_job_runtime_snapshot
from article_api.jobs import (
    cleanup_job,
    clear_jobs,
    create_job,
    get_job,
    get_job_result,
    inspect_job,
    job_count,
    list_jobs,
    retry_job,
    sweep_job_retention,
    wait_for_job,
)
from article_api.uploads import store_uploaded_docx

from .conftest import RULE_MUTATORS, make_compliant_doc


def _build_mutated_doc(tmp_docx, *, filename: str, rule_ids: tuple[str, ...]) -> Path:
    path = tmp_docx(make_compliant_doc, filename=filename)
    doc = Document(path)
    for rule_id in rule_ids:
        RULE_MUTATORS[rule_id](doc)
    doc.save(path)
    return path


def _make_high_risk_lnu_doc(source_path: Path) -> Path:
    doc = Document()
    heading = doc.add_paragraph("1 绪论")
    heading.style = doc.styles["Heading 1"]
    doc.add_paragraph("这是正文示例。")
    table = doc.add_table(rows=5, cols=1)
    values = [
        "1,4-Diaminobutane",
        "2-Amino-2-Deoxy-D-Glucopyranose",
        "2.80E+08",
        "3.20E+01",
        "2.79E+08",
    ]
    for row, value in zip(table.rows, values):
        row.cells[0].text = value
    doc.save(source_path)
    return source_path


def _make_style_conflict_lnu_doc(source_path: Path) -> Path:
    doc = Document()
    heading = doc.add_paragraph("1 绪论")
    heading.style = doc.styles["Heading 1"]
    conflict = doc.add_paragraph("3.6 分子对接验证结果")
    conflict.style = doc.styles["Heading 1"]
    doc.add_paragraph("这是正文示例。")
    doc.save(source_path)
    return source_path


def _create_and_wait_job(operation: str, payload: dict[str, object]) -> dict[str, object]:
    job = create_job(operation, payload)
    wait_for_job(job["job_id"])
    return job


def _future_timestamp(timestamp: str, *, seconds: float) -> str:
    return (datetime.fromisoformat(timestamp.replace("Z", "+00:00")) + timedelta(seconds=seconds)).astimezone(
        timezone.utc
    ).isoformat().replace("+00:00", "Z")


def test_create_verify_job_completes_and_returns_result(tmp_docx):
    clear_jobs()
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_verify.docx",
        rule_ids=("H02",),
    )

    job = _create_and_wait_job(
        "verify",
        {
            "file_path": str(source_path),
            "profile_path": None,
            "scopes": ["headings"],
            "strict_profile": False,
        },
    )
    status = get_job(job["job_id"])
    result = get_job_result(job["job_id"])

    assert job["status"] == "queued"
    assert status["status"] == "succeeded"
    assert status["summary"]["business_status"] == "needs_fix"
    assert status["summary"]["readiness"] == "needs-fix"
    assert status["summary"]["selected_scopes"] == ["headings"]
    assert status["runtime"]["cleanup_policy"] == "manual"
    assert status["runtime"]["stage_input_enabled"] is False
    assert status["runtime"]["workspace_present"] is False
    assert status["runtime"]["worker_model"] == "single"
    assert status["runtime"]["lease_state"] == "released"
    assert status["runtime"]["last_heartbeat_at"] is not None
    assert status["runtime"]["heartbeat_count"] >= 2
    assert [item["event"] for item in status["runtime"]["events"][:2]] == ["job_queued", "worker_started"]
    assert status["runtime"]["events"][-1]["event"] == "job_succeeded"
    assert result["result"]["document"]["name"] == "article_job_verify.docx"
    assert result["result"]["overall_status"] == "needs_fix"
    assert result["result"]["readiness"] == "needs-fix"


def test_apply_job_keeps_task_status_separate_from_business_status(tmp_path):
    clear_jobs()
    source_path = _make_high_risk_lnu_doc(Path(tmp_path) / "article_job_apply_manual_review.docx")

    job = _create_and_wait_job(
        "apply",
        {
            "file_path": str(source_path),
            "profile_path": "lnu",
            "scopes": ["headings"],
            "renumber_headings": True,
        },
    )
    result = get_job_result(job["job_id"])

    assert job["status"] == "queued"
    assert result["status"] == "succeeded"
    assert result["summary"]["business_status"] == "manual_review"
    assert result["summary"]["readiness"] == "manual-review-required"
    assert result["result"]["readiness"] == "manual-review-required"
    assert result["result"]["verification"]["overall_status"] == "manual_review"


def test_normalize_job_completes_and_writes_output(tmp_path):
    clear_jobs()
    source_path = _make_style_conflict_lnu_doc(Path(tmp_path) / "article_job_normalize.docx")
    output_path = Path(tmp_path) / "article_job_normalized.docx"

    job = _create_and_wait_job(
        "normalize",
        {
            "file_path": str(source_path),
            "output_path": str(output_path),
            "profile_path": "lnu",
        },
    )
    status = get_job(job["job_id"])
    result = get_job_result(job["job_id"])

    assert status["status"] == "succeeded"
    assert status["summary"]["business_status"] == "warning"
    assert status["summary"]["output_path"] == str(output_path)
    assert status["summary"]["changed"] is True
    assert output_path.exists()
    assert result["result"]["document"]["name"] == "article_job_normalize.docx"
    assert result["result"]["output"]["name"] == "article_job_normalized.docx"
    assert result["result"]["after"]["preflight_status"] == "warning"
    assert any(item["role"] == "output" and item["written"] is True for item in result["artifacts"])


def test_batch_job_operation_is_not_supported():
    clear_jobs()

    with pytest.raises(ValueError, match="Unsupported job operation: batch"):
        create_job("batch", {"input_path": "/tmp/article-batch", "operation": "audit"})


def test_apply_job_result_is_frozen_after_output_changes(tmp_docx):
    clear_jobs()
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_apply_freeze.docx",
        rule_ids=("H02",),
    )

    job = _create_and_wait_job(
        "apply",
        {
            "file_path": str(source_path),
            "profile_path": None,
            "scopes": ["headings"],
        },
    )
    first_result = get_job_result(job["job_id"])
    output_path = Path(first_result["result"]["output"]["path"])

    doc = Document(output_path)
    doc.paragraphs[0].runs[0].text = "附加修改，模拟任务完成后文件再次被改动"
    doc.save(output_path)

    second_result = get_job_result(job["job_id"])

    assert first_result == second_result
    assert first_result["artifacts"][0]["written"] is True
    assert first_result["artifacts"][0]["exists_at_completion"] is True


def test_apply_job_captures_blocked_guard_as_failed_job(tmp_path):
    clear_jobs()
    source_path = _make_style_conflict_lnu_doc(Path(tmp_path) / "article_job_apply_fail.docx")

    job = _create_and_wait_job(
        "apply",
        {
            "file_path": str(source_path),
            "profile_path": "lnu",
            "scopes": ["headings"],
            "renumber_headings": True,
        },
    )
    status = get_job(job["job_id"])
    result = get_job_result(job["job_id"])

    assert job["status"] == "queued"
    assert status["status"] == "failed"
    assert status["summary"]["error_code"] == "apply_guard_blocked"
    assert status["summary"]["guard_blocked"] is True
    assert status["summary"]["guard_warning_count"] >= 1
    assert result["error"]["type"] == "ApplyGuardBlockedError"
    assert result["error"]["code"] == "apply_guard_blocked"
    assert result["error"]["http_status"] == 409
    assert result["error"]["guard"]["blocked"] is True
    assert result["error"]["guard"]["checked"] is True
    assert result["error"]["guard"]["warnings"]
    assert result["resolved_request"]["scopes"] == ["headings"]
    assert result["resolved_request"]["output_path"].endswith("article_job_apply_fail_headings.docx")
    assert result["artifacts"][0]["written"] is False
    assert result["artifacts"][0]["exists_at_completion"] is False
    assert "Apply blocked by structural risk" in result["error"]["message"]
    assert result["result"] is None


def test_create_job_rejects_invalid_file_path_without_storing_job(tmp_path):
    clear_jobs()
    missing_path = tmp_path / "missing.docx"

    with pytest.raises(ValueError, match="文件不存在"):
        create_job("verify", {"file_path": str(missing_path)})

    assert job_count() == 0


def test_get_job_raises_for_missing_job():
    clear_jobs()

    with pytest.raises(LookupError, match="Job not found"):
        get_job("missing-job")


def test_inspect_job_returns_lightweight_backend_view(tmp_docx):
    clear_jobs()
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_inspect.docx",
        rule_ids=("H02",),
    )

    job = _create_and_wait_job(
        "verify",
        {
            "file_path": str(source_path),
            "scopes": ["headings"],
        },
    )
    inspection = inspect_job(job["job_id"])

    assert inspection["job_id"] == job["job_id"]
    assert inspection["operation"] == "verify"
    assert inspection["status"] == "succeeded"
    assert inspection["summary"]["business_status"] == "needs_fix"
    assert inspection["runtime"]["cleanup_policy"] == "manual"
    assert inspection["result_available"] is True
    assert "request" not in inspection
    assert "result" not in inspection


def test_list_jobs_returns_latest_first(tmp_docx):
    clear_jobs()
    first_source = _build_mutated_doc(
        tmp_docx,
        filename="article_job_list_first.docx",
        rule_ids=("H02",),
    )
    second_source = _build_mutated_doc(
        tmp_docx,
        filename="article_job_list_second.docx",
        rule_ids=("H02",),
    )

    first = _create_and_wait_job("verify", {"file_path": str(first_source), "scopes": ["headings"]})
    second = _create_and_wait_job("verify", {"file_path": str(second_source), "scopes": ["headings"]})
    jobs = list_jobs()

    assert [item["job_id"] for item in jobs[:2]] == [second["job_id"], first["job_id"]]


def test_build_job_inspection_reduces_job_payload_to_backend_view(tmp_docx):
    clear_jobs()
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_build_inspection.docx",
        rule_ids=("H02",),
    )
    job = _create_and_wait_job("verify", {"file_path": str(source_path), "scopes": ["headings"]})
    job_payload = get_job(job["job_id"])

    inspection = build_job_inspection(job_payload)

    assert inspection["job_id"] == job["job_id"]
    assert inspection["summary"]["business_status"] == "needs_fix"
    assert inspection["request_paths"]["source_file_path"] is None
    assert "request" not in inspection
    assert "resolved_request" not in inspection


def test_get_job_inspection_preserves_staging_paths_and_stays_stable(tmp_docx, tmp_path):
    clear_jobs()
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_get_inspection.docx",
        rule_ids=("H02",),
    )
    job = _create_and_wait_job(
        "verify",
        {
            "file_path": str(source_path),
            "scopes": ["headings"],
            "stage_input": True,
            "runtime_root": str(tmp_path / "runtime"),
        },
    )
    first = get_job_inspection(job["job_id"])
    Path(first["request_paths"]["staged_input_path"]).unlink()
    second = get_job_inspection(job["job_id"])

    assert first == second
    assert first["request_paths"]["source_file_path"] == str(Path(source_path).resolve())
    assert first["request_paths"]["workspace_root"] == first["runtime"]["workspace_root"]


def test_get_job_runtime_snapshot_returns_runtime_focused_view(tmp_docx, tmp_path):
    clear_jobs()
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_runtime_snapshot.docx",
        rule_ids=("H02",),
    )
    job = _create_and_wait_job(
        "apply",
        {
            "file_path": str(source_path),
            "scopes": ["headings"],
            "dry_run": True,
            "stage_input": True,
            "runtime_root": str(tmp_path / "runtime"),
        },
    )
    snapshot = get_job_runtime_snapshot(job["job_id"])

    assert snapshot["job_id"] == job["job_id"]
    assert snapshot["summary"]["result_mode"] == "preview"
    assert snapshot["runtime"]["workspace_present"] is True
    assert len(snapshot["artifacts"]) == 2
    assert "created_at" not in snapshot


def test_runtime_snapshot_reports_recovered_failed_jobs(monkeypatch, tmp_path):
    monkeypatch.setenv(storage_module.STATE_ROOT_ENV_VAR, str(tmp_path / "state"))
    clear_jobs()
    payload = {
        "job_id": "snapshot-stale-job",
        "operation": "verify",
        "status": "running",
        "mode": "background",
        "created_at": "2026-04-19T00:00:00Z",
        "updated_at": "2026-04-19T00:01:00Z",
        "started_at": "2026-04-19T00:00:30Z",
        "finished_at": None,
        "request": {"file_path": "/tmp/stale.docx", "scopes": ["headings"]},
        "resolved_request": {"file_path": "/tmp/stale.docx", "scopes": ["headings"], "strict_profile": False},
        "workspace": None,
        "runtime": {
            "cleanup_policy": "manual",
            "runtime_root": None,
            "source_upload_id": None,
            "stage_input_enabled": False,
            "workspace_present": False,
            "workspace_root": None,
            "workspace_inputs": None,
            "workspace_outputs": None,
            "source_file_path": None,
            "staged_input_path": None,
            "output_path": None,
            "attempt_count": 0,
            "attempts": [],
            "events": [],
            "worker_model": "single",
            "lease_state": "active",
            "last_heartbeat_at": "2026-04-19T00:01:00Z",
            "heartbeat_count": 1,
        },
        "result_available": False,
        "summary": None,
        "artifacts": [],
        "error": None,
        "result": None,
    }
    storage_module.upsert_job(payload)

    snapshot = jobs_module.runtime_snapshot()

    assert snapshot["recovery"]["recovered_failed_count"] == 1
    assert snapshot["recovery"]["recovered_failed_job_ids"] == ["snapshot-stale-job"]
    assert get_job("snapshot-stale-job")["status"] == "failed"


def test_runtime_snapshot_flags_running_job_with_stale_heartbeat(monkeypatch, tmp_path):
    monkeypatch.setenv(storage_module.STATE_ROOT_ENV_VAR, str(tmp_path / "state"))
    monkeypatch.setenv(jobs_module.HEARTBEAT_STALE_SECONDS_ENV, "1")
    clear_jobs()
    payload = {
        "job_id": "stale-heartbeat-job",
        "operation": "verify",
        "status": "running",
        "mode": "background",
        "created_at": "2026-04-19T00:00:00Z",
        "updated_at": "2026-04-19T00:00:00Z",
        "started_at": "2026-04-19T00:00:00Z",
        "finished_at": None,
        "request": {"file_path": "/tmp/stale-heartbeat.docx", "scopes": ["headings"]},
        "resolved_request": {"file_path": "/tmp/stale-heartbeat.docx", "scopes": ["headings"], "strict_profile": False},
        "workspace": None,
        "runtime": {
            "cleanup_policy": "manual",
            "runtime_root": None,
            "source_upload_id": None,
            "stage_input_enabled": False,
            "workspace_present": False,
            "workspace_root": None,
            "workspace_inputs": None,
            "workspace_outputs": None,
            "source_file_path": None,
            "staged_input_path": None,
            "output_path": None,
            "attempt_count": 0,
            "attempts": [],
            "events": [],
            "worker_model": "single",
            "lease_state": "active",
            "last_heartbeat_at": "2026-04-19T00:00:00Z",
            "heartbeat_count": 1,
        },
        "result_available": False,
        "summary": None,
        "artifacts": [],
        "error": None,
        "result": None,
    }
    storage_module.upsert_job(payload)
    future = Future()
    with jobs_module._JOB_LOCK:
        jobs_module._ACTIVE_FUTURES[payload["job_id"]] = future
    try:
        snapshot = jobs_module.runtime_snapshot()
    finally:
        future.set_result(None)
        with jobs_module._JOB_LOCK:
            jobs_module._ACTIVE_FUTURES.pop(payload["job_id"], None)

    assert snapshot["health"]["stale_heartbeat_count"] == 1
    assert snapshot["health"]["stale_heartbeat_job_ids"] == ["stale-heartbeat-job"]
    assert snapshot["recovery"]["heartbeat_stale_after_seconds"] == 1.0


def test_runtime_snapshot_reports_pending_recovery_jobs_within_grace(monkeypatch, tmp_path):
    monkeypatch.setenv(storage_module.STATE_ROOT_ENV_VAR, str(tmp_path / "state"))
    monkeypatch.setenv(jobs_module.RECOVERY_GRACE_SECONDS_ENV, "3600")
    clear_jobs()
    payload = {
        "job_id": "pending-recovery-job",
        "operation": "verify",
        "status": "queued",
        "mode": "background",
        "created_at": "2099-01-01T00:00:00Z",
        "updated_at": "2099-01-01T00:00:00Z",
        "started_at": None,
        "finished_at": None,
        "request": {"file_path": "/tmp/pending.docx", "scopes": ["headings"]},
        "resolved_request": {"file_path": "/tmp/pending.docx", "scopes": ["headings"], "strict_profile": False},
        "workspace": None,
        "runtime": {
            "cleanup_policy": "manual",
            "runtime_root": None,
            "source_upload_id": None,
            "stage_input_enabled": False,
            "workspace_present": False,
            "workspace_root": None,
            "workspace_inputs": None,
            "workspace_outputs": None,
            "source_file_path": None,
            "staged_input_path": None,
            "output_path": None,
            "attempt_count": 0,
            "attempts": [],
            "events": [],
            "worker_model": "single",
            "lease_state": "pending",
            "last_heartbeat_at": None,
            "heartbeat_count": 0,
        },
        "result_available": False,
        "summary": None,
        "artifacts": [],
        "error": None,
        "result": None,
    }
    storage_module.upsert_job(payload)

    snapshot = jobs_module.runtime_snapshot()

    assert snapshot["recovery"]["grace_seconds"] == 3600.0
    assert snapshot["recovery"]["pending_recovery_count"] == 1
    assert snapshot["recovery"]["pending_recovery_job_ids"] == ["pending-recovery-job"]
    assert get_job("pending-recovery-job")["status"] == "queued"


def test_get_job_inspection_raises_for_missing_job():
    clear_jobs()

    with pytest.raises(LookupError, match="Job not found"):
        get_job_inspection("missing-job")


def test_multiple_jobs_keep_frozen_summaries_isolated(tmp_docx, tmp_path):
    clear_jobs()
    verify_source = _build_mutated_doc(
        tmp_docx,
        filename="article_job_verify_isolation.docx",
        rule_ids=("H02",),
    )
    apply_source = _build_mutated_doc(
        tmp_docx,
        filename="article_job_apply_isolation.docx",
        rule_ids=("H02",),
    )

    verify_job = _create_and_wait_job(
        "verify",
        {
            "file_path": str(verify_source),
            "scopes": ["headings"],
        },
    )
    apply_job = _create_and_wait_job(
        "apply",
        {
            "file_path": str(apply_source),
            "output_path": str(Path(tmp_path) / "article_job_apply_isolation_fixed.docx"),
            "scopes": ["headings"],
        },
    )

    verify_status = get_job(verify_job["job_id"])
    apply_status = get_job(apply_job["job_id"])

    assert verify_status["summary"]["document_name"] == "article_job_verify_isolation.docx"
    assert verify_status["summary"]["business_status"] == "needs_fix"
    assert "output_path" not in verify_status["summary"]
    assert apply_status["summary"]["document_name"] == "article_job_apply_isolation.docx"
    assert apply_status["summary"]["business_status"] == "verified"
    assert apply_status["summary"]["output_path"].endswith("article_job_apply_isolation_fixed.docx")


def test_apply_dry_run_job_has_preview_summary_and_no_verification(tmp_docx, tmp_path):
    clear_jobs()
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_apply_preview.docx",
        rule_ids=("H02",),
    )
    output_path = Path(tmp_path) / "article_job_apply_preview_fixed.docx"

    job = _create_and_wait_job(
        "apply",
        {
            "file_path": str(source_path),
            "output_path": str(output_path),
            "scopes": ["headings"],
            "dry_run": True,
        },
    )
    result = get_job_result(job["job_id"])

    assert job["status"] == "queued"
    assert result["summary"]["dry_run"] is True
    assert result["summary"]["result_mode"] == "preview"
    assert result["summary"]["business_status"] == "preview"
    assert result["result"]["mode"] == "preview"
    assert result["artifacts"][0]["result_mode"] == "preview"
    assert result["artifacts"][0]["written"] is False
    assert result["artifacts"][0]["exists_at_completion"] is False
    assert "verification" not in result["result"]


def test_apply_job_freezes_default_output_path_when_not_provided(tmp_docx):
    clear_jobs()
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_default_output.docx",
        rule_ids=("H02",),
    )

    job = _create_and_wait_job(
        "apply",
        {
            "file_path": str(source_path),
            "scopes": ["headings"],
        },
    )
    status = get_job(job["job_id"])

    assert status["resolved_request"]["output_path"].endswith("article_job_default_output_headings.docx")
    assert status["artifacts"][0]["path"].endswith("article_job_default_output_headings.docx")


def test_apply_job_with_display_name_uses_next_versioned_output(monkeypatch, tmp_docx, tmp_path):
    clear_jobs()
    monkeypatch.setattr(output_naming, "DEFAULT_OUTPUT_DIR", tmp_path / "versioned-outputs")
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="runtime_upload_name.docx",
        rule_ids=("H02",),
    )
    output_dir = output_naming.DEFAULT_OUTPUT_DIR
    output_dir.mkdir(parents=True, exist_ok=True)
    existing = output_dir / "article_job_versioned_格式修复_V01.docx"
    existing.write_bytes(b"existing")

    job = _create_and_wait_job(
        "apply",
        {
            "file_path": str(source_path),
            "source_display_name": "article_job_versioned.docx",
            "scopes": ["headings"],
        },
    )
    status = get_job(job["job_id"])

    assert status["resolved_request"]["output_path"].endswith("article_job_versioned_格式修复_V02.docx")


def test_staged_apply_job_with_display_name_keeps_versioned_basename_in_workspace(monkeypatch, tmp_docx, tmp_path):
    clear_jobs()
    monkeypatch.setattr(output_naming, "DEFAULT_OUTPUT_DIR", tmp_path / "versioned-outputs")
    runtime_root = tmp_path / "runtime"
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="runtime_upload_staged_name.docx",
        rule_ids=("H02",),
    )

    job = _create_and_wait_job(
        "apply",
        {
            "file_path": str(source_path),
            "source_display_name": "article_job_upload_display.docx",
            "scopes": ["headings"],
            "stage_input": True,
            "runtime_root": str(runtime_root),
        },
    )
    status = get_job(job["job_id"])

    assert status["resolved_request"]["output_path"].startswith(status["workspace"]["outputs"])
    assert status["resolved_request"]["output_path"].endswith("article_job_upload_display_格式修复_V01.docx")


def test_apply_job_strips_existing_version_suffix_from_display_name(monkeypatch, tmp_docx, tmp_path):
    clear_jobs()
    monkeypatch.setattr(output_naming, "DEFAULT_OUTPUT_DIR", tmp_path / "versioned-outputs")
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="runtime_upload_versioned.docx",
        rule_ids=("H02",),
    )

    job = _create_and_wait_job(
        "apply",
        {
            "file_path": str(source_path),
            "source_display_name": "article_job_versioned_input_格式修复_V01.docx",
            "scopes": ["headings"],
        },
    )
    status = get_job(job["job_id"])

    assert status["resolved_request"]["output_path"].endswith("article_job_versioned_input_格式修复_V01.docx")
    assert "V01_格式修复" not in status["resolved_request"]["output_path"]


def test_normalize_job_freezes_default_output_path_when_not_provided(tmp_path):
    clear_jobs()
    source_path = _make_style_conflict_lnu_doc(Path(tmp_path) / "article_job_default_normalize.docx")

    job = _create_and_wait_job(
        "normalize",
        {
            "file_path": str(source_path),
            "profile_path": "lnu",
        },
    )
    status = get_job(job["job_id"])

    assert status["resolved_request"]["output_path"].endswith("article_job_default_normalize_normalized.docx")
    assert status["artifacts"][0]["path"].endswith("article_job_default_normalize_normalized.docx")


def test_stage_input_keeps_original_request_and_freezes_staged_input_snapshot(tmp_docx, tmp_path):
    clear_jobs()
    runtime_root = tmp_path / "runtime"
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_stage_verify.docx",
        rule_ids=("H02",),
    )

    job = _create_and_wait_job(
        "verify",
        {
            "file_path": str(source_path),
            "scopes": ["headings"],
            "stage_input": True,
            "runtime_root": str(runtime_root),
        },
    )
    status = get_job(job["job_id"])
    result = get_job_result(job["job_id"])

    assert status["request"]["file_path"] == str(source_path)
    assert status["resolved_request"]["source_file_path"] == str(Path(source_path).resolve())
    assert status["resolved_request"]["file_path"] != str(source_path)
    assert status["resolved_request"]["file_path"].startswith(status["workspace"]["inputs"])
    assert Path(status["resolved_request"]["file_path"]).exists()
    assert status["runtime"]["cleanup_policy"] == "manual"
    assert status["runtime"]["workspace_present"] is True
    assert status["runtime"]["workspace_root"] == status["workspace"]["root"]
    assert status["runtime"]["staged_input_path"] == status["resolved_request"]["file_path"]
    assert status["runtime"]["workspace_exists_at_completion"] is True
    assert status["runtime"]["staged_input_exists_at_completion"] is True
    assert result["result"]["document"]["path"] == str(Path(source_path).resolve())
    assert result["result"]["document"]["name"] == "article_job_stage_verify.docx"
    assert result["artifacts"][0]["role"] == "input"
    assert result["artifacts"][0]["staged"] is True


def test_stage_input_default_output_path_still_uses_original_document_name(tmp_docx, tmp_path):
    clear_jobs()
    runtime_root = tmp_path / "runtime"
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_stage_apply.docx",
        rule_ids=("H02",),
    )

    job = _create_and_wait_job(
        "apply",
        {
            "file_path": str(source_path),
            "scopes": ["headings"],
            "stage_input": True,
            "runtime_root": str(runtime_root),
        },
    )
    status = get_job(job["job_id"])

    assert status["resolved_request"]["output_path"].endswith("article_job_stage_apply_headings.docx")
    assert status["resolved_request"]["output_path"].startswith(status["workspace"]["outputs"])
    assert "article_job_stage_apply.docx" in status["resolved_request"]["source_file_path"]
    assert status["runtime"]["output_path"] == status["resolved_request"]["output_path"]
    assert status["runtime"]["workspace_outputs"] == status["workspace"]["outputs"]
    assert status["runtime"]["output_exists_at_completion"] is True


def test_stage_input_result_remains_stable_after_staged_file_is_deleted(tmp_docx, tmp_path):
    clear_jobs()
    runtime_root = tmp_path / "runtime"
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_stage_freeze.docx",
        rule_ids=("H02",),
    )

    job = _create_and_wait_job(
        "verify",
        {
            "file_path": str(source_path),
            "scopes": ["headings"],
            "stage_input": True,
            "runtime_root": str(runtime_root),
        },
    )
    before = get_job_result(job["job_id"])
    staged_input_path = Path(before["resolved_request"]["file_path"])
    staged_input_path.unlink()
    after = get_job_result(job["job_id"])

    assert before == after
    assert before["artifacts"][0]["exists_at_completion"] is True
    assert before["runtime"]["staged_input_exists_at_completion"] is True


def test_stage_input_dry_run_keeps_input_and_output_artifacts_distinct(tmp_docx, tmp_path):
    clear_jobs()
    runtime_root = tmp_path / "runtime"
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_stage_preview.docx",
        rule_ids=("H02",),
    )

    job = _create_and_wait_job(
        "apply",
        {
            "file_path": str(source_path),
            "scopes": ["headings"],
            "dry_run": True,
            "stage_input": True,
            "runtime_root": str(runtime_root),
        },
    )
    result = get_job_result(job["job_id"])
    input_artifact = next(item for item in result["artifacts"] if item["role"] == "input")
    output_artifact = next(item for item in result["artifacts"] if item["role"] == "output")

    assert result["result"]["mode"] == "preview"
    assert input_artifact["staged"] is True
    assert output_artifact["result_mode"] == "preview"
    assert output_artifact["written"] is False
    assert result["runtime"]["output_exists_at_completion"] is False


def test_stage_input_invalid_path_does_not_store_job(tmp_path):
    clear_jobs()
    runtime_root = tmp_path / "runtime"
    missing_path = tmp_path / "missing.docx"

    with pytest.raises(ValueError, match="文件不存在"):
        create_job(
            "verify",
            {
                "file_path": str(missing_path),
                "stage_input": True,
                "runtime_root": str(runtime_root),
            },
        )

    assert job_count() == 0


def test_cleanup_job_removes_managed_runtime_files_and_preserves_frozen_result(tmp_docx, tmp_path):
    clear_jobs()
    runtime_root = tmp_path / "runtime"
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_cleanup_upload.docx",
        rule_ids=("H02",),
    )

    class FakeUpload:
        filename = "article_job_cleanup_upload.docx"

        def __init__(self, payload: bytes):
            from io import BytesIO

            self.file = BytesIO(payload)

    upload = store_uploaded_docx(FakeUpload(Path(source_path).read_bytes()), runtime_root=runtime_root)
    job = _create_and_wait_job(
        "apply",
        {
            "file_path": upload.stored_path,
            "upload_id": upload.upload_id,
            "scopes": ["headings"],
            "stage_input": True,
            "runtime_root": str(runtime_root),
        },
    )
    before = get_job_result(job["job_id"])
    workspace_root = Path(before["runtime"]["workspace_root"])
    upload_source_path = Path(before["runtime"]["source_file_path"])

    cleanup_status = cleanup_job(job["job_id"])
    after = get_job_result(job["job_id"])

    assert cleanup_status["cleanup"]["state"] == "cleaned"
    assert cleanup_status["cleanup"]["attempt_count"] == 1
    assert not workspace_root.exists()
    assert upload_source_path.exists()
    assert before["summary"] == after["summary"]
    assert before["result"] == after["result"]
    assert after["cleanup"]["state"] == "cleaned"
    assert after["runtime"] == before["runtime"]
    assert after["artifacts"] == before["artifacts"]


def test_cleanup_job_is_idempotent_for_completed_job(tmp_docx, tmp_path):
    clear_jobs()
    runtime_root = tmp_path / "runtime"
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_cleanup_idempotent.docx",
        rule_ids=("H02",),
    )
    job = _create_and_wait_job(
        "apply",
        {
            "file_path": str(source_path),
            "scopes": ["headings"],
            "stage_input": True,
            "runtime_root": str(runtime_root),
        },
    )

    first = cleanup_job(job["job_id"])
    second = cleanup_job(job["job_id"])

    assert first["cleanup"]["attempt_count"] == 1
    assert second["cleanup"]["attempt_count"] == 2
    assert second["cleanup"]["state"] == "noop"


def test_cleanup_job_rejects_job_before_completion(monkeypatch, tmp_docx):
    clear_jobs()
    source_path = tmp_docx(make_compliant_doc, filename="article_job_cleanup_running.docx")

    def slow_verify(*args, **kwargs):
        time.sleep(0.2)
        return jobs_module.verify_document(*args, **kwargs)

    monkeypatch.setitem(jobs_module._JOB_HANDLERS, "verify", slow_verify)
    job = create_job("verify", {"file_path": str(source_path), "scopes": ["headings"]})

    with pytest.raises(RuntimeError, match="only available after completion"):
        cleanup_job(job["job_id"])

    wait_for_job(job["job_id"])


def test_get_job_result_rejects_running_job(monkeypatch, tmp_docx):
    clear_jobs()
    source_path = tmp_docx(make_compliant_doc, filename="article_job_result_running.docx")

    def slow_verify(*args, **kwargs):
        time.sleep(0.2)
        return jobs_module.verify_document(*args, **kwargs)

    monkeypatch.setitem(jobs_module._JOB_HANDLERS, "verify", slow_verify)
    job = create_job("verify", {"file_path": str(source_path), "scopes": ["headings"]})

    with pytest.raises(RuntimeError, match="not ready yet"):
        get_job_result(job["job_id"])

    wait_for_job(job["job_id"])


def test_running_job_exposes_active_single_worker_lease(monkeypatch, tmp_docx):
    clear_jobs()
    source_path = tmp_docx(make_compliant_doc, filename="article_job_runtime_lease_active.docx")

    def slow_attempt(_operation: str, resolved_request: dict[str, object]) -> dict[str, object]:
        time.sleep(0.2)
        return {
            "ok": True,
            "result": jobs_module.verify_document(
                file_path=str(resolved_request["file_path"]),
                profile_path=resolved_request.get("profile_path"),
                scopes=resolved_request.get("scopes"),
                strict_profile=bool(resolved_request.get("strict_profile")),
            ),
        }

    monkeypatch.setattr(jobs_module, "_execute_job_attempt", slow_attempt)
    job = create_job("verify", {"file_path": str(source_path), "scopes": ["headings"]})

    deadline = time.time() + 1.0
    running = None
    while time.time() < deadline:
        current = get_job(job["job_id"])
        if current["status"] == "running":
            running = current
            break
        time.sleep(0.01)

    assert running is not None
    assert running["runtime"]["worker_model"] == "single"
    assert running["runtime"]["lease_state"] == "active"
    assert running["runtime"]["last_heartbeat_at"] is not None
    assert running["runtime"]["heartbeat_count"] >= 2

    finished = wait_for_job(job["job_id"])

    assert finished["runtime"]["lease_state"] == "released"


def test_wait_for_job_timeout_does_not_mutate_persistent_status(monkeypatch, tmp_docx):
    clear_jobs()
    source_path = tmp_docx(make_compliant_doc, filename="article_job_wait_timeout.docx")

    def slow_attempt(_operation: str, resolved_request: dict[str, object]) -> dict[str, object]:
        time.sleep(0.2)
        return {
            "ok": True,
            "result": jobs_module.verify_document(
                file_path=str(resolved_request["file_path"]),
                profile_path=resolved_request.get("profile_path"),
                scopes=resolved_request.get("scopes"),
                strict_profile=bool(resolved_request.get("strict_profile")),
            ),
        }

    monkeypatch.setattr(jobs_module, "_execute_job_attempt", slow_attempt)
    job = create_job("verify", {"file_path": str(source_path), "scopes": ["headings"]})

    with pytest.raises(TimeoutError, match="Timed out waiting for job completion"):
        wait_for_job(job["job_id"], timeout=0.01, poll_interval=0.005)

    pending = get_job(job["job_id"])

    assert pending["status"] in {"queued", "running"}

    finished = wait_for_job(job["job_id"])
    assert finished["status"] == "succeeded"


def test_verify_job_retries_transient_internal_error_before_succeeding(monkeypatch, tmp_docx):
    clear_jobs()
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_retry_success.docx",
        rule_ids=("H02",),
    )
    attempts = {"count": 0}

    def fake_attempt(operation: str, resolved_request: dict[str, object]) -> dict[str, object]:
        attempts["count"] += 1
        if attempts["count"] == 1:
            return {
                "ok": False,
                "error": {
                    "code": "internal_error",
                    "type": "RuntimeError",
                    "message": "transient backend failure",
                    "http_status": 500,
                },
            }
        return {
            "ok": True,
            "result": jobs_module.verify_document(
                file_path=str(resolved_request["file_path"]),
                profile_path=resolved_request.get("profile_path"),
                scopes=resolved_request.get("scopes"),
                strict_profile=bool(resolved_request.get("strict_profile")),
            ),
        }

    monkeypatch.setattr(jobs_module, "_execute_job_attempt", fake_attempt)
    job = create_job(
        "verify",
        {
            "file_path": str(source_path),
            "scopes": ["headings"],
            "max_attempts": 2,
        },
    )
    wait_for_job(job["job_id"])
    result = get_job_result(job["job_id"])

    assert attempts["count"] == 2
    assert result["status"] == "succeeded"
    assert result["summary"]["attempt_count"] == 2
    assert result["summary"]["max_attempts"] == 2
    assert result["runtime"]["attempt_count"] == 2
    assert result["runtime"]["attempts"][0]["error_code"] == "internal_error"
    assert result["runtime"]["attempts"][0]["retryable"] is True
    assert result["runtime"]["attempts"][1]["status"] == "succeeded"
    assert [item["event"] for item in result["runtime"]["events"] if item["event"] == "retry_scheduled"] == [
        "retry_scheduled"
    ]


def test_verify_job_exhausts_retry_budget_on_timeout(monkeypatch, tmp_docx):
    clear_jobs()
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_retry_timeout.docx",
        rule_ids=("H02",),
    )
    attempts = {"count": 0}

    def fake_attempt(_operation: str, _resolved_request: dict[str, object]) -> dict[str, object]:
        attempts["count"] += 1
        return {
            "ok": False,
            "error": {
                "code": "worker_timeout",
                "type": "TimeoutExpired",
                "message": "Job attempt exceeded timeout of 0.01 seconds",
                "http_status": 504,
            },
        }

    monkeypatch.setattr(jobs_module, "_execute_job_attempt", fake_attempt)
    job = create_job(
        "verify",
        {
            "file_path": str(source_path),
            "scopes": ["headings"],
            "max_attempts": 2,
            "timeout_seconds": 0.01,
        },
    )
    wait_for_job(job["job_id"])
    result = get_job_result(job["job_id"])

    assert attempts["count"] == 2
    assert result["status"] == "failed"
    assert result["error"]["code"] == "worker_timeout"
    assert result["summary"]["error_code"] == "worker_timeout"
    assert result["summary"]["attempt_count"] == 2
    assert result["runtime"]["attempt_count"] == 2
    assert result["runtime"]["timeout_seconds"] == 0.01
    assert [item["retryable"] for item in result["runtime"]["attempts"]] == [True, True]


def test_unexpected_orchestration_exception_is_persisted_as_failed_job(monkeypatch, tmp_docx):
    clear_jobs()
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_orchestration_failure.docx",
        rule_ids=("H02",),
    )

    def fake_attempt(operation: str, resolved_request: dict[str, object]) -> dict[str, object]:
        return {
            "ok": True,
            "result": jobs_module.verify_document(
                file_path=str(resolved_request["file_path"]),
                profile_path=resolved_request.get("profile_path"),
                scopes=resolved_request.get("scopes"),
                strict_profile=bool(resolved_request.get("strict_profile")),
            ),
        }

    def broken_summary(_operation: str, _result: dict[str, object], _resolved_request: dict[str, object]) -> dict[str, object]:
        raise RuntimeError("summary exploded")

    monkeypatch.setattr(jobs_module, "_execute_job_attempt", fake_attempt)
    monkeypatch.setattr(jobs_module, "_result_summary", broken_summary)

    job = create_job(
        "verify",
        {
            "file_path": str(source_path),
            "scopes": ["headings"],
        },
    )
    wait_for_job(job["job_id"])
    result = get_job_result(job["job_id"])

    assert result["status"] == "failed"
    assert result["error"]["code"] == "internal_error"
    assert result["error"]["type"] == "RuntimeError"
    assert "Unexpected job orchestration failure: summary exploded" in result["error"]["message"]
    assert result["summary"]["error_code"] == "internal_error"
    assert result["result"] is None
    assert result["finished_at"] is not None
    assert result["runtime"]["lease_state"] == "released"
    assert result["runtime"]["events"][-1]["event"] == "job_failed"


def test_retry_job_recreates_finished_local_job_with_linked_origin(tmp_docx):
    clear_jobs()
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_retry_local.docx",
        rule_ids=("H02",),
    )

    first = _create_and_wait_job(
        "verify",
        {
            "file_path": str(source_path),
            "scopes": ["headings"],
        },
    )
    retried = retry_job(first["job_id"])
    wait_for_job(retried["job_id"])
    status = get_job(retried["job_id"])

    assert retried["job_id"] != first["job_id"]
    assert retried["request"]["file_path"] == str(source_path)
    assert retried["request"]["retry_of_job_id"] == first["job_id"]
    assert status["status"] == "succeeded"
    assert status["runtime"]["retry_of_job_id"] == first["job_id"]


def test_retry_job_rejects_missing_local_source_file(tmp_docx):
    clear_jobs()
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_retry_missing_local.docx",
        rule_ids=("H02",),
    )

    first = _create_and_wait_job(
        "verify",
        {
            "file_path": str(source_path),
            "scopes": ["headings"],
        },
    )
    Path(source_path).unlink()

    with pytest.raises(ValueError, match="文件不存在"):
        retry_job(first["job_id"])


def test_retry_job_rejects_missing_upload_record(tmp_docx, tmp_path):
    clear_jobs()
    runtime_root = tmp_path / "runtime"
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_retry_missing_upload_record.docx",
        rule_ids=("H02",),
    )

    class FakeUpload:
        filename = "article_job_retry_missing_upload_record.docx"

        def __init__(self, payload: bytes):
            from io import BytesIO

            self.file = BytesIO(payload)

    upload = store_uploaded_docx(FakeUpload(Path(source_path).read_bytes()), runtime_root=runtime_root)
    storage_module.upsert_upload(
        {
            "upload_id": upload.upload_id,
            "file_name": upload.file_name,
            "stored_path": upload.stored_path,
            "workspace_dir": upload.workspace_dir,
            "runtime_root": str(runtime_root),
            "size_bytes": upload.size_bytes,
            "created_at": "2026-04-19T00:00:00Z",
        }
    )
    first = _create_and_wait_job(
        "verify",
        {
            "file_path": upload.stored_path,
            "upload_id": upload.upload_id,
            "source_display_name": upload.file_name,
            "scopes": ["headings"],
            "runtime_root": str(runtime_root),
        },
    )
    storage_module.clear_uploads()

    with pytest.raises(RuntimeError, match="Upload not found for retry"):
        retry_job(first["job_id"])


def test_retry_job_rejects_upload_record_with_missing_file(tmp_docx, tmp_path):
    clear_jobs()
    runtime_root = tmp_path / "runtime"
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_retry_missing_upload_file.docx",
        rule_ids=("H02",),
    )

    class FakeUpload:
        filename = "article_job_retry_missing_upload_file.docx"

        def __init__(self, payload: bytes):
            from io import BytesIO

            self.file = BytesIO(payload)

    upload = store_uploaded_docx(FakeUpload(Path(source_path).read_bytes()), runtime_root=runtime_root)
    storage_module.upsert_upload(
        {
            "upload_id": upload.upload_id,
            "file_name": upload.file_name,
            "stored_path": upload.stored_path,
            "workspace_dir": upload.workspace_dir,
            "runtime_root": str(runtime_root),
            "size_bytes": upload.size_bytes,
            "created_at": "2026-04-19T00:00:00Z",
        }
    )
    first = _create_and_wait_job(
        "verify",
        {
            "file_path": upload.stored_path,
            "upload_id": upload.upload_id,
            "source_display_name": upload.file_name,
            "scopes": ["headings"],
            "runtime_root": str(runtime_root),
        },
    )
    Path(upload.stored_path).unlink()

    with pytest.raises(RuntimeError, match="Uploaded file is unavailable for retry"):
        retry_job(first["job_id"])


def test_retry_upload_backed_job_still_works_after_job_cleanup(tmp_docx, tmp_path):
    clear_jobs()
    runtime_root = tmp_path / "runtime"
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_retry_upload_after_cleanup.docx",
        rule_ids=("H02",),
    )

    class FakeUpload:
        filename = "article_job_retry_upload_after_cleanup.docx"

        def __init__(self, payload: bytes):
            from io import BytesIO

            self.file = BytesIO(payload)

    upload = store_uploaded_docx(FakeUpload(Path(source_path).read_bytes()), runtime_root=runtime_root)
    storage_module.upsert_upload(
        {
            "upload_id": upload.upload_id,
            "file_name": upload.file_name,
            "stored_path": upload.stored_path,
            "workspace_dir": upload.workspace_dir,
            "runtime_root": str(runtime_root),
            "size_bytes": upload.size_bytes,
            "created_at": "2026-04-19T00:00:00Z",
        }
    )
    first = _create_and_wait_job(
        "verify",
        {
            "file_path": upload.stored_path,
            "upload_id": upload.upload_id,
            "source_display_name": upload.file_name,
            "scopes": ["headings"],
            "runtime_root": str(runtime_root),
        },
    )

    cleanup_status = cleanup_job(first["job_id"])
    retried = retry_job(first["job_id"])
    wait_for_job(retried["job_id"])
    status = get_job(retried["job_id"])

    assert cleanup_status["cleanup"]["state"] in {"cleaned", "noop"}
    assert Path(upload.stored_path).exists()
    assert retried["request"]["upload_id"] == upload.upload_id
    assert status["status"] == "succeeded"
    assert status["runtime"]["source_upload_id"] == upload.upload_id
    assert status["runtime"]["retry_of_job_id"] == first["job_id"]


def test_create_job_rejects_invalid_worker_controls_without_storing_job(tmp_docx):
    clear_jobs()
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_invalid_worker_controls.docx",
        rule_ids=("H02",),
    )

    with pytest.raises(ValueError, match="max_attempts must be >= 1"):
        create_job(
            "verify",
            {
                "file_path": str(source_path),
                "scopes": ["headings"],
                "max_attempts": 0,
            },
        )

    with pytest.raises(ValueError, match="retry_delay_seconds must be >= 0"):
        create_job(
            "verify",
            {
                "file_path": str(source_path),
                "scopes": ["headings"],
                "retry_delay_seconds": -0.1,
            },
        )

    with pytest.raises(ValueError, match="timeout_seconds must be > 0"):
        create_job(
            "verify",
            {
                "file_path": str(source_path),
                "scopes": ["headings"],
                "timeout_seconds": 0,
            },
        )

    assert job_count() == 0


def test_active_future_is_not_reconciled_as_stale_running(monkeypatch, tmp_docx):
    clear_jobs()
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_active_future_not_stale.docx",
        rule_ids=("H02",),
    )

    def slow_attempt(_operation: str, resolved_request: dict[str, object]) -> dict[str, object]:
        time.sleep(0.2)
        return {
            "ok": True,
            "result": jobs_module.verify_document(
                file_path=str(resolved_request["file_path"]),
                profile_path=resolved_request.get("profile_path"),
                scopes=resolved_request.get("scopes"),
                strict_profile=bool(resolved_request.get("strict_profile")),
            ),
        }

    monkeypatch.setattr(jobs_module, "_execute_job_attempt", slow_attempt)
    job = create_job("verify", {"file_path": str(source_path), "scopes": ["headings"]})

    first_view = get_job(job["job_id"])
    listed = list_jobs()

    assert first_view["status"] in {"queued", "running"}
    assert first_view["summary"] is None
    assert first_view["error"] is None
    assert listed[0]["job_id"] == job["job_id"]
    assert listed[0]["status"] in {"queued", "running"}
    assert listed[0]["error"] is None

    with pytest.raises(RuntimeError, match="not ready yet"):
        get_job_result(job["job_id"])

    finished = wait_for_job(job["job_id"])
    assert finished["status"] == "succeeded"


def test_upload_id_respects_runtime_root_isolation(tmp_docx, tmp_path):
    clear_jobs()
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_upload_runtime_isolation.docx",
        rule_ids=("H02",),
    )

    class FakeUpload:
        filename = "article_job_upload_runtime_isolation.docx"

        def __init__(self, payload: bytes):
            from io import BytesIO

            self.file = BytesIO(payload)

    upload_a = store_uploaded_docx(FakeUpload(Path(source_path).read_bytes()), runtime_root=tmp_path / "runtime-a")
    upload_b = store_uploaded_docx(FakeUpload(Path(source_path).read_bytes()), runtime_root=tmp_path / "runtime-b")

    job_a = _create_and_wait_job(
        "verify",
        {
            "file_path": upload_a.stored_path,
            "upload_id": upload_a.upload_id,
            "source_display_name": upload_a.file_name,
            "scopes": ["headings"],
            "runtime_root": str(tmp_path / "runtime-a"),
        },
    )
    job_b = _create_and_wait_job(
        "verify",
        {
            "file_path": upload_b.stored_path,
            "upload_id": upload_b.upload_id,
            "source_display_name": upload_b.file_name,
            "scopes": ["headings"],
            "runtime_root": str(tmp_path / "runtime-b"),
        },
    )

    result_a = get_job_result(job_a["job_id"])
    result_b = get_job_result(job_b["job_id"])

    assert result_a["runtime"]["runtime_root"] == str((tmp_path / "runtime-a").resolve())
    assert result_b["runtime"]["runtime_root"] == str((tmp_path / "runtime-b").resolve())
    assert result_a["runtime"]["source_upload_id"] == upload_a.upload_id
    assert result_b["runtime"]["source_upload_id"] == upload_b.upload_id


def test_cleanup_job_skips_output_path_outside_runtime_root(tmp_docx, tmp_path):
    clear_jobs()
    runtime_root = tmp_path / "runtime"
    external_output = tmp_path / "external-output.docx"
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_cleanup_external_output.docx",
        rule_ids=("H02",),
    )

    job = _create_and_wait_job(
        "apply",
        {
            "file_path": str(source_path),
            "output_path": str(external_output),
            "scopes": ["headings"],
            "stage_input": True,
            "runtime_root": str(runtime_root),
        },
    )
    result = get_job_result(job["job_id"])
    assert external_output.exists()

    cleanup_status = cleanup_job(job["job_id"])

    assert external_output.exists()
    assert str(external_output.resolve()) in cleanup_status["cleanup"]["skipped_paths"]
    assert result["summary"] == get_job_result(job["job_id"])["summary"]


def test_completed_job_keeps_runtime_files_until_manual_cleanup(tmp_docx, tmp_path):
    clear_jobs()
    runtime_root = tmp_path / "runtime"
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_manual_retention.docx",
        rule_ids=("H02",),
    )

    job = _create_and_wait_job(
        "apply",
        {
            "file_path": str(source_path),
            "scopes": ["headings"],
            "stage_input": True,
            "runtime_root": str(runtime_root),
        },
    )

    status = get_job(job["job_id"])
    inspection = inspect_job(job["job_id"])
    listed = list_jobs()
    result = get_job_result(job["job_id"])
    workspace_root = Path(result["runtime"]["workspace_root"])
    output_path = Path(result["runtime"]["output_path"])

    assert status["runtime"]["cleanup_policy"] == "manual"
    assert status["cleanup"] is None
    assert inspection["cleanup"] is None
    assert listed[0]["job_id"] == job["job_id"]
    assert listed[0]["cleanup"] is None
    assert result["cleanup"] is None
    assert workspace_root.exists()
    assert output_path.exists()


def test_job_retention_sweep_dry_run_reports_expired_job_without_materializing_cleanup(tmp_docx, tmp_path):
    clear_jobs()
    runtime_root = tmp_path / "runtime"
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_retention_dry_run.docx",
        rule_ids=("H02",),
    )

    job = _create_and_wait_job(
        "apply",
        {
            "file_path": str(source_path),
            "scopes": ["headings"],
            "stage_input": True,
            "runtime_root": str(runtime_root),
        },
    )
    before = get_job_result(job["job_id"])
    workspace_root = Path(before["runtime"]["workspace_root"])
    report = sweep_job_retention(
        1,
        now=_future_timestamp(before["finished_at"], seconds=5),
        dry_run=True,
    )

    after = get_job_result(job["job_id"])

    assert report["policy"] == "retention"
    assert report["dry_run"] is True
    assert report["eligible_count"] == 1
    assert report["cleaned_count"] == 0
    assert report["noop_count"] == 0
    assert report["items"][0]["job_id"] == job["job_id"]
    assert report["items"][0]["action"] == "would_clean"
    assert after["cleanup"] is None
    assert after["summary"] == before["summary"]
    assert after["result"] == before["result"]
    assert workspace_root.exists()


def test_job_retention_sweep_cleans_expired_job_and_preserves_frozen_snapshot(tmp_docx, tmp_path):
    clear_jobs()
    runtime_root = tmp_path / "runtime"
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_retention_clean.docx",
        rule_ids=("H02",),
    )

    job = _create_and_wait_job(
        "apply",
        {
            "file_path": str(source_path),
            "scopes": ["headings"],
            "stage_input": True,
            "runtime_root": str(runtime_root),
        },
    )
    before = get_job_result(job["job_id"])
    workspace_root = Path(before["runtime"]["workspace_root"])
    report = sweep_job_retention(
        1,
        now=_future_timestamp(before["finished_at"], seconds=5),
    )

    after = get_job_result(job["job_id"])

    assert report["policy"] == "retention"
    assert report["dry_run"] is False
    assert report["eligible_count"] == 1
    assert report["cleaned_count"] == 1
    assert report["items"][0]["job_id"] == job["job_id"]
    assert report["items"][0]["action"] == "cleaned"
    assert report["items"][0]["cleanup"]["policy"] == "retention"
    assert after["cleanup"]["policy"] == "retention"
    assert after["cleanup"]["state"] == "cleaned"
    assert before["summary"] == after["summary"]
    assert before["result"] == after["result"]
    assert not workspace_root.exists()


def test_active_job_polling_does_not_materialize_retention_side_effects(monkeypatch, tmp_docx, tmp_path):
    clear_jobs()
    runtime_root = tmp_path / "runtime"
    source_path = _build_mutated_doc(
        tmp_docx,
        filename="article_job_polling_retention_side_effects.docx",
        rule_ids=("H02",),
    )

    def slow_attempt(_operation: str, resolved_request: dict[str, object]) -> dict[str, object]:
        time.sleep(0.2)
        return {
            "ok": True,
            "result": jobs_module.apply_fix(
                file_path=str(resolved_request["file_path"]),
                output_path=str(resolved_request["output_path"]),
                profile_path=resolved_request.get("profile_path"),
                scopes=resolved_request.get("scopes"),
                strict_profile=bool(resolved_request.get("strict_profile")),
                toc=bool(resolved_request.get("toc")),
                renumber_headings=bool(resolved_request.get("renumber_headings")),
                layout_rebalance=bool(resolved_request.get("layout_rebalance")),
                dry_run=bool(resolved_request.get("dry_run")),
                force=bool(resolved_request.get("force")),
            ),
        }

    monkeypatch.setattr(jobs_module, "_execute_job_attempt", slow_attempt)
    job = create_job(
        "apply",
        {
            "file_path": str(source_path),
            "scopes": ["headings"],
            "stage_input": True,
            "runtime_root": str(runtime_root),
        },
    )

    status = get_job(job["job_id"])
    inspection = inspect_job(job["job_id"])
    listed = list_jobs()
    workspace_root = Path(status["runtime"]["workspace_root"])

    assert status["status"] in {"queued", "running"}
    assert status["runtime"]["cleanup_policy"] == "manual"
    assert status["cleanup"] is None
    assert status["error"] is None
    assert inspection["cleanup"] is None
    assert listed[0]["job_id"] == job["job_id"]
    assert listed[0]["cleanup"] is None
    assert workspace_root.exists()

    finished = wait_for_job(job["job_id"])

    assert finished["status"] == "succeeded"
    assert get_job_result(job["job_id"])["cleanup"] is None
    assert workspace_root.exists()


def test_stale_running_recovery_preserves_cleanup_metadata(tmp_path, monkeypatch):
    monkeypatch.setenv("ARTICLE_API_STATE_ROOT", str(tmp_path / "state"))
    clear_jobs()
    payload = {
        "job_id": "stale-running-cleanup-job",
        "operation": "verify",
        "status": "running",
        "mode": "background",
        "created_at": "2026-04-19T00:00:00Z",
        "updated_at": "2026-04-19T00:05:00Z",
        "started_at": "2026-04-19T00:01:00Z",
        "finished_at": None,
        "request": {"file_path": "/tmp/stale-cleanup.docx", "scopes": ["headings"]},
        "resolved_request": {
            "file_path": "/tmp/stale-cleanup.docx",
            "scopes": ["headings"],
            "strict_profile": False,
        },
        "workspace": None,
        "runtime": {
            "cleanup_policy": "manual",
            "runtime_root": None,
            "source_upload_id": None,
            "stage_input_enabled": False,
            "workspace_present": False,
            "workspace_root": None,
            "workspace_inputs": None,
            "workspace_outputs": None,
            "source_file_path": None,
            "staged_input_path": None,
            "output_path": None,
            "retry_of_job_id": None,
            "max_attempts": 1,
            "retry_delay_seconds": 0.0,
            "timeout_seconds": None,
            "attempt_count": 0,
            "attempts": [],
        },
        "result_available": False,
        "summary": None,
        "artifacts": [],
        "error": None,
        "result": None,
    }
    cleanup_payload = {
        "policy": "manual",
        "state": "cleaned",
        "cleaned_at": "2026-04-19T00:06:00Z",
        "attempt_count": 1,
        "removed_paths": ["/tmp/runtime/stale-running-cleanup-job"],
        "missing_paths": [],
        "skipped_paths": [],
    }
    storage_module.upsert_job(payload)
    storage_module.upsert_job_cleanup(payload["job_id"], cleanup_payload)

    recovered = get_job(payload["job_id"])
    listed = list_jobs()
    result = get_job_result(payload["job_id"])

    assert recovered["status"] == "failed"
    assert recovered["summary"]["error_code"] == "worker_recovery_failed"
    assert recovered["cleanup"] == cleanup_payload
    assert listed[0]["job_id"] == payload["job_id"]
    assert listed[0]["cleanup"] == cleanup_payload
    assert result["error"]["code"] == "worker_recovery_failed"
    assert result["cleanup"] == cleanup_payload
