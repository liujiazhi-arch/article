from __future__ import annotations

import importlib.util
import io
from pathlib import Path
import time
import zipfile

from docx import Document
import pytest

from .conftest import RULE_MUTATORS, make_compliant_doc


_HTTP_RUNTIME_READY = (
    importlib.util.find_spec("fastapi") is not None
    and importlib.util.find_spec("httpx") is not None
)
pytestmark = pytest.mark.skipif(
    not _HTTP_RUNTIME_READY,
    reason="requires fastapi + httpx in the active interpreter",
)

if _HTTP_RUNTIME_READY:  # pragma: no branch
    from fastapi.testclient import TestClient
    import article_api.app as app_module
    from article_api.app import create_app


def _make_style_conflict_lnu_doc(source_path: Path) -> Path:
    doc = Document()
    heading = doc.add_paragraph("1 绪论")
    heading.style = doc.styles["Heading 1"]
    conflict = doc.add_paragraph("3.6 分子对接验证结果")
    conflict.style = doc.styles["Heading 1"]
    doc.add_paragraph("这是正文示例。")
    doc.save(source_path)
    return source_path


def _make_client(monkeypatch, state_root: Path) -> TestClient:
    monkeypatch.setenv("ARTICLE_API_STATE_ROOT", str(state_root))
    return TestClient(create_app())


def _wait_for_job_completion(client: TestClient, job_id: str) -> dict[str, object]:
    for _ in range(150):
        status_response = client.get(f"/jobs/{job_id}")
        assert status_response.status_code == 200
        status_payload = status_response.json()
        if status_payload["status"] in {"succeeded", "failed"}:
            return status_payload
        time.sleep(0.05)
    raise TimeoutError(f"Timed out waiting for HTTP job completion: {job_id}")


def _upload_docx(client: TestClient, source_path: Path, runtime_root: Path) -> dict[str, object]:
    with open(source_path, "rb") as handle:
        upload_response = client.post(
            "/uploads/docx",
            params={"runtime_root": str(runtime_root)},
            files={
                "file": (
                    source_path.name,
                    handle,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
    assert upload_response.status_code == 201
    return upload_response.json()


def _upload_pdf(client: TestClient, runtime_root: Path) -> dict[str, object]:
    upload_response = client.post(
        "/uploads/pdf",
        params={"runtime_root": str(runtime_root)},
        files={"file": ("20221303306-刘佳轾-排版复核.pdf", b"%PDF-1.7\nfake-pdf-binary", "application/pdf")},
    )
    assert upload_response.status_code == 201
    return upload_response.json()


def test_live_http_upload_apply_result_download_and_cleanup(monkeypatch, tmp_docx, tmp_path):
    monkeypatch.delenv("ARTICLE_LOCAL_RELEASE_API_URL", raising=False)
    source_path = tmp_docx(make_compliant_doc, filename="article_http_smoke.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    doc.save(source_path)
    runtime_root = tmp_path / "runtime"
    client = _make_client(monkeypatch, tmp_path / "state")

    health_response = client.get("/health")
    console_response = client.get("/")
    ready_response = client.get("/ready")
    version_response = client.get("/version")
    update_response = client.get("/updates/latest")
    profiles_response = client.get("/profiles")
    runtime_response = client.get("/ops/runtime")
    storage_response = client.get("/ops/storage")
    emblem_response = client.get("/assets/lnu-emblem.jpg")
    assert console_response.status_code == 200
    assert health_response.status_code == 200
    assert ready_response.status_code == 200
    assert version_response.status_code == 200
    assert update_response.status_code == 200
    assert profiles_response.status_code == 200
    assert runtime_response.status_code == 200
    assert storage_response.status_code == 200
    assert emblem_response.status_code == 200
    assert emblem_response.headers["content-type"].startswith("image/jpeg")
    assert "论文格式本地控制台" in console_response.text
    assert "辽宁大学毕业论文" in console_response.text
    assert "brand-mark" in console_response.text
    assert "assets/lnu-emblem.jpg" in console_response.text
    assert "选择 Word 论文" in console_response.text
    assert "辽宁大学毕业论文格式" in console_response.text
    assert "格式检查与修复" in console_response.text
    assert "上传 DOCX，生成修复稿；再上传导出的 PDF 复审。" not in console_response.text
    assert "一键生成修复稿" not in console_response.text
    assert "生成修复方案" in console_response.text
    assert "按所选范围修复" in console_response.text
    assert "上传 PDF 复审" in console_response.text
    assert "按这些问题生成下一版 DOCX" in console_response.text
    assert "任务查询" in console_response.text
    assert "最近任务" in console_response.text
    assert "refreshRecentJobs" in console_response.text
    assert "/jobs?limit=10" in console_response.text
    assert "高级设置" not in console_response.text
    assert "分步操作" not in console_response.text
    assert "advanced-details" not in console_response.text
    assert "single-profile" not in console_response.text
    assert "cn-common" not in console_response.text
    assert "profile: 'lnu'" not in console_response.text
    assert "可选动作" not in console_response.text
    assert "处理选项" not in console_response.text
    assert "PDF 复审" in console_response.text
    assert "检查导出的 PDF" in console_response.text
    assert "开始复审" in console_response.text
    assert "data-workflow-mode=\"default_user\"" not in console_response.text
    assert "data-workflow-mode=\"advanced_word\"" not in console_response.text
    assert "selectedWorkflowMode: 'default_user'" in console_response.text
    assert "function userFacingError" in console_response.text
    assert "function jobErrorText" in console_response.text
    assert "jobErrorText(detail.error || detail)" in console_response.text
    assert "jobErrorText(payload.error || payload)" in console_response.text
    assert "renderWorkflowStatusItems" in console_response.text
    assert "formatRenderFinding" in console_response.text
    assert "renderFindingItems" in console_response.text
    assert "verifyIssueItems" in console_response.text
    assert "formatRuleSummary" in console_response.text
    assert "pollAgentCandidateJob" in console_response.text
    assert "下一版 DOCX 已提交后端任务" in console_response.text
    assert "后端任务仍在运行，不是页面卡死" in console_response.text
    assert "正在转换 PDF 页面；30 页左右可能需要 1-3 分钟，不是页面卡住" in console_response.text
    assert "需要版式复核原因" in console_response.text
    assert "技术详情" in console_response.text
    assert "预计下一版路径" in console_response.text
    assert "等待结果整理" in console_response.text
    assert "const AGENT_CANDIDATE_PROGRESS = { submitted: 30, running: 55, finalizing: 80, finished: 100 };" in console_response.text
    assert "PDF 复审不修改 DOCX" in console_response.text
    assert "详细报告" in console_response.text
    assert "已向后端发送高级模式请求" not in console_response.text
    assert "后端没有拿到 Word 导出的 render_verify_word.pdf" not in console_response.text
    assert "PDF 复审完成" in console_response.text
    assert "PDF 复审" in console_response.text
    assert "上传 PDF" in console_response.text
    assert "render-pdf-upload-zone" in console_response.text
    assert "render-pdf-file" in console_response.text
    assert "/uploads/pdf" in console_response.text
    assert "高级模式" not in console_response.text
    assert "Agent 候选稿模式" not in console_response.text
    assert "排障模式工作台" not in console_response.text
    assert "进入排障模式" not in console_response.text
    assert "const PIPELINE_STEP_IDS = ['preflight', 'plan', 'apply', 'verify'];" not in console_response.text
    assert "for (const id of ['preflight', 'plan'])" in console_response.text
    assert "guardedRenderWorkflow('advanced_word')" not in console_response.text
    assert "原文不会被覆盖" in console_response.text
    assert "总体结论" in console_response.text
    assert "report-action-button" in console_response.text
    assert "report-progress" in console_response.text
    assert "正在修复" in console_response.text
    assert "修复稿已生成" in console_response.text
    assert "结构复查完成" in console_response.text
    assert "source-summary" in console_response.text
    assert "displayFileName" in console_response.text
    assert "sourceDisplayName" in console_response.text
    assert "outputFolderLabel" in console_response.text
    assert "outputSummary" in console_response.text
    assert "复核后排障工具" not in console_response.text
    assert "先完成 PDF 版式复核后再使用" not in console_response.text
    assert "查看路径" in console_response.text
    assert "桌面/论文格式修复输出" in console_response.text
    assert "任务状态" in console_response.text
    assert "renderJobReportTrace" in console_response.text
    assert "查看审查报告" in console_response.text
    assert "检查新版本" in console_response.text
    assert "只检查软件版本，不上传论文" in console_response.text
    assert "/updates/latest" in console_response.text
    assert "对应文件" in console_response.text
    assert "outputArtifact.available === true" in console_response.text
    assert "文件已清理，不能直接下载" in console_response.text
    assert "/artifacts/report/download" in console_response.text
    assert "心跳间隔" not in console_response.text
    assert "阶段:" not in console_response.text
    assert "超时阈值" not in console_response.text
    assert "等待选择任务" not in console_response.text
    assert "历史任务" not in console_response.text
    assert "批量任务" not in console_response.text
    assert "适合发给学弟学妹使用" not in console_response.text
    assert health_response.json()["service"] == "article-api"
    assert ready_response.json()["status"] == "ready"
    assert ready_response.json()["checks"]["runtime_root"]["status"] == "ok"
    assert version_response.json()["api_version"] == "v0"
    assert version_response.json()["update_check"]["mode"] == "manual"
    assert update_response.json()["status"] == "not_configured"
    assert update_response.json()["auto_update"] is False
    profiles_payload = profiles_response.json()
    assert profiles_payload["summary"]["profile_count"] == 1
    assert profiles_payload["summary"]["default_profile_id"] == "lnu-checker-2026"
    assert [item["id"] for item in profiles_payload["profiles"]] == ["lnu-checker-2026"]
    assert "cn-common" not in str(profiles_payload).lower()
    assert "课程作业" not in str(profiles_payload)
    assert "综述" not in str(profiles_payload)
    assert [item["id"] for item in profiles_payload["summary"]["support_scenarios"]] == ["school_degree_thesis"]
    lnu_profile = profiles_payload["profiles"][0]
    assert any(item["label"] == "学校学位论文" for item in lnu_profile["support_scenarios"])
    assert runtime_response.json()["runtime"]["worker_model"] == "single"
    assert runtime_response.json()["runtime"]["recovery"]["strategy"] == "fail_unfinished_without_active_worker_future"
    assert storage_response.json()["storage"]["integrity_check"] == "ok"

    preflight_source = _make_style_conflict_lnu_doc(tmp_path / "article_http_preflight.docx")
    preflight_response = client.post(
        "/preflight",
        json={
            "file_path": str(preflight_source),
            "profile": "lnu",
        },
    )
    assert preflight_response.status_code == 200
    preflight_payload = preflight_response.json()
    assert preflight_payload["operation"] == "preflight"
    assert preflight_payload["preflight_status"] == "blocked"
    assert preflight_payload["summary"]["wild_doc_detected"] is True
    assert any(item["id"] == "style_text_conflicts" for item in preflight_payload["wild_doc"]["signals"])

    normalize_response = client.post(
        "/normalize",
        json={
            "file_path": str(preflight_source),
            "output_path": str(tmp_path / "article_http_preflight_normalized.docx"),
            "profile": "lnu",
        },
    )
    assert normalize_response.status_code == 200
    normalize_payload = normalize_response.json()
    assert normalize_payload["operation"] == "normalize"
    assert normalize_payload["changed"] is True
    assert normalize_payload["summary"]["before_preflight_status"] == "blocked"
    assert normalize_payload["summary"]["after_preflight_status"] == "warning"
    assert normalize_payload["wild_doc"]["after"]["style_conflict_count"] == 0

    normalize_job_response = client.post(
        "/jobs/normalize",
        json={
            "file_path": str(preflight_source),
            "output_path": str(tmp_path / "article_http_preflight_job_normalized.docx"),
            "profile": "lnu",
        },
    )
    assert normalize_job_response.status_code == 201
    normalize_job_payload = normalize_job_response.json()
    normalize_job_status = _wait_for_job_completion(client, normalize_job_payload["job_id"])
    assert normalize_job_status["status"] == "succeeded"
    normalize_job_result = client.get(f"/jobs/{normalize_job_payload['job_id']}/result")
    assert normalize_job_result.status_code == 200
    assert normalize_job_result.json()["summary"]["business_status"] == "warning"
    assert normalize_job_result.json()["result"]["after"]["preflight_status"] == "warning"

    monkeypatch.setattr(
        app_module,
        "render_verify_document",
        lambda *args, **kwargs: {
            "document": {"path": str(source_path), "name": source_path.name},
            "profile": {"id": "lnu-checker-2026", "requested": "lnu", "fallback_used": False, "display": "lnu"},
            "output_dir": str(tmp_path / "http-render-proof"),
            "render_engine": "manual-pdf",
            "evidence_source": "manual-pdf",
            "evidence_trust": "authoritative",
            "evidence_authoritative": True,
            "layout_decision_eligible": True,
            "render_fallback_used": False,
            "page_count": 1,
            "page_images": [str(tmp_path / "http-render-proof" / "page-1.png")],
            "render_findings": [],
            "render_summary": {"finding_count": 0, "highest_severity": None},
            "selected_scopes": ["toc"],
            "overall_status": "verified",
            "readiness": "render-check-required",
            "manual_review_rule_ids": ["LNU_TOC03"],
            "unsupported_rule_ids": [],
            "review_items": ["目录需要刷新后复核页码。"],
            "report_path": str(tmp_path / "http-render-proof" / "render_verify_report.md"),
        },
    )
    render_verify_response = client.post(
        "/render-verify",
        json={
            "file_path": str(source_path),
            "profile": "lnu",
            "scopes": ["toc"],
            "rendered_pdf": str(tmp_path / "manual-word-export.pdf"),
            "workflow_mode": "default_user",
        },
    )
    assert render_verify_response.status_code == 200
    render_verify_payload = render_verify_response.json()
    assert render_verify_payload["operation"] == "render-verify"
    assert render_verify_payload["page_count"] == 1
    assert render_verify_payload["summary"]["manual_review_rule_count"] == 1
    assert render_verify_payload["selected_scopes"] == ["toc"]

    upload_payload = _upload_docx(client, source_path, runtime_root)
    pdf_upload_payload = _upload_pdf(client, runtime_root)
    assert upload_payload["file_name"] == "article_http_smoke.docx"
    assert pdf_upload_payload["file_name"] == "20221303306-刘佳轾-排版复核.pdf"
    assert pdf_upload_payload["available"] is True
    assert Path(pdf_upload_payload["stored_path"]).exists()

    uploads_response = client.get("/uploads")
    upload_view_response = client.get(f"/uploads/{upload_payload['upload_id']}")
    assert uploads_response.status_code == 200
    assert upload_view_response.status_code == 200
    assert uploads_response.json()[0]["upload_id"] == upload_payload["upload_id"]
    assert upload_view_response.json()["available"] is True

    normalize_upload_response = client.post(
        f"/uploads/{upload_payload['upload_id']}/jobs/normalize",
        json={},
    )
    assert normalize_upload_response.status_code == 201
    normalize_upload_payload = normalize_upload_response.json()
    normalize_upload_status = _wait_for_job_completion(client, normalize_upload_payload["job_id"])
    assert normalize_upload_status["status"] == "succeeded"
    normalize_upload_result = client.get(f"/jobs/{normalize_upload_payload['job_id']}/result")
    assert normalize_upload_result.status_code == 200
    assert normalize_upload_result.json()["result"]["document"]["name"] == "article_http_smoke.docx"
    assert normalize_upload_result.json()["runtime"]["source_upload_id"] == upload_payload["upload_id"]

    create_response = client.post(
        f"/uploads/{upload_payload['upload_id']}/jobs/apply",
        json={
            "scopes": ["headings"],
            "stage_input": True,
        },
    )
    assert create_response.status_code == 201
    create_payload = create_response.json()
    assert create_payload["request"]["upload_id"] == upload_payload["upload_id"]
    assert "file_path" not in create_payload["request"]

    job_id = create_payload["job_id"]
    status_payload = _wait_for_job_completion(client, job_id)
    assert status_payload["status"] == "succeeded"

    result_response = client.get(f"/jobs/{job_id}/result")
    assert result_response.status_code == 200
    result_payload = result_response.json()
    assert result_payload["result"]["document"]["name"] == "article_http_smoke.docx"
    assert result_payload["runtime"]["source_upload_id"] == upload_payload["upload_id"]
    assert next(item for item in result_payload["artifacts"] if item["role"] == "output")["available"] is True
    report_artifact = next(item for item in result_payload["artifacts"] if item["role"] == "report")
    assert report_artifact["available"] is True
    assert report_artifact["download_name"] == "job_report.md"

    download_response = client.get(f"/jobs/{job_id}/artifacts/output/download")
    assert download_response.status_code == 200
    assert download_response.content
    assert "article_http_smoke_%E6%A0%BC%E5%BC%8F%E4%BF%AE%E5%A4%8D_V01.docx" in download_response.headers.get("content-disposition", "")
    report_download_response = client.get(f"/jobs/{job_id}/artifacts/report/download")
    assert report_download_response.status_code == 200
    assert report_download_response.headers["content-type"].startswith("text/markdown")
    assert f"任务编号: {job_id}".encode("utf-8") in report_download_response.content
    feedback_response = client.get("/feedback/download")
    assert feedback_response.status_code == 200
    assert feedback_response.headers["content-type"].startswith("application/zip")
    assert "feedback" in feedback_response.headers.get("content-disposition", "").lower() or "%E5%8F%8D%E9%A6%88%E5%8C%85" in feedback_response.headers.get("content-disposition", "")
    with zipfile.ZipFile(io.BytesIO(feedback_response.content)) as archive:
        feedback_names = set(archive.namelist())
        feedback_manifest = archive.read("manifest.json").decode("utf-8")
    assert {"manifest.json", "doctor.json", "jobs.json", "uploads.json", "runtime_files.json"} <= feedback_names
    assert not any(name.endswith((".docx", ".pdf", ".png")) for name in feedback_names)
    assert str(tmp_path) not in feedback_manifest

    cleanup_response = client.post(f"/jobs/{job_id}/cleanup")
    assert cleanup_response.status_code == 200
    assert cleanup_response.json()["cleanup"]["state"] == "cleaned"

    after_response = client.get(f"/jobs/{job_id}/result")
    assert after_response.status_code == 200
    after_payload = after_response.json()
    assert after_payload["summary"] == result_payload["summary"]
    assert after_payload["result"] == result_payload["result"]
    assert next(item for item in after_payload["artifacts"] if item["role"] == "output")["available"] is False
    assert next(item for item in after_payload["artifacts"] if item["role"] == "report")["available"] is True

    missing_download_response = client.get(f"/jobs/{job_id}/artifacts/output/download")
    assert missing_download_response.status_code == 409
    missing_download_detail = missing_download_response.json()["detail"]
    assert missing_download_detail["code"] == "artifact_unavailable"
    assert missing_download_detail["user_message"] == "下载文件已经不可用。"
    assert "重新运行修复任务" in missing_download_detail["next_action"]

    summary_response = client.get("/ops/summary")
    assert summary_response.status_code == 200
    summary_payload = summary_response.json()
    assert summary_payload["jobs"]["total"] >= 1
    assert summary_payload["jobs"]["succeeded"] >= 1
    assert summary_payload["runtime"]["worker_model"] == "single"
    assert summary_payload["storage"]["index_count"] >= 1
    assert summary_payload["runtime"]["pending_recovery_count"] == 0


def test_live_http_inspect_and_retry_upload_backed_job(monkeypatch, tmp_docx, tmp_path):
    source_path = tmp_docx(make_compliant_doc, filename="article_http_retry.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    doc.save(source_path)
    runtime_root = tmp_path / "runtime"
    client = _make_client(monkeypatch, tmp_path / "state")
    upload_payload = _upload_docx(client, source_path, runtime_root)

    create_response = client.post(
        f"/uploads/{upload_payload['upload_id']}/jobs/verify",
        json={
            "scopes": ["headings"],
            "stage_input": True,
        },
    )
    assert create_response.status_code == 201
    create_payload = create_response.json()
    first_job_id = create_payload["job_id"]
    first_status = _wait_for_job_completion(client, first_job_id)
    assert first_status["status"] == "succeeded"

    inspect_response = client.get(f"/jobs/{first_job_id}/inspect")
    assert inspect_response.status_code == 200
    inspect_payload = inspect_response.json()
    assert inspect_payload["job_id"] == first_job_id
    assert inspect_payload["status"] == "succeeded"
    assert inspect_payload["result_available"] is True
    assert inspect_payload["runtime"]["source_upload_id"] == upload_payload["upload_id"]
    assert "request" not in inspect_payload
    assert "result" not in inspect_payload

    retry_response = client.post(f"/jobs/{first_job_id}/retry")
    assert retry_response.status_code == 201
    retry_payload = retry_response.json()
    assert retry_payload["request"]["upload_id"] == upload_payload["upload_id"]
    assert retry_payload["request"]["retry_of_job_id"] == first_job_id
    assert "file_path" not in retry_payload["request"]

    retried_status = _wait_for_job_completion(client, retry_payload["job_id"])
    assert retried_status["status"] == "succeeded"
    assert retried_status["runtime"]["retry_of_job_id"] == first_job_id

    missing_retry_response = client.post("/jobs/missing-job/retry")
    assert missing_retry_response.status_code == 404
    assert "Job not found" in missing_retry_response.json()["detail"]


def test_live_http_upload_cleanup_and_missing_upload_mapping(monkeypatch, tmp_docx, tmp_path):
    source_path = tmp_docx(make_compliant_doc, filename="article_http_upload_cleanup.docx")
    runtime_root = tmp_path / "runtime"
    client = _make_client(monkeypatch, tmp_path / "state")
    upload_payload = _upload_docx(client, source_path, runtime_root)

    cleanup_response = client.post(f"/uploads/{upload_payload['upload_id']}/cleanup")
    assert cleanup_response.status_code == 200
    cleanup_payload = cleanup_response.json()
    assert cleanup_payload["upload_id"] == upload_payload["upload_id"]
    assert cleanup_payload["cleanup"]["state"] == "cleaned"
    assert cleanup_payload["cleanup"]["attempt_count"] == 1
    assert cleanup_payload["available"] is False

    upload_status_response = client.get(f"/uploads/{upload_payload['upload_id']}")
    assert upload_status_response.status_code == 200
    assert upload_status_response.json()["available"] is False

    blocked_job_response = client.post(
        f"/uploads/{upload_payload['upload_id']}/jobs/verify",
        json={"scopes": ["headings"]},
    )
    assert blocked_job_response.status_code == 409
    blocked_detail = blocked_job_response.json()["detail"]
    assert blocked_detail["code"] == "upload_unavailable"
    assert blocked_detail["user_message"] == "上传文件已经不可用。"
    assert blocked_detail["next_action"] == "请重新上传论文后再试。"
    assert blocked_detail["retryable"] is False

    missing_upload_cleanup_response = client.post("/uploads/missing-upload/cleanup")
    assert missing_upload_cleanup_response.status_code == 404
    missing_detail = missing_upload_cleanup_response.json()["detail"]
    assert missing_detail["code"] == "upload_not_found"
    assert missing_detail["user_message"] == "没有找到这个上传记录。"
    assert missing_detail["next_action"] == "请重新上传论文后再试。"
    assert missing_detail["retryable"] is False


def test_live_http_pdf_upload_rejects_non_pdf_payload_with_chinese_guidance(monkeypatch, tmp_path):
    runtime_root = tmp_path / "runtime"
    client = _make_client(monkeypatch, tmp_path / "state")

    response = client.post(
        "/uploads/pdf",
        params={"runtime_root": str(runtime_root)},
        files={"file": ("paper.pdf", b"not-a-pdf", "application/pdf")},
    )

    assert response.status_code == 400
    detail = response.json()["detail"]
    assert detail["code"] == "invalid_pdf"
    assert detail["user_message"] == "上传的文件不是有效的 PDF 文档。"
    assert detail["next_action"] == "请从 Word/WPS 重新导出 PDF 后再上传。"
    assert not list((runtime_root / "uploads").glob("*.pdf"))


def test_live_http_retention_sweep_cleans_expired_job_and_upload(monkeypatch, tmp_docx, tmp_path):
    source_path = tmp_docx(make_compliant_doc, filename="article_http_retention.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    doc.save(source_path)
    runtime_root = tmp_path / "runtime"
    client = _make_client(monkeypatch, tmp_path / "state")
    upload_payload = _upload_docx(client, source_path, runtime_root)

    create_response = client.post(
        f"/uploads/{upload_payload['upload_id']}/jobs/apply",
        json={
            "scopes": ["headings"],
            "stage_input": True,
        },
    )
    assert create_response.status_code == 201
    create_payload = create_response.json()
    job_id = create_payload["job_id"]

    status_payload = _wait_for_job_completion(client, job_id)
    assert status_payload["status"] == "succeeded"

    monkeypatch.setattr(app_module, "_utcnow", lambda: "2099-01-01T00:00:00Z")
    sweep_response = client.post(
        "/ops/retention/sweep",
        json={
            "job_max_age_seconds": 1,
            "upload_max_age_seconds": 1,
        },
    )
    assert sweep_response.status_code == 200
    sweep_payload = sweep_response.json()
    assert sweep_payload["job_retention"]["cleaned_count"] == 1
    assert sweep_payload["upload_retention"]["cleaned_count"] == 1

    result_response = client.get(f"/jobs/{job_id}/result")
    upload_response = client.get(f"/uploads/{upload_payload['upload_id']}")
    assert result_response.status_code == 200
    assert upload_response.status_code == 200
    assert result_response.json()["cleanup"]["policy"] == "retention"
    assert upload_response.json()["cleanup"]["policy"] == "retention"
    assert upload_response.json()["available"] is False


def test_live_http_run_default_retention_uses_env_thresholds(monkeypatch, tmp_docx, tmp_path):
    monkeypatch.setenv("ARTICLE_API_JOB_RETENTION_SECONDS", "1")
    monkeypatch.setenv("ARTICLE_API_UPLOAD_RETENTION_SECONDS", "1")
    source_path = tmp_docx(make_compliant_doc, filename="article_http_default_retention.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    doc.save(source_path)
    runtime_root = tmp_path / "runtime"
    client = _make_client(monkeypatch, tmp_path / "state")
    upload_payload = _upload_docx(client, source_path, runtime_root)

    create_response = client.post(
        f"/uploads/{upload_payload['upload_id']}/jobs/apply",
        json={
            "scopes": ["headings"],
            "stage_input": True,
        },
    )
    assert create_response.status_code == 201
    create_payload = create_response.json()
    job_id = create_payload["job_id"]
    status_payload = _wait_for_job_completion(client, job_id)
    assert status_payload["status"] == "succeeded"

    monkeypatch.setattr(app_module, "_utcnow", lambda: "2099-01-01T00:00:00Z")
    sweep_response = client.post("/ops/retention/run-defaults", json={})
    assert sweep_response.status_code == 200
    sweep_payload = sweep_response.json()
    assert sweep_payload["job_retention"]["cleaned_count"] == 1
    assert sweep_payload["upload_retention"]["cleaned_count"] == 1

    summary_response = client.get("/ops/summary")
    assert summary_response.status_code == 200
    assert summary_response.json()["retention"]["state"]["last_trigger"] == "manual-defaults"


def test_live_http_download_missing_artifact_maps_to_404(monkeypatch, tmp_docx, tmp_path):
    source_path = tmp_docx(make_compliant_doc, filename="article_http_missing_artifact.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    doc.save(source_path)
    client = _make_client(monkeypatch, tmp_path / "state")

    create_response = client.post(
        "/jobs/verify",
        json={
            "file_path": str(source_path),
            "scopes": ["headings"],
            "stage_input": True,
            "runtime_root": str(tmp_path / "runtime"),
        },
    )
    assert create_response.status_code == 201
    create_payload = create_response.json()
    status_payload = _wait_for_job_completion(client, create_payload["job_id"])
    assert status_payload["status"] == "succeeded"

    missing_artifact_response = client.get(
        f"/jobs/{create_payload['job_id']}/artifacts/missing/download"
    )
    assert missing_artifact_response.status_code == 404
    missing_artifact_detail = missing_artifact_response.json()["detail"]
    assert missing_artifact_detail["code"] == "artifact_not_found"
    assert missing_artifact_detail["user_message"] == "没有找到这个下载项。"
    assert "刷新任务结果" in missing_artifact_detail["next_action"]
