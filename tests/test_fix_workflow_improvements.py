from __future__ import annotations

import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import audit_thesis
import pytest
import fix_thesis
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from thesis_fix import toc as toc_fix
from thesis_tool.workflow import apply_scoped_fix, build_scoped_fix_preview, render_scoped_fix_preview

from .conftest import audit_rule_status, make_compliant_doc, make_violating_doc


def _paragraph_by_prefix(doc: Document, prefix: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith(prefix):
            return paragraph
    raise AssertionError(f"Paragraph not found: {prefix}")


def _normalize_spaces(text: str) -> str:
    return " ".join((text or "").split())


def _paragraph_by_normalized_prefix(doc: Document, prefix: str):
    normalized_prefix = _normalize_spaces(prefix)
    for paragraph in doc.paragraphs:
        if _normalize_spaces(paragraph.text).startswith(normalized_prefix):
            return paragraph
    raise AssertionError(f"Paragraph not found: {prefix}")


def _profile_cfg(profile_path: str = "lnu") -> dict:
    return fix_thesis.build_fix_runtime(profile_path=profile_path).cfg


def _add_mock_drawing(paragraph) -> None:
    run = paragraph.add_run()
    run._r.append(OxmlElement("w:drawing"))


def _add_anchor_drawing(paragraph, *, drawing_id: int) -> None:
    drawing = OxmlElement("w:drawing")
    anchor = OxmlElement("wp:anchor")
    for name in ("distT", "distB", "distL", "distR"):
        anchor.set(name, "0")
    extent = OxmlElement("wp:extent")
    extent.set("cx", "360000")
    extent.set("cy", "360000")
    doc_pr = OxmlElement("wp:docPr")
    doc_pr.set("id", str(drawing_id))
    doc_pr.set("name", f"Anchor {drawing_id}")
    anchor.extend((extent, doc_pr, OxmlElement("wp:cNvGraphicFramePr"), OxmlElement("a:graphic")))
    drawing.append(anchor)
    paragraph.add_run()._r.append(drawing)


def _make_cover_and_body_anchor_docx(path: Path) -> None:
    doc = Document()
    cover = doc.add_paragraph("封面图")
    _add_anchor_drawing(cover, drawing_id=1)
    doc.add_paragraph("第1章 绪论").style = doc.styles["Heading 1"]
    body = doc.add_paragraph("正文图")
    _add_anchor_drawing(body, drawing_id=2)
    doc.save(path)


def _assert_cover_anchor_and_body_inline(path: Path) -> None:
    doc = Document(path)
    cover = _paragraph_by_prefix(doc, "封面图")
    body = _paragraph_by_prefix(doc, "正文图")
    assert cover._p.xpath(".//wp:anchor")
    assert not cover._p.xpath(".//wp:inline")
    assert not body._p.xpath(".//wp:anchor")
    assert body._p.xpath(".//wp:inline")


def _inject_equation_layout_table(
    docx_path: Path,
    eq_number: str = "(1.1)",
    *,
    additional_numbers: tuple[str, ...] = (),
) -> None:
    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    m_ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"

    def _w(tag: str) -> str:
        return f"{{{w_ns}}}{tag}"

    def _m(tag: str) -> str:
        return f"{{{m_ns}}}{tag}"

    with zipfile.ZipFile(docx_path, "r") as zf:
        parts = {name: zf.read(name) for name in zf.namelist()}

    root = ET.fromstring(parts["word/document.xml"])
    body = root.find(_w("body"))
    assert body is not None
    sect_pr = body.find(_w("sectPr"))
    assert sect_pr is not None

    tbl = ET.Element(_w("tbl"))
    tbl_pr = ET.SubElement(tbl, _w("tblPr"))
    tbl_borders = ET.SubElement(tbl_pr, _w("tblBorders"))
    for name in ("top", "bottom", "left", "right", "insideH", "insideV"):
        border = ET.SubElement(tbl_borders, _w(name))
        border.set(_w("val"), "single")
        border.set(_w("sz"), "18")
        border.set(_w("color"), "000000")

    for current_number in (eq_number, *additional_numbers):
        tr = ET.SubElement(tbl, _w("tr"))
        ET.SubElement(ET.SubElement(tr, _w("tc")), _w("p"))

        math_tc = ET.SubElement(tr, _w("tc"))
        math_p = ET.SubElement(math_tc, _w("p"))
        omath_para = ET.SubElement(math_p, _m("oMathPara"))
        omath = ET.SubElement(omath_para, _m("oMath"))
        mr = ET.SubElement(omath, _m("r"))
        mt = ET.SubElement(mr, _m("t"))
        mt.text = "E=mc2"

        num_tc = ET.SubElement(tr, _w("tc"))
        num_p = ET.SubElement(num_tc, _w("p"))
        run = ET.SubElement(num_p, _w("r"))
        text = ET.SubElement(run, _w("t"))
        text.text = current_number

    body.insert(list(body).index(sect_pr), tbl)
    parts["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)

    with zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, payload in parts.items():
            zf.writestr(name, payload)


def _inject_text_equation_layout_table(
    docx_path: Path,
    *,
    eq_number: str = "(1.1)",
    formula_text: str = "样品得率（%）=(m2-m1)/m0×100",
) -> None:
    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    def _w(tag: str) -> str:
        return f"{{{w_ns}}}{tag}"

    with zipfile.ZipFile(docx_path, "r") as zf:
        parts = {name: zf.read(name) for name in zf.namelist()}

    root = ET.fromstring(parts["word/document.xml"])
    body = root.find(_w("body"))
    assert body is not None
    sect_pr = body.find(_w("sectPr"))
    assert sect_pr is not None

    tbl = ET.Element(_w("tbl"))
    tbl_pr = ET.SubElement(tbl, _w("tblPr"))
    tbl_borders = ET.SubElement(tbl_pr, _w("tblBorders"))
    for name in ("top", "bottom", "left", "right", "insideH", "insideV"):
        border = ET.SubElement(tbl_borders, _w(name))
        border.set(_w("val"), "single" if name in {"top", "bottom"} else "none")
        border.set(_w("sz"), "18" if name in {"top", "bottom"} else "0")
        border.set(_w("color"), "000000" if name in {"top", "bottom"} else "auto")
        border.set(_w("space"), "0")

    tr = ET.SubElement(tbl, _w("tr"))
    for cell_text in ("", formula_text, eq_number):
        tc = ET.SubElement(tr, _w("tc"))
        tc_pr = ET.SubElement(tc, _w("tcPr"))
        tc_borders = ET.SubElement(tc_pr, _w("tcBorders"))
        for name in ("top", "left", "bottom", "right"):
            border = ET.SubElement(tc_borders, _w(name))
            border.set(_w("val"), "single" if name == "bottom" else "nil")
            border.set(_w("sz"), "6" if name == "bottom" else "0")
            border.set(_w("color"), "000000" if name == "bottom" else "auto")
            border.set(_w("space"), "0")
        p = ET.SubElement(tc, _w("p"))
        if cell_text:
            r = ET.SubElement(p, _w("r"))
            t = ET.SubElement(r, _w("t"))
            t.text = cell_text

    body.insert(list(body).index(sect_pr), tbl)
    parts["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)

    with zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, payload in parts.items():
            zf.writestr(name, payload)


def _inject_separate_equation_number_paragraph(docx_path: Path, eq_number: str = "（1.1）") -> None:
    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    m_ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"

    def _w(tag: str) -> str:
        return f"{{{w_ns}}}{tag}"

    def _m(tag: str) -> str:
        return f"{{{m_ns}}}{tag}"

    with zipfile.ZipFile(docx_path, "r") as zf:
        parts = {name: zf.read(name) for name in zf.namelist()}

    root = ET.fromstring(parts["word/document.xml"])
    body = root.find(_w("body"))
    assert body is not None
    sect_pr = body.find(_w("sectPr"))
    assert sect_pr is not None

    formula_p = ET.Element(_w("p"))
    formula_p_pr = ET.SubElement(formula_p, _w("pPr"))
    formula_jc = ET.SubElement(formula_p_pr, _w("jc"))
    formula_jc.set(_w("val"), "center")
    formula_run = ET.SubElement(formula_p, _w("r"))
    omath = ET.SubElement(formula_run, _m("oMath"))
    math_run = ET.SubElement(omath, _w("r"))
    math_text = ET.SubElement(math_run, _w("t"))
    math_text.text = "x"

    number_p = ET.Element(_w("p"))
    number_p_pr = ET.SubElement(number_p, _w("pPr"))
    number_jc = ET.SubElement(number_p_pr, _w("jc"))
    number_jc.set(_w("val"), "center")
    number_run = ET.SubElement(number_p, _w("r"))
    number_text = ET.SubElement(number_run, _w("t"))
    number_text.text = eq_number

    insert_at = list(body).index(sect_pr)
    body.insert(insert_at, formula_p)
    body.insert(insert_at + 1, number_p)
    parts["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)

    with zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, payload in parts.items():
            zf.writestr(name, payload)


def _inject_inline_math_into_paragraph(docx_path: Path, paragraph_text: str) -> None:
    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    m_ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    ns = {"w": w_ns, "m": m_ns}

    def _w(tag: str) -> str:
        return f"{{{w_ns}}}{tag}"

    def _m(tag: str) -> str:
        return f"{{{m_ns}}}{tag}"

    with zipfile.ZipFile(docx_path, "r") as zf:
        parts = {name: zf.read(name) for name in zf.namelist()}

    root = ET.fromstring(parts["word/document.xml"])
    target = None
    for p_elem in root.findall(".//w:p", ns):
        texts = [t.text or "" for t in p_elem.findall(".//w:t", ns)]
        if "".join(texts) == paragraph_text:
            target = p_elem
            break
    assert target is not None

    run = ET.SubElement(target, _w("r"))
    omath = ET.SubElement(run, _m("oMath"))
    mr = ET.SubElement(omath, _m("r"))
    mt = ET.SubElement(mr, _m("t"))
    mt.text = "m"

    parts["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, payload in parts.items():
            zf.writestr(name, payload)


def _spacing_attrs(paragraph):
    spacing = paragraph._p.pPr.find(qn("w:spacing")) if paragraph._p.pPr is not None else None
    if spacing is None:
        return None, None, None
    return spacing.get(qn("w:before")), spacing.get(qn("w:after")), spacing.get(qn("w:line"))


def _inject_existing_footer_page_field(docx_path: Path):
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    office_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    ns = {"w": w_ns}

    with zipfile.ZipFile(docx_path, "r") as zf:
        parts = {name: zf.read(name) for name in zf.namelist()}

    rels_root = ET.fromstring(parts["word/_rels/document.xml.rels"])
    rel = ET.SubElement(rels_root, f"{{{rel_ns}}}Relationship")
    rel.set("Id", "rIdFooterTest")
    rel.set("Type", f"{office_rel_ns}/footer")
    rel.set("Target", "footer1.xml")
    parts["word/_rels/document.xml.rels"] = ET.tostring(rels_root, encoding="utf-8", xml_declaration=True)

    document_root = ET.fromstring(parts["word/document.xml"])
    sect_pr = document_root.find("w:body/w:sectPr", ns)
    assert sect_pr is not None
    footer_ref = ET.SubElement(sect_pr, f"{{{w_ns}}}footerReference")
    footer_ref.set(f"{{{w_ns}}}type", "default")
    footer_ref.set(f"{{{office_rel_ns}}}id", "rIdFooterTest")
    parts["word/document.xml"] = ET.tostring(document_root, encoding="utf-8", xml_declaration=True)

    footer_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:ftr xmlns:w="{w_ns}">
  <w:p>
    <w:pPr/>
    <w:r><w:fldChar w:fldCharType="begin"/></w:r>
    <w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>
    <w:r><w:fldChar w:fldCharType="separate"/></w:r>
    <w:r><w:t>1</w:t></w:r>
    <w:r><w:fldChar w:fldCharType="end"/></w:r>
  </w:p>
</w:ftr>"""
    parts["word/footer1.xml"] = footer_xml.encode("utf-8")

    content_types_root = ET.fromstring(parts["[Content_Types].xml"])
    override = ET.SubElement(content_types_root, f"{{{ct_ns}}}Override")
    override.set("PartName", "/word/footer1.xml")
    override.set("ContentType", "application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml")
    parts["[Content_Types].xml"] = ET.tostring(content_types_root, encoding="utf-8", xml_declaration=True)

    with zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, payload in parts.items():
            zf.writestr(name, payload)


def _inject_hidden_page_field_body_paragraph(docx_path: Path, display_text: str = "-"):
    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    ns = {"w": w_ns}

    def _w(tag: str) -> str:
        return f"{{{w_ns}}}{tag}"

    with zipfile.ZipFile(docx_path, "r") as zf:
        parts = {name: zf.read(name) for name in zf.namelist()}

    root = ET.fromstring(parts["word/document.xml"])
    body = root.find("w:body", ns)
    assert body is not None
    sect_pr = body.find(_w("sectPr"))
    assert sect_pr is not None

    paragraph = ET.Element(_w("p"))
    p_pr = ET.SubElement(paragraph, _w("pPr"))
    spacing = ET.SubElement(p_pr, _w("spacing"))
    spacing.set(_w("line"), "360")
    spacing.set(_w("lineRule"), "auto")

    def _hidden_run():
        run = ET.SubElement(paragraph, _w("r"))
        r_pr = ET.SubElement(run, _w("rPr"))
        ET.SubElement(r_pr, _w("vanish")).set(_w("val"), "1")
        return run

    fld_begin = ET.SubElement(_hidden_run(), _w("fldChar"))
    fld_begin.set(_w("fldCharType"), "begin")
    instr = ET.SubElement(_hidden_run(), _w("instrText"))
    instr.text = " PAGE "
    fld_end = ET.SubElement(_hidden_run(), _w("fldChar"))
    fld_end.set(_w("fldCharType"), "end")
    ET.SubElement(_hidden_run(), _w("t")).text = display_text

    body.insert(list(body).index(sect_pr), paragraph)
    parts["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)

    with zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, payload in parts.items():
            zf.writestr(name, payload)


def test_scoped_fix_dry_run_returns_preview_without_writing_output(tmp_docx, tmp_path):
    source_path = tmp_docx(make_compliant_doc, filename="scoped_fix_preview_source.docx")
    preview_path = Path(tmp_path) / "scoped_fix_preview_output.docx"

    preview = build_scoped_fix_preview(
        str(source_path),
        output_path=str(preview_path),
        profile_path="cn-common",
        scopes=["headings"],
    )

    assert preview["selected_scopes"] == ["headings"]
    assert preview["renumber_headings"] is False
    assert not preview_path.exists()

    rendered = render_scoped_fix_preview(preview)
    assert "dry-run: 是" in rendered
    assert "修复范围: headings" in rendered
    assert "预计会触达的模块" in rendered


def test_cleanup_docx_hidden_page_number_artifacts_removes_body_page_field_paragraph(tmp_docx):
    source_path = tmp_docx(make_compliant_doc, filename="cleanup_hidden_page_field_source.docx")
    _inject_hidden_page_field_body_paragraph(source_path, "-")

    removed = fix_thesis._cleanup_docx_hidden_page_number_artifacts(str(source_path))

    assert removed == 1
    doc = Document(source_path)
    texts = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
    assert "-" not in texts


def test_fix_docx_keeps_heading_numbers_when_renumber_not_enabled(tmp_docx, tmp_path):
    source_path = tmp_docx(make_compliant_doc, filename="renumber_default_source.docx")
    doc = Document(source_path)
    _paragraph_by_prefix(doc, "第一章").runs[0].text = "0 引言"
    _paragraph_by_prefix(doc, "1.1 研究背景").runs[0].text = "0.1 研究背景"
    _paragraph_by_prefix(doc, "1.1.1 研究现状").runs[0].text = "0.1.1 研究现状"
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "renumber_default_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="cn-common",
        scopes=["headings"],
    )

    fixed_doc = Document(fixed_path)
    assert _paragraph_by_prefix(fixed_doc, "0 引言").text == "0 引言"
    assert _paragraph_by_prefix(fixed_doc, "0.1 研究背景").text == "0.1 研究背景"
    assert _paragraph_by_prefix(fixed_doc, "0.1.1 研究现状").text == "0.1.1 研究现状"


def test_fix_docx_can_renumber_headings_when_opted_in(tmp_docx, tmp_path):
    source_path = tmp_docx(make_compliant_doc, filename="renumber_optin_source.docx")
    doc = Document(source_path)
    _paragraph_by_prefix(doc, "第一章").runs[0].text = "0 引言"
    _paragraph_by_prefix(doc, "1.1 研究背景").runs[0].text = "0.1 研究背景"
    _paragraph_by_prefix(doc, "1.1.1 研究现状").runs[0].text = "0.1.1 研究现状"
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "renumber_optin_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="cn-common",
        scopes=["headings"],
        renumber_headings=True,
    )

    fixed_doc = Document(fixed_path)
    assert _paragraph_by_prefix(fixed_doc, "1 引言").text == "1 引言"
    assert _paragraph_by_prefix(fixed_doc, "1.1 研究背景").text == "1.1 研究背景"
    assert _paragraph_by_prefix(fixed_doc, "1.1.1 研究现状").text == "1.1.1 研究现状"


def test_lnu_headings_scope_formats_h1_as_di_chapter(tmp_path):
    source_path = Path(tmp_path) / "lnu_h1_chapter_format_source.docx"
    doc = Document()
    heading = doc.add_paragraph("1 材料与方法")
    heading.style = doc.styles["Heading 1"]
    doc.add_paragraph("这是正文内容。")
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "lnu_h1_chapter_format_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["headings"],
    )

    fixed_doc = Document(fixed_path)
    assert _paragraph_by_prefix(fixed_doc, "第1章 材料与方法").text == "第1章 材料与方法"


def test_lnu_heading_fix_preserves_split_run_text_across_repeated_apply(tmp_path):
    source_path = Path(tmp_path) / "split_run_heading_source.docx"
    first_path = Path(tmp_path) / "split_run_heading_first.docx"
    second_path = Path(tmp_path) / "split_run_heading_second.docx"
    expected = "0.1 鹿茸的研究进展"

    doc = Document()
    preface = doc.add_paragraph("序  言")
    preface.style = doc.styles["Heading 1"]
    heading = doc.add_paragraph()
    heading.style = doc.styles["Heading 2"]
    heading.add_run("0.1 鹿茸的")
    heading.add_run("研究进展")
    doc.save(source_path)

    for input_path, output_path in ((source_path, first_path), (first_path, second_path)):
        apply_scoped_fix(
            str(input_path),
            str(output_path),
            profile_path="lnu",
            scopes=["headings"],
        )

    actual = [
        next(paragraph.text for paragraph in Document(path).paragraphs if paragraph.text.startswith("0.1"))
        for path in (first_path, second_path)
    ]
    assert actual == [expected, expected]


def test_lnu_figures_scope_uses_zero_chapter_caption_number_for_preface(tmp_path):
    source_path = Path(tmp_path) / "lnu_preface_caption_zero_source.docx"
    doc = Document()
    preface = doc.add_paragraph("序  言")
    preface.style = doc.styles["Heading 1"]
    figure = doc.add_paragraph()
    _add_mock_drawing(figure)
    caption = doc.add_paragraph("图1.1 绪论框架图")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("第1章 材料与方法").style = doc.styles["Heading 1"]
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "lnu_preface_caption_zero_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["figures_tables", "headings"],
    )

    fixed_doc = Document(fixed_path)
    assert _paragraph_by_normalized_prefix(fixed_doc, "图0.1 绪论框架图").text.startswith("图0.1")


def test_heading_style_prepass_identifies_and_repairs_heading_like_paragraphs(tmp_docx, tmp_path):
    source_path = tmp_docx(make_compliant_doc, filename="heading_prepass_source.docx")
    doc = Document(source_path)
    heading = _paragraph_by_prefix(doc, "1.1 研究背景")
    heading.style = doc.styles["Normal"]
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(source_path)

    preview_before = build_scoped_fix_preview(
        str(source_path),
        profile_path="cn-common",
        scopes=["headings"],
    )
    assert preview_before["heading_style_candidates"] >= 1

    fixed_path = Path(tmp_path) / "heading_prepass_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="cn-common",
        scopes=["headings"],
    )

    h02_status = audit_rule_status(fixed_path, "H02", profile_path="cn-common")
    assert h02_status["rule"]["passed"], h02_status["rule"]["issues"]

    preview_after = build_scoped_fix_preview(
        str(fixed_path),
        profile_path="cn-common",
        scopes=["headings"],
    )
    assert preview_after["heading_style_candidates"] == 0


def test_body_scope_does_not_promote_heading_like_paragraphs_via_heading_prepass(tmp_docx, tmp_path):
    source_path = tmp_docx(make_compliant_doc, filename="body_scope_no_heading_prepass_source.docx")
    doc = Document(source_path)
    heading = _paragraph_by_prefix(doc, "1.1 研究背景")
    heading.style = doc.styles["Normal"]
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "body_scope_no_heading_prepass_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="cn-common",
        scopes=["body_paragraphs"],
    )

    fixed_doc = Document(fixed_path)
    fixed_heading = _paragraph_by_prefix(fixed_doc, "1.1 研究背景")
    h02_status = audit_rule_status(fixed_path, "H02")

    assert fixed_heading.style.name == "Normal"
    assert not h02_status["rule"]["passed"]


def test_body_scope_demotes_misstyled_heading_paragraph_to_body_style(tmp_path):
    source_path = Path(tmp_path) / "body_scope_demote_heading_style.docx"
    doc = Document()
    heading = doc.add_paragraph("1 绪论")
    heading.style = doc.styles["Heading 1"]
    paragraph = doc.add_paragraph("这是一个被误套用标题样式的正文句子，包含完整语义和句号。")
    paragraph.style = doc.styles["Heading 1"]
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "body_scope_demote_heading_style_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["body_paragraphs"],
    )

    fixed_doc = Document(fixed_path)
    fixed_paragraph = fixed_doc.paragraphs[1]
    assert fixed_paragraph.style.name == "Normal"


def test_body_scope_clears_residual_page_break_before_on_demoted_body_paragraph(tmp_path):
    source_path = Path(tmp_path) / "body_scope_clear_page_break_source.docx"
    doc = Document()
    heading = doc.add_paragraph("1 绪论")
    heading.style = doc.styles["Heading 1"]
    paragraph = doc.add_paragraph("这是一个被误套用标题样式的正文句子。")
    paragraph.style = doc.styles["Heading 1"]
    paragraph.paragraph_format.page_break_before = True
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "body_scope_clear_page_break_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["body_paragraphs"],
    )

    fixed_doc = Document(fixed_path)
    fixed_paragraph = fixed_doc.paragraphs[1]
    assert fixed_paragraph.style.name == "Normal"
    assert fixed_paragraph.paragraph_format.page_break_before is not True


def test_body_scope_preserves_figure_note_paragraph_formatting(tmp_path):
    source_path = Path(tmp_path) / "body_scope_preserve_figure_note_source.docx"
    doc = Document()
    heading = doc.add_paragraph("2.1.1 溶胀特性分析")
    heading.style = doc.styles["Heading 3"]

    figure = doc.add_paragraph()
    figure.add_run("[mock figure]")

    caption = doc.add_paragraph("图2.1 不同改性条件筛选组样品的溶胀率与溶失率")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    note = doc.add_paragraph("注：（A）0.5～24 h 溶胀率变化曲线；（B）24 h 溶胀率。")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.first_line_indent = 0
    note.paragraph_format.space_before = 0
    note.paragraph_format.space_after = 0
    note.paragraph_format.line_spacing = 1.0

    body = doc.add_paragraph("这是正文段落。")
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "body_scope_preserve_figure_note_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["body_paragraphs"],
    )

    fixed_doc = Document(fixed_path)
    fixed_note = next(paragraph for paragraph in fixed_doc.paragraphs if paragraph.text.startswith("注：（A）0.5"))

    assert fixed_note.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert fixed_note.paragraph_format.first_line_indent in (0, None)
    spacing = fixed_note._p.pPr.find(qn("w:spacing"))
    assert spacing is not None
    assert spacing.get(qn("w:line")) in {"240", None}

    results, _score, _report = audit_thesis.audit_docx(str(fixed_path), profile_path="lnu")
    t03_status = next(item for item in results if item["id"] == "T03")
    t04_status = next(item for item in results if item["id"] == "T04")
    assert t03_status["passed"], t03_status["issues"]
    assert t04_status["passed"], t04_status["issues"]


def test_figures_scope_moves_post_figure_analysis_before_figure_block(tmp_path):
    source_path = Path(tmp_path) / "lnu_figure_reorder_source.docx"
    doc = Document()
    doc.add_paragraph("第2章 实验结果与分析").style = doc.styles["Heading 1"]
    doc.add_paragraph("2.1 不同改性技术鹿皮明胶改性条件的筛选").style = doc.styles["Heading 2"]
    doc.add_paragraph("2.1.1 溶胀特性分析").style = doc.styles["Heading 3"]

    figure = doc.add_paragraph()
    _add_mock_drawing(figure)

    caption = doc.add_paragraph("图2.1 不同改性条件筛选组样品的溶胀率与溶失率")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    note = doc.add_paragraph("注：（A）0.5～24 h 溶胀率变化曲线；（B）24 h 溶胀率；（C）24 h 溶失率。")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("不同改性条件对鹿皮明胶冻干凝胶溶胀行为具有明显影响，具体结果如图2.1 A、B 所示。")
    doc.add_paragraph("2.1.2 溶失特性分析").style = doc.styles["Heading 3"]
    doc.add_paragraph("各组样品 24 h 溶失率结果如图2.1 C 所示。")
    doc.add_paragraph("2.2 不同改性技术对鹿皮明胶功能特性的影响").style = doc.styles["Heading 2"]
    doc.save(source_path)

    before_results, _score, _report = audit_thesis.audit_docx(str(source_path), profile_path="lnu")
    before_f03 = next(item for item in before_results if item["id"] == "LNU_F03")
    before_f06 = next(item for item in before_results if item["id"] == "LNU_F06")
    assert not before_f03["passed"]
    assert not before_f06["passed"]

    fixed_path = Path(tmp_path) / "lnu_figure_reorder_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["figures_tables"],
    )

    fixed_doc = Document(fixed_path)
    paragraph_info = [
        (idx, paragraph.text.strip(), bool(paragraph._p.xpath(".//w:drawing")))
        for idx, paragraph in enumerate(fixed_doc.paragraphs)
    ]

    idx_analysis_1 = next(idx for idx, text, has_drawing in paragraph_info if text.startswith("不同改性条件对鹿皮明胶冻干凝胶溶胀行为"))
    idx_heading_212 = next(idx for idx, text, has_drawing in paragraph_info if text == "2.1.2 溶失特性分析")
    idx_analysis_2 = next(idx for idx, text, has_drawing in paragraph_info if text.startswith("各组样品 24 h 溶失率结果如图2.1 C 所示"))
    idx_drawing = next(idx for idx, text, has_drawing in paragraph_info if has_drawing)
    idx_caption = next(idx for idx, text, has_drawing in paragraph_info if text.startswith("图2.1 "))
    idx_note = next(idx for idx, text, has_drawing in paragraph_info if text.startswith("注"))
    idx_next_h2 = next(idx for idx, text, has_drawing in paragraph_info if text == "2.2 不同改性技术对鹿皮明胶功能特性的影响")

    assert idx_analysis_1 < idx_heading_212 < idx_analysis_2 < idx_drawing < idx_caption < idx_note < idx_next_h2
    assert not any(text.startswith("见图") or text.startswith("见表") for _, text, _ in paragraph_info)

    drawing_paragraph = fixed_doc.paragraphs[idx_drawing]
    caption_paragraph = fixed_doc.paragraphs[idx_caption]
    note_paragraph = fixed_doc.paragraphs[idx_note]
    drawing_before, drawing_after, drawing_line = _spacing_attrs(drawing_paragraph)
    caption_before, caption_after, caption_line = _spacing_attrs(caption_paragraph)
    note_before, note_after, note_line = _spacing_attrs(note_paragraph)

    assert drawing_before == "360"
    assert drawing_after in {None, "0"}
    assert caption_line == "360"
    assert note_line == "240"
    assert note_after == "360"

    results, _score, _report = audit_thesis.audit_docx(str(fixed_path), profile_path="lnu")
    lnu_f03 = next(item for item in results if item["id"] == "LNU_F03")
    lnu_f06 = next(item for item in results if item["id"] == "LNU_F06")
    assert lnu_f03["passed"], lnu_f03["issues"]
    assert lnu_f06["passed"], lnu_f06["issues"]


def test_layout_rebalance_is_opt_in_and_does_not_change_default_flow(tmp_path):
    source_path = Path(tmp_path) / "layout_rebalance_opt_in_source.docx"
    doc = Document()
    doc.add_paragraph("第2章 实验结果与分析").style = doc.styles["Heading 1"]
    doc.add_paragraph("2.2 基因组组装结果").style = doc.styles["Heading 2"]
    doc.add_paragraph("组装结果显示连续性良好，图2.1 与表2.1共同支持该判断。")
    table_caption = doc.add_paragraph("表2.1 基因组组装统计")
    table_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "层次"
    table.cell(0, 1).text = "N50"
    table.cell(1, 0).text = "Contig"
    table.cell(1, 1).text = "522958"
    figure = doc.add_paragraph()
    _add_mock_drawing(figure)
    caption = doc.add_paragraph("图2.1 基因组组装评估结果图")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "layout_rebalance_opt_in_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["figures_tables"],
    )

    fixed_doc = Document(fixed_path)
    paragraph_info = [
        (idx, paragraph.text.strip(), bool(paragraph._p.xpath(".//w:drawing")))
        for idx, paragraph in enumerate(fixed_doc.paragraphs)
    ]
    idx_table_caption = next(idx for idx, text, has_drawing in paragraph_info if text.startswith("表2.1 "))
    idx_figure_caption = next(idx for idx, text, has_drawing in paragraph_info if text.startswith("图2.1 "))
    assert idx_table_caption < idx_figure_caption


def test_layout_rebalance_moves_figure_block_after_existing_reference(tmp_path):
    source_path = Path(tmp_path) / "layout_rebalance_enabled_source.docx"
    doc = Document()
    doc.add_paragraph("第2章 实验结果与分析").style = doc.styles["Heading 1"]
    doc.add_paragraph("2.2 基因组组装结果").style = doc.styles["Heading 2"]
    doc.add_paragraph("组装结果显示连续性良好，图2.1 与表2.1共同支持该判断。")
    table_caption = doc.add_paragraph("表2.1 基因组组装统计")
    table_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "层次"
    table.cell(0, 1).text = "N50"
    table.cell(1, 0).text = "Contig"
    table.cell(1, 1).text = "522958"
    figure = doc.add_paragraph()
    _add_mock_drawing(figure)
    caption = doc.add_paragraph("图2.1 基因组组装评估结果图")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "layout_rebalance_enabled_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["figures_tables"],
        layout_rebalance=True,
    )

    fixed_doc = Document(fixed_path)
    paragraph_info = [
        (idx, paragraph.text.strip(), bool(paragraph._p.xpath(".//w:drawing")))
        for idx, paragraph in enumerate(fixed_doc.paragraphs)
    ]
    idx_ref = next(idx for idx, text, has_drawing in paragraph_info if text.startswith("组装结果显示连续性良好"))
    idx_table_caption = next(idx for idx, text, has_drawing in paragraph_info if text.startswith("表2.1 "))
    idx_drawing = next(idx for idx, text, has_drawing in paragraph_info if has_drawing)
    idx_figure_caption = next(idx for idx, text, has_drawing in paragraph_info if text.startswith("图2.1 "))

    assert idx_ref < idx_drawing < idx_figure_caption < idx_table_caption
    assert not any(text.startswith("相关结果如图") for _, text, _ in paragraph_info if text)


def test_layout_rebalance_prefers_earliest_reference_in_same_subsection(tmp_path):
    source_path = Path(tmp_path) / "layout_rebalance_earliest_anchor_source.docx"
    doc = Document()
    doc.add_paragraph("第2章 实验结果与分析").style = doc.styles["Heading 1"]
    doc.add_paragraph("2.1 结果分析").style = doc.styles["Heading 2"]
    doc.add_paragraph("不同改性条件结果如图2.1A、B所示。")
    doc.add_paragraph("进一步地，24 h结果如图2.1C所示。")
    figure = doc.add_paragraph()
    _add_mock_drawing(figure)
    caption = doc.add_paragraph("图2.1 不同改性条件筛选组样品的溶胀率与溶失率")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note = doc.add_paragraph("注：（A）0.5～24 h 溶胀率变化曲线；（B）24 h 溶胀率；（C）24 h 溶失率。")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "layout_rebalance_earliest_anchor_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["figures_tables"],
        layout_rebalance=True,
    )

    fixed_doc = Document(fixed_path)
    paragraph_info = [
        (idx, paragraph.text.strip(), bool(paragraph._p.xpath(".//w:drawing")))
        for idx, paragraph in enumerate(fixed_doc.paragraphs)
    ]
    idx_first_ref = next(idx for idx, text, has_drawing in paragraph_info if text.startswith("不同改性条件结果如图2.1A"))
    idx_second_ref = next(idx for idx, text, has_drawing in paragraph_info if text.startswith("进一步地，24 h结果如图2.1C"))
    idx_drawing = next(idx for idx, text, has_drawing in paragraph_info if has_drawing)
    idx_caption = next(idx for idx, text, has_drawing in paragraph_info if text.startswith("图2.1 "))

    assert idx_first_ref < idx_drawing < idx_caption < idx_second_ref


def test_layout_rebalance_can_cross_h3_when_reference_is_earlier_in_same_h2(tmp_path):
    source_path = Path(tmp_path) / "layout_rebalance_cross_h3_source.docx"
    doc = Document()
    doc.add_paragraph("第2章 实验结果与分析").style = doc.styles["Heading 1"]
    doc.add_paragraph("2.1 结果分析").style = doc.styles["Heading 2"]
    doc.add_paragraph("2.1.2 溶胀特性分析").style = doc.styles["Heading 3"]
    doc.add_paragraph("具体结果如图2.1A、B所示。")
    doc.add_paragraph("2.1.3 溶失特性分析").style = doc.styles["Heading 3"]
    doc.add_paragraph("24 h结果如图2.1C所示。")
    figure = doc.add_paragraph()
    _add_mock_drawing(figure)
    caption = doc.add_paragraph("图2.1 不同改性条件筛选组样品的溶胀率与溶失率")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note = doc.add_paragraph("注：（A）0.5～24 h 溶胀率变化曲线；（B）24 h 溶胀率；（C）24 h 溶失率。")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "layout_rebalance_cross_h3_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["figures_tables"],
        layout_rebalance=True,
    )

    fixed_doc = Document(fixed_path)
    paragraph_info = [
        (idx, paragraph.text.strip(), bool(paragraph._p.xpath(".//w:drawing")))
        for idx, paragraph in enumerate(fixed_doc.paragraphs)
    ]
    idx_first_ref = next(idx for idx, text, has_drawing in paragraph_info if text.startswith("具体结果如图2.1A"))
    idx_h3 = next(idx for idx, text, has_drawing in paragraph_info if text.startswith("2.1.3 溶失特性分析"))
    idx_drawing = next(idx for idx, text, has_drawing in paragraph_info if has_drawing)
    idx_caption = next(idx for idx, text, has_drawing in paragraph_info if text.startswith("图2.1 "))

    assert idx_first_ref < idx_drawing < idx_caption < idx_h3


def test_layout_rebalance_does_not_insert_synthetic_lead_when_spaced_reference_already_exists(tmp_path):
    source_path = Path(tmp_path) / "layout_rebalance_spaced_reference_source.docx"
    doc = Document()
    doc.add_paragraph("第2章 实验结果与分析").style = doc.styles["Heading 1"]
    doc.add_paragraph("2.1 结果分析").style = doc.styles["Heading 2"]
    doc.add_paragraph("2.1.2 溶胀特性分析").style = doc.styles["Heading 3"]
    doc.add_paragraph("具体结果如图 2.1 A、B 所示。")
    figure = doc.add_paragraph()
    _add_mock_drawing(figure)
    caption = doc.add_paragraph("图2.1 不同改性条件筛选组样品的溶胀率与溶失率")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note = doc.add_paragraph("注：（A）0.5～24 h 溶胀率变化曲线；（B）24 h 溶胀率。")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "layout_rebalance_spaced_reference_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["figures_tables"],
        layout_rebalance=True,
    )

    fixed_doc = Document(fixed_path)
    paragraph_texts = [paragraph.text.strip() for paragraph in fixed_doc.paragraphs if paragraph.text.strip()]
    assert "相关结果如图2.1所示。" not in paragraph_texts


def test_layout_rebalance_does_not_touch_non_figure_scopes(tmp_path):
    source_path = Path(tmp_path) / "layout_rebalance_non_figure_scope_source.docx"
    doc = Document()
    doc.add_paragraph("第2章 实验结果与分析").style = doc.styles["Heading 1"]
    doc.add_paragraph("2.2 基因组组装结果").style = doc.styles["Heading 2"]
    doc.add_paragraph("组装结果显示连续性良好，图2.1 与表2.1共同支持该判断。")
    table_caption = doc.add_paragraph("表2.1 基因组组装统计")
    table_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "层次"
    table.cell(0, 1).text = "N50"
    table.cell(1, 0).text = "Contig"
    table.cell(1, 1).text = "522958"
    figure = doc.add_paragraph()
    _add_mock_drawing(figure)
    caption = doc.add_paragraph("图2.1 基因组组装评估结果图")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "layout_rebalance_non_figure_scope_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["headings"],
        layout_rebalance=True,
    )

    fixed_doc = Document(fixed_path)
    paragraph_texts = [paragraph.text.strip() for paragraph in fixed_doc.paragraphs if paragraph.text.strip()]
    assert paragraph_texts[-2:] == ["表2.1 基因组组装统计", "图2.1 基因组组装评估结果图"]


def test_layout_rebalance_can_move_table_block_to_reference_anchor(tmp_path):
    source_path = Path(tmp_path) / "layout_rebalance_table_anchor_source.docx"
    doc = Document()
    doc.add_paragraph("第2章 实验结果与分析").style = doc.styles["Heading 1"]
    doc.add_paragraph("2.3 基因组组分分析").style = doc.styles["Heading 2"]
    doc.add_paragraph("相关统计见表2.1。")
    figure = doc.add_paragraph()
    _add_mock_drawing(figure)
    figure_caption = doc.add_paragraph("图2.1 基因组组分分析结果图")
    figure_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_caption = doc.add_paragraph("表2.1 基因组组分统计")
    table_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "类型"
    table.cell(0, 1).text = "数量"
    table.cell(1, 0).text = "CDS"
    table.cell(1, 1).text = "7118"
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "layout_rebalance_table_anchor_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["figures_tables"],
        layout_rebalance=True,
    )

    fixed_doc = Document(fixed_path)
    paragraph_info = [
        (idx, paragraph.text.strip(), bool(paragraph._p.xpath(".//w:drawing")))
        for idx, paragraph in enumerate(fixed_doc.paragraphs)
    ]
    idx_ref = next(idx for idx, text, has_drawing in paragraph_info if text.startswith("相关统计见表2.1"))
    idx_table_caption = next(idx for idx, text, has_drawing in paragraph_info if text.startswith("表2.1 "))
    idx_figure_caption = next(idx for idx, text, has_drawing in paragraph_info if text.startswith("图2.1 "))

    assert idx_ref < idx_table_caption < idx_figure_caption


def test_layout_rebalance_does_not_insert_lead_when_only_later_section_mentions_table(tmp_path):
    source_path = Path(tmp_path) / "layout_rebalance_later_reference_source.docx"
    doc = Document()
    doc.add_paragraph("第2章 实验结果与分析").style = doc.styles["Heading 1"]
    doc.add_paragraph("2.1 基因组组装结果").style = doc.styles["Heading 2"]
    doc.add_paragraph("本节先给出表格。")
    table_caption = doc.add_paragraph("表2.1 基因组组装统计")
    table_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "层次"
    table.cell(0, 1).text = "N50"
    table.cell(1, 0).text = "Contig"
    table.cell(1, 1).text = "522958"
    doc.add_paragraph("2.2 讨论").style = doc.styles["Heading 2"]
    doc.add_paragraph("后文将结合表2.1进一步讨论。")
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "layout_rebalance_later_reference_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["figures_tables"],
        layout_rebalance=True,
    )

    fixed_doc = Document(fixed_path)
    paragraph_texts = [paragraph.text.strip() for paragraph in fixed_doc.paragraphs if paragraph.text.strip()]
    idx_caption = next(i for i, text in enumerate(paragraph_texts) if _normalize_spaces(text) == "表2.1 基因组组装统计")
    idx_heading = paragraph_texts.index("2.2 讨论")

    assert "相关结果如表2.1所示。" not in paragraph_texts
    assert idx_caption < idx_heading


def test_figures_scope_removes_existing_synthetic_caption_reference_lead(tmp_path):
    source_path = Path(tmp_path) / "synthetic_caption_reference_lead_source.docx"
    doc = Document()
    doc.add_paragraph("第2章 实验结果与分析").style = doc.styles["Heading 1"]
    doc.add_paragraph("2.1 基因组组装结果").style = doc.styles["Heading 2"]
    doc.add_paragraph("相关结果如表2.1所示。")
    table_caption = doc.add_paragraph("表2.1 基因组组装统计")
    table_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "层次"
    table.cell(0, 1).text = "N50"
    table.cell(1, 0).text = "Contig"
    table.cell(1, 1).text = "522958"
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "synthetic_caption_reference_lead_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["figures_tables"],
    )

    fixed_doc = Document(fixed_path)
    paragraph_texts = [paragraph.text.strip() for paragraph in fixed_doc.paragraphs if paragraph.text.strip()]

    assert "相关结果如表2.1所示。" not in paragraph_texts
    assert any(_normalize_spaces(text) == "表2.1 基因组组装统计" for text in paragraph_texts)


def test_figures_scope_clears_borders_on_equation_layout_tables(tmp_path):
    source_path = Path(tmp_path) / "equation_layout_table_source.docx"
    doc = Document()
    doc.add_paragraph("第1章 绪论").style = doc.styles["Heading 1"]
    doc.add_paragraph("其中公式如下。")
    doc.save(source_path)
    _inject_equation_layout_table(source_path, "(1.1)")

    fixed_path = Path(tmp_path) / "equation_layout_table_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["figures_tables"],
    )

    with zipfile.ZipFile(fixed_path, "r") as zf:
        root = ET.fromstring(zf.read("word/document.xml"))

    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    equation_tables = [
        tbl
        for tbl in root.findall(".//w:tbl", ns)
        if audit_thesis.is_equation_layout_table(tbl)
    ]
    assert len(equation_tables) == 1

    tbl_borders = equation_tables[0].find("w:tblPr/w:tblBorders", ns)
    assert tbl_borders is not None
    for border_name in ("top", "bottom", "left", "right", "insideH", "insideV"):
        border = tbl_borders.find(f"w:{border_name}", ns)
        assert border is not None
        assert border.get(qn("w:val")) == "nil"
        assert border.get(qn("w:sz")) == "0"


def test_figures_scope_clears_borders_on_grouped_equation_layout_table(tmp_path):
    source_path = Path(tmp_path) / "grouped_equation_layout_table_source.docx"
    doc = Document()
    doc.add_paragraph("第1章 绪论").style = doc.styles["Heading 1"]
    doc.add_paragraph("其中两条公式如下。")
    doc.save(source_path)
    _inject_equation_layout_table(
        source_path,
        "（1.3）",
        additional_numbers=("（1.4）",),
    )

    fixed_path = Path(tmp_path) / "grouped_equation_layout_table_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["figures_tables"],
    )

    with zipfile.ZipFile(fixed_path, "r") as zf:
        root = ET.fromstring(zf.read("word/document.xml"))

    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    table = root.find(".//w:tbl", ns)
    assert table is not None
    assert len(table.findall("w:tr", ns)) == 2
    borders = table.find("w:tblPr/w:tblBorders", ns)
    assert borders is not None
    for border_name in ("top", "bottom", "left", "right", "insideH", "insideV"):
        border = borders.find(f"w:{border_name}", ns)
        assert border is not None
        assert border.get(qn("w:val")) == "nil"
        assert border.get(qn("w:sz")) == "0"


def test_body_scope_right_aligns_separate_equation_number_paragraphs(tmp_path):
    source_path = Path(tmp_path) / "separate_equation_number_source.docx"
    doc = Document()
    doc.add_paragraph("第1章 绪论").style = doc.styles["Heading 1"]
    doc.add_paragraph("样品溶胀率按式（1.1）计算：")
    doc.save(source_path)
    _inject_separate_equation_number_paragraph(source_path, "（1.1）")

    assert not audit_rule_status(source_path, "EQ02", profile_path="lnu")["rule"]["passed"]

    fixed_path = Path(tmp_path) / "separate_equation_number_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["body_paragraphs"],
    )

    assert audit_rule_status(fixed_path, "EQ02", profile_path="lnu")["rule"]["passed"]

    with zipfile.ZipFile(fixed_path, "r") as zf:
        root = ET.fromstring(zf.read("word/document.xml"))

    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    number_paragraph = next(
        p for p in root.findall(".//w:p", ns)
        if "".join(t.text or "" for t in p.findall(".//w:t", ns)) == "（1.1）"
    )
    assert number_paragraph.find("w:pPr/w:jc", ns).get(qn("w:val")) == "right"


def test_body_scope_fixes_equation_explanation_variable_subscripts(tmp_path):
    source_path = Path(tmp_path) / "equation_explanation_source.docx"
    doc = Document()
    doc.add_paragraph("第1章 绪论").style = doc.styles["Heading 1"]
    doc.add_paragraph("式中，W0为样品初始干质量，Wt为样品在浸泡t时刻质量。")
    doc.save(source_path)

    assert not audit_rule_status(source_path, "LNU_EQ05", profile_path="lnu")["rule"]["passed"]

    fixed_path = Path(tmp_path) / "equation_explanation_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["body_paragraphs"],
    )

    assert audit_rule_status(fixed_path, "LNU_EQ05", profile_path="lnu")["rule"]["passed"]


def test_figures_scope_clears_borders_on_text_equation_layout_tables(tmp_path):
    source_path = Path(tmp_path) / "text_equation_layout_table_source.docx"
    doc = Document()
    doc.add_paragraph("第1章 绪论").style = doc.styles["Heading 1"]
    doc.add_paragraph("筛选阶段样品得率按式（1.1）计算：")
    doc.save(source_path)
    _inject_text_equation_layout_table(source_path, eq_number="(1.1)")

    fixed_path = Path(tmp_path) / "text_equation_layout_table_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["figures_tables"],
    )

    with zipfile.ZipFile(fixed_path, "r") as zf:
        root = ET.fromstring(zf.read("word/document.xml"))

    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    equation_tables = [
        tbl
        for tbl in root.findall(".//w:tbl", ns)
        if audit_thesis.is_equation_layout_table(tbl)
    ]
    assert len(equation_tables) == 1

    tbl_borders = equation_tables[0].find("w:tblPr/w:tblBorders", ns)
    assert tbl_borders is not None
    for border_name in ("top", "bottom", "left", "right", "insideH", "insideV"):
        border = tbl_borders.find(f"w:{border_name}", ns)
        assert border is not None
        assert border.get(qn("w:val")) == "nil"
        assert border.get(qn("w:sz")) == "0"

    for tc_borders in equation_tables[0].findall(".//w:tcPr/w:tcBorders", ns):
        for border_name in ("top", "bottom", "left", "right"):
            border = tc_borders.find(f"w:{border_name}", ns)
            assert border is not None
            assert border.get(qn("w:val")) == "nil"
            assert border.get(qn("w:sz")) == "0"


def test_figures_scope_converts_existing_blank_separator_to_structured_spacing(tmp_path):
    source_path = Path(tmp_path) / "lnu_figure_blank_separator_source.docx"
    doc = Document()
    doc.add_paragraph("第2章 实验结果与分析").style = doc.styles["Heading 1"]
    doc.add_paragraph("2.1.1 溶胀特性分析").style = doc.styles["Heading 3"]
    doc.add_paragraph("")
    figure = doc.add_paragraph()
    _add_mock_drawing(figure)
    caption = doc.add_paragraph("图2.1 不同改性条件筛选组样品的溶胀率与溶失率")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note = doc.add_paragraph("注：（A）0.5～24 h 溶胀率变化曲线；（B）24 h 溶胀率。")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")
    doc.add_paragraph("下文正文从这里开始。")
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "lnu_figure_blank_separator_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["figures_tables"],
    )

    fixed_doc = Document(fixed_path)
    paragraph_info = [
        (idx, paragraph.text.strip(), bool(paragraph._p.xpath(".//w:drawing")))
        for idx, paragraph in enumerate(fixed_doc.paragraphs)
    ]
    idx_drawing = next(idx for idx, text, has_drawing in paragraph_info if has_drawing)
    idx_note = next(idx for idx, text, has_drawing in paragraph_info if text.startswith("注"))
    drawing_before, _, _ = _spacing_attrs(fixed_doc.paragraphs[idx_drawing])
    _, note_after, _ = _spacing_attrs(fixed_doc.paragraphs[idx_note])

    assert drawing_before == "360"
    assert note_after == "360"


def test_figures_scope_normalizes_table_block_spacing_and_caption_attachment(tmp_path):
    source_path = Path(tmp_path) / "lnu_table_block_spacing_source.docx"
    doc = Document()
    doc.add_paragraph("第2章 实验结果与分析").style = doc.styles["Heading 1"]
    doc.add_paragraph("2.1 基因组组装结果").style = doc.styles["Heading 2"]
    doc.add_paragraph("组装统计结果如下所示。")

    caption = doc.add_paragraph("表2.1 基因组组装统计")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = 6
    doc.add_paragraph("")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "层次"
    table.cell(0, 1).text = "N50"
    table.cell(1, 0).text = "Contig"
    table.cell(1, 1).text = "522958"

    doc.add_paragraph("2.2 基因组组分分析").style = doc.styles["Heading 2"]
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "lnu_table_block_spacing_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["figures_tables"],
    )

    fixed_doc = Document(fixed_path)
    fixed_caption = _paragraph_by_normalized_prefix(fixed_doc, "表2.1 基因组组装统计")
    fixed_next_heading = _paragraph_by_prefix(fixed_doc, "2.2 基因组组分分析")
    caption_before, caption_after, _ = _spacing_attrs(fixed_caption)
    next_before, _, _ = _spacing_attrs(fixed_next_heading)
    cfg = _profile_cfg("lnu")

    body_children = list(fixed_doc.element.body.iterchildren())
    caption_idx = next(
        i
        for i, elem in enumerate(body_children)
        if elem.tag == qn("w:p") and _normalize_spaces("".join(elem.itertext())).startswith("表2.1 基因组组装统计")
    )

    assert body_children[caption_idx + 1].tag == qn("w:tbl")
    assert caption_before == "360"
    assert caption_after == "0"
    if cfg.get("table_blank_line_mode") == "blank_paragraph":
        assert body_children[caption_idx + 2].tag == qn("w:p")
        assert not "".join(body_children[caption_idx + 2].itertext()).strip()
        assert next_before in (None, "0")
    else:
        assert next_before == "360"

    results, _score, _report = audit_thesis.audit_docx(str(fixed_path), profile_path="lnu")
    lnu_tb04 = next(item for item in results if item["id"] == "LNU_TB04")
    assert lnu_tb04["passed"], lnu_tb04["issues"]


def test_body_scope_splits_inline_citation_in_protected_math_paragraph(tmp_path):
    source_path = Path(tmp_path) / "protected_math_citation_source.docx"
    paragraph_text = "该方法的计算过程参考相关文献[45]。"

    doc = Document()
    doc.add_paragraph("第2章 实验结果与分析").style = doc.styles["Heading 1"]
    doc.add_paragraph(paragraph_text)
    doc.save(source_path)
    _inject_inline_math_into_paragraph(source_path, paragraph_text)

    fixed_path = Path(tmp_path) / "protected_math_citation_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["body_paragraphs"],
    )

    fixed_doc = Document(fixed_path)
    paragraph = next(p for p in fixed_doc.paragraphs if "该方法的计算过程参考相关文献" in p.text)
    citation_runs = [run for run in paragraph.runs if run.text == "[45]"]
    assert len(citation_runs) == 1
    assert citation_runs[0].font.superscript is True

    results, _score, _report = audit_thesis.audit_docx(str(fixed_path), profile_path="lnu")
    c04 = next(item for item in results if item["id"] == "C04")
    assert c04["passed"], c04["issues"]


def test_scoped_fix_normalizes_inline_citation_groups_and_reference_layout(tmp_path):
    source_path = Path(tmp_path) / "citation_reference_layout_source.docx"
    doc = Document()
    doc.add_paragraph("第1章 绪论").style = doc.styles["Heading 1"]
    doc.add_paragraph("综述显示。[1][2]连续研究[1,2,3]表明。")
    doc.add_paragraph("参考文献").style = doc.styles["Heading 1"]
    doc.add_paragraph("[1] Some reference text.")
    doc.add_paragraph("[2] Another reference text.")
    doc.add_paragraph("[3] Third reference text.")
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "citation_reference_layout_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["body_paragraphs", "references"],
    )

    fixed_doc = Document(fixed_path)
    body_paragraph = next(p for p in fixed_doc.paragraphs if p.text.startswith("综述显示"))
    assert body_paragraph.text == "综述显示[1,2]。连续研究[1-3]表明。"
    citation_runs = [run for run in body_paragraph.runs if run.text.startswith("[")]
    assert [run.text for run in citation_runs] == ["[1,2]", "[1-3]"]
    assert all(run.font.superscript is True for run in citation_runs)

    with zipfile.ZipFile(fixed_path, "r") as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    ref_paragraph = next(
        p_elem
        for p_elem in root.findall(".//w:body/w:p", ns)
        if "".join(t.text or "" for t in p_elem.findall(".//w:t", ns)).startswith("[1]")
    )
    p_pr = ref_paragraph.find("w:pPr", ns)
    assert p_pr is not None
    assert ref_paragraph.find(".//w:tab", ns) is not None
    jc = p_pr.find("w:jc", ns)
    spacing = p_pr.find("w:spacing", ns)
    assert jc is not None and jc.get(qn("w:val")) == "both"
    assert spacing is not None and spacing.get(qn("w:line")) == "360"
    assert p_pr.find("w:suppressAutoHyphens", ns) is not None
    first_run = ref_paragraph.find("w:r", ns)
    assert first_run is not None
    assert first_run.find("w:rPr/w:vertAlign", ns) is None

    results, _score, _report = audit_thesis.audit_docx(str(fixed_path), profile_path="lnu")
    by_id = {item["id"]: item for item in results}
    for rule_id in ("C03", "C04", "R04", "LNU_REF03"):
        assert by_id[rule_id]["passed"], by_id[rule_id]["issues"]


def test_figures_scope_adds_spacing_between_caption_note_prefix_and_number(tmp_path):
    source_path = Path(tmp_path) / "caption_note_num_spacing_source.docx"
    doc = Document()
    doc.add_paragraph("第2章 实验结果与分析").style = doc.styles["Heading 1"]
    figure = doc.add_paragraph()
    _add_mock_drawing(figure)
    caption = doc.add_paragraph("图2.1 明胶样品的起泡性与起泡稳定性")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note = doc.add_paragraph("注1) 起泡性数据为 3 次平行测定结果。")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("下文继续讨论起泡稳定性结果。")
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "caption_note_num_spacing_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["figures_tables"],
    )

    fixed_doc = Document(fixed_path)
    fixed_note = next(p for p in fixed_doc.paragraphs if p.text.startswith("注"))
    assert fixed_note.text.startswith("注1)")


def test_figures_scope_treats_spaced_colon_table_note_as_caption_note(tmp_path):
    source_path = Path(tmp_path) / "caption_note_spaced_colon_source.docx"
    doc = Document()
    doc.add_paragraph("第2章 实验结果与分析").style = doc.styles["Heading 1"]
    doc.add_paragraph("2.1 正式实验样品得率").style = doc.styles["Heading 2"]
    doc.add_paragraph("正式实验结果见表2.1。")

    caption = doc.add_paragraph("表2.1 正式实验各组样品得率")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "组别"
    table.cell(0, 1).text = "得率"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "85.0"

    note = doc.add_paragraph("注 ：正式实验各组样品初始投料质量均为 1.00 g，数据以平均值 ± 标准差表示（n = 3）。")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("2.2 后续分析").style = doc.styles["Heading 2"]
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "caption_note_spaced_colon_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["figures_tables", "body_paragraphs", "headings"],
    )

    fixed_doc = Document(fixed_path)
    fixed_note = next(p for p in fixed_doc.paragraphs if p.text.startswith("注"))
    note_before, _note_after, _ = _spacing_attrs(fixed_note)
    assert fixed_note.text.startswith("注：")
    assert note_before == "0"

    results, _score, _report = audit_thesis.audit_docx(str(fixed_path), profile_path="lnu")
    s03 = next(item for item in results if item["id"] == "S03")
    lnu_tb04 = next(item for item in results if item["id"] == "LNU_TB04")
    assert s03["passed"], s03["issues"]
    assert lnu_tb04["passed"], lnu_tb04["issues"]


def test_figures_scope_formats_english_caption_and_explanatory_note_separately(tmp_path):
    source_path = Path(tmp_path) / "english_caption_note_source.docx"
    doc = Document()
    doc.add_paragraph("第2章 实验结果与分析").style = doc.styles["Heading 1"]
    figure = doc.add_paragraph()
    _add_mock_drawing(figure)
    caption = doc.add_paragraph("图2.1 明胶凝胶强度")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    english_caption = doc.add_paragraph("Fig. 2.1A  Gel strength of gelatin samples.")
    english_caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
    note = doc.add_paragraph("注：图2.1B 为实验组。")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("下文继续分析。")
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "english_caption_note_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["figures_tables"],
    )

    fixed_doc = Document(fixed_path)
    fixed_english_caption = _paragraph_by_prefix(fixed_doc, "Fig.")
    fixed_note = _paragraph_by_prefix(fixed_doc, "注：")
    _, _, en_line = _spacing_attrs(fixed_english_caption)
    note_before, _, note_line = _spacing_attrs(fixed_note)

    assert fixed_english_caption.text.startswith("Fig. 2.1(A)")
    assert fixed_english_caption.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert en_line == "240"
    assert fixed_note.text.startswith("注：图2.1(B)")
    assert fixed_note.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert note_before == "0"
    assert note_line == "240"

    results, _score, _report = audit_thesis.audit_docx(str(fixed_path), profile_path="lnu")
    lnu_f06 = next(item for item in results if item["id"] == "LNU_F06")
    assert lnu_f06["passed"], lnu_f06["issues"]


def test_heading_and_table_scopes_keep_heading_spacing_after_table_block(tmp_path):
    source_path = Path(tmp_path) / "lnu_heading_table_spacing_source.docx"
    doc = Document()
    doc.add_paragraph("第2章 实验结果与分析").style = doc.styles["Heading 1"]
    doc.add_paragraph("2.1 基因组组装结果").style = doc.styles["Heading 2"]
    doc.add_paragraph("组装统计结果如下所示。")

    caption = doc.add_paragraph("表2.1 基因组组装统计")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "层次"
    table.cell(0, 1).text = "N50"
    table.cell(1, 0).text = "Contig"
    table.cell(1, 1).text = "522958"

    next_heading = doc.add_paragraph("2.2 基因组组分分析")
    next_heading.style = doc.styles["Heading 2"]
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "lnu_heading_table_spacing_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["headings", "figures_tables"],
    )

    fixed_doc = Document(fixed_path)
    fixed_next_heading = _paragraph_by_prefix(fixed_doc, "2.2 基因组组分分析")
    next_before, next_after, _ = _spacing_attrs(fixed_next_heading)
    cfg = _profile_cfg("lnu")
    body_children = list(fixed_doc.element.body.iterchildren())
    caption_idx = next(
        i
        for i, elem in enumerate(body_children)
        if elem.tag == qn("w:p") and _normalize_spaces("".join(elem.itertext())).startswith("表2.1 基因组组装统计")
    )

    assert body_children[caption_idx + 1].tag == qn("w:tbl")
    if cfg.get("table_blank_line_mode") == "blank_paragraph":
        assert body_children[caption_idx + 2].tag == qn("w:p")
        assert not "".join(body_children[caption_idx + 2].itertext()).strip()
        assert next_before == "120"
    else:
        assert next_before == "360"
    assert next_after == "120"

    results, _score, _report = audit_thesis.audit_docx(str(fixed_path), profile_path="lnu")
    s01 = next(item for item in results if item["id"] == "S01")
    lnu_tb04 = next(item for item in results if item["id"] == "LNU_TB04")
    assert s01["passed"], s01["issues"]
    assert lnu_tb04["passed"], lnu_tb04["issues"]


def test_figures_scope_converges_for_adjacent_table_blocks(tmp_path):
    source_path = Path(tmp_path) / "adjacent_tables_source.docx"
    first_path = Path(tmp_path) / "adjacent_tables_first.docx"
    second_path = Path(tmp_path) / "adjacent_tables_second.docx"
    doc = Document()
    doc.add_paragraph("第2章 实验结果与分析").style = doc.styles["Heading 1"]
    for number in (1, 2):
        doc.add_paragraph(f"表2.{number} 统计结果").alignment = WD_ALIGN_PARAGRAPH.CENTER
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "组别"
        table.cell(0, 1).text = "结果"
        table.cell(1, 0).text = "实验组"
        table.cell(1, 1).text = str(number)
        if number == 1:
            doc.add_paragraph("")
    doc.add_paragraph("下文继续分析。")
    doc.save(source_path)

    for input_path, output_path in ((source_path, first_path), (first_path, second_path)):
        apply_scoped_fix(
            str(input_path),
            str(output_path),
            profile_path="lnu",
            scopes=["figures_tables"],
        )

    fixed_doc = Document(second_path)
    children = list(fixed_doc.element.body.iterchildren())
    tables = [index for index, child in enumerate(children) if child.tag == qn("w:tbl")]
    between = children[tables[0] + 1 : tables[1]]
    assert sum(child.tag == qn("w:p") and not "".join(child.itertext()).strip() for child in between) == 1


def test_figures_scope_allows_table_block_at_document_end(tmp_path):
    source_path = Path(tmp_path) / "lnu_table_block_at_document_end_source.docx"
    doc = Document()
    doc.add_paragraph("第2章 实验结果与分析").style = doc.styles["Heading 1"]
    doc.add_paragraph("2.1 基因组组装结果").style = doc.styles["Heading 2"]
    paragraph = doc.add_paragraph("组装统计结果如下所示。")
    paragraph.paragraph_format.space_after = Pt(18)

    caption = doc.add_paragraph("表2.1 基因组组装统计")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = 6

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "层次"
    table.cell(0, 1).text = "N50"
    table.cell(1, 0).text = "Contig"
    table.cell(1, 1).text = "522958"
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "lnu_table_block_at_document_end_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["figures_tables"],
    )

    results, _score, _report = audit_thesis.audit_docx(str(fixed_path), profile_path="lnu")
    lnu_tb04 = next(item for item in results if item["id"] == "LNU_TB04")
    assert lnu_tb04["passed"], lnu_tb04["issues"]


def test_figures_scope_sets_repeat_header_for_long_tables(tmp_path):
    source_path = Path(tmp_path) / "tb03_repeat_header_source.docx"
    doc = Document()
    doc.add_paragraph("第2章 实验结果与分析").style = doc.styles["Heading 1"]
    doc.add_paragraph("2.1 基因组组装结果").style = doc.styles["Heading 2"]
    doc.add_paragraph("长表结果如下所示。")
    doc.add_paragraph("表2.1 长表统计")
    table = doc.add_table(rows=6, cols=2)
    for row_idx in range(6):
        table.cell(row_idx, 0).text = f"项目{row_idx + 1}"
        table.cell(row_idx, 1).text = f"数值{row_idx + 1}"
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "tb03_repeat_header_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["figures_tables"],
    )

    with zipfile.ZipFile(fixed_path) as zf:
        document_xml = ET.fromstring(zf.read("word/document.xml"))
    first_table = document_xml.find(".//w:tbl", {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"})
    assert first_table is not None
    first_row = first_table.find("w:tr", {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"})
    assert first_row is not None
    assert first_row.find("w:trPr/w:tblHeader", {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}) is not None

    results, _score, _report = audit_thesis.audit_docx(str(fixed_path), profile_path="lnu")
    tb03 = next(item for item in results if item["id"] == "TB03")
    assert tb03["passed"], tb03["issues"]


def test_heading_and_figure_scopes_share_gap_budget_without_over_spacing(tmp_path):
    source_path = Path(tmp_path) / "lnu_heading_figure_spacing_source.docx"
    doc = Document()
    doc.add_paragraph("第2章 实验结果与分析").style = doc.styles["Heading 1"]
    doc.add_paragraph("2.1.1 溶胀特性分析").style = doc.styles["Heading 3"]
    figure = doc.add_paragraph()
    _add_mock_drawing(figure)
    caption = doc.add_paragraph("图2.1 不同改性条件筛选组样品的溶胀率与溶失率")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note = doc.add_paragraph("注：（A）0.5～24 h 溶胀率变化曲线；（B）24 h 溶胀率。")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("2.2 不同改性技术对鹿皮明胶功能特性的影响").style = doc.styles["Heading 2"]
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "lnu_heading_figure_spacing_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["headings", "figures_tables"],
    )

    fixed_doc = Document(fixed_path)
    paragraph_info = [
        (idx, paragraph.text.strip(), bool(paragraph._p.xpath(".//w:drawing")))
        for idx, paragraph in enumerate(fixed_doc.paragraphs)
    ]
    idx_drawing = next(idx for idx, text, has_drawing in paragraph_info if has_drawing)
    idx_note = next(idx for idx, text, has_drawing in paragraph_info if text.startswith("注"))
    previous_heading = next(paragraph for paragraph in fixed_doc.paragraphs if paragraph.text.startswith("2.1.1 "))
    _, previous_after, _ = _spacing_attrs(previous_heading)
    drawing_before, _, _ = _spacing_attrs(fixed_doc.paragraphs[idx_drawing])
    _, note_after, _ = _spacing_attrs(fixed_doc.paragraphs[idx_note])

    assert int(previous_after or 0) + int(drawing_before or 0) == 360
    next_heading = next(paragraph for paragraph in fixed_doc.paragraphs if paragraph.text.startswith("2.2 "))
    next_before, _, _ = _spacing_attrs(next_heading)
    assert int(note_after or 0) + int(next_before or 0) == 360

    results, _score, _report = audit_thesis.audit_docx(str(fixed_path), profile_path="lnu")
    lnu_f03 = next(item for item in results if item["id"] == "LNU_F03")
    lnu_f06 = next(item for item in results if item["id"] == "LNU_F06")
    assert lnu_f03["passed"], lnu_f03["issues"]
    assert lnu_f06["passed"], lnu_f06["issues"]


def test_page_scope_repairs_existing_footer_page_field(tmp_docx, tmp_path):
    source_path = Path(tmp_docx(make_compliant_doc, filename="existing_footer_page_source.docx"))
    _inject_existing_footer_page_field(source_path)

    fixed_path = Path(tmp_path) / "existing_footer_page_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["page"],
    )

    results, _score, _report = audit_thesis.audit_docx(str(fixed_path), profile_path="lnu")
    pg01 = next(item for item in results if item["id"] == "PG01")
    p03 = next(item for item in results if item["id"] == "P03")
    assert pg01["passed"], pg01["issues"]
    assert p03["passed"], p03["issues"]


def test_fix_docx_writes_reopenable_docx(tmp_docx, tmp_path):
    source_path = tmp_docx(make_compliant_doc, filename="atomic_write_reopen_source.docx")
    fixed_path = Path(tmp_path) / "atomic_write_reopen_fixed.docx"

    fix_thesis.fix_docx(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["headings"],
    )

    assert fixed_path.exists()
    with zipfile.ZipFile(fixed_path, "r") as zip_handle:
        assert zip_handle.testzip() is None
    reopened = Document(fixed_path)
    assert len(reopened.paragraphs) > 0


def test_fix_docx_is_part_idempotent_after_first_toc_and_page_fix(tmp_path):
    source_path = Path(tmp_path) / "idempotent_source.docx"
    first_path = Path(tmp_path) / "idempotent_first.docx"
    second_path = Path(tmp_path) / "idempotent_second.docx"
    doc = Document()
    doc.add_paragraph("封面信息")
    toc_title = doc.add_paragraph("Table of contents")
    toc_title._p.get_or_add_pPr().append(OxmlElement("w:sectPr"))
    doc.add_paragraph("第1章 绪论")
    doc.add_paragraph("1.1 研究背景")
    doc.add_paragraph("这是正文。")
    doc.save(source_path)

    fix_thesis.fix_docx(
        str(source_path),
        str(first_path),
        profile_path="lnu",
        scopes=["toc", "page"],
    )
    fix_thesis.fix_docx(
        str(first_path),
        str(second_path),
        profile_path="lnu",
        scopes=["toc", "page"],
    )

    with zipfile.ZipFile(first_path, "r") as first_zip, zipfile.ZipFile(second_path, "r") as second_zip:
        assert first_zip.read("word/document.xml") == second_zip.read("word/document.xml")
        assert first_zip.read("word/styles.xml") == second_zip.read("word/styles.xml")


def test_toc_refresh_preserves_frontmatter_page_section_across_repeated_apply(tmp_path):
    source_path = Path(tmp_path) / "frontmatter_sections_source.docx"
    first_path = Path(tmp_path) / "frontmatter_sections_first.docx"
    second_path = Path(tmp_path) / "frontmatter_sections_second.docx"

    doc = Document()
    cover = doc.add_paragraph("封面信息")
    cover._p.get_or_add_pPr().append(OxmlElement("w:sectPr"))
    doc.add_paragraph("摘  要")
    doc.add_paragraph("摘要正文")
    doc.add_paragraph("关键词：测试；修复")
    manual_break = doc.add_paragraph()
    page_break = OxmlElement("w:br")
    page_break.set(qn("w:type"), "page")
    manual_break.add_run()._r.append(page_break)
    toc_title = doc.add_paragraph("目  录")
    toc_title._p.get_or_add_pPr().append(OxmlElement("w:sectPr"))
    doc.add_paragraph("序  言").style = doc.styles["Heading 1"]
    doc.add_paragraph("这是正文内容。")
    doc.save(source_path)

    for input_path, output_path in ((source_path, first_path), (first_path, second_path)):
        apply_scoped_fix(
            str(input_path),
            str(output_path),
            profile_path="lnu",
            scopes=["toc", "page", "acknowledgement"],
            toc=True,
        )

        with zipfile.ZipFile(output_path) as zf:
            root = ET.fromstring(zf.read("word/document.xml"))
        paragraphs = root.findall("w:body/w:p", fix_thesis.NSMAP)
        toc_heading = next(
            paragraph
            for paragraph in paragraphs
            if paragraph.find("w:pPr/w:pStyle", fix_thesis.NSMAP) is not None
            and paragraph.find("w:pPr/w:pStyle", fix_thesis.NSMAP).get(qn("w:val")) == "TOCHeading"
        )
        toc_heading_index = paragraphs.index(toc_heading)
        assert toc_heading.find("w:pPr/w:pageBreakBefore", fix_thesis.NSMAP) is not None
        assert not any(
            paragraph.find('.//w:br[@w:type="page"]', fix_thesis.NSMAP) is not None
            for paragraph in paragraphs[:toc_heading_index]
        )
        sections = root.findall(".//w:sectPr", fix_thesis.NSMAP)
        assert len(sections) == 3
        assert sections[0].find("w:footerReference", fix_thesis.NSMAP) is None
        for section, expected_format in zip(sections[1:], ("upperRoman", "decimal")):
            page_number = section.find("w:pgNumType", fix_thesis.NSMAP)
            assert page_number is not None
            assert page_number.get(qn("w:fmt")) == expected_format
            assert page_number.get(qn("w:start")) == "1"
            assert section.find("w:footerReference", fix_thesis.NSMAP) is not None


def test_auto_toc_bookmark_excludes_hidden_page_marker():
    document_root = ET.Element(f"{{{fix_thesis.W_NS}}}document")
    body = ET.SubElement(document_root, f"{{{fix_thesis.W_NS}}}body")
    heading = ET.SubElement(body, f"{{{fix_thesis.W_NS}}}p")
    content = ET.SubElement(body, f"{{{fix_thesis.W_NS}}}p")
    page_marker = ET.SubElement(body, f"{{{fix_thesis.W_NS}}}p")
    page_run = ET.SubElement(page_marker, f"{{{fix_thesis.W_NS}}}r")
    page_run_properties = ET.SubElement(page_run, f"{{{fix_thesis.W_NS}}}rPr")
    ET.SubElement(page_run_properties, f"{{{fix_thesis.W_NS}}}vanish")
    page_instruction = ET.SubElement(page_run, f"{{{fix_thesis.W_NS}}}instrText")
    page_instruction.text = " PAGE "

    assert toc_fix._add_body_toc_bookmark(
        document_root,
        list(body),
        0,
        toc_fix.BODY_TOC_BOOKMARK_NAME,
        fix_thesis.NSMAP,
        fix_thesis.W_NS,
        fix_thesis.set_attr,
    )
    assert heading.find("w:bookmarkStart", fix_thesis.NSMAP) is not None
    assert content.find("w:bookmarkEnd", fix_thesis.NSMAP) is not None
    assert page_marker.find("w:bookmarkEnd", fix_thesis.NSMAP) is None


def test_fix_docx_is_part_idempotent_after_inserting_acknowledgement(tmp_path):
    source_path = Path(tmp_path) / "ack_idempotent_source.docx"
    first_path = Path(tmp_path) / "ack_idempotent_first.docx"
    second_path = Path(tmp_path) / "ack_idempotent_second.docx"
    doc = Document()
    doc.add_paragraph("1 绪论")
    doc.add_paragraph("1.1 BATMAN-TCM 数据库").style = doc.styles["Heading 2"]
    doc.add_paragraph("这是正文。")
    doc.save(source_path)

    for input_path, output_path in ((source_path, first_path), (first_path, second_path)):
        fix_thesis.fix_docx(
            str(input_path),
            str(output_path),
            profile_path="lnu",
            toc=True,
            scopes=["toc", "headings", "acknowledgement"],
        )

    with zipfile.ZipFile(first_path, "r") as first_zip, zipfile.ZipFile(second_path, "r") as second_zip:
        assert first_zip.read("word/document.xml") == second_zip.read("word/document.xml")
        assert first_zip.read("word/styles.xml") == second_zip.read("word/styles.xml")

    first_doc = Document(first_path)
    assert any(
        paragraph.style.style_id == "TOC1" and paragraph.text.strip() == "致  谢"
        for paragraph in first_doc.paragraphs
    )


def test_toc_scope_inserts_visible_toc_without_word_field(tmp_path):
    source_path = Path(tmp_path) / "visible_toc_source.docx"
    fixed_path = Path(tmp_path) / "visible_toc_fixed.docx"
    doc = Document()
    doc.add_paragraph("封面信息")
    doc.add_paragraph("第1章 绪论")
    doc.add_paragraph("1.1 研究背景")
    doc.add_paragraph("这是正文。")
    doc.save(source_path)

    fix_thesis.fix_docx(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["toc"],
    )

    fixed_doc = Document(fixed_path)
    texts = [paragraph.text.strip() for paragraph in fixed_doc.paragraphs]
    with zipfile.ZipFile(fixed_path, "r") as zf:
        document_xml = zf.read("word/document.xml").decode("utf-8")

    assert "目  录" in texts
    assert texts.index("目  录") < texts.index("第1章 绪论")
    assert any(text.startswith("第1章 绪论") for text in texts)
    assert any(text.startswith("1.1") and "研究背景" in text for text in texts)
    assert not any("待核对" in text for text in texts)
    assert 'TOC \\o "1-3"' not in document_xml


def test_toc_scope_replaces_existing_toc_content_control(tmp_path):
    source_path = Path(tmp_path) / "sdt_toc_source.docx"
    first_path = Path(tmp_path) / "sdt_toc_first.docx"
    second_path = Path(tmp_path) / "sdt_toc_second.docx"
    doc = Document()
    doc.add_paragraph("封面信息")
    heading = doc.add_paragraph("第1章 绪论")
    heading.style = doc.styles["Heading 1"]
    bookmark_start = OxmlElement("w:bookmarkStart")
    bookmark_start.set(qn("w:id"), "42")
    bookmark_start.set(qn("w:name"), "_TocFixture")
    bookmark_end = OxmlElement("w:bookmarkEnd")
    bookmark_end.set(qn("w:id"), "42")
    heading._p.insert(1, bookmark_start)
    heading._p.append(bookmark_end)
    doc.add_paragraph("这是正文。")

    sdt = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    for text, instruction in (
        ("目  录", None),
        (None, ' TOC \\o "1-3" \\h \\z \\u '),
        ("第1章 绪论\t1", None),
    ):
        paragraph = OxmlElement("w:p")
        run = OxmlElement("w:r")
        node = OxmlElement("w:instrText" if instruction else "w:t")
        node.text = instruction or text
        run.append(node)
        if text == "第1章 绪论\t1":
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("w:anchor"), "_TocFixture")
            hyperlink.append(run)
            paragraph.append(hyperlink)
        else:
            paragraph.append(run)
        content.append(paragraph)
    sdt.append(content)
    doc.element.body.insert(1, sdt)
    doc.save(source_path)

    for input_path, output_path in ((source_path, first_path), (first_path, second_path)):
        apply_scoped_fix(
            str(input_path),
            str(output_path),
            profile_path="lnu",
            scopes=["toc"],
            toc=True,
        )
        with zipfile.ZipFile(output_path) as zf:
            root = ET.fromstring(zf.read("word/document.xml"))
        toc_fields = [
            instr
            for instr in root.findall(".//w:instrText", fix_thesis.NSMAP)
            if "TOC" in (instr.text or "").upper()
        ]
        toc_titles = [
            paragraph
            for paragraph in root.findall(".//w:p", fix_thesis.NSMAP)
            if _normalize_spaces(fix_thesis.get_paragraph_text(paragraph)) == "目 录"
        ]
        toc_links = root.findall(".//w:hyperlink", fix_thesis.NSMAP)
        assert len(toc_fields) == 1
        assert len(toc_titles) == 1
        assert [link.get(qn("w:anchor")) for link in toc_links] == ["_TocFixture"]


def test_toc_scope_ignores_frontmatter_that_only_looks_like_a_heading(tmp_path):
    source_path = Path(tmp_path) / "toc_false_heading_source.docx"
    fixed_path = Path(tmp_path) / "toc_false_heading_fixed.docx"
    doc = Document()
    doc.add_paragraph("封面信息")

    address = doc.add_paragraph()
    address.alignment = WD_ALIGN_PARAGRAPH.CENTER
    address_run = address.add_run("123 Research Road, Leiden, NL")
    address_run.bold = True
    address_run.font.size = Pt(14)

    abstract = doc.add_paragraph()
    abstract.alignment = WD_ALIGN_PARAGRAPH.CENTER
    abstract_run = abstract.add_run("Abstract")
    abstract_run.bold = True
    abstract_run.font.size = Pt(14)
    doc.add_paragraph("This is the abstract body.")
    doc.add_paragraph("第1章 绪论")
    doc.add_paragraph("1.1 研究背景")
    doc.add_paragraph("这是正文。")
    doc.save(source_path)

    fix_thesis.fix_docx(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["toc"],
    )

    fixed_doc = Document(fixed_path)
    toc_title_index = next(
        index for index, paragraph in enumerate(fixed_doc.paragraphs)
        if paragraph.text.strip() == "目  录"
    )
    toc_entries = []
    for paragraph in fixed_doc.paragraphs[toc_title_index + 1:]:
        if not paragraph.style.style_id.startswith("TOC"):
            break
        if paragraph.style.style_id in {"TOC1", "TOC2", "TOC3"}:
            toc_entries.append(paragraph.text.strip())

    assert toc_entries == ["第1章 绪论", "1.1 研究背景"]


def test_abstract_scope_does_not_move_keywords_from_an_unrelated_english_toc(tmp_path):
    source_path = Path(tmp_path) / "english_toc_keywords_source.docx"
    fixed_path = Path(tmp_path) / "english_toc_keywords_fixed.docx"
    doc = Document()
    doc.add_paragraph("Journal manuscript")
    doc.add_paragraph("Table of contents")
    doc.add_paragraph("Abstract: 8")
    doc.add_paragraph("Keywords/Subject terms: 8")
    doc.add_paragraph("Abstract")
    doc.add_paragraph("This is the abstract body.")
    doc.add_paragraph("Keywords/Subject terms: alpha; beta")
    doc.add_paragraph("第1章 绪论")
    doc.save(source_path)

    fix_thesis.fix_docx(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["abstract"],
    )

    texts = [paragraph.text.strip() for paragraph in Document(fixed_path).paragraphs]

    assert texts.index("Table of contents") < texts.index("Keywords/Subject terms: 8")
    assert texts.index("Keywords/Subject terms: 8") < texts.index("Abstract")
    assert texts.index("Abstract") < texts.index("Keywords/Subject terms: alpha; beta")


def test_toc_scope_preserves_existing_visible_toc_page_numbers(tmp_path):
    source_path = Path(tmp_path) / "visible_toc_existing_pages_source.docx"
    fixed_path = Path(tmp_path) / "visible_toc_existing_pages_fixed.docx"
    doc = Document()
    doc.add_paragraph("封面信息")
    doc.add_paragraph("目  录")
    doc.add_paragraph("第1章 绪论\t3")
    doc.add_paragraph("1.1 研究背景\t4")
    doc.add_paragraph("第1章 绪论")
    doc.add_paragraph("1.1 研究背景")
    doc.add_paragraph("这是正文。")
    doc.save(source_path)

    fix_thesis.fix_docx(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["toc"],
    )

    texts = [paragraph.text.strip() for paragraph in Document(fixed_path).paragraphs]

    assert any(text.startswith("第1章 绪论") and text.endswith("3") for text in texts)
    assert any(text.startswith("1.1 研究背景") and text.endswith("4") for text in texts)
    assert not any("待核对" in text for text in texts)


def test_fix_docx_keeps_page_break_before_reference_heading_in_default_flow(tmp_docx, tmp_path):
    source_path = tmp_docx(make_violating_doc, filename="default_fix_keeps_reference_break_source.docx", rule_id="T01")
    fixed_path = Path(tmp_path) / "default_fix_keeps_reference_break_fixed.docx"

    fix_thesis.fix_docx(str(source_path), str(fixed_path))

    results, _score, _report = audit_thesis.audit_docx(str(fixed_path))
    s02 = next(item for item in results if item["id"] == "S02")
    assert s02["passed"], s02["issues"]


def test_fix_docx_validation_failure_keeps_existing_output(tmp_docx, tmp_path, monkeypatch):
    source_path = tmp_docx(make_compliant_doc, filename="atomic_write_failure_source.docx")
    output_path = Path(tmp_path) / "atomic_write_failure_output.docx"
    original_payload = b"previous-good-file"
    output_path.write_bytes(original_payload)

    def _raise_reopen_error(_path):
        raise RuntimeError("forced reopen failure")

    monkeypatch.setattr(fix_thesis, "Document", _raise_reopen_error)

    with pytest.raises(RuntimeError, match="python-docx"):
        fix_thesis.fix_docx(
            str(source_path),
            str(output_path),
            profile_path="lnu",
            scopes=["headings"],
        )

    assert output_path.read_bytes() == original_payload


def test_figures_scope_renumbers_split_run_lnu_captions(tmp_path):
    source_path = Path(tmp_path) / "split_run_caption_source.docx"
    doc = Document()
    heading = doc.add_paragraph("2 结果")
    heading.style = doc.styles["Heading 1"]

    caption = doc.add_paragraph()
    caption.style = "Caption"
    caption.add_run("图 ")
    caption.add_run("")
    caption.add_run("1")
    caption.add_run(" 鹿茸成分靶点与肝毒性靶点韦恩图")
    doc.save(source_path)

    fixed_path = Path(tmp_path) / "split_run_caption_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["figures_tables"],
    )

    fixed_doc = Document(fixed_path)
    texts = [paragraph.text for paragraph in fixed_doc.paragraphs]
    assert "图2.1 鹿茸成分靶点与肝毒性靶点韦恩图" in texts


def test_figures_scope_preserves_cover_anchor_and_inlines_body_anchor(tmp_path):
    source_path = Path(tmp_path) / "anchor_scope_source.docx"
    fixed_path = Path(tmp_path) / "anchor_scope_fixed.docx"
    _make_cover_and_body_anchor_docx(source_path)

    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["figures_tables"],
    )

    _assert_cover_anchor_and_body_inline(fixed_path)


def test_normalize_preserves_cover_anchor_and_inlines_body_anchor(tmp_path):
    source_path = Path(tmp_path) / "anchor_normalize_source.docx"
    fixed_path = Path(tmp_path) / "anchor_normalize_fixed.docx"
    _make_cover_and_body_anchor_docx(source_path)

    fix_thesis.normalize_docx(str(source_path), str(fixed_path), profile_path="lnu")

    _assert_cover_anchor_and_body_inline(fixed_path)
