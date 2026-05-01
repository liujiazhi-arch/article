from pathlib import Path

import audit_thesis
from docx import Document
from docx.oxml import OxmlElement

from thesis_workbench import default_output_path
from thesis_tool.scopes import scope_for_rule
from thesis_tool.workflow import (
    apply_scoped_fix,
    build_document_diagnostics,
    build_document_normalize,
    build_document_preflight,
    build_scope_plan,
    build_scope_verify,
    render_document_diagnostics_compact,
    render_document_diagnostics,
    render_document_normalize,
    render_document_normalize_compact,
    render_document_preflight,
    render_document_preflight_compact,
    render_scope_plan,
    render_scope_verify,
)

from .conftest import RULE_MUTATORS, _add_heading, _set_heading, audit_rule_status, make_compliant_doc


def _build_multi_violation_doc(tmp_docx, filename: str, rule_ids: tuple[str, ...]) -> Path:
    path = tmp_docx(make_compliant_doc, filename=filename)
    doc = Document(path)
    for rule_id in rule_ids:
        RULE_MUTATORS[rule_id](doc)
    doc.save(path)
    return path


def _build_lnu_f05_fail_doc(path: Path) -> Path:
    doc = Document()
    heading = doc.add_paragraph("第一章 绪论")
    heading.style = "Heading 1"
    doc.add_paragraph("本节介绍实验装置。")
    doc.add_paragraph("图1.1 实验装置示意图")
    doc.save(path)
    return path


def _build_lnu_f03_fail_doc(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("第2章 实验结果与分析").style = "Heading 1"
    doc.add_paragraph("2.1.1 溶胀特性分析").style = "Heading 3"
    figure = doc.add_paragraph()
    figure.add_run()._r.append(OxmlElement("w:drawing"))
    doc.add_paragraph("图2.1 不同改性条件筛选组样品的溶胀率与溶失率")
    doc.add_paragraph("注：（A）0.5～24 h 溶胀率变化曲线；（B）24 h 溶胀率。")
    doc.add_paragraph("图后正文。")
    doc.save(path)
    return path


def _build_abstract_pu01_scope_doc(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("摘 要")
    doc.add_paragraph("这是中文,摘要.内容")
    doc.add_paragraph("关键词：测试；流程；稳定性")
    doc.add_paragraph("Abstract")
    doc.add_paragraph("This is the English abstract.")
    _add_heading(doc, "第一章 绪论", level=1, size=30)
    doc.add_paragraph("这是正文,内容.保留")
    doc.save(path)
    return path


def test_scope_plan_groups_failed_rules_by_scope(tmp_docx):
    docx_path = _build_multi_violation_doc(
        tmp_docx,
        filename="scope_plan_grouping.docx",
        rule_ids=("H02", "KW01"),
    )

    plan = build_scope_plan(str(docx_path))
    scopes = {scope["id"]: scope for scope in plan["scopes"]}

    assert "H02" in scopes["headings"]["failed_rules"]
    assert scopes["headings"]["status"] == "autofix_ready"
    assert scopes["headings"]["failed_count"] >= 1

    assert "KW01" in scopes["abstract"]["failed_rules"]
    assert scopes["abstract"]["status"] == "manual_review"
    assert scopes["abstract"]["failed_count"] >= 1

    rendered = render_scope_plan(plan)
    assert "正文标题（headings）" in rendered
    assert "摘要（abstract）" in rendered


def test_scope_plan_exposes_action_buckets(tmp_docx):
    docx_path = _build_multi_violation_doc(
        tmp_docx,
        filename="scope_plan_actions.docx",
        rule_ids=("H02", "KW01"),
    )

    plan = build_scope_plan(str(docx_path))
    scopes = {scope["id"]: scope for scope in plan["scopes"]}

    assert scopes["headings"]["status"] == "autofix_ready"
    assert scopes["headings"]["autofixable_count"] >= 1
    assert scopes["headings"]["unsupported_count"] == 0

    assert scopes["abstract"]["status"] == "manual_review"
    assert scopes["abstract"]["autofixable_count"] == 0
    assert scopes["abstract"]["manual_review_count"] >= 1
    assert scopes["abstract"]["unsupported_count"] == 0

    rendered = render_scope_plan(plan)
    assert "可自动修复" in rendered
    assert "当前不支持" in rendered


def test_scoped_fix_only_repairs_selected_scope(tmp_docx, tmp_path):
    source_path = _build_multi_violation_doc(
        tmp_docx,
        filename="scope_fix_source.docx",
        rule_ids=("H02", "T01"),
    )
    before_h02 = audit_rule_status(source_path, "H02")
    before_t01 = audit_rule_status(source_path, "T01")
    assert not before_h02["rule"]["passed"]
    assert not before_t01["rule"]["passed"]

    fixed_path = Path(tmp_path) / "scope_fix_headings_only.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        scopes=["headings"],
    )

    after_h02 = audit_rule_status(fixed_path, "H02")
    after_t01 = audit_rule_status(fixed_path, "T01")
    assert after_h02["rule"]["passed"], after_h02["rule"]["issues"]
    assert not after_t01["rule"]["passed"]


def test_scope_verify_can_focus_on_selected_scope(tmp_docx, tmp_path):
    source_path = _build_multi_violation_doc(
        tmp_docx,
        filename="scope_verify_source.docx",
        rule_ids=("T01", "H02"),
    )
    fixed_path = Path(tmp_path) / "scope_verify_body_only.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        scopes=["body_paragraphs"],
    )

    body_verify = build_scope_verify(str(fixed_path), scopes=["body_paragraphs"])
    assert body_verify["overall_status"] == "verified"
    assert body_verify["readiness"] == "structure-ready"
    assert body_verify["failed_count"] == 0
    assert body_verify["selected_scopes"] == ["body_paragraphs"]

    headings_verify = build_scope_verify(str(fixed_path), scopes=["headings"])
    assert headings_verify["overall_status"] in {"needs_fix", "unsupported"}
    assert headings_verify["readiness"] == "needs-fix"
    assert headings_verify["failed_count"] >= 1

    rendered = render_scope_verify(headings_verify)
    assert "验证状态:" in rendered
    assert "可提交状态: needs-fix" in rendered
    assert "正文标题（headings）" in rendered


def test_scope_verify_exposes_manual_review_rule_summary_for_lnu_f05(tmp_path):
    source_path = _build_lnu_f05_fail_doc(Path(tmp_path) / "scope_verify_lnu_f05_fail.docx")
    verification = build_scope_verify(str(source_path), profile_path="lnu", scopes=["figures"])

    assert verification["overall_status"] in {"manual_review", "needs_fix"}
    assert verification["readiness"] == "manual-review-required"
    assert "LNU_F05" in verification["manual_review_rule_ids"]
    assert verification["unsupported_rule_ids"] == []
    lnu_f05_item = next(item for item in verification["manual_review_rules"] if item["id"] == "LNU_F05")
    assert lnu_f05_item["check_level"] == "Semi"

    rendered = render_scope_verify(verification)
    assert "可提交状态: manual-review-required" in rendered
    assert "仍需人工复核：" in rendered
    assert "LNU_F05" in rendered
    assert "评分较高不等于可直接提交" in rendered


def test_scope_verify_treats_lnu_f03_as_autofixable_not_manual_review(tmp_path):
    source_path = _build_lnu_f03_fail_doc(Path(tmp_path) / "scope_verify_lnu_f03_fail.docx")
    verification = build_scope_verify(str(source_path), profile_path="lnu", scopes=["figures"])

    assert verification["overall_status"] == "needs_fix"
    assert "LNU_F03" not in verification["manual_review_rule_ids"]
    assert "LNU_F03" not in verification["unsupported_rule_ids"]
    figures_scope = next(scope for scope in verification["scopes"] if scope["id"] == "figures_tables")
    assert "LNU_F03" in figures_scope["failed_rules"]
    assert figures_scope["autofixable_count"] >= 1


def test_scope_verify_exposes_unsupported_rule_summary_ids(tmp_docx):
    source_path = _build_multi_violation_doc(
        tmp_docx,
        filename="scope_verify_unsupported_summary.docx",
        rule_ids=("KW01",),
    )
    verification = build_scope_verify(str(source_path), profile_path="cn-common", scopes=["abstract"])

    assert verification["overall_status"] == "manual_review"
    assert verification["readiness"] == "manual-review-required"
    assert "KW01" in verification["manual_review_rule_ids"]
    assert verification["unsupported_rule_ids"] == []

    rendered = render_scope_verify(verification)
    assert "可提交状态: manual-review-required" in rendered
    assert "仍需人工复核：" in rendered
    assert "KW01" in rendered


def test_default_output_path_normalizes_comma_separated_scopes():
    output_path = default_output_path("/tmp/sample.docx", ["headings,body_paragraphs"])
    assert output_path.endswith("/tmp/sample_body_paragraphs_headings.docx")


def test_lnu_s03_is_grouped_under_headings_scope(tmp_docx):
    source_path = tmp_docx(make_compliant_doc, filename="lnu_s03_scope_source.docx")
    doc = Document(source_path)
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("参考文献"):
            p_pr = paragraph._p.pPr
            if p_pr is not None:
                page_break = p_pr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pageBreakBefore")
                if page_break is not None:
                    p_pr.remove(page_break)
            break
    doc.save(source_path)

    plan = build_scope_plan(str(source_path), profile_path="lnu")
    scopes = {scope["id"]: scope for scope in plan["scopes"]}

    assert "LNU_S03" in scopes["headings"]["failed_rules"]
    assert plan["unscoped_failed"] == []


def test_lnu_toc01_is_grouped_under_toc_scope():
    assert scope_for_rule("LNU_TOC01") == "toc"


def test_abstract_scope_repairs_lnu_abs04_without_touching_body_pu01(tmp_path):
    source_path = _build_abstract_pu01_scope_doc(Path(tmp_path) / "abstract_pu01_scope_source.docx")

    before_results, _score, _report = audit_thesis.audit_docx(str(source_path), profile_path="lnu")
    assert any(item["id"] == "LNU_ABS04" and not item["passed"] for item in before_results)
    assert any(item["id"] == "PU01" and not item["passed"] for item in before_results)

    fixed_path = Path(tmp_path) / "abstract_pu01_scope_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["abstract"],
    )

    after_results, _score, _report = audit_thesis.audit_docx(str(fixed_path), profile_path="lnu")
    assert any(item["id"] == "LNU_ABS04" and item["passed"] for item in after_results)
    assert any(item["id"] == "PU01" and not item["passed"] for item in after_results)

    fixed_doc = Document(fixed_path)
    assert any(paragraph.text == "这是中文，摘要。内容" for paragraph in fixed_doc.paragraphs)
    assert any(paragraph.text == "这是正文,内容.保留" for paragraph in fixed_doc.paragraphs)


def test_headings_scope_also_repairs_acknowledgement_titles(tmp_docx, tmp_path):
    source_path = tmp_docx(make_compliant_doc, filename="ack_headings_scope_source.docx")
    doc = Document(source_path)
    doc.add_paragraph("致  谢")
    doc.add_paragraph("1.1 致谢说明")
    doc.add_paragraph("感谢内容")
    doc.save(source_path)

    before_h01 = audit_rule_status(source_path, "H01")
    before_h02 = audit_rule_status(source_path, "H02")
    assert not before_h01["rule"]["passed"]
    assert not before_h02["rule"]["passed"]

    fixed_path = Path(tmp_path) / "ack_headings_scope_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="cn-common",
        scopes=["headings"],
    )

    after_h01 = audit_rule_status(fixed_path, "H01")
    after_h02 = audit_rule_status(fixed_path, "H02")
    assert after_h01["rule"]["passed"], after_h01["rule"]["issues"]
    assert after_h02["rule"]["passed"], after_h02["rule"]["issues"]


def test_body_scope_repairs_lnu_unit_spacing_in_acknowledgement_paragraph(tmp_docx, tmp_path):
    source_path = tmp_docx(make_compliant_doc, filename="ack_unit_scope_source.docx")
    doc = Document(source_path)
    doc.add_paragraph("致  谢")
    doc.add_paragraph("页边距保持上2.5cm，下2.5cm。")
    doc.save(source_path)

    before_results, _score, _report = audit_thesis.audit_docx(str(source_path), profile_path="lnu")
    before_unit = next(item for item in before_results if item["id"] == "LNU_UNIT01")
    assert not before_unit["passed"]

    fixed_path = Path(tmp_path) / "ack_unit_scope_fixed.docx"
    apply_scoped_fix(
        str(source_path),
        str(fixed_path),
        profile_path="lnu",
        scopes=["body_paragraphs"],
    )

    after_results, _score, _report = audit_thesis.audit_docx(str(fixed_path), profile_path="lnu")
    after_unit = next(item for item in after_results if item["id"] == "LNU_UNIT01")
    assert after_unit["passed"], after_unit["issues"]


def test_c01_without_any_superscript_citations_is_manual_review(tmp_docx):
    source_path = _build_multi_violation_doc(
        tmp_docx,
        filename="scope_plan_c01_manual_review.docx",
        rule_ids=("C01",),
    )

    plan = build_scope_plan(str(source_path))
    scopes = {scope["id"]: scope for scope in plan["scopes"]}
    c01_item = next(item for item in scopes["body_paragraphs"]["failed_items"] if item["id"] == "C01")

    assert c01_item["action"] == "manual_review"
    assert scopes["body_paragraphs"]["manual_review_count"] >= 1


def test_build_document_diagnostics_reports_heading_risks(tmp_path):
    source_path = Path(tmp_path) / "diagnostics_heading_risks.docx"
    doc = Document()
    _add_heading(doc, "1 绪论", level=1, size=30)
    doc.add_paragraph("这是正文段落。")

    conflict = doc.add_paragraph("3.6 分子对接验证结果")
    _set_heading(conflict, level=1, size=30)

    table = doc.add_table(rows=1, cols=1)
    cell_paragraph = table.cell(0, 0).paragraphs[0]
    cell_paragraph.text = "3 1,4-Diaminobutane"
    _set_heading(cell_paragraph, level=1, size=30)
    doc.save(source_path)

    diagnostics = build_document_diagnostics(str(source_path), profile_path="lnu")

    assert diagnostics["toc"]["status"] == "no_toc"
    assert diagnostics["heading_renumber_guard"]["status"] == "block"
    assert diagnostics["heading_renumber_guard"]["reason"] == "style_conflict"
    assert any(item["text"] == "3.6 分子对接验证结果" for item in diagnostics["style_text_conflicts"])
    assert any("1,4-Diaminobutane" in item["text"] for item in diagnostics["table_heading_candidates"])
    assert diagnostics["table_heading_candidate_summary"]["comma_compound"] >= 1
    assert any("headings scope" in action for action in diagnostics["recommended_actions"])
    assert any("表格中存在伪标题风险" in action for action in diagnostics["recommended_actions"])

    rendered = render_document_diagnostics(diagnostics)
    assert "目录状态: no_toc" in rendered
    assert "正文重编号守卫: block (style_conflict)" in rendered
    assert "表格内伪标题候选：" in rendered
    assert "风险摘要:" in rendered
    assert "reason=comma_compound" in rendered
    assert "样式/文本层级冲突：" in rendered
    assert "建议动作：" in rendered


def test_build_document_diagnostics_marks_table_risk_as_warn_not_block(tmp_path):
    source_path = Path(tmp_path) / "diagnostics_table_risk_warn.docx"
    doc = Document()
    _add_heading(doc, "1 绪论", level=1, size=30)
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).paragraphs[0].text = "1,4-Diaminobutane"
    table.cell(1, 0).paragraphs[0].text = "2.80E+08"
    doc.save(source_path)

    diagnostics = build_document_diagnostics(str(source_path), profile_path="lnu")

    assert diagnostics["heading_renumber_guard"]["status"] == "warn"
    assert diagnostics["heading_renumber_guard"]["reason"] == "table_risk"

    compact = render_document_diagnostics_compact(diagnostics)
    assert "heading_renumber_guard_status=warn" in compact
    assert "heading_renumber_guard_reason=table_risk" in compact


def test_build_document_preflight_marks_style_conflict_as_blocked(tmp_path):
    source_path = Path(tmp_path) / "preflight_blocked.docx"
    doc = Document()
    _add_heading(doc, "1 绪论", level=1, size=30)
    conflict = doc.add_paragraph("3.6 分子对接验证结果")
    _set_heading(conflict, level=1, size=30)
    doc.save(source_path)

    preflight = build_document_preflight(str(source_path), profile_path="lnu")

    assert preflight["preflight_status"] == "blocked"
    assert preflight["heading_renumber_guard"]["status"] == "block"
    rendered = render_document_preflight(preflight)
    assert "预检状态: blocked" in rendered
    compact = render_document_preflight_compact(preflight)
    assert "preflight_status=blocked" in compact


def test_build_document_preflight_marks_table_risk_as_warning(tmp_path):
    source_path = Path(tmp_path) / "preflight_warning.docx"
    doc = Document()
    _add_heading(doc, "1 绪论", level=1, size=30)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].text = "1,4-Diaminobutane"
    doc.save(source_path)

    preflight = build_document_preflight(str(source_path), profile_path="lnu")

    assert preflight["preflight_status"] == "warning"
    assert preflight["table_heading_risk_count"] >= 1
    rendered = render_document_preflight(preflight)
    assert "预检状态: warning" in rendered
    assert "建议动作：" in rendered
    compact = render_document_preflight_compact(preflight)
    assert "preflight_status=warning" in compact


def test_build_document_normalize_reduces_style_conflict_risk(tmp_path):
    source_path = Path(tmp_path) / "normalize_style_conflict.docx"
    doc = Document()
    heading = doc.add_paragraph("1 绪论")
    heading.style = doc.styles["Heading 1"]
    conflict = doc.add_paragraph("3.6 分子对接验证结果")
    conflict.style = doc.styles["Heading 1"]
    doc.add_paragraph("这是正文示例。")
    doc.save(source_path)
    output_path = Path(tmp_path) / "normalize_style_conflict_output.docx"

    normalize = build_document_normalize(
        str(source_path),
        output_path=str(output_path),
        profile_path="lnu",
    )

    assert output_path.exists()
    assert normalize["changed"] is True
    assert normalize["summary"]["before_preflight_status"] == "blocked"
    assert normalize["summary"]["after_preflight_status"] == "warning"
    assert normalize["summary"]["before_style_conflict_count"] >= 1
    assert normalize["summary"]["after_style_conflict_count"] == 0
    assert any(item["id"] == "heading_styles" for item in normalize["operations"])

    rendered = render_document_normalize(normalize)
    assert "预规整改动: 有" in rendered
    assert "预检变化: blocked -> warning" in rendered

    compact = render_document_normalize_compact(normalize)
    assert "file=normalize_style_conflict.docx" in compact
    assert "after_style_conflict_count=0" in compact


def test_build_document_normalize_normalizes_manual_toc_title_and_entries(tmp_path):
    from docx.oxml.ns import qn

    source_path = Path(tmp_path) / "normalize_manual_toc.docx"
    doc = Document()

    toc_title = doc.add_paragraph("目录")
    toc_title_pr = toc_title._p.get_or_add_pPr()
    toc_title_jc = OxmlElement("w:jc")
    toc_title_jc.set(qn("w:val"), "left")
    toc_title_pr.append(toc_title_jc)

    toc_field = doc.add_paragraph("")
    toc_field_pr = toc_field._p.get_or_add_pPr()
    toc_field_style = OxmlElement("w:pStyle")
    toc_field_style.set(qn("w:val"), "TOCField")
    toc_field_pr.append(toc_field_style)
    toc_instr = OxmlElement("w:instrText")
    toc_instr.text = ' TOC \\\\o "1-3" \\\\h \\\\z \\\\u '
    toc_field.add_run("")._r.append(toc_instr)

    toc_entry = doc.add_paragraph("第一章 绪论\t1")
    toc_entry_pr = toc_entry._p.get_or_add_pPr()
    toc_entry_style = OxmlElement("w:pStyle")
    toc_entry_style.set(qn("w:val"), "TOC1")
    toc_entry_pr.append(toc_entry_style)
    toc_entry_spacing = OxmlElement("w:spacing")
    toc_entry_spacing.set(qn("w:after"), "0")
    toc_entry_pr.append(toc_entry_spacing)

    _add_heading(doc, "第一章 绪论", level=1, size=30)
    doc.save(source_path)

    output_path = Path(tmp_path) / "normalize_manual_toc_output.docx"
    normalize = build_document_normalize(
        str(source_path),
        output_path=str(output_path),
        profile_path="lnu",
    )

    assert output_path.exists()
    assert normalize["changed"] is True
    operation_ids = {item["id"] for item in normalize["operations"]}
    assert "toc_title" in operation_ids
    assert "toc_entries" in operation_ids

    normalized_doc = Document(output_path)
    normalized_title = next(paragraph for paragraph in normalized_doc.paragraphs if paragraph.text.strip() == "目录")
    title_pr = normalized_title._p.pPr
    assert title_pr is not None
    title_style = title_pr.find(qn("w:pStyle"))
    assert title_style is not None
    assert title_style.get(qn("w:val")) == "TOCHeading"
    title_jc = title_pr.find(qn("w:jc"))
    assert title_jc is not None
    assert title_jc.get(qn("w:val")) == "center"

    normalized_entry = next(paragraph for paragraph in normalized_doc.paragraphs if "\t1" in paragraph.text)
    entry_pr = normalized_entry._p.pPr
    assert entry_pr is not None
    entry_spacing = entry_pr.find(qn("w:spacing"))
    assert entry_spacing is not None
    assert entry_spacing.get(qn("w:after")) == "100"

    assert any(paragraph.text.strip() == "第一章 绪论" for paragraph in normalized_doc.paragraphs)


def test_build_document_normalize_keeps_manual_toc_in_warning_state_without_field(tmp_path):
    from docx.oxml.ns import qn

    source_path = Path(tmp_path) / "normalize_manual_only_toc.docx"
    doc = Document()

    toc_title = doc.add_paragraph("目录")
    toc_title_pr = toc_title._p.get_or_add_pPr()
    toc_title_jc = OxmlElement("w:jc")
    toc_title_jc.set(qn("w:val"), "left")
    toc_title_pr.append(toc_title_jc)

    toc_entry = doc.add_paragraph("第一章 绪论\t1")
    toc_entry_pr = toc_entry._p.get_or_add_pPr()
    toc_entry_style = OxmlElement("w:pStyle")
    toc_entry_style.set(qn("w:val"), "TOC1")
    toc_entry_pr.append(toc_entry_style)
    toc_entry_spacing = OxmlElement("w:spacing")
    toc_entry_spacing.set(qn("w:after"), "0")
    toc_entry_pr.append(toc_entry_spacing)

    _add_heading(doc, "第一章 绪论", level=1, size=30)
    doc.save(source_path)

    output_path = Path(tmp_path) / "normalize_manual_only_toc_output.docx"
    normalize = build_document_normalize(
        str(source_path),
        output_path=str(output_path),
        profile_path="lnu",
    )

    assert output_path.exists()
    assert normalize["changed"] is True
    assert normalize["summary"]["before_toc_status"] == "manual_toc"
    assert normalize["summary"]["after_toc_status"] == "manual_toc"
    assert normalize["summary"]["before_preflight_status"] == "warning"
    assert normalize["summary"]["after_preflight_status"] == "warning"
    assert any(item["id"] == "toc_title" for item in normalize["operations"])
    assert any(item["id"] == "toc_entries" for item in normalize["operations"])
    assert any("建议先 verify 关键 scope" in step for step in normalize["next_steps"])


def test_build_document_normalize_removes_duplicate_body_toc_title(tmp_path):
    from docx.oxml.ns import qn

    source_path = Path(tmp_path) / "normalize_duplicate_toc_title.docx"
    doc = Document()

    toc_title = doc.add_paragraph("目  录")
    toc_title_pr = toc_title._p.get_or_add_pPr()
    toc_title_style = OxmlElement("w:pStyle")
    toc_title_style.set(qn("w:val"), "TOCHeading")
    toc_title_pr.append(toc_title_style)

    toc_field = doc.add_paragraph("")
    toc_field_pr = toc_field._p.get_or_add_pPr()
    toc_field_style = OxmlElement("w:pStyle")
    toc_field_style.set(qn("w:val"), "TOCField")
    toc_field_pr.append(toc_field_style)
    toc_instr = OxmlElement("w:instrText")
    toc_instr.text = ' TOC \\\\o "1-3" \\\\h \\\\z \\\\u '
    toc_field.add_run("")._r.append(toc_instr)

    toc_entry = doc.add_paragraph("第一章 绪论\t1")
    toc_entry_pr = toc_entry._p.get_or_add_pPr()
    toc_entry_style = OxmlElement("w:pStyle")
    toc_entry_style.set(qn("w:val"), "TOC1")
    toc_entry_pr.append(toc_entry_style)
    doc.add_paragraph("目录")
    _add_heading(doc, "第一章 绪论", level=1, size=30)
    doc.add_paragraph("正文内容。")
    doc.save(source_path)

    output_path = Path(tmp_path) / "normalize_duplicate_toc_title_output.docx"
    normalize = build_document_normalize(
        str(source_path),
        output_path=str(output_path),
        profile_path="lnu",
    )

    assert output_path.exists()
    assert normalize["changed"] is True
    assert any(item["id"] == "duplicate_toc_title" for item in normalize["operations"])

    normalized_doc = Document(output_path)
    texts = [paragraph.text.strip() for paragraph in normalized_doc.paragraphs if paragraph.text.strip()]
    assert texts.count("目录") == 0
    assert texts.count("目  录") == 1
    assert "第一章 绪论" in texts
    assert "正文内容。" in texts


def test_build_document_diagnostics_reports_field_only_toc(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    source_path = Path(tmp_path) / "diagnostics_field_only_toc.docx"
    doc = Document()
    title = doc.add_paragraph("目  录")
    title_p_pr = title._p.get_or_add_pPr()
    title_style = OxmlElement("w:pStyle")
    title_style.set(qn("w:val"), "TOCHeading")
    title_p_pr.append(title_style)
    title_jc = OxmlElement("w:jc")
    title_jc.set(qn("w:val"), "center")
    title_p_pr.append(title_jc)

    field = doc.add_paragraph("")
    field_p_pr = field._p.get_or_add_pPr()
    field_style = OxmlElement("w:pStyle")
    field_style.set(qn("w:val"), "TOCField")
    field_p_pr.append(field_style)
    run = field.add_run("")
    instr = OxmlElement("w:instrText")
    instr.text = ' TOC \\\\o "1-3" \\\\h \\\\z \\\\u '
    run._r.append(instr)

    _add_heading(doc, "1 绪论", level=1, size=30)
    doc.save(source_path)

    diagnostics = build_document_diagnostics(str(source_path), profile_path="lnu")

    assert diagnostics["toc"]["status"] == "field_only"
    assert diagnostics["toc"]["title_count"] == 1
    assert diagnostics["toc"]["field_count"] == 1


def test_build_scope_verify_marks_field_only_toc_as_render_check_required(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    source_path = Path(tmp_path) / "verify_field_only_toc.docx"
    doc = Document()
    title = doc.add_paragraph("目  录")
    title_p_pr = title._p.get_or_add_pPr()
    title_style = OxmlElement("w:pStyle")
    title_style.set(qn("w:val"), "TOCHeading")
    title_p_pr.append(title_style)
    title_jc = OxmlElement("w:jc")
    title_jc.set(qn("w:val"), "center")
    title_p_pr.append(title_jc)

    field = doc.add_paragraph("")
    field_p_pr = field._p.get_or_add_pPr()
    field_style = OxmlElement("w:pStyle")
    field_style.set(qn("w:val"), "TOCField")
    field_p_pr.append(field_style)
    run = field.add_run("")
    instr = OxmlElement("w:instrText")
    instr.text = ' TOC \\\\o "1-3" \\\\h \\\\z \\\\u '
    run._r.append(instr)

    _add_heading(doc, "1 绪论", level=1, size=30)
    doc.save(source_path)

    verification = build_scope_verify(str(source_path), profile_path="lnu", scopes=["toc"])

    assert verification["overall_status"] == "verified"
    assert verification["readiness"] == "render-check-required"
    assert verification["render_check_rule_ids"] == ["TOC_REFRESH_REQUIRED"]

    rendered = render_scope_verify(verification)
    assert "结果: 所选范围结构规则已通过，但仍需做渲染复核。" in rendered
    assert "TOC_REFRESH_REQUIRED" in rendered


def test_build_document_diagnostics_detects_lnu_preface_zero_based_mismatch(tmp_path):
    source_path = Path(tmp_path) / "diagnostics_preface_status.docx"
    doc = Document()
    _add_heading(doc, "1 序言", level=1, size=30)
    _add_heading(doc, "1.1 研究背景", level=2, size=30)
    _add_heading(doc, "2 材料与方法", level=1, size=30)
    doc.save(source_path)

    diagnostics = build_document_diagnostics(str(source_path), profile_path="lnu")

    assert diagnostics["preface_status"] == "zero_based_mismatch"
    assert any("--renumber-headings" in action for action in diagnostics["recommended_actions"])


def test_build_document_diagnostics_reports_zero_based_lnu_preface(tmp_path):
    source_path = Path(tmp_path) / "diagnostics_lnu_preface.docx"
    doc = Document()
    _add_heading(doc, "序  言", level=1, size=30)
    _add_heading(doc, "0.1 研究背景", level=2, size=28)
    _add_heading(doc, "0.1.1 研究现状", level=3, size=24)
    _add_heading(doc, "1 材料与方法", level=1, size=30)
    doc.save(source_path)

    diagnostics = build_document_diagnostics(str(source_path), profile_path="lnu")

    assert diagnostics["preface_status"] == "zero_based_ok"

    rendered = render_document_diagnostics(diagnostics)
    assert "辽大序言编号: zero_based_ok" in rendered


def test_render_document_diagnostics_compact_exposes_stable_summary_keys(tmp_path):
    source_path = Path(tmp_path) / "diagnostics_compact_summary.docx"
    doc = Document()
    _add_heading(doc, "1 序言", level=1, size=30)
    _add_heading(doc, "1.1 研究背景", level=2, size=30)
    doc.save(source_path)

    diagnostics = build_document_diagnostics(str(source_path), profile_path="lnu")
    rendered = render_document_diagnostics_compact(diagnostics)

    assert "file=diagnostics_compact_summary.docx" in rendered
    assert "profile=lnu-checker-2026 (requested: lnu)" in rendered
    assert "toc_status=no_toc" in rendered
    assert "preface_status=zero_based_mismatch" in rendered
    assert "recommended_action_count=" in rendered
