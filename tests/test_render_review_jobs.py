from types import SimpleNamespace
import sys
import time
from pathlib import Path

from fastapi.testclient import TestClient
from docx import Document
from PIL import Image
import pytest
import release_smoke

from article_api.app import create_app
from article_api import storage
from article_api.job_artifacts import build_artifacts, refresh_artifact_availability
from article_api.job_execution import handler_request
from article_api.jobs import _render_retry_options
from article_api.render_review_jobs import build_render_review_job_kwargs


DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _write_docx(path):
    doc = Document()
    doc.add_heading("1 绪论", level=1)
    for paragraph in release_smoke.SMOKE_BODY_ANCHORS:
        doc.add_paragraph(paragraph)
    doc.save(path)


def _write_pdf(path, page_count):
    if page_count == 1:
        release_smoke._build_smoke_pdf(path, python_executable=Path(sys.executable))
        return
    pages = [Image.new("RGB", (120, 160), "white") for _ in range(page_count)]
    pages[0].save(path, "PDF", save_all=True, append_images=pages[1:], resolution=72)
    for page in pages:
        page.close()


def _wait_for_job(client, job_id):
    deadline = time.monotonic() + 20
    while time.monotonic() < deadline:
        payload = client.get(f"/jobs/{job_id}").json()
        if payload["status"] in {"succeeded", "failed"}:
            return payload
        time.sleep(0.05)
    raise AssertionError(f"job did not finish: {job_id}")


def _stored_render_job(job_id, *, evidence_source, evidence_trust, content_match=None):
    summary = {
        "render_evidence_status": "render-evidence-ready",
        "evidence_source": evidence_source,
        "evidence_trust": evidence_trust,
        "evidence_authoritative": evidence_trust == "authoritative",
        "layout_decision_eligible": True,
    }
    if content_match is not None:
        summary.update(content_match)
    return {
        "job_id": job_id,
        "operation": "render-verify",
        "status": "succeeded",
        "mode": "background",
        "created_at": "2026-06-16T00:00:00Z",
        "updated_at": "2026-06-16T00:00:01Z",
        "request": {},
        "resolved_request": {"workflow_mode": "default_user"},
        "workspace": None,
        "runtime": None,
        "summary": dict(summary),
        "artifacts": [],
        "result": {
            **summary,
            "summary": dict(summary),
            "evidence_items": [{"page": 1, "rule_id": "render.object_flow"}],
        },
        "error": None,
    }


def test_legacy_manual_pdf_job_history_fails_closed_without_rewriting_storage(tmp_path, monkeypatch):
    monkeypatch.setenv("ARTICLE_API_STATE_ROOT", str(tmp_path / "state"))
    toc_output = tmp_path / "legacy-static-toc.docx"
    toc_output.write_bytes(b"unsafe legacy toc")
    stored = _stored_render_job(
        "legacy-manual-pdf",
        evidence_source="manual-pdf",
        evidence_trust="user-confirmed",
    )
    stored["summary"]["toc_output_available"] = True
    stored["artifacts"] = [{"role": "toc-output", "path": str(toc_output)}]
    stored["result"]["summary"]["toc_output_available"] = True
    stored["result"]["toc_finalization"] = {
        "status": "generated",
        "available": True,
        "output_path": str(toc_output),
    }
    storage.upsert_job(stored)

    with TestClient(create_app()) as client:
        status_payload = client.get("/jobs/legacy-manual-pdf").json()
        list_payload = client.get("/jobs").json()
        result_payload = client.get("/jobs/legacy-manual-pdf/result").json()
        download_response = client.get("/jobs/legacy-manual-pdf/artifacts/toc-output/download")

    listed = next(item for item in list_payload if item["job_id"] == "legacy-manual-pdf")
    for summary in (status_payload["summary"], listed["summary"], result_payload["result"]["summary"]):
        assert summary["render_evidence_status"] == "unsupported-evidence"
        assert summary["layout_decision_eligible"] is False
        assert summary["evidence_trust"] == "unverified"
        assert summary["pdf_content_match_status"] == "not-verified"
        assert summary["pdf_content_matched"] is False
        assert summary["toc_output_available"] is False

    assert all(artifact["role"] != "toc-output" for artifact in status_payload["artifacts"])
    assert result_payload["result"]["toc_finalization"]["available"] is False
    assert download_response.status_code == 404

    persisted = storage.get_job("legacy-manual-pdf", include_result=True)
    assert persisted["summary"]["layout_decision_eligible"] is True
    assert persisted["result"]["summary"]["render_evidence_status"] == "render-evidence-ready"
    assert persisted["result"]["toc_finalization"]["available"] is True


@pytest.mark.parametrize(
    ("job_id", "evidence_source", "evidence_trust", "content_match"),
    [
        ("authoritative-word", "word-pdf", "authoritative", None),
        (
            "matched-manual",
            "manual-pdf",
            "user-confirmed",
            {"pdf_content_match_status": "matched", "pdf_content_matched": True},
        ),
    ],
)
def test_trusted_historical_render_evidence_remains_eligible(
    tmp_path,
    monkeypatch,
    job_id,
    evidence_source,
    evidence_trust,
    content_match,
):
    monkeypatch.setenv("ARTICLE_API_STATE_ROOT", str(tmp_path / "state"))
    storage.upsert_job(
        _stored_render_job(
            job_id,
            evidence_source=evidence_source,
            evidence_trust=evidence_trust,
            content_match=content_match,
        )
    )

    with TestClient(create_app()) as client:
        payload = client.get(f"/jobs/{job_id}/result").json()["result"]

    assert payload["summary"]["render_evidence_status"] == "render-evidence-ready"
    assert payload["summary"]["layout_decision_eligible"] is True
    assert payload["summary"]["evidence_trust"] == evidence_trust


def test_render_review_job_requires_pdf_upload(tmp_path, monkeypatch):
    monkeypatch.setenv("ARTICLE_RUNTIME_ROOT", str(tmp_path / "runtime"))
    client = TestClient(create_app())
    docx_path = tmp_path / "demo.docx"
    _write_docx(docx_path)

    docx_response = client.post(
        "/uploads/docx",
        files={"file": ("demo.docx", docx_path.read_bytes(), DOCX_MEDIA_TYPE)},
    )
    docx_upload_id = docx_response.json()["upload_id"]

    response = client.post(
        f"/uploads/{docx_upload_id}/render-review-jobs",
        json={"pdf_upload_id": "missing-pdf"},
    )

    assert response.status_code == 404
    assert response.json()["detail"]["user_message"] == "没有找到这个 PDF 文件。"


def test_render_review_job_creates_background_render_verify_job(tmp_path, monkeypatch):
    monkeypatch.setenv("ARTICLE_RUNTIME_ROOT", str(tmp_path / "runtime"))
    client = TestClient(create_app())
    docx_path = tmp_path / "demo.docx"
    pdf_path = tmp_path / "demo.pdf"
    _write_docx(docx_path)
    pdf_path.write_bytes(b"%PDF-1.4\n%%EOF\n")

    docx_response = client.post(
        "/uploads/docx",
        files={"file": ("demo.docx", docx_path.read_bytes(), DOCX_MEDIA_TYPE)},
    )
    pdf_response = client.post(
        "/uploads/pdf",
        files={"file": ("demo.pdf", pdf_path.read_bytes(), "application/pdf")},
    )

    response = client.post(
        f"/uploads/{docx_response.json()['upload_id']}/render-review-jobs",
        json={
            "pdf_upload_id": pdf_response.json()["upload_id"],
            "pdf_matches_docx_confirmed": True,
        },
    )

    assert response.status_code == 201
    payload = response.json()
    assert payload["operation"] == "render-verify"
    assert payload["status"] == "queued"
    assert payload["resolved_request"]["workflow_mode"] == "default_user"
    assert payload["request"]["pdf_matches_docx_confirmed"] is True
    assert payload["request"]["generate_static_toc"] is True
    assert payload["resolved_request"]["pdf_matches_docx_confirmed"] is True
    assert payload["resolved_request"]["generate_static_toc"] is True
    assert payload["resolved_request"]["rendered_pdf"].endswith(".pdf")
    assert payload["workspace"] is not None
    assert f"jobs/{payload['job_id']}/inputs" in payload["resolved_request"]["file_path"]


def test_build_render_review_job_kwargs_links_docx_and_pdf_uploads():
    uploads = {
        "docx-1": {
            "file_name": "论文.docx",
            "stored_path": "/tmp/runtime/uploads/paper.docx",
            "runtime_root": "/tmp/runtime",
        },
        "pdf-1": {
            "file_name": "论文.pdf",
            "stored_path": "/tmp/runtime/uploads/paper.pdf",
            "runtime_root": "/tmp/runtime",
        },
    }
    request = SimpleNamespace(
        pdf_upload_id="pdf-1",
        pdf_matches_docx_confirmed=True,
        profile="lnu",
        strict_profile=None,
        scopes=None,
        max_attempts=1,
        retry_delay_seconds=0.0,
        timeout_seconds=None,
        generate_static_toc=True,
    )

    payload = build_render_review_job_kwargs(
        "docx-1",
        request,
        resolve_upload_fn=lambda upload_id: uploads[upload_id],
    )

    assert payload["file_path"] == "/tmp/runtime/uploads/paper.docx"
    assert payload["rendered_pdf"] == "/tmp/runtime/uploads/paper.pdf"
    assert payload["pdf_matches_docx_confirmed"] is True
    assert payload["generate_static_toc"] is True
    assert payload["workflow_mode"] == "default_user"
    assert payload["source_display_name"] == "论文.docx"
    assert payload["pdf_display_name"] == "论文.pdf"
    assert payload["runtime_root"] == "/tmp/runtime"
    assert payload["stage_input"] is True
    assert payload["_public_request"] == {
        "docx_upload_id": "docx-1",
        "pdf_upload_id": "pdf-1",
        "pdf_matches_docx_confirmed": True,
        "generate_static_toc": True,
    }


def test_render_review_worker_request_omits_display_metadata():
    request = handler_request(
        {
            "file_path": "/tmp/paper.docx",
            "rendered_pdf": "/tmp/paper.pdf",
            "docx_upload_id": "docx-1",
            "pdf_upload_id": "pdf-1",
            "workflow_mode": "default_user",
            "pdf_display_name": "论文.pdf",
            "pdf_matches_docx_confirmed": True,
            "generate_static_toc": True,
        }
    )

    assert request == {
        "file_path": "/tmp/paper.docx",
        "rendered_pdf": "/tmp/paper.pdf",
        "pdf_matches_docx_confirmed": True,
        "generate_static_toc": True,
    }


def test_render_retry_options_preserve_pdf_docx_confirmation_and_static_toc_request():
    options = _render_retry_options(
        "render-verify",
        {"pdf_matches_docx_confirmed": True, "generate_static_toc": True},
    )

    assert options == {"pdf_matches_docx_confirmed": True, "generate_static_toc": True}


def test_render_review_artifacts_register_generated_static_toc(tmp_path):
    output_path = tmp_path / "static_toc.docx"
    output_path.write_bytes(b"docx")
    report_path = tmp_path / "render_verify_report.md"
    report_path.write_text("# PDF 复核报告\n", encoding="utf-8")

    artifacts = build_artifacts(
        "render-verify",
        {},
        None,
        result={
            "report_path": str(report_path),
            "toc_finalization": {
                "available": True,
                "output_path": str(output_path),
            }
        },
        status="succeeded",
    )
    refreshed = refresh_artifact_availability(artifacts)
    by_role = {artifact["role"]: artifact for artifact in refreshed}

    assert set(by_role) == {"report", "toc-output"}
    assert by_role["report"]["kind"] == "markdown"
    assert by_role["report"]["path"] == str(report_path)
    assert by_role["report"]["available"] is True
    assert by_role["toc-output"]["kind"] == "docx"
    assert by_role["toc-output"]["path"] == str(output_path)
    assert by_role["toc-output"]["available"] is True


def test_build_render_review_job_kwargs_rejects_non_pdf_upload():
    uploads = {
        "docx-1": {"file_name": "论文.docx", "stored_path": "/tmp/paper.docx"},
        "pdf-1": {"file_name": "截图.png", "stored_path": "/tmp/page.png"},
    }
    request = SimpleNamespace(
        pdf_upload_id="pdf-1",
        pdf_matches_docx_confirmed=False,
        profile="lnu",
        strict_profile=None,
        scopes=None,
        max_attempts=1,
        retry_delay_seconds=0.0,
        timeout_seconds=None,
    )

    with pytest.raises(ValueError, match="PDF"):
        build_render_review_job_kwargs(
            "docx-1",
            request,
            resolve_upload_fn=lambda upload_id: uploads[upload_id],
        )


def test_render_review_jobs_keep_page_evidence_in_separate_workspaces(tmp_path, monkeypatch):
    monkeypatch.setenv("ARTICLE_API_STATE_ROOT", str(tmp_path / "state"))
    monkeypatch.setenv("ARTICLE_API_RUNTIME_ROOT", str(tmp_path / "runtime"))
    docx_path = tmp_path / "paper.docx"
    first_pdf = tmp_path / "first.pdf"
    second_pdf = tmp_path / "second.pdf"
    _write_docx(docx_path)
    _write_pdf(first_pdf, 2)
    _write_pdf(second_pdf, 1)

    with TestClient(create_app()) as client:
        docx_upload = client.post(
            "/uploads/docx",
            files={"file": (docx_path.name, docx_path.read_bytes(), DOCX_MEDIA_TYPE)},
        ).json()
        results = []
        for pdf_path in (first_pdf, second_pdf):
            pdf_upload = client.post(
                "/uploads/pdf",
                files={"file": (pdf_path.name, pdf_path.read_bytes(), "application/pdf")},
            ).json()
            job = client.post(
                f"/uploads/{docx_upload['upload_id']}/render-review-jobs",
                json={"pdf_upload_id": pdf_upload["upload_id"]},
            ).json()
            assert _wait_for_job(client, job["job_id"])["status"] == "succeeded"
            results.append(client.get(f"/jobs/{job['job_id']}/result").json()["result"])

    assert results[0]["page_count"] == 2
    assert results[1]["page_count"] == 1
    assert results[0]["output_dir"] != results[1]["output_dir"]
    assert all(Path(page_path).exists() for page_path in results[0]["page_images"])


def test_render_review_job_retry_restores_both_uploads(tmp_path, monkeypatch):
    monkeypatch.setenv("ARTICLE_API_STATE_ROOT", str(tmp_path / "state"))
    monkeypatch.setenv("ARTICLE_API_RUNTIME_ROOT", str(tmp_path / "runtime"))
    docx_path = tmp_path / "paper.docx"
    pdf_path = tmp_path / "paper.pdf"
    _write_docx(docx_path)
    _write_pdf(pdf_path, 1)

    with TestClient(create_app()) as client:
        docx_upload = client.post(
            "/uploads/docx",
            files={"file": (docx_path.name, docx_path.read_bytes(), DOCX_MEDIA_TYPE)},
        ).json()
        pdf_upload = client.post(
            "/uploads/pdf",
            files={"file": (pdf_path.name, pdf_path.read_bytes(), "application/pdf")},
        ).json()
        first = client.post(
            f"/uploads/{docx_upload['upload_id']}/render-review-jobs",
            json={
                "pdf_upload_id": pdf_upload["upload_id"],
                "pdf_matches_docx_confirmed": True,
                "scopes": ["toc"],
            },
        ).json()
        first_status = _wait_for_job(client, first["job_id"])
        assert first_status["status"] == "succeeded"
        assert first_status["summary"]["pdf_matches_docx_confirmed"] is True
        assert first_status["summary"]["pdf_name"] == "paper.pdf"
        assert first_status["summary"]["business_status"] == "render-review-required"
        assert first_status["summary"]["render_evidence_status"] == "render-review-required"
        assert first_status["summary"]["evidence_trust"] == "user-confirmed"
        assert first_status["summary"]["layout_decision_eligible"] is True
        assert first_status["summary"]["pdf_content_match_status"] == "matched"
        assert first_status["summary"]["pdf_content_matched"] is True
        first_result = client.get(f"/jobs/{first['job_id']}/result").json()["result"]
        assert first_result["pdf_matches_docx_confirmed"] is True
        assert first_result["summary"]["pdf_matches_docx_confirmed"] is True
        assert first_result["evidence_trust"] == "user-confirmed"
        assert first_result["evidence_authoritative"] is False
        assert first_result["layout_decision_eligible"] is True
        assert first_result["render_evidence_status"] == "render-review-required"

        retry_response = client.post(f"/jobs/{first['job_id']}/retry")

        assert retry_response.status_code == 201
        retried = retry_response.json()
        retried_status = _wait_for_job(client, retried["job_id"])
        assert retried_status["status"] == "succeeded", client.get(
            f"/jobs/{retried['job_id']}/result"
        ).json()
        assert retried["request"]["docx_upload_id"] == docx_upload["upload_id"]
        assert retried["request"]["pdf_upload_id"] == pdf_upload["upload_id"]
        assert retried["request"]["pdf_matches_docx_confirmed"] is True
        assert retried["request"]["generate_static_toc"] is True
        assert retried["workspace"] is not None
        assert retried["resolved_request"]["stage_input"] is True
        assert retried["resolved_request"]["scopes"] == ["toc"]
        assert retried["resolved_request"]["pdf_matches_docx_confirmed"] is True
        assert retried["resolved_request"]["generate_static_toc"] is True
        assert retried_status["summary"]["pdf_matches_docx_confirmed"] is True
        assert retried_status["summary"]["pdf_name"] == "paper.pdf"
        retried_result = client.get(f"/jobs/{retried['job_id']}/result").json()["result"]
        assert retried_result["pdf_matches_docx_confirmed"] is True
        assert retried_result["summary"]["pdf_matches_docx_confirmed"] is True
        assert retried_result["output_dir"] != first_result["output_dir"]
