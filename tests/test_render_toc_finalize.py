from pathlib import Path

from docx import Document

from thesis_tool.render_toc_evidence import verify_docx_pdf_content_match
from thesis_tool.render_toc_finalize import build_static_toc_finalization


TOC_MISMATCH = [{"id": "toc_page_number_mismatch"}]


def _matching_docx_and_pdf(tmp_path):
    input_docx = tmp_path / "paper.docx"
    document = Document()
    document.add_heading("1 绪论", level=1)
    anchors = [
        "论文格式检查工具用于识别目录页码和正文标题",
        "系统通过多项规则复核论文格式问题并输出证据",
        "实验结果表明目录映射能够覆盖正文主要章节",
    ]
    for anchor in anchors:
        document.add_paragraph(anchor)
    document.save(input_docx)
    return input_docx, {1: "目录\n1 绪论 ...... 2", 2: "\n".join(anchors)}


def test_static_toc_finalization_is_opt_in(tmp_path):
    result = build_static_toc_finalization(
        "paper.docx",
        tmp_path,
        {},
        requested=False,
        pdf_matches_docx_confirmed=True,
        toc_findings=TOC_MISMATCH,
    )

    assert result == {
        "status": "not-requested",
        "available": False,
        "message": "本次只完成 PDF 复核",
        "next_action": "需要时再生成静态目录版",
    }


def test_static_toc_finalization_requires_confirmed_pdf(tmp_path):
    result = build_static_toc_finalization(
        "paper.docx",
        tmp_path,
        {1: "目录"},
        requested=True,
        pdf_matches_docx_confirmed=False,
        toc_findings=TOC_MISMATCH,
    )

    assert result["status"] == "blocked"
    assert result["available"] is False
    assert "确认" in result["message"]


def test_static_toc_finalization_skips_when_current_pdf_has_no_toc_issue(tmp_path):
    result = build_static_toc_finalization(
        "paper.docx",
        tmp_path,
        {1: "目录"},
        requested=True,
        pdf_matches_docx_confirmed=True,
        toc_findings=[],
    )

    assert result["status"] == "not-needed"
    assert result["available"] is False


def test_static_toc_finalization_exposes_generated_output(tmp_path):
    input_docx, page_texts = _matching_docx_and_pdf(tmp_path)
    captured = {}

    def fake_finalize(input_docx, output_path, page_texts):
        captured.update(input_docx=str(input_docx), output_path=str(output_path), page_texts=page_texts)
        Path(output_path).write_bytes(b"docx")
        return {
            "status": "finalized",
            "complete": True,
            "entry_count": 12,
            "mapped_count": 12,
            "entries": [],
            "unmatched_titles": [],
            "written": True,
            "output_path": str(output_path),
        }

    result = build_static_toc_finalization(
        input_docx,
        tmp_path,
        page_texts,
        requested=True,
        pdf_matches_docx_confirmed=True,
        toc_findings=TOC_MISMATCH,
        finalize_fn=fake_finalize,
    )

    assert result["status"] == "generated"
    assert result["available"] is True
    assert result["entry_count"] == 12
    assert result["output_path"] == str((tmp_path / "static_toc.docx").resolve())
    assert captured["input_docx"] == str(input_docx)
    assert captured["page_texts"] == page_texts


def test_static_toc_finalization_does_not_expose_incomplete_output(tmp_path):
    input_docx, page_texts = _matching_docx_and_pdf(tmp_path)

    def fake_finalize(*_args):
        return {
            "status": "incomplete",
            "complete": False,
            "entry_count": 4,
            "mapped_count": 3,
            "entries": [],
            "unmatched_titles": ["参考文献"],
            "written": False,
            "output_path": None,
        }

    result = build_static_toc_finalization(
        input_docx,
        tmp_path,
        page_texts,
        requested=True,
        pdf_matches_docx_confirmed=True,
        toc_findings=TOC_MISMATCH,
        finalize_fn=fake_finalize,
    )

    assert result["status"] == "incomplete"
    assert result["available"] is False
    assert result["unmatched_titles"] == ["参考文献"]
    assert "不会写入" in result["message"]


def test_static_toc_finalization_blocks_mismatched_docx_and_pdf(tmp_path):
    input_docx, _page_texts = _matching_docx_and_pdf(tmp_path)
    called = False

    def fake_finalize(*_args):
        nonlocal called
        called = True
        return {}

    result = build_static_toc_finalization(
        input_docx,
        tmp_path,
        {
            1: "目录\n1 绪论 ...... 2",
            2: "这是一篇主题完全不同的论文正文",
            3: "其研究对象和实验结论均不相同",
        },
        requested=True,
        pdf_matches_docx_confirmed=True,
        toc_findings=TOC_MISMATCH,
        finalize_fn=fake_finalize,
    )

    assert result["status"] == "blocked"
    assert result["available"] is False
    assert result["reason"] == "content-mismatch"
    assert result["content_match"]["matched_anchor_count"] == 0
    assert called is False


def test_static_toc_finalization_downgrades_finalize_error(tmp_path):
    input_docx, page_texts = _matching_docx_and_pdf(tmp_path)

    def failing_finalize(*_args):
        raise ValueError("DOCX 中未找到目录位置")

    result = build_static_toc_finalization(
        input_docx,
        tmp_path,
        page_texts,
        requested=True,
        pdf_matches_docx_confirmed=True,
        toc_findings=TOC_MISMATCH,
        finalize_fn=failing_finalize,
    )

    assert result["status"] == "unavailable"
    assert result["available"] is False
    assert result["reason"] == "finalization-error"
    assert "静态目录不可生成" in result["message"]
    assert "DOCX 中未找到目录位置" in result["detail"]


def test_content_match_requires_distinct_docx_anchors(tmp_path):
    input_docx = tmp_path / "repeated.docx"
    document = Document()
    document.add_heading("1 绪论", level=1)
    repeated = "论文格式检查工具用于识别目录页码和正文标题"
    for _index in range(3):
        document.add_paragraph(repeated)
    document.save(input_docx)

    result = verify_docx_pdf_content_match(input_docx, {1: repeated})

    assert result["status"] == "insufficient-evidence"
    assert result["matched"] is False
    assert result["anchor_count"] == 1


def test_content_match_does_not_count_toc_entries_as_body_anchors(tmp_path):
    input_docx = tmp_path / "toc_entries.docx"
    document = Document()
    document.add_heading("目录", level=1)
    toc_entries = [
        "1 论文格式检查研究背景\t1",
        "2 论文格式检查系统设计\t5",
        "3 论文格式检查实验结果\t9",
    ]
    for entry in toc_entries:
        document.add_paragraph(entry)
    document.add_heading("1 绪论", level=1)
    body_anchor = "论文格式检查工具用于识别目录页码和正文标题"
    document.add_paragraph(body_anchor)
    document.save(input_docx)

    result = verify_docx_pdf_content_match(
        input_docx,
        {1: "目录\n" + "\n".join(toc_entries), 2: "1 绪论\n" + body_anchor},
    )

    assert result["status"] == "insufficient-evidence"
    assert result["matched"] is False
    assert result["anchor_count"] == 1
