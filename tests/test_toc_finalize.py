from __future__ import annotations

import io
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement

from thesis_tool.toc_finalize import analyze_static_toc, finalize_static_toc


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
NSMAP = {"w": W_NS}


def _make_toc_docx(path: Path) -> Path:
    document = Document()
    document.styles.add_style("TOC1", WD_STYLE_TYPE.PARAGRAPH)
    document.styles.add_style("TOC2", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("目  录")
    document.add_paragraph("第1章 绪论\t8", style="TOC1")
    document.add_paragraph("1.1 研究背景\t9", style="TOC2")
    document.add_paragraph("第1章 绪论", style="Heading 1")
    document.add_paragraph("1.1 研究背景", style="Heading 2")
    document.save(path)
    return path


def _add_field_paragraph(document: Document, field_type: str, *, styled: bool = True) -> None:
    style_name = "TOCField" if field_type == "begin" else "TOCEnd"
    if styled and style_name not in document.styles:
        document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    paragraph = document.add_paragraph(style=style_name if styled else None)
    field = OxmlElement("w:fldChar")
    field.set(f"{{{W_NS}}}fldCharType", field_type)
    paragraph.add_run()._r.append(field)
    if field_type == "begin":
        instruction = OxmlElement("w:instrText")
        instruction.text = ' TOC \\o "1-3" '
        paragraph.add_run()._r.append(instruction)
        separator = OxmlElement("w:fldChar")
        separator.set(f"{{{W_NS}}}fldCharType", "separate")
        paragraph.add_run()._r.append(separator)


def _make_field_toc_docx(path: Path, *, styled_end: bool = True) -> Path:
    document = Document()
    for style_name in ("TOC1", "TOC2"):
        document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("封面信息")
    document.add_paragraph("目  录")
    _add_field_paragraph(document, "begin")
    document.add_paragraph("第1章 绪论\t8", style="TOC1")
    document.add_paragraph("1.1 研究背景\t9", style="TOC2")
    _add_field_paragraph(document, "end", styled=styled_end)
    document.add_paragraph("第1章 绪论", style="Heading 1")
    document.add_paragraph("正文保持不变")
    document.add_picture(str(Path(__file__).parents[1] / "scripts/article_api/static/assets/lnu-emblem.jpg"))
    document.add_paragraph("1.1 研究背景", style="Heading 2")
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(f"{{{W_NS}}}val", "true")
    document.settings.element.append(update_fields)
    document.save(path)
    return path


def _make_empty_field_toc_docx(path: Path) -> Path:
    document = Document()
    for style_name in ("TOC1", "TOC2"):
        document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("目  录")
    _add_field_paragraph(document, "begin")
    _add_field_paragraph(document, "end")
    document.add_paragraph("第1章 绪论", style="Heading 1")
    document.add_paragraph("1.1 研究背景", style="Heading 2")
    document.save(path)
    return path


def _make_embedded_field_toc_docx(path: Path) -> Path:
    document = Document()
    for style_name in ("TOC1", "TOC2"):
        document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("目  录")
    first = document.add_paragraph(style="TOC1")
    begin = OxmlElement("w:fldChar")
    begin.set(f"{{{W_NS}}}fldCharType", "begin")
    first.add_run()._r.append(begin)
    instruction = OxmlElement("w:instrText")
    instruction.text = ' TOC \\o "1-3" '
    first.add_run()._r.append(instruction)
    separate = OxmlElement("w:fldChar")
    separate.set(f"{{{W_NS}}}fldCharType", "separate")
    first.add_run()._r.append(separate)
    first.add_run("第1章 绪论\t8")
    last = document.add_paragraph("1.1 研究背景\t9", style="TOC2")
    end = OxmlElement("w:fldChar")
    end.set(f"{{{W_NS}}}fldCharType", "end")
    last.add_run()._r.append(end)
    document.add_paragraph("第1章 绪论", style="Heading 1")
    document.add_paragraph("1.1 研究背景", style="Heading 2")
    document.save(path)
    return path


def test_analyze_static_toc_maps_visible_entries_to_printed_pdf_pages(tmp_path: Path):
    source_path = _make_toc_docx(tmp_path / "source.docx")

    result = analyze_static_toc(
        source_path,
        {
            1: "目 录\n第1章 绪论 ...... 8\n1.1 研究背景 ...... 9",
            4: "第1章 绪论\n正文\n\u2010 1 \u2010",
            5: "1.1 研究背景\n正文\n\u2212 2 \u2212",
        },
    )

    assert result["status"] == "ready"
    assert result["complete"] is True
    assert result["entry_count"] == 2
    assert result["mapped_count"] == 2
    assert result["entries"] == [
        {"title": "第1章 绪论", "level": 1, "pdf_page": 4, "page": 1},
        {"title": "1.1 研究背景", "level": 2, "pdf_page": 5, "page": 2},
    ]
    assert result["unmatched_titles"] == []


def test_analyze_static_toc_merges_partial_body_headings_with_existing_toc_entries(tmp_path: Path):
    source_path = tmp_path / "source.docx"
    document = Document()
    document.styles.add_style("TOC1", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("目  录")
    document.add_paragraph("第1章 绪论\t1", style="TOC1")
    document.add_paragraph("参考文献\t2", style="TOC1")
    document.add_paragraph("致谢\t3", style="TOC1")
    document.add_paragraph("第1章 绪论", style="Heading 1")
    document.add_paragraph("参考文献")
    document.add_paragraph("致谢")
    document.save(source_path)

    result = analyze_static_toc(
        source_path,
        {
            2: "第1章 绪论\n- 1 -",
            3: "参考文献\n- 2 -",
            4: "致谢\n- 3 -",
        },
    )

    assert result["complete"] is True
    assert [(entry["title"], entry["level"]) for entry in result["entries"]] == [
        ("第1章 绪论", 1),
        ("参考文献", 1),
        ("致谢", 1),
    ]


def test_finalize_static_toc_preserves_repeated_heading_titles(tmp_path: Path):
    source_path = tmp_path / "source.docx"
    output_path = tmp_path / "output.docx"
    document = Document()
    for style_name in ("TOC1", "TOC2"):
        document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("目  录")
    entries = [
        ("第1章 方法", "TOC1", "Heading 1"),
        ("1.1 本章小结", "TOC2", "Heading 2"),
        ("第2章 结果", "TOC1", "Heading 1"),
        ("1.1 本章小结", "TOC2", "Heading 2"),
    ]
    for index, (title, toc_style, _) in enumerate(entries, start=1):
        document.add_paragraph(f"{title}\t{index}", style=toc_style)
    for title, _, heading_style in entries:
        document.add_paragraph(title, style=heading_style)
    document.save(source_path)
    page_texts = {
        4: "第1章 方法\n- 1 -",
        5: "1.1 本章小结\n- 2 -",
        6: "第2章 结果\n- 3 -",
        7: "1.1 本章小结\n- 4 -",
    }

    result = finalize_static_toc(source_path, output_path, page_texts)

    assert result["complete"] is True
    assert result["entry_count"] == 4
    assert result["mapped_count"] == 4
    output = Document(output_path)
    assert [p.text for p in output.paragraphs if p.style.style_id in {"TOC1", "TOC2"}] == [
        "第1章 方法\t1",
        "1.1 本章小结\t2",
        "第2章 结果\t3",
        "1.1 本章小结\t4",
    ]


def test_finalize_static_toc_writes_plain_entries_and_preserves_body_and_media(tmp_path: Path):
    source_path = _make_field_toc_docx(tmp_path / "source.docx")
    output_path = tmp_path / "finalized.docx"
    page_texts = {
        1: "目 录\n第1章 绪论 ...... 8\n1.1 研究背景 ...... 9",
        4: "第1章 绪论\n正文保持不变\n- 1 -",
        5: "1.1 研究背景\n- 2 -",
    }

    result = finalize_static_toc(source_path, output_path, page_texts)

    assert result["status"] == "finalized"
    assert result["written"] is True
    assert output_path.exists()
    output = Document(output_path)
    toc_texts = [p.text for p in output.paragraphs if p.style.style_id in {"TOC1", "TOC2", "TOC3"}]
    assert toc_texts == ["第1章 绪论\t1", "1.1 研究背景\t2"]
    assert "正文保持不变" in [paragraph.text for paragraph in output.paragraphs]

    with zipfile.ZipFile(source_path) as source_zip, zipfile.ZipFile(output_path) as output_zip:
        source_media = {name: source_zip.read(name) for name in source_zip.namelist() if name.startswith("word/media/")}
        output_media = {name: output_zip.read(name) for name in output_zip.namelist() if name.startswith("word/media/")}
        document_root = ET.fromstring(output_zip.read("word/document.xml"))
        settings_root = ET.fromstring(output_zip.read("word/settings.xml"))

    assert output_media == source_media
    assert document_root.findall(".//w:instrText", NSMAP) == []
    assert document_root.findall(".//w:fldChar", NSMAP) == []
    update_fields = settings_root.find("w:updateFields", NSMAP)
    assert update_fields is not None
    assert update_fields.get(f"{{{W_NS}}}val") == "true"


def test_finalize_static_toc_does_not_write_when_mapping_is_incomplete(tmp_path: Path):
    source_path = _make_field_toc_docx(tmp_path / "source.docx")
    output_path = tmp_path / "existing.docx"
    output_path.write_bytes(b"keep-existing-output")

    result = finalize_static_toc(
        source_path,
        output_path,
        {
            1: "目 录\n第1章 绪论 ...... 8\n1.1 研究背景 ...... 9",
            4: "第1章 绪论\n正文保持不变\n- 1 -",
        },
    )

    assert result["status"] == "incomplete"
    assert result["complete"] is False
    assert result["written"] is False
    assert result["mapped_count"] == 1
    assert result["unmatched_titles"] == ["1.1 研究背景"]
    assert output_path.read_bytes() == b"keep-existing-output"


def test_finalize_static_toc_rejects_resolved_output_path_that_overwrites_input(tmp_path: Path):
    source_path = _make_toc_docx(tmp_path / "source.docx")
    (tmp_path / "nested").mkdir()
    equivalent_output_path = tmp_path / "nested" / ".." / source_path.name
    source_bytes = source_path.read_bytes()

    with pytest.raises(ValueError, match="输入.*输出"):
        finalize_static_toc(
            source_path,
            equivalent_output_path,
            {
                4: "第1章 绪论\n正文\n- 1 -",
                5: "1.1 研究背景\n正文\n- 2 -",
            },
        )

    assert source_path.read_bytes() == source_bytes


def test_finalize_static_toc_is_idempotent(tmp_path: Path):
    source_path = _make_field_toc_docx(tmp_path / "source.docx")
    first_path = tmp_path / "first.docx"
    second_path = tmp_path / "second.docx"
    page_texts = {
        1: "目 录\n第1章 绪论 ...... 8\n1.1 研究背景 ...... 9",
        4: "第1章 绪论\n正文保持不变\n- 1 -",
        5: "1.1 研究背景\n- 2 -",
    }

    finalize_static_toc(source_path, first_path, page_texts)
    result = finalize_static_toc(first_path, second_path, page_texts)

    with zipfile.ZipFile(first_path) as first_zip, zipfile.ZipFile(second_path) as second_zip:
        first_parts = {name: first_zip.read(name) for name in first_zip.namelist()}
        second_parts = {name: second_zip.read(name) for name in second_zip.namelist()}

    assert result["status"] == "finalized"
    assert result["written"] is True
    assert second_parts == first_parts


def test_finalize_static_toc_builds_entries_when_word_field_cache_is_empty(tmp_path: Path):
    source_path = _make_empty_field_toc_docx(tmp_path / "source.docx")
    output_path = tmp_path / "output.docx"
    page_texts = {
        1: "目 录\n第1章 绪论 ...... 8\n1.1 研究背景 ...... 9",
        4: "第1章 绪论\n- 1 -",
        5: "1.1 研究背景\n- 2 -",
    }

    result = finalize_static_toc(source_path, output_path, page_texts)

    assert result["complete"] is True
    assert result["mapped_count"] == 2
    output = Document(output_path)
    assert [p.text for p in output.paragraphs if p.style.style_id in {"TOC1", "TOC2"}] == [
        "第1章 绪论\t1",
        "1.1 研究背景\t2",
    ]


def test_finalize_static_toc_preserves_entries_when_field_markers_share_their_paragraphs(tmp_path: Path):
    source_path = _make_embedded_field_toc_docx(tmp_path / "source.docx")
    output_path = tmp_path / "output.docx"

    result = finalize_static_toc(
        source_path,
        output_path,
        {
            1: "目 录\n第1章 绪论 ...... 8\n1.1 研究背景 ...... 9",
            4: "第1章 绪论\n- 1 -",
            5: "1.1 研究背景\n- 2 -",
        },
    )

    assert result["complete"] is True
    output = Document(output_path)
    assert [p.text for p in output.paragraphs if p.style.style_id in {"TOC1", "TOC2"}] == [
        "第1章 绪论\t1",
        "1.1 研究背景\t2",
    ]


def test_finalize_static_toc_keeps_ignorable_prefixes_declared(tmp_path: Path):
    source_path = _make_field_toc_docx(tmp_path / "source.docx")
    output_path = tmp_path / "output.docx"

    finalize_static_toc(
        source_path,
        output_path,
        {
            4: "第1章 绪论\n正文保持不变\n- 1 -",
            5: "1.1 研究背景\n- 2 -",
        },
    )

    with zipfile.ZipFile(output_path) as output_zip:
        document_xml = output_zip.read("word/document.xml")
    declared_prefixes = {
        prefix
        for _, (prefix, _) in ET.iterparse(io.BytesIO(document_xml), events=("start-ns",))
        if prefix
    }
    document_root = ET.fromstring(document_xml)
    ignorable_prefixes = set(document_root.get(f"{{{MC_NS}}}Ignorable", "").split())
    assert ignorable_prefixes <= declared_prefixes


@pytest.mark.parametrize("hyphen", ["-", "\u2010", "\u2011", "\u2012", "\u2013", "\u2014", "\u2212"])
def test_analyze_static_toc_accepts_supported_page_number_hyphens(tmp_path: Path, hyphen: str):
    source_path = _make_toc_docx(tmp_path / "source.docx")

    result = analyze_static_toc(
        source_path,
        {
            1: "目 录\n第1章 绪论 ...... 8\n1.1 研究背景 ...... 9",
            4: f"第1章 绪论\n1.1 研究背景\n正文\n{hyphen} 12 {hyphen}",
        },
    )

    assert result["complete"] is True
    assert [(entry["pdf_page"], entry["page"]) for entry in result["entries"]] == [(4, 12), (4, 12)]


def test_finalize_static_toc_removes_unstyled_toc_field_end_paragraph(tmp_path: Path):
    source_path = _make_field_toc_docx(tmp_path / "source.docx", styled_end=False)
    output_path = tmp_path / "output.docx"

    finalize_static_toc(
        source_path,
        output_path,
        {
            1: "目 录\n第1章 绪论 ...... 8\n1.1 研究背景 ...... 9",
            4: "第1章 绪论\n正文保持不变\n- 1 -",
            5: "1.1 研究背景\n- 2 -",
        },
    )

    with zipfile.ZipFile(output_path) as output_zip:
        document_root = ET.fromstring(output_zip.read("word/document.xml"))
    assert document_root.findall(".//w:fldChar", NSMAP) == []


@pytest.mark.parametrize("unsafe_footer", ["2026", "- 7", "7 -"])
def test_analyze_static_toc_rejects_unwrapped_or_one_sided_page_numbers(
    tmp_path: Path,
    unsafe_footer: str,
):
    source_path = _make_toc_docx(tmp_path / "source.docx")

    result = analyze_static_toc(
        source_path,
        {
            1: "目 录\n第1章 绪论 ...... 8\n1.1 研究背景 ...... 9",
            4: f"第1章 绪论\n正文\n{unsafe_footer}",
            5: "1.1 研究背景\n正文\n- 2 -",
        },
    )

    assert result["complete"] is False
    assert result["unmatched_titles"] == ["第1章 绪论"]


def test_analyze_static_toc_rejects_title_prefix_matches(tmp_path: Path):
    source_path = _make_toc_docx(tmp_path / "source.docx")

    result = analyze_static_toc(
        source_path,
        {
            1: "目 录\n第1章 绪论 ...... 8\n1.1 研究背景 ...... 9",
            4: "第1章 绪论与展望\n正文\n- 1 -",
            5: "1.1 研究背景\n正文\n- 2 -",
        },
    )

    assert result["complete"] is False
    assert result["unmatched_titles"] == ["第1章 绪论"]
