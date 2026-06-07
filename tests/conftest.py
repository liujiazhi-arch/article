from __future__ import annotations

import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS_DIR = PROJECT_ROOT / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import audit_thesis  # noqa: E402
import fix_thesis  # noqa: E402


@pytest.fixture(autouse=True)
def isolate_article_runtime_roots(monkeypatch, request, tmp_path):
    if not request.node.nodeid.startswith("tests/test_article"):
        yield
        return

    monkeypatch.setenv("ARTICLE_API_STATE_ROOT", str((tmp_path / "article-state").resolve()))
    monkeypatch.setenv("ARTICLE_API_RUNTIME_ROOT", str((tmp_path / "article-runtime").resolve()))
    yield


TARGET_RULE_IDS = (
    "P01",
    "T01",
    "T02",
    "T03",
    "T04",
    "T05",
    "T06",
    "H01",
    "H02",
    "H03",
    "H04",
    "F01",
    "F02",
    "F07",
    "TB01",
    "C01",
    "R01",
    "KW01",
)

FIX_ROUNDTRIP_RULE_IDS = ("P01", "T01", "H02", "F07", "TB01", "PU01")


def _find_or_create(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def _set_run_format(
    run,
    *,
    east_asia: str = "宋体",
    ascii_font: str = "Times New Roman",
    size: int = 24,
    bold: bool = False,
    superscript: bool = False,
):
    r_pr = run._r.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    r_fonts.set(qn("w:ascii"), ascii_font)
    r_fonts.set(qn("w:hAnsi"), ascii_font)

    sz = r_pr.sz
    if sz is None:
        sz = OxmlElement("w:sz")
        r_pr.append(sz)
    sz.set(qn("w:val"), str(size))

    b = r_pr.b
    if b is None:
        b = OxmlElement("w:b")
        r_pr.append(b)
    b.set(qn("w:val"), "1" if bold else "0")

    vert_align = r_pr.find(qn("w:vertAlign"))
    if superscript:
        if vert_align is None:
            vert_align = OxmlElement("w:vertAlign")
            r_pr.append(vert_align)
        vert_align.set(qn("w:val"), "superscript")
    elif vert_align is not None:
        r_pr.remove(vert_align)


def _set_paragraph_properties(
    paragraph,
    *,
    align: str = "both",
    first_line: int | None = 480,
    line: int = 360,
    before: int = 0,
    after: int = 0,
    auto_space: bool = True,
):
    p_pr = paragraph._p.get_or_add_pPr()

    jc = _find_or_create(p_pr, "w:jc")
    jc.set(qn("w:val"), align)

    ind = _find_or_create(p_pr, "w:ind")
    if first_line is None:
        ind.attrib.pop(qn("w:firstLine"), None)
    else:
        ind.set(qn("w:firstLine"), str(first_line))

    spacing = _find_or_create(p_pr, "w:spacing")
    spacing.set(qn("w:line"), str(line))
    spacing.set(qn("w:lineRule"), "auto")
    spacing.set(qn("w:before"), str(before))
    spacing.set(qn("w:after"), str(after))

    if auto_space:
        for tag in ("w:autoSpaceDE", "w:autoSpaceDN"):
            elem = _find_or_create(p_pr, tag)
            elem.set(qn("w:val"), "0")


def _set_page_break_before(paragraph, enabled: bool = True):
    p_pr = paragraph._p.get_or_add_pPr()
    elem = _find_or_create(p_pr, "w:pageBreakBefore")
    elem.set(qn("w:val"), "1" if enabled else "0")


def _set_heading(paragraph, *, level: int, size: int, page_break_before: bool = False):
    paragraph.style = f"Heading {level}"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    before_after = {
        1: (120, 120),
        2: (120, 120),
        3: (120, 120),
        4: (60, 60),
    }
    before, after = before_after[level]
    _set_paragraph_properties(
        paragraph,
        align="center" if level == 1 else "left",
        first_line=0,
        line=360,
        before=before,
        after=after,
        auto_space=False,
    )
    _set_page_break_before(paragraph, enabled=page_break_before)
    font_cn = "宋体" if level == 4 else "黑体"
    for run in paragraph.runs:
        _set_run_format(run, east_asia=font_cn, size=size, bold=False)


def _add_heading(doc: Document, text: str, *, level: int, size: int, page_break_before: bool = False):
    paragraph = doc.add_paragraph()
    paragraph.add_run(text)
    _set_heading(paragraph, level=level, size=size, page_break_before=page_break_before)
    return paragraph


def _set_reference_format(paragraph, *, left: int = 560, hanging: int = 560, tab_pos: int = 560):
    p_pr = paragraph._p.get_or_add_pPr()
    ind = _find_or_create(p_pr, "w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    ind.attrib.pop(qn("w:firstLine"), None)

    spacing = _find_or_create(p_pr, "w:spacing")
    spacing.set(qn("w:line"), "240")
    spacing.set(qn("w:lineRule"), "auto")

    jc = _find_or_create(p_pr, "w:jc")
    jc.set(qn("w:val"), "left")

    tabs = _find_or_create(p_pr, "w:tabs")
    for old_tab in list(tabs):
        tabs.remove(old_tab)
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "left")
    tab.set(qn("w:pos"), str(tab_pos))
    tabs.append(tab)


def _set_caption(paragraph, text: str):
    paragraph.clear()
    run = paragraph.add_run(text)
    _set_paragraph_properties(
        paragraph,
        align="center",
        first_line=0,
        line=360,
        before=0,
        after=0,
        auto_space=False,
    )
    _set_run_format(run, size=21)


def _set_table_borders(table, *, valid: bool = True):
    tbl_pr = table._tbl.tblPr
    tbl_borders = tbl_pr.first_child_found_in("w:tblBorders")
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)

    specs = {
        "top": ("single", "12" if valid else "4"),
        "bottom": ("single", "12"),
        "left": ("none", "0"),
        "right": ("none", "0"),
        "insideV": ("none", "0"),
    }
    for name, (val, size) in specs.items():
        elem = tbl_borders.find(qn(f"w:{name}"))
        if elem is None:
            elem = OxmlElement(f"w:{name}")
            tbl_borders.append(elem)
        elem.set(qn("w:val"), val)
        elem.set(qn("w:sz"), size)
        elem.set(qn("w:color"), "000000")

    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_borders = tc_pr.first_child_found_in("w:tcBorders")
        if tc_borders is None:
            tc_borders = OxmlElement("w:tcBorders")
            tc_pr.append(tc_borders)
        bottom = tc_borders.find(qn("w:bottom"))
        if bottom is None:
            bottom = OxmlElement("w:bottom")
            tc_borders.append(bottom)
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), "000000")


def _build_base_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, attr, Cm(2.54))

    doc.add_paragraph("封面标题")

    _add_heading(doc, "第一章 绪论", level=1, size=30)

    body = doc.add_paragraph()
    _set_paragraph_properties(body)
    _set_run_format(body.add_run("这是正文示例"), size=24)
    _set_run_format(
        body.add_run("[1]"),
        east_asia="Times New Roman",
        ascii_font="Times New Roman",
        size=24,
        superscript=True,
    )
    _set_run_format(body.add_run("。"), size=24)

    _add_heading(doc, "1.1 研究背景", level=2, size=30)
    _add_heading(doc, "1.1.1 研究现状", level=3, size=24)
    _add_heading(doc, "1.1.1.1 术语定义", level=4, size=24)

    keywords = doc.add_paragraph()
    _set_paragraph_properties(keywords)
    _set_run_format(keywords.add_run("关键词"), east_asia="黑体", size=24)
    _set_run_format(keywords.add_run("：测试；审计；修复"), east_asia="宋体", size=24)

    figure_caption = doc.add_paragraph()
    _set_caption(figure_caption, "图1-1 实验装置")

    table_caption = doc.add_paragraph()
    _set_caption(table_caption, "表1-1 实验结果")

    _add_heading(doc, "参考文献", level=1, size=30, page_break_before=True)

    reference = doc.add_paragraph()
    _set_run_format(reference.add_run("[1]\t张三. 论文格式检查研究。"), size=24)
    _set_reference_format(reference)

    table = doc.add_table(rows=2, cols=2)
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.text = f"{row_index}{col_index}"
            for paragraph in cell.paragraphs:
                _set_paragraph_properties(paragraph)
                for run in paragraph.runs:
                    _set_run_format(run, size=24)
    _set_table_borders(table, valid=True)

    return doc


def _inject_existing_footer_page_field(docx_path: Path):
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    office_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    ns = {"w": w_ns}

    with zipfile.ZipFile(docx_path, "r") as zf:
        parts = {name: zf.read(name) for name in zf.namelist()}

    rels_root = ET.fromstring(parts["word/_rels/document.xml.rels"])
    rel = ET.SubElement(rels_root, f"{{{rel_ns}}}Relationship")
    rel.set("Id", "rIdFooterTest")
    rel.set("Type", f"{office_rel_ns}/footer")
    rel.set("Target", "footer1.xml")
    parts["word/_rels/document.xml.rels"] = ET.tostring(rels_root, encoding="utf-8", xml_declaration=True)

    document_root = ET.fromstring(parts["word/document.xml"])
    sect_pr = document_root.find("w:body/w:sectPr", ns)
    assert sect_pr is not None
    footer_ref = ET.SubElement(sect_pr, f"{{{w_ns}}}footerReference")
    footer_ref.set(f"{{{w_ns}}}type", "default")
    footer_ref.set(f"{{{office_rel_ns}}}id", "rIdFooterTest")
    parts["word/document.xml"] = ET.tostring(document_root, encoding="utf-8", xml_declaration=True)

    footer_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:ftr xmlns:w="{w_ns}">
  <w:p>
    <w:pPr><w:jc w:val="center"/></w:pPr>
    <w:r><w:fldChar w:fldCharType="begin"/></w:r>
    <w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>
    <w:r><w:fldChar w:fldCharType="separate"/></w:r>
    <w:r><w:t>1</w:t></w:r>
    <w:r><w:fldChar w:fldCharType="end"/></w:r>
  </w:p>
</w:ftr>"""
    parts["word/footer1.xml"] = footer_xml.encode("utf-8")

    content_types_root = ET.fromstring(parts["[Content_Types].xml"])
    override = ET.SubElement(content_types_root, f"{{{ct_ns}}}Override")
    override.set("PartName", "/word/footer1.xml")
    override.set("ContentType", "application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml")
    parts["[Content_Types].xml"] = ET.tostring(content_types_root, encoding="utf-8", xml_declaration=True)

    with zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, payload in parts.items():
            zf.writestr(name, payload)


def _find_paragraph(doc: Document, prefix: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith(prefix):
            return paragraph
    raise AssertionError(f"Paragraph not found: {prefix}")


def _mutate_p01(doc: Document):
    doc.sections[0].left_margin = Cm(1.0)


def _mutate_t01(doc: Document):
    paragraph = _find_paragraph(doc, "这是正文示例")
    _set_run_format(paragraph.runs[0], east_asia="黑体", size=24)


def _mutate_t02(doc: Document):
    paragraph = _find_paragraph(doc, "这是正文示例")
    _set_run_format(paragraph.runs[0], ascii_font="Arial", size=24)


def _mutate_t03(doc: Document):
    paragraph = _find_paragraph(doc, "这是正文示例")
    _set_run_format(paragraph.runs[0], size=18)


def _mutate_t04(doc: Document):
    paragraph = _find_paragraph(doc, "这是正文示例")
    _set_paragraph_properties(paragraph, line=240)


def _mutate_t05(doc: Document):
    paragraph = _find_paragraph(doc, "这是正文示例")
    _set_paragraph_properties(paragraph, first_line=0)


def _mutate_t06(doc: Document):
    paragraph = _find_paragraph(doc, "这是正文示例")
    _set_paragraph_properties(paragraph, align="left")


def _mutate_h01(doc: Document):
    paragraph = _find_paragraph(doc, "第一章 绪论")
    _set_paragraph_properties(paragraph, align="left", first_line=0, line=360, before=240, after=120, auto_space=False)


def _mutate_h02(doc: Document):
    paragraph = _find_paragraph(doc, "1.1 研究背景")
    _set_paragraph_properties(paragraph, align="center", first_line=0, line=360, before=120, after=60, auto_space=False)


def _mutate_h03(doc: Document):
    paragraph = _find_paragraph(doc, "1.1.1 研究现状")
    _set_paragraph_properties(paragraph, align="center", first_line=0, line=360, before=60, after=60, auto_space=False)


def _mutate_h04(doc: Document):
    paragraph = _find_paragraph(doc, "1.1.1.1 术语定义")
    _set_paragraph_properties(paragraph, align="center", first_line=0, line=360, before=60, after=60, auto_space=False)


def _mutate_f01(doc: Document):
    paragraph = _find_paragraph(doc, "图1-1")
    _set_paragraph_properties(paragraph, align="left", first_line=0, line=360, before=0, after=0, auto_space=False)


def _mutate_f02(doc: Document):
    paragraph = _find_paragraph(doc, "表1-1")
    _set_caption(paragraph, "表1-1 实验结果")
    _set_run_format(paragraph.runs[0], size=24)


def _mutate_f07(doc: Document):
    paragraph = _find_paragraph(doc, "图1-1")
    paragraph.runs[0].text += "。"


def _mutate_tb01(doc: Document):
    _set_table_borders(doc.tables[0], valid=False)


def _mutate_c01(doc: Document):
    paragraph = _find_paragraph(doc, "这是正文示例")
    _set_run_format(
        paragraph.runs[1],
        east_asia="Times New Roman",
        ascii_font="Times New Roman",
        size=24,
        superscript=False,
    )


def _mutate_r01(doc: Document):
    paragraph = _find_paragraph(doc, "[1]\t")
    _set_reference_format(paragraph, left=0, hanging=0, tab_pos=560)


def _mutate_kw01(doc: Document):
    paragraph = _find_paragraph(doc, "关键词：")
    paragraph.clear()
    run = paragraph.add_run("关键词：测试；审计")
    _set_paragraph_properties(paragraph)
    _set_run_format(run, size=24)


def _mutate_pu01(doc: Document):
    paragraph = _find_paragraph(doc, "这是正文示例")
    paragraph.runs[0].text = "这是正文,示例.段落"


RULE_MUTATORS = {
    "P01": _mutate_p01,
    "T01": _mutate_t01,
    "T02": _mutate_t02,
    "T03": _mutate_t03,
    "T04": _mutate_t04,
    "T05": _mutate_t05,
    "T06": _mutate_t06,
    "H01": _mutate_h01,
    "H02": _mutate_h02,
    "H03": _mutate_h03,
    "H04": _mutate_h04,
    "F01": _mutate_f01,
    "F02": _mutate_f02,
    "F07": _mutate_f07,
    "TB01": _mutate_tb01,
    "C01": _mutate_c01,
    "R01": _mutate_r01,
    "KW01": _mutate_kw01,
    "PU01": _mutate_pu01,
}


def make_compliant_doc(path: str | Path) -> Path:
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    _build_base_document().save(output_path)
    _inject_existing_footer_page_field(output_path)
    return output_path


def make_violating_doc(path: str | Path, rule_id: str) -> Path:
    if rule_id not in RULE_MUTATORS:
        raise KeyError(f"Unsupported rule_id: {rule_id}")
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = _build_base_document()
    RULE_MUTATORS[rule_id](doc)
    doc.save(output_path)
    _inject_existing_footer_page_field(output_path)
    return output_path


def audit_rule_status(docx_path: str | Path, rule_id: str, *, profile_path: str | None = None) -> dict:
    results, score, report = audit_thesis.audit_docx(str(docx_path), profile_path=profile_path)
    result = next(item for item in results if item["id"] == rule_id)
    return {
        "rule": result,
        "results": results,
        "score": score,
        "report": report,
        "failed_ids": [item["id"] for item in results if not item["passed"]],
    }


@pytest.fixture
def tmp_docx(tmp_path):
    def _build(builder=make_compliant_doc, *, filename: str = "sample.docx", rule_id: str | None = None):
        path = tmp_path / filename
        if builder is make_violating_doc:
            if rule_id is None:
                raise ValueError("rule_id is required for make_violating_doc")
            return make_violating_doc(path, rule_id)
        return builder(path)

    return _build
