from pathlib import Path

from docx import Document
import pytest

import thesis_tool.render_verify as render_verify_module
import thesis_tool.render_sources as render_sources_module


def test_build_pdf_toc_page_number_report_detects_mismatch(monkeypatch, tmp_path):
    rendered_pdf = tmp_path / "toc_mismatch.pdf"
    rendered_pdf.write_bytes(b"%PDF")

    def fake_extract_pdf_page_texts(pdf_path: str | None, *, page_count: int | None = None):
        assert pdf_path == str(rendered_pdf)
        assert page_count is None
        return (
            {
                1: "目录\n第1章 绪论 ........ 1",
                2: "第1章 绪论\n这是正文\n- 2 -",
            },
            {
                "source": "pdf",
                "available": True,
                "page_text_available_count": 2,
                "page_text_extraction_warning_count": 0,
                "warnings": [],
            },
        )

    monkeypatch.setattr(render_verify_module, "_extract_pdf_page_texts", fake_extract_pdf_page_texts)

    report = render_verify_module.build_pdf_toc_page_number_report(str(rendered_pdf))

    assert report["summary"]["toc_page_number_mismatch_count"] == 1
    assert report["findings"][0]["id"] == "toc_page_number_mismatch"
    assert "第1章 绪论" in report["findings"][0]["message"]

def test_build_render_verify_report_flags_toc_declared_page_mismatch(monkeypatch, tmp_path):
    source_path = tmp_path / "render_verify_toc_source.docx"
    doc = Document()
    doc.add_heading("Render Verify", level=1)
    doc.save(source_path)

    def fake_run_render_engine(input_docx: str, output_dir: str, renderer: str) -> dict:
        page_dir = Path(output_dir) / "word_pdf_pages"
        page_dir.mkdir(parents=True, exist_ok=True)
        for page_number in range(1, 5):
            (page_dir / f"page-{page_number}.png").write_bytes(b"png")
        pdf_path = Path(output_dir) / "render_verify_word.pdf"
        pdf_path.write_bytes(b"pdf")
        return {
            "engine": "word-pdf",
            "page_dir": str(page_dir),
            "pdf_path": str(pdf_path),
            "fallback_used": False,
            "warnings": [],
        }

    monkeypatch.setattr(render_sources_module, "_run_render_engine", fake_run_render_engine)
    monkeypatch.setattr(
        render_verify_module,
        "_extract_pdf_page_texts",
        lambda pdf_path, *, page_count: (
            {
                1: "目 录\n摘要 ...... 2\n1 绪论 ...... 2\n2 材料与方法",
                2: "摘要\n本文研究论文格式检查工具。",
                3: "1 绪论\n研究背景与意义。\n- 3 -",
                4: "2 材料与方法\n实验设置。\n- 4 -",
            },
            {
                "source": "pdf",
                "available": True,
                "page_text_available_count": 4,
                "page_text_extraction_warning_count": 0,
                "warnings": [],
            },
        ),
        raising=False,
    )
    monkeypatch.setattr(
        render_verify_module,
        "analyze_page_images",
        lambda page_images, **kwargs: {
            "findings": [],
            "summary": {
                "finding_count": 0,
                "highest_severity": None,
                "blank_page_count": 0,
                "large_blank_count": 0,
                "render_suspect_count": 0,
            },
            "layout_score": {"score": 100, "penalty": 0, "actionable_finding_count": 0},
        },
    )
    monkeypatch.setattr(
        render_verify_module,
        "build_document_preflight",
        lambda *args, **kwargs: {
            "preflight_status": "ready",
            "toc_status": "generated_toc",
            "preface_status": "not_detected",
            "heading_renumber_guard": {"status": "clear"},
            "style_conflict_count": 0,
            "table_heading_risk_count": 0,
            "recommended_actions": [],
            "diagnostics": {"toc": {"status": "generated_toc"}, "heading_renumber_guard": {"status": "clear"}},
        },
    )
    monkeypatch.setattr(
        render_verify_module,
        "build_scope_verify",
        lambda *args, **kwargs: {
            "profile_id": "lnu-checker-2026",
            "requested_profile": "lnu",
            "fallback_used": False,
            "profile_display": "lnu-checker-2026 (requested: lnu)",
            "selected_scopes": ["toc"],
            "overall_status": "verified",
            "readiness": "structure-ready",
            "manual_review_rule_ids": [],
            "unsupported_rule_ids": [],
        },
    )

    report = render_verify_module.build_render_verify_report(str(source_path), profile_path="lnu", scopes=["toc"])

    mismatch = next(item for item in report["render_findings"] if item["id"] == "toc_page_number_mismatch")
    assert mismatch["page"] == 1
    assert mismatch["declared_page"] == 2
    assert mismatch["actual_page"] == 3
    assert mismatch["toc_entry"] == "1 绪论"
    assert "目录页码不一致" in mismatch["message"]
    assert report["summary"]["render_finding_count"] == 2
    assert any(item["rule_id"] == "render.toc_page_number_mismatch" for item in report["evidence_items"])

    unconfirmed = next(item for item in report["render_findings"] if item["id"] == "toc_page_number_unconfirmed")
    assert unconfirmed["toc_entry"] == "2 材料与方法"
    assert unconfirmed["severity"] == "warning"
    assert "缺少页码" in unconfirmed["message"]

def test_toc_page_number_check_uses_printed_page_number_not_pdf_index():
    findings = render_verify_module._build_toc_page_number_findings(
        {
            1: "目 录\n1 绪论 ...... 1",
            2: "摘 要\n摘要正文\nI",
            3: "1 绪论\n研究背景与意义。\n- 1 -",
        }
    )

    assert findings == []

def test_toc_page_number_check_includes_acknowledgement_heading():
    findings = render_verify_module._build_toc_page_number_findings(
        {
            1: "目 录\n致 谢 ...... 22",
            2: "致 谢\n感谢导师和同学。\n- 25 -",
        }
    )

    assert len(findings) == 1
    assert findings[0]["id"] == "toc_page_number_mismatch"
    assert findings[0]["toc_entry"] == "致 谢"
    assert findings[0]["declared_page"] == 22
    assert findings[0]["actual_page"] == 25

def test_toc_page_number_check_reads_wps_nonbreaking_hyphen_footer():
    findings = render_verify_module._build_toc_page_number_findings(
        {
            1: "目 录\n1 绪论 ...... 2",
            2: "1 绪论\n研究背景与意义。\n‑1‑",
        }
    )

    assert len(findings) == 1
    assert findings[0]["id"] == "toc_page_number_mismatch"
    assert findings[0]["actual_page"] == 1

def test_toc_page_number_check_requires_exact_normalized_heading_line():
    findings = render_verify_module._build_toc_page_number_findings(
        {
            1: "目 录\n1 绪论 ...... 2",
            2: "1 绪论研究背景\n正文内容。\n- 7 -",
        }
    )

    assert len(findings) == 1
    assert findings[0]["id"] == "toc_page_number_unconfirmed"
    assert findings[0]["actual_page"] is None
    assert "未能在后续 PDF 页面中确认" in findings[0]["message"]

@pytest.mark.parametrize("footer", ["2024", "- 7", "7 -"])
def test_toc_page_number_check_rejects_unreliable_footer_number(footer):
    findings = render_verify_module._build_toc_page_number_findings(
        {
            1: "目 录\n1 绪论 ...... 2",
            2: f"1 绪论\n正文内容。\n{footer}",
        }
    )

    assert len(findings) == 1
    assert findings[0]["id"] == "toc_page_number_unconfirmed"
    assert findings[0]["actual_page"] is None
    assert "未读取到该页显示页码" in findings[0]["message"]

def test_toc_page_number_check_reads_entries_from_continued_toc_page():
    findings = render_verify_module._build_toc_page_number_findings(
        {
            1: "目 录\n1 绪论 ...... 1",
            2: "2 材料与方法 ...... 3\n3 结果 ...... 5",
            3: "1 绪论\n研究背景。\n- 1 -",
            4: "2 材料与方法\n实验设计。\n- 4 -",
            5: "3 结果\n实验结果。\n- 5 -",
        }
    )

    assert [(item["toc_entry"], item["declared_page"], item["actual_page"]) for item in findings] == [
        ("2 材料与方法", 3, 4),
    ]

def test_toc_page_number_check_attaches_unique_real_line_bbox():
    findings = render_verify_module._build_toc_page_number_findings(
        {
            1: "目 录\n1 绪论 ...... 2",
            2: "1 绪论\n研究背景。\n- 1 -",
        },
        line_boxes_by_page={
            1: [
                {
                    "text": "1绪论......2",
                    "bbox": {"x": 0.12, "y": 0.31, "w": 0.76, "h": 0.025},
                }
            ]
        },
    )

    assert len(findings) == 1
    assert findings[0]["bbox"] == {"x": 0.12, "y": 0.31, "w": 0.76, "h": 0.025}

def test_toc_page_number_check_does_not_treat_heading_number_as_printed_page_number():
    findings = render_verify_module._build_toc_page_number_findings(
        {
            1: "目 录\n1 绪论 ...... 1",
            2: "1 绪论\n研究背景与意义。",
        }
    )

    assert findings == [
        {
            "id": "toc_page_number_unconfirmed",
            "rule_id": "render.toc_page_number_unconfirmed",
            "classification": "render_integrity",
            "severity": "warning",
            "page": 1,
            "toc_entry": "1 绪论",
            "declared_page": 1,
            "actual_page": None,
            "message": "目录条目“1 绪论”已定位到 PDF 第 2 页，但未读取到该页显示页码，需人工复核。",
            "actionable": False,
            "suggested_scope": "toc",
            "suggested_action": "打开 PDF 对照目录和正文标题页码。",
        }
    ]

def test_toc_page_number_check_warns_when_heading_not_located():
    findings = render_verify_module._build_toc_page_number_findings(
        {
            1: "目录\n1 绪论 ...... 2",
            2: "摘要\n本文研究论文格式检查工具。",
        }
    )

    assert findings == [
        {
            "id": "toc_page_number_unconfirmed",
            "rule_id": "render.toc_page_number_unconfirmed",
            "classification": "render_integrity",
            "severity": "warning",
            "page": 1,
            "toc_entry": "1 绪论",
            "declared_page": 2,
            "actual_page": None,
            "message": "目录条目“1 绪论”未能在后续 PDF 页面中确认对应标题，需人工复核。",
            "actionable": False,
            "suggested_scope": "toc",
            "suggested_action": "打开 PDF 对照目录和正文标题页码。",
        }
    ]

def test_build_render_verify_report_connects_static_toc_finalization(monkeypatch, tmp_path):
    source_path = tmp_path / "paper.docx"
    Document().save(source_path)
    pdf_path = tmp_path / "paper.pdf"
    pdf_path.write_bytes(b"%PDF")
    page_dir = tmp_path / "pages"
    page_dir.mkdir()
    page_image = page_dir / "page-1.png"
    page_image.write_bytes(b"png")
    page_texts = {1: "目录\n第1章 绪论 ...... 1", 2: "第1章 绪论\n- 1 -"}
    toc_findings = [{"id": "toc_page_number_mismatch", "page": 1, "severity": "warning"}]
    content_match = {"status": "matched", "matched": True}
    captured = {}

    monkeypatch.setattr(
        render_verify_module,
        "_resolve_render_metadata",
        lambda *_args, **_kwargs: {
            "engine": "manual-pdf",
            "pdf_path": str(pdf_path),
            "page_dir": str(page_dir),
            "warnings": [],
            "fallback_used": False,
        },
    )
    monkeypatch.setattr(
        render_verify_module,
        "_extract_pdf_page_texts",
        lambda *_args, **_kwargs: (
            page_texts,
            {"available": True, "page_text_available_count": 2, "page_text_extraction_warning_count": 0, "warnings": []},
        ),
    )
    monkeypatch.setattr(render_verify_module, "_extract_pdf_line_boxes", lambda *_args: {})
    monkeypatch.setattr(
        render_verify_module,
        "verify_docx_pdf_content_match",
        lambda *_args, **_kwargs: content_match,
    )
    monkeypatch.setattr(render_verify_module, "_build_toc_page_number_findings", lambda *_args, **_kwargs: toc_findings)
    monkeypatch.setattr(
        render_verify_module,
        "analyze_page_images",
        lambda *_args, **_kwargs: {"findings": [], "summary": {"finding_count": 0}, "layout_score": {}},
    )
    monkeypatch.setattr(
        render_verify_module,
        "build_document_preflight",
        lambda *_args, **_kwargs: {
            "preflight_status": "ready",
            "heading_renumber_guard": {},
            "diagnostics": {},
            "recommended_actions": [],
        },
    )
    monkeypatch.setattr(
        render_verify_module,
        "build_scope_verify",
        lambda *_args, **_kwargs: {
            "profile_id": "lnu-checker-2026",
            "selected_scopes": ["toc"],
            "overall_status": "verified",
            "readiness": "structure-ready",
            "manual_review_rule_ids": [],
            "unsupported_rule_ids": [],
        },
    )

    def fake_finalize(input_docx, output_dir, current_page_texts, **kwargs):
        captured.update(
            input_docx=input_docx,
            output_dir=output_dir,
            page_texts=current_page_texts,
            kwargs=kwargs,
        )
        return {
            "status": "generated",
            "available": True,
            "entry_count": 26,
            "mapped_count": 26,
            "output_path": str(tmp_path / "static_toc.docx"),
        }

    monkeypatch.setattr(render_verify_module, "build_static_toc_finalization", fake_finalize, raising=False)

    report = render_verify_module.build_render_verify_report(
        str(source_path),
        output_dir=str(tmp_path / "proof"),
        rendered_pdf=str(pdf_path),
        pdf_matches_docx_confirmed=True,
        generate_static_toc=True,
    )

    assert report["toc_finalization"]["status"] == "generated"
    assert report["summary"]["toc_finalization_status"] == "generated"
    assert report["summary"]["toc_output_available"] is True
    assert report["summary"]["toc_entry_count"] == 26
    assert report["summary"]["toc_mapped_count"] == 26
    assert captured["input_docx"] == str(source_path)
    assert captured["page_texts"] == page_texts
    assert captured["kwargs"] == {
        "requested": True,
        "pdf_matches_docx_confirmed": True,
        "content_match": content_match,
        "toc_findings": toc_findings,
    }
