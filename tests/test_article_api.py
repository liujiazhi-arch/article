from io import BytesIO
import zipfile
from pathlib import Path
import pytest
from docx import Document

import article_api.app as app_module
from article_api import app_ops
import article_api.jobs as jobs_module
import article_api.output_naming as output_naming
import article_api.storage as storage
from article_api.app import (
    ApplyRequest,
    AuditRequest,
    NormalizeRequest,
    NormalizeJobRequest,
    PreflightRequest,
    RenderVerifyRequest,
    RetentionSweepRequest,
    UploadApplyRequest,
    UploadNormalizeRequest,
    UploadVerifyRequest,
    VerifyRequest,
)
from article_api import create_app, fastapi_available
from article_api.jobs import clear_jobs, job_count, wait_for_job
from article_api.storage import clear_uploads, upload_count

from .conftest import RULE_MUTATORS, make_compliant_doc


def test_render_workflow_modes_do_not_advertise_large_blank_public_findings():
    from article_api.response_payloads import RENDER_WORKFLOW_MODES

    text = "\n".join(str(mode) for mode in RENDER_WORKFLOW_MODES)

    assert "大块空白" not in text
    assert "大面积空白" not in text
    assert "页底空白" not in text


class FakeHTTPException(Exception):
    def __init__(self, status_code: int, detail: str):
        super().__init__(detail)
        self.status_code = status_code
        self.detail = detail


class FakeRoute:
    def __init__(self, path: str, endpoint, status_code: int):
        self.path = path
        self.endpoint = endpoint
        self.status_code = status_code


class FakeFastAPI:
    def __init__(self, **_kwargs):
        self.routes = []

    def get(self, path: str, **kwargs):
        status_code = kwargs.get("status_code", 200)

        def decorator(func):
            self.routes.append(FakeRoute(path, func, status_code))
            return func

        return decorator

    def post(self, path: str, **kwargs):
        status_code = kwargs.get("status_code", 200)

        def decorator(func):
            self.routes.append(FakeRoute(path, func, status_code))
            return func

        return decorator


def _build_fake_app(monkeypatch):
    app_module._RETENTION_STATE.update(
        {
            "last_run_at": None,
            "last_trigger": None,
            "last_result": None,
            "last_error": None,
            "last_monotonic": None,
        }
    )
    monkeypatch.setattr(app_module, "FastAPI", FakeFastAPI)
    monkeypatch.setattr(app_module, "HTTPException", FakeHTTPException)
    monkeypatch.setattr(app_module, "FileResponse", None)
    return app_module.create_app()


def _routes_by_path(app):
    return {route.path: route for route in app.routes}


def _make_style_conflict_lnu_doc(source_path):
    doc = Document()
    heading = doc.add_paragraph("1 绪论")
    heading.style = doc.styles["Heading 1"]
    conflict = doc.add_paragraph("3.6 分子对接验证结果")
    conflict.style = doc.styles["Heading 1"]
    doc.add_paragraph("这是正文示例。")
    doc.save(source_path)
    return source_path


def test_create_app_handles_missing_fastapi_dependency():
    if fastapi_available():
        app = create_app()
        route_paths = {route.path for route in app.routes}
        assert "/" in route_paths
        assert "/static/{asset_path:path}" in route_paths
        assert "/assets/{asset_path:path}" not in route_paths
        assert "/health" in route_paths
        assert "/ready" in route_paths
        assert "/version" in route_paths
        assert "/updates/latest" in route_paths
        assert "/profiles" in route_paths
        assert "/render-workflow-modes" in route_paths
        assert "/render-evidence/screenshot/{token}" in route_paths
        assert "/render-evidence/screenshot" not in route_paths
        assert "/audit" in route_paths
        assert "/plan" in route_paths
        assert "/preflight" in route_paths
        assert "/normalize" in route_paths
        assert "/render-verify" in route_paths
        assert "/verify" in route_paths
        assert "/apply" in route_paths
        assert "/jobs/normalize" in route_paths
        assert "/jobs/verify" in route_paths
        assert "/jobs/apply" in route_paths
        assert "/batch" not in route_paths
        assert "/jobs/batch" not in route_paths
        assert "/jobs/batches/recent" not in route_paths
        assert "/uploads/docx" in route_paths
        assert "/uploads/pdf" in route_paths
        assert "/assets/lnu-emblem.jpg" not in route_paths
        assert "/uploads" in route_paths
        assert "/uploads/{upload_id}" in route_paths
        assert "/uploads/{upload_id}/cleanup" in route_paths
        assert "/uploads/{upload_id}/jobs/normalize" in route_paths
        assert "/uploads/{upload_id}/jobs/verify" in route_paths
        assert "/uploads/{upload_id}/jobs/apply" in route_paths
        assert "/jobs" in route_paths
        assert "/jobs/{job_id}" in route_paths
        assert "/jobs/{job_id}/inspect" in route_paths
        assert "/jobs/{job_id}/result" in route_paths
        assert "/jobs/{job_id}/cleanup" in route_paths
        assert "/jobs/{job_id}/retry" in route_paths
        assert "/jobs/{job_id}/artifacts/{artifact_role}/download" in route_paths
        assert "/feedback/download" in route_paths
        assert "/ops/summary" in route_paths
        assert "/ops/storage" in route_paths
        assert "/ops/runtime" in route_paths
        assert "/ops/retention/run-defaults" in route_paths
        assert "/ops/retention/sweep" in route_paths
        return

    with pytest.raises(RuntimeError, match="FastAPI API prototype is unavailable"):
        create_app()


def test_apply_request_supports_force_flag():
    request = ApplyRequest(file_path="demo.docx", scopes=["headings"], force=True)

    assert request.force is True


def test_request_models_default_to_auto_strict_profile():
    apply_request = ApplyRequest(file_path="demo.docx")
    preflight_request = PreflightRequest(file_path="demo.docx")
    normalize_request = NormalizeRequest(file_path="demo.docx")
    normalize_job_request = NormalizeJobRequest(file_path="demo.docx")
    render_verify_request = RenderVerifyRequest(file_path="demo.docx")
    verify_request = VerifyRequest(file_path="demo.docx")

    assert apply_request.strict_profile is None
    assert preflight_request.strict_profile is None
    assert normalize_request.strict_profile is None
    assert normalize_job_request.strict_profile is None
    assert render_verify_request.strict_profile is None
    assert render_verify_request.renderer == "auto"
    assert render_verify_request.rendered_pdf is None
    assert render_verify_request.page_images_dir is None
    assert verify_request.strict_profile is None


def test_apply_request_supports_candidate_mode():
    request = ApplyRequest(file_path="demo.docx", candidate_mode="compact_candidate")

    assert request.candidate_mode == "compact_candidate"


def test_render_verify_request_rejects_artifact_tool_renderer():
    with pytest.raises(Exception):
        RenderVerifyRequest(file_path="demo.docx", renderer="artifact-tool")


def test_render_workflow_modes_payload_describes_pdf_and_candidate_modes():
    payload = app_module.build_render_workflow_modes_payload()

    mode_ids = [item["id"] for item in payload["modes"]]
    assert mode_ids == ["default_user", "agent_candidate"]
    assert payload["recommended_mode"] == "default_user"
    assert payload["modes"][0]["requires_manual_pdf"] is True
    assert payload["modes"][0]["uses_automation"] is False
    assert payload["modes"][1]["creates_candidate_docx"] is True
    assert payload["modes"][1]["candidate_modes"] == ["fast_candidate", "compact_candidate"]
    assert any("Word/WPS" in issue for issue in payload["render_layer_issues"])


def test_render_verify_request_supports_workflow_mode():
    request = RenderVerifyRequest(file_path="demo.docx", workflow_mode="default_user", rendered_pdf="/tmp/demo.pdf")

    assert request.workflow_mode == "default_user"


def test_verify_request_supports_staging_fields():
    request = VerifyRequest(file_path="demo.docx", stage_input=True, runtime_root="/tmp/article-runtime")

    assert request.stage_input is True
    assert request.runtime_root == "/tmp/article-runtime"


def test_upload_apply_stages_input_by_default():
    assert UploadApplyRequest().stage_input is True


def test_verify_request_supports_worker_control_fields():
    request = VerifyRequest(
        file_path="demo.docx",
        max_attempts=3,
        retry_delay_seconds=0.25,
        timeout_seconds=4.0,
    )

    assert request.max_attempts == 3
    assert request.retry_delay_seconds == 0.25
    assert request.timeout_seconds == 4.0


def test_audit_endpoint_returns_fields_required_for_rule_matrix(monkeypatch, tmp_docx):
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = tmp_docx(make_compliant_doc, filename="article_api_rule_matrix.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    RULE_MUTATORS["KW01"](doc)
    doc.save(source_path)

    payload = routes["/audit"].endpoint(AuditRequest(file_path=str(source_path)))

    assert payload["summary"]["total_rules"] == len(payload["results"])
    assert payload["summary"]["failed_rules"] == len(payload["failed_results"])
    result_ids = {item["id"] for item in payload["results"]}
    assert {item["id"] for item in payload["failed_results"]}.issubset(result_ids)
    required_keys = {"id", "name", "severity", "passed", "issues", "affected", "action"}
    for item in payload["results"]:
        assert required_keys.issubset(item)
        assert isinstance(item["issues"], list)
        assert isinstance(item["affected"], list)
    failed_by_id = {item["id"]: item for item in payload["failed_results"]}
    assert failed_by_id["H02"]["action"] == "autofix"
    assert failed_by_id["KW01"]["action"] == "manual_review"


def test_retention_sweep_request_supports_optional_thresholds():
    request = RetentionSweepRequest(job_max_age_seconds=60, upload_max_age_seconds=120, dry_run=True)

    assert request.job_max_age_seconds == 60
    assert request.upload_max_age_seconds == 120
    assert request.dry_run is True


def test_fake_app_health_ready_version_and_summary_endpoints(monkeypatch, tmp_docx, tmp_path):
    monkeypatch.delenv("ARTICLE_LOCAL_RELEASE_API_URL", raising=False)
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = tmp_docx(make_compliant_doc, filename="article_api_summary.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    doc.save(source_path)

    health_payload = routes["/health"].endpoint()
    version_payload = routes["/version"].endpoint()
    updates_payload = routes["/updates/latest"].endpoint()
    profiles_payload = routes["/profiles"].endpoint()
    preflight_payload = routes["/preflight"].endpoint(
        PreflightRequest(
            file_path=str(source_path),
        )
    )
    ready_payload = routes["/ready"].endpoint()
    runtime_payload = routes["/ops/runtime"].endpoint()
    storage_payload = routes["/ops/storage"].endpoint()
    create_payload = routes["/jobs/verify"].endpoint(
        VerifyRequest(
            file_path=str(source_path),
            scopes=["headings"],
        )
    )
    wait_for_job(create_payload["job_id"])
    summary_payload = routes["/ops/summary"].endpoint()

    assert "/" in routes
    assert "/static/{asset_path:path}" in routes
    assert health_payload["service"] == "article-api"
    assert health_payload["status"] == "ok"
    assert version_payload["version"] == "0.1.2"
    assert version_payload["api_version"] == "v0"
    assert version_payload["update_check"]["mode"] == "manual"
    assert version_payload["update_check"]["configured"] is False
    assert version_payload["update_check"]["auto_update"] is False
    assert "只检查软件版本，不上传论文" in version_payload["update_check"]["privacy"]
    assert updates_payload["status"] == "not_configured"
    assert updates_payload["current_version"] == "0.1.2"
    assert updates_payload["update_available"] is False
    assert updates_payload["auto_update"] is False
    assert "ARTICLE_LOCAL_RELEASE_API_URL" in updates_payload["next_action"]
    assert preflight_payload["operation"] == "preflight"
    assert preflight_payload["document"]["name"] == "article_api_summary.docx"
    assert preflight_payload["profile"]["id"] == "lnu-checker-2026"
    assert preflight_payload["preflight_status"] in {"ready", "warning", "blocked"}
    assert "diagnostics" in preflight_payload
    assert preflight_payload["summary"]["wild_doc_signal_count"] >= 0
    assert profiles_payload["summary"]["default_profile_id"] == "lnu-checker-2026"
    assert profiles_payload["summary"]["profile_count"] == 1
    assert profiles_payload["summary"]["support_scenario_count"] == 1
    assert [item["id"] for item in profiles_payload["summary"]["support_scenarios"]] == ["school_degree_thesis"]
    assert not any(item["id"] == "ams-graduate" for item in profiles_payload["profiles"])
    assert [item["id"] for item in profiles_payload["profiles"]] == ["lnu-checker-2026"]
    assert "cn-common" not in str(profiles_payload).lower()
    assert "课程作业" not in str(profiles_payload)
    assert "综述" not in str(profiles_payload)
    default_profile = profiles_payload["profiles"][0]
    assert default_profile["support_level_label"] == "一等支持"
    assert any(item["label"] == "学校学位论文" for item in default_profile["support_scenarios"])
    assert default_profile["document_types"] == ["本科毕业论文"]
    assert ready_payload["status"] == "ready"
    assert ready_payload["checks"]["storage"]["status"] == "ok"
    assert ready_payload["checks"]["runtime_root"]["status"] == "ok"
    assert runtime_payload["status"] == "ok"
    assert runtime_payload["runtime"]["worker_model"] == "single"
    assert runtime_payload["summary"]["worker_model"] == "single"
    assert runtime_payload["summary"]["status"] == "ok"
    assert runtime_payload["runtime"]["executor"]["max_workers"] == 1
    assert runtime_payload["runtime"]["recovery"]["strategy"] == "fail_unfinished_without_active_worker_future"
    assert storage_payload["status"] == "ok"
    assert storage_payload["storage"]["status"] == "ok"
    assert storage_payload["storage"]["integrity_check"] == "ok"
    assert storage_payload["summary"]["schema_version"] >= 1
    assert storage_payload["summary"]["index_count"] >= 1
    assert storage_payload["runtime_root"]["managed_directory_count"] >= 0
    assert summary_payload["status"] == "ok"
    assert summary_payload["checks"]["storage"]["status"] == "ok"
    assert summary_payload["checks"]["runtime_worker"]["status"] == "ok"
    assert summary_payload["jobs"]["total"] >= 1
    assert summary_payload["jobs"]["succeeded"] >= 1
    assert summary_payload["uploads"]["total"] == 0
    assert summary_payload["storage"]["schema_version"] >= 1
    assert summary_payload["storage"]["index_count"] >= 1
    assert summary_payload["cleanup"]["jobs_cleaned"] == 0
    assert summary_payload["runtime"]["worker_model"] == "single"
    assert summary_payload["runtime"]["recovered_failed_count"] == 0
    assert summary_payload["runtime"]["pending_recovery_count"] == 0
    assert summary_payload["retention"]["defaults"]["autorun_enabled"] is False


def test_fake_app_profiles_route_uses_app_module_profile_catalog(monkeypatch):
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    calls = []

    def fake_profile_catalog(**kwargs):
        calls.append(kwargs)
        return {"profiles": [{"id": "patched-profile"}], "summary": {"profile_count": 1}}

    monkeypatch.setattr(app_module, "build_profile_catalog", fake_profile_catalog)

    payload = routes["/profiles"].endpoint()

    assert payload["profiles"] == [{"id": "patched-profile"}]
    assert calls == [
        {
            "service_name": app_module.SERVICE_NAME,
            "stage": app_module.SERVICE_STAGE,
            "version": app_module.SERVICE_VERSION,
            "api_version": app_module.API_VERSION,
        }
    ]


def test_update_check_payload_fetches_configured_github_release_without_document_data(monkeypatch):
    monkeypatch.setenv(
        "ARTICLE_LOCAL_RELEASE_API_URL",
        "https://api.github.com/repos/example/article/releases/latest",
    )
    calls = []

    def fake_fetch(url, *, timeout_seconds):
        calls.append((url, timeout_seconds))
        return {
            "tag_name": "v0.2.0",
            "name": "v0.2.0",
            "html_url": "https://github.com/example/article/releases/tag/v0.2.0",
            "published_at": "2026-06-06T00:00:00Z",
            "assets": [
                {
                    "name": "lnu-thesis-local-windows.zip",
                    "browser_download_url": "https://github.com/example/article/releases/download/v0.2.0/lnu-thesis-local-windows.zip",
                },
                {
                    "name": "lnu-thesis-local-windows.zip.sha256",
                    "browser_download_url": "https://github.com/example/article/releases/download/v0.2.0/lnu-thesis-local-windows.zip.sha256",
                },
            ],
        }

    payload = app_ops.build_latest_update_payload(fetch_release_fn=fake_fetch, current_version="0.1.0")

    assert calls == [("https://api.github.com/repos/example/article/releases/latest", 5.0)]
    assert payload["status"] == "ok"
    assert payload["configured"] is True
    assert payload["current_version"] == "0.1.0"
    assert payload["latest_version"] == "0.2.0"
    assert payload["latest_tag"] == "v0.2.0"
    assert payload["update_available"] is True
    assert payload["release_url"] == "https://github.com/example/article/releases/tag/v0.2.0"
    assert payload["download_url"].endswith("/lnu-thesis-local-windows.zip")
    assert payload["auto_update"] is False
    assert "只检查软件版本，不上传论文" in payload["privacy"]
    assert "file_path" not in payload
    assert "api_key" not in str(payload).lower()


def test_update_check_payload_accepts_github_release_tag_api_url(monkeypatch):
    release_api_url = "https://api.github.com/repos/example/article/releases/tags/v0.1.0-beta"
    monkeypatch.setenv("ARTICLE_LOCAL_RELEASE_API_URL", release_api_url)
    calls = []

    def fake_fetch(url, *, timeout_seconds):
        calls.append((url, timeout_seconds))
        return {
            "tag_name": "v0.1.0-beta",
            "name": "v0.1.0-beta",
            "html_url": "https://github.com/example/article/releases/tag/v0.1.0-beta",
            "published_at": "2026-06-07T14:20:15Z",
            "assets": [
                {
                    "name": "lnu-thesis-local-windows.zip",
                    "browser_download_url": "https://github.com/example/article/releases/download/v0.1.0-beta/lnu-thesis-local-windows.zip",
                },
            ],
        }

    payload = app_ops.build_latest_update_payload(fetch_release_fn=fake_fetch, current_version="0.1.0")

    assert calls == [(release_api_url, 5.0)]
    assert payload["status"] == "ok"
    assert payload["configured"] is True
    assert payload["latest_version"] == "0.1.0-beta"
    assert payload["latest_tag"] == "v0.1.0-beta"
    assert payload["update_available"] is False
    assert payload["release_url"] == "https://github.com/example/article/releases/tag/v0.1.0-beta"
    assert payload["download_url"].endswith("/lnu-thesis-local-windows.zip")
    assert payload["auto_update"] is False
    assert "只检查软件版本，不上传论文" in payload["privacy"]


def test_update_check_payload_does_not_fetch_when_release_url_is_unconfigured(monkeypatch):
    monkeypatch.delenv("ARTICLE_LOCAL_RELEASE_API_URL", raising=False)

    def fail_fetch(_url, *, timeout_seconds):
        raise AssertionError("update check must not contact the network until configured")

    payload = app_ops.build_latest_update_payload(fetch_release_fn=fail_fetch, current_version="0.1.0")

    assert payload["status"] == "not_configured"
    assert payload["configured"] is False
    assert payload["update_available"] is False
    assert payload["auto_update"] is False
    assert "ARTICLE_LOCAL_RELEASE_API_URL" in payload["next_action"]


def test_update_check_payload_rejects_non_github_release_api_url(monkeypatch):
    monkeypatch.setenv("ARTICLE_LOCAL_RELEASE_API_URL", "https://example.com/not-github/releases/latest")

    def fail_fetch(_url, *, timeout_seconds):
        raise AssertionError("update check must reject non-GitHub Release API URLs before fetch")

    payload = app_ops.build_latest_update_payload(fetch_release_fn=fail_fetch, current_version="0.1.0")

    assert payload["status"] == "invalid_config"
    assert payload["configured"] is False
    assert payload["update_available"] is False
    assert payload["auto_update"] is False
    assert "api.github.com/repos/<owner>/<repo>/releases/latest" in payload["next_action"]


def test_update_check_payload_handles_malformed_release_response(monkeypatch):
    monkeypatch.setenv(
        "ARTICLE_LOCAL_RELEASE_API_URL",
        "https://api.github.com/repos/example/article/releases/latest",
    )

    payload = app_ops.build_latest_update_payload(
        fetch_release_fn=lambda _url, *, timeout_seconds: [],
        current_version="0.1.0",
    )

    assert payload["status"] == "error"
    assert payload["update_available"] is False
    assert payload["auto_update"] is False
    assert "GitHub Release" in payload["next_action"]


def test_update_check_payload_does_not_announce_update_for_unparseable_tag(monkeypatch):
    monkeypatch.setenv(
        "ARTICLE_LOCAL_RELEASE_API_URL",
        "https://api.github.com/repos/example/article/releases/latest",
    )

    payload = app_ops.build_latest_update_payload(
        fetch_release_fn=lambda _url, *, timeout_seconds: {
            "tag_name": "release-0.1.0",
            "html_url": "https://github.com/example/article/releases/tag/release-0.1.0",
            "assets": [],
        },
        current_version="0.1.0",
    )

    assert payload["status"] == "ok"
    assert payload["latest_version"] == "release-0.1.0"
    assert payload["update_available"] is False
    assert "当前已是最新版本" in payload["next_action"]


def test_fake_app_pdf_upload_endpoint_stores_pdf(tmp_path, monkeypatch):
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)

    class FakeUpload:
        filename = "20221303306-刘佳轾-排版复核.pdf"

        def __init__(self, payload: bytes):
            self.file = BytesIO(payload)

    upload_response = routes["/uploads/pdf"].endpoint(
        FakeUpload(b"%PDF-1.7\nfake-pdf-binary"),
        runtime_root=str(tmp_path / "runtime"),
    )

    assert routes["/uploads/pdf"].status_code == 201
    assert upload_response["file_name"] == "20221303306-刘佳轾-排版复核.pdf"
    assert upload_response["size_bytes"] == len(b"%PDF-1.7\nfake-pdf-binary")
    assert upload_response["available"] is True
    assert Path(upload_response["stored_path"]).exists()
    assert Path(upload_response["stored_path"]).read_bytes() == b"%PDF-1.7\nfake-pdf-binary"


def test_fake_app_preflight_endpoint_detects_wild_doc_signals(monkeypatch, tmp_path):
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = _make_style_conflict_lnu_doc(tmp_path / "article_api_preflight.docx")

    payload = routes["/preflight"].endpoint(
        PreflightRequest(
            file_path=str(source_path),
            profile="lnu",
        )
    )

    assert payload["status"] == "ok"
    assert payload["operation"] == "preflight"
    assert payload["preflight_status"] == "blocked"
    assert payload["summary"]["wild_doc_detected"] is True
    assert payload["summary"]["wild_doc_signal_count"] >= 1
    assert any(item["id"] == "style_text_conflicts" for item in payload["wild_doc"]["signals"])
    assert payload["diagnostics"]["style_text_conflicts"]


def test_build_normalize_payload_wraps_engine_result(monkeypatch):
    monkeypatch.setattr(
        app_module,
        "normalize_document",
        lambda *args, **kwargs: {
            "document": {"path": "/tmp/demo.docx", "name": "demo.docx"},
            "output": {"path": "/tmp/demo_normalized.docx", "name": "demo_normalized.docx"},
            "profile": {"id": "lnu-checker-2026", "requested": "lnu", "fallback_used": False, "display": "lnu"},
            "changed": True,
            "operations": [{"id": "heading_styles", "label": "扶正高置信度标题样式", "count": 1}],
            "summary": {
                "operation_count": 1,
                "before_preflight_status": "blocked",
                "after_preflight_status": "warning",
            },
            "before": {
                "preflight_status": "blocked",
                "toc_status": "no_toc",
                "style_conflict_count": 1,
                "table_heading_risk_count": 0,
            },
            "after": {
                "preflight_status": "warning",
                "toc_status": "no_toc",
                "style_conflict_count": 0,
                "table_heading_risk_count": 0,
            },
            "next_steps": ["继续 verify。"],
        },
    )

    payload = app_module.build_normalize_payload(
        file_path="/tmp/demo.docx",
        output_path="/tmp/demo_normalized.docx",
        profile_path="lnu",
    )

    assert payload["status"] == "ok"
    assert payload["operation"] == "normalize"
    assert payload["changed"] is True
    assert payload["summary"]["before_preflight_status"] == "blocked"
    assert payload["wild_doc"]["after"]["style_conflict_count"] == 0


def test_fake_app_normalize_endpoint_returns_wild_doc_delta(monkeypatch):
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    monkeypatch.setattr(
        app_module,
        "normalize_document",
        lambda *args, **kwargs: {
            "document": {"path": "/tmp/demo.docx", "name": "demo.docx"},
            "output": {"path": "/tmp/demo_normalized.docx", "name": "demo_normalized.docx"},
            "profile": {"id": "lnu-checker-2026", "requested": "lnu", "fallback_used": False, "display": "lnu"},
            "changed": True,
            "operations": [{"id": "heading_styles", "label": "扶正高置信度标题样式", "count": 1}],
            "summary": {
                "operation_count": 1,
                "before_preflight_status": "blocked",
                "after_preflight_status": "warning",
            },
            "before": {
                "preflight_status": "blocked",
                "toc_status": "no_toc",
                "style_conflict_count": 1,
                "table_heading_risk_count": 0,
            },
            "after": {
                "preflight_status": "warning",
                "toc_status": "no_toc",
                "style_conflict_count": 0,
                "table_heading_risk_count": 0,
            },
            "next_steps": ["继续 verify。"],
        },
    )

    payload = routes["/normalize"].endpoint(
        NormalizeRequest(
            file_path="/tmp/demo.docx",
            profile="lnu",
        )
    )

    assert payload["status"] == "ok"
    assert payload["operation"] == "normalize"
    assert payload["changed"] is True
    assert payload["wild_doc"]["before"]["preflight_status"] == "blocked"
    assert payload["wild_doc"]["after"]["style_conflict_count"] == 0


def test_build_render_verify_payload_wraps_engine_result(monkeypatch, tmp_path):
    output_dir = tmp_path / "render-proof"
    output_dir.mkdir()
    (output_dir / "page-1.png").write_bytes(b"png")
    (output_dir / "page-2.png").write_bytes(b"png")

    monkeypatch.setattr(
        app_module,
        "render_verify_document",
        lambda *args, **kwargs: {
            "document": {"path": "/tmp/demo.docx", "name": "demo.docx"},
            "profile": {"id": "lnu-checker-2026", "requested": "lnu", "fallback_used": False, "display": "lnu"},
            "output_dir": str(output_dir),
            "render_engine": "word-pdf",
            "evidence_source": "word-pdf",
            "evidence_trust": "authoritative",
            "evidence_authoritative": True,
            "layout_decision_eligible": True,
            "render_fallback_used": False,
            "page_count": 2,
            "page_images": [str(output_dir / "page-1.png"), str(output_dir / "page-2.png")],
            "evidence_items": [
                {
                    "page": 2,
                    "screenshot_path": str(output_dir / "page-2.png"),
                    "screenshot_url": "/render-evidence/screenshot/test-token",
                    "rule_id": "render.object_flow",
                    "bbox": {"x": 0.12, "y": 0.64, "w": 0.72, "h": 0.18},
                    "message": "页底留白需要人工复核",
                    "severity": "warning",
                    "next_action": "回到 WPS/Word 调整图片大小或分页设置后重新导出 PDF",
                }
            ],
            "render_findings": [
                {
                    "id": "isolated_punctuation",
                    "rule_id": "render.isolated_punctuation",
                    "severity": "warning",
                    "page": 2,
                    "message": "页面文本存在单独成行的标点，疑似换行或排版挤压导致。",
                }
            ],
            "render_summary": {
                "finding_count": 1,
                "highest_severity": "warning",
                "actionable_finding_count": 1,
                "expected_blank_count": 0,
                "object_flow_issue_count": 0,
                "heading_break_issue_count": 0,
                "isolated_punctuation_count": 1,
            },
            "layout_score": {
                "score": 100,
                "penalty": 0,
                "expected_blank_count": 0,
                "actionable_finding_count": 0,
                "object_flow_issue_count": 1,
                "heading_break_issue_count": 0,
                "render_integrity_issue_count": 0,
            },
            "render_text_summary": {
                "source": "pdf",
                "available": True,
                "page_text_available_count": 2,
                "page_text_extraction_warning_count": 0,
                "warnings": [],
            },
            "selected_scopes": ["toc"],
            "overall_status": "verified",
            "readiness": "render-check-required",
            "manual_review_rule_ids": ["LNU_TOC03"],
            "unsupported_rule_ids": [],
            "review_items": ["目录需复核页码、层级和可见目录结果。"],
            "report_path": str(output_dir / "render_verify_report.md"),
        },
    )

    payload = app_module.build_render_verify_payload(
        file_path="/tmp/demo.docx",
        output_dir=str(output_dir),
        profile_path="lnu",
        scopes=["toc"],
        renderer="auto",
        workflow_mode="default_user",
        rendered_pdf="/tmp/demo.pdf",
    )

    assert payload["status"] == "ok"
    assert payload["operation"] == "render-verify"
    assert payload["summary"]["page_count"] == 2
    assert payload["summary"]["render_engine"] == "word-pdf"
    assert payload["summary"]["evidence_source"] == "word-pdf"
    assert payload["summary"]["evidence_trust"] == "authoritative"
    assert payload["summary"]["layout_decision_eligible"] is True
    assert payload["summary"]["render_fallback_used"] is False
    assert payload["summary"]["render_finding_count"] == 1
    assert payload["summary"]["render_highest_severity"] == "warning"
    assert payload["summary"]["layout_score"] == 100
    assert payload["summary"]["layout_penalty"] == 0
    assert payload["summary"]["actionable_finding_count"] == 1
    assert payload["summary"]["isolated_punctuation_count"] == 1
    assert payload["summary"]["page_text_available_count"] == 2
    assert payload["summary"]["evidence_item_count"] == 1
    assert payload["evidence_items"][0]["rule_id"] == "render.object_flow"
    assert payload["evidence_items"][0]["bbox"]["x"] == 0.12
    assert payload["evidence_items"][0]["screenshot_url"].startswith("/render-evidence/screenshot/")
    assert "path=" not in payload["evidence_items"][0]["screenshot_url"]
    assert payload["summary"]["review_item_count"] == 1
    assert payload["summary"]["manual_review_rule_count"] == 1
    assert payload["render_workflow_mode"]["id"] == "default_user"
    assert payload["selected_scopes"] == ["toc"]


def test_render_evidence_screenshot_route_uses_registered_token(monkeypatch, tmp_path):
    output_dir = tmp_path / "render-proof"
    output_dir.mkdir()
    screenshot_path = output_dir / "page-1.png"
    screenshot_path.write_bytes(b"png")
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)

    monkeypatch.setattr(
        app_module,
        "render_verify_document",
        lambda *args, **kwargs: {
            "document": {"path": "/tmp/demo.docx", "name": "demo.docx"},
            "profile": {"id": "lnu-checker-2026", "requested": "lnu", "fallback_used": False, "display": "lnu"},
            "output_dir": str(output_dir),
            "render_engine": "manual-pdf",
            "evidence_source": "manual-pdf",
            "evidence_trust": "authoritative",
            "evidence_authoritative": True,
            "layout_decision_eligible": True,
            "render_fallback_used": False,
            "page_count": 1,
            "page_images": [str(screenshot_path)],
            "evidence_items": [
                {
                    "page": 1,
                    "screenshot_path": str(screenshot_path),
                    "rule_id": "render.object_flow",
                    "bbox": {"x": 0.1, "y": 0.1, "w": 0.2, "h": 0.2},
                    "message": "页底留白需要人工复核",
                    "severity": "warning",
                    "next_action": "回到 WPS/Word 调整图片大小或分页设置后重新导出 PDF",
                }
            ],
            "render_findings": [],
            "render_summary": {"finding_count": 0, "highest_severity": None, "actionable_finding_count": 0},
            "layout_score": {"score": 100, "penalty": 0},
            "render_text_summary": {"page_text_available_count": 0, "page_text_extraction_warning_count": 0},
            "selected_scopes": ["toc"],
            "overall_status": "verified",
            "readiness": "render-check-required",
            "manual_review_rule_ids": [],
            "unsupported_rule_ids": [],
            "review_items": [],
            "report_path": str(output_dir / "render_verify_report.md"),
        },
    )

    payload = app_module.build_render_verify_payload(
        file_path="/tmp/demo.docx",
        profile_path="lnu",
        workflow_mode="default_user",
        rendered_pdf="/tmp/demo.pdf",
    )
    screenshot_url = payload["evidence_items"][0]["screenshot_url"]
    token = screenshot_url.rsplit("/", 1)[-1]

    response = routes["/render-evidence/screenshot/{token}"].endpoint(token)
    assert response["path"] == str(screenshot_path.resolve())
    assert response["media_type"] == "image/png"
    assert "path=" not in screenshot_url
    with pytest.raises(FakeHTTPException) as exc_info:
        routes["/render-evidence/screenshot/{token}"].endpoint(str(screenshot_path))
    assert exc_info.value.status_code == 404


def test_render_verify_job_result_registers_evidence_screenshot_from_stored_path(tmp_path, monkeypatch):
    from article_api.jobs import get_job_result
    from article_api import render_evidence

    monkeypatch.setenv("ARTICLE_API_STATE_ROOT", str(tmp_path / "state"))
    clear_jobs()
    screenshot = tmp_path / "page-1.png"
    screenshot.write_bytes(b"png")
    storage.upsert_job(
        {
            "job_id": "job-render-1",
            "operation": "render-verify",
            "status": "succeeded",
            "mode": "background",
            "created_at": "2026-06-16T00:00:00Z",
            "updated_at": "2026-06-16T00:00:01Z",
            "request": {},
            "resolved_request": {"file_path": "/tmp/demo.docx", "workflow_mode": "default_user"},
            "workspace": None,
            "runtime": None,
            "summary": {},
            "artifacts": [],
            "result": {
                "evidence_items": [
                    {
                        "page": 1,
                        "screenshot_path": str(screenshot),
                        "rule_id": "render.isolated_punctuation",
                        "bbox": {"x": 0.1, "y": 0.2, "w": 0.3, "h": 0.4},
                        "message": "需要确认。",
                        "severity": "warning",
                        "next_action": "重新导出 PDF。",
                    }
                ]
            },
            "error": None,
        }
    )

    payload = get_job_result("job-render-1")
    url = payload["result"]["evidence_items"][0]["screenshot_url"]
    token = url.rsplit("/", 1)[-1]

    assert render_evidence.resolve_render_evidence_screenshot(token) == str(screenshot.resolve())
    render_evidence._SCREENSHOT_PATHS_BY_TOKEN.clear()
    assert render_evidence.resolve_render_evidence_screenshot(token) == str(screenshot.resolve())


def test_build_render_verify_payload_requires_pdf_for_default_user_mode(monkeypatch):
    monkeypatch.setattr(app_module, "render_verify_document", lambda *args, **kwargs: {})

    with pytest.raises(ValueError, match="PDF 版式复核"):
        app_module.build_render_verify_payload(
            file_path="/tmp/demo.docx",
            profile_path="lnu",
            workflow_mode="default_user",
        )


def test_fake_app_render_verify_endpoint_returns_proof_summary(monkeypatch):
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    monkeypatch.setattr(
        app_module,
        "render_verify_document",
        lambda *args, **kwargs: {
            "document": {"path": "/tmp/demo.docx", "name": "demo.docx"},
            "profile": {"id": "lnu-checker-2026", "requested": "lnu", "fallback_used": False, "display": "lnu"},
            "output_dir": "/tmp/render-proof",
            "render_engine": "word-pdf",
            "evidence_source": "word-pdf",
            "evidence_trust": "authoritative",
            "evidence_authoritative": True,
            "layout_decision_eligible": True,
            "render_fallback_used": False,
            "page_count": 1,
            "page_images": ["/tmp/render-proof/page-1.png"],
            "render_findings": [],
            "render_summary": {"finding_count": 0, "highest_severity": None},
            "selected_scopes": ["figures_tables"],
            "overall_status": "verified",
            "readiness": "render-check-required",
            "manual_review_rule_ids": [],
            "unsupported_rule_ids": [],
            "review_items": ["逐页检查图表是否与题注分离。"],
            "report_path": "/tmp/render-proof/render_verify_report.md",
        },
    )

    payload = routes["/render-verify"].endpoint(
        RenderVerifyRequest(
            file_path="/tmp/demo.docx",
            profile="lnu",
            scopes=["figures_tables"],
            renderer="word-pdf",
            page_images_dir="/tmp/wps-pages",
        )
    )

    assert payload["status"] == "ok"
    assert payload["operation"] == "render-verify"
    assert payload["page_count"] == 1
    assert payload["summary"]["render_engine"] == "word-pdf"
    assert payload["summary"]["render_finding_count"] == 0
    assert payload["summary"]["review_item_count"] == 1
    assert payload["selected_scopes"] == ["figures_tables"]


def test_fake_app_ready_maps_storage_failure_to_503(monkeypatch):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    monkeypatch.setattr(app_module.storage, "init_storage", lambda *_args, **_kwargs: (_ for _ in ()).throw(RuntimeError("db offline")))

    with pytest.raises(FakeHTTPException) as exc_info:
        routes["/ready"].endpoint()

    assert exc_info.value.status_code == 503


def test_fake_app_apply_endpoint_maps_guard_block_to_409(monkeypatch, tmp_path):
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = _make_style_conflict_lnu_doc(tmp_path / "article_api_guard_block.docx")

    with pytest.raises(FakeHTTPException) as exc_info:
        routes["/apply"].endpoint(
            ApplyRequest(
                file_path=str(source_path),
                profile="lnu",
                scopes=["headings"],
                renumber_headings=True,
            )
        )

    assert exc_info.value.status_code == 409
    assert "Apply blocked by structural risk" in exc_info.value.detail


def test_fake_app_job_endpoints_return_created_status_and_lookup(monkeypatch, tmp_docx):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = tmp_docx(make_compliant_doc, filename="article_api_job_verify.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    doc.save(source_path)

    create_response = routes["/jobs/verify"].endpoint(
        VerifyRequest(
            file_path=str(source_path),
            scopes=["headings"],
        )
    )
    wait_for_job(create_response["job_id"])
    list_response = routes["/jobs"].endpoint()
    status_response = routes["/jobs/{job_id}"].endpoint(create_response["job_id"])
    inspect_response = routes["/jobs/{job_id}/inspect"].endpoint(create_response["job_id"])
    result_response = routes["/jobs/{job_id}/result"].endpoint(create_response["job_id"])

    assert routes["/jobs/verify"].status_code == 201
    assert create_response["status"] == "queued"
    assert any(item["job_id"] == create_response["job_id"] for item in list_response)
    assert status_response["status"] == "succeeded"
    assert status_response["summary"]["readiness"] == "manual-review-required"
    assert inspect_response["status"] == "succeeded"
    assert "request" not in inspect_response
    assert "result" not in inspect_response
    assert result_response["result"]["overall_status"] == "needs_fix"
    assert result_response["result"]["readiness"] == "manual-review-required"


def test_fake_app_normalize_job_endpoint_runs_and_returns_output(monkeypatch, tmp_path):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = _make_style_conflict_lnu_doc(tmp_path / "article_api_job_normalize.docx")
    output_path = tmp_path / "article_api_job_normalized.docx"

    create_response = routes["/jobs/normalize"].endpoint(
        NormalizeJobRequest(
            file_path=str(source_path),
            output_path=str(output_path),
            profile="lnu",
        )
    )
    wait_for_job(create_response["job_id"])
    result_response = routes["/jobs/{job_id}/result"].endpoint(create_response["job_id"])

    assert routes["/jobs/normalize"].status_code == 201
    assert result_response["status"] == "succeeded"
    assert result_response["summary"]["business_status"] == "warning"
    assert result_response["summary"]["output_path"] == str(output_path)
    assert result_response["result"]["changed"] is True
    assert result_response["result"]["after"]["preflight_status"] == "warning"


def test_fake_app_jobs_endpoint_supports_operation_status_and_limit_filters(monkeypatch, tmp_docx):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    verify_source = tmp_docx(make_compliant_doc, filename="article_api_jobs_filter_verify.docx")
    verify_doc = Document(verify_source)
    RULE_MUTATORS["H02"](verify_doc)
    verify_doc.save(verify_source)
    normalize_output = tmp_docx(make_compliant_doc, filename="article_api_jobs_filter_normalized.docx")
    normalize_output.unlink()

    verify_job = routes["/jobs/verify"].endpoint(
        VerifyRequest(
            file_path=str(verify_source),
            scopes=["headings"],
        )
    )
    normalize_job = routes["/jobs/normalize"].endpoint(
        NormalizeJobRequest(
            file_path=str(verify_source),
            output_path=str(normalize_output),
            profile="lnu",
        )
    )
    wait_for_job(verify_job["job_id"])
    wait_for_job(normalize_job["job_id"])

    filtered = routes["/jobs"].endpoint(operation="verify", status="succeeded", limit=1)

    assert len(filtered) == 1
    assert filtered[0]["operation"] == "verify"
    assert filtered[0]["status"] == "succeeded"


def test_fake_app_job_status_maps_missing_job_to_404(monkeypatch):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)

    with pytest.raises(FakeHTTPException) as exc_info:
        routes["/jobs/{job_id}"].endpoint("missing-job")

    assert exc_info.value.status_code == 404


def test_fake_app_job_inspect_maps_missing_job_to_404(monkeypatch):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)

    with pytest.raises(FakeHTTPException) as exc_info:
        routes["/jobs/{job_id}/inspect"].endpoint("missing-job")

    assert exc_info.value.status_code == 404


def test_fake_app_job_submit_invalid_path_maps_to_400_and_does_not_store(monkeypatch, tmp_path):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    missing_path = tmp_path / "missing.docx"

    with pytest.raises(FakeHTTPException) as exc_info:
        routes["/jobs/verify"].endpoint(VerifyRequest(file_path=str(missing_path)))

    assert exc_info.value.status_code == 400
    assert job_count() == 0


def test_fake_app_job_apply_returns_failed_job_payload_on_guard_block(monkeypatch, tmp_path):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = _make_style_conflict_lnu_doc(tmp_path / "article_api_job_apply_fail.docx")

    response = routes["/jobs/apply"].endpoint(
        ApplyRequest(
            file_path=str(source_path),
            profile="lnu",
            scopes=["headings"],
            renumber_headings=True,
        )
    )
    wait_for_job(response["job_id"])
    status = routes["/jobs/{job_id}"].endpoint(response["job_id"])

    assert routes["/jobs/apply"].status_code == 201
    assert response["status"] == "queued"
    assert status["status"] == "failed"
    assert status["summary"]["error_code"] == "apply_guard_blocked"
    assert status["error"]["code"] == "apply_guard_blocked"


def test_fake_app_job_result_maps_missing_job_to_404(monkeypatch):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)

    with pytest.raises(FakeHTTPException) as exc_info:
        routes["/jobs/{job_id}/result"].endpoint("missing-job")

    assert exc_info.value.status_code == 404


def test_fake_app_job_verify_can_stage_input(monkeypatch, tmp_docx, tmp_path):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = tmp_docx(make_compliant_doc, filename="article_api_job_stage_verify.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    doc.save(source_path)

    response = routes["/jobs/verify"].endpoint(
        VerifyRequest(
            file_path=str(source_path),
            scopes=["headings"],
            stage_input=True,
            runtime_root=str(tmp_path / "runtime"),
        )
    )
    wait_for_job(response["job_id"])
    result = routes["/jobs/{job_id}/result"].endpoint(response["job_id"])

    assert response["resolved_request"]["file_path"].startswith(response["workspace"]["inputs"])
    assert result["result"]["document"]["name"] == "article_api_job_stage_verify.docx"


def test_fake_app_upload_endpoint_stores_docx_and_job_downloads_output(monkeypatch, tmp_docx, tmp_path):
    clear_jobs()
    clear_uploads()
    monkeypatch.setattr(output_naming, "DEFAULT_OUTPUT_DIR", tmp_path / "versioned-outputs")
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = tmp_docx(make_compliant_doc, filename="article_api_upload_apply.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    doc.save(source_path)

    class FakeUpload:
        filename = "article_api_upload_apply.docx"

        def __init__(self, payload: bytes):
            self.file = BytesIO(payload)

    upload_response = routes["/uploads/docx"].endpoint(
        FakeUpload(Path(source_path).read_bytes()),
        runtime_root=str(tmp_path / "runtime"),
    )
    uploads_response = routes["/uploads"].endpoint()
    upload_status = routes["/uploads/{upload_id}"].endpoint(upload_response["upload_id"])
    create_response = routes["/uploads/{upload_id}/jobs/apply"].endpoint(
        upload_response["upload_id"],
        UploadApplyRequest(
            scopes=["headings"],
        ),
    )
    wait_for_job(create_response["job_id"])
    status = routes["/jobs/{job_id}"].endpoint(create_response["job_id"])
    result = routes["/jobs/{job_id}/result"].endpoint(create_response["job_id"])
    download_response = routes["/jobs/{job_id}/artifacts/{artifact_role}/download"].endpoint(
        create_response["job_id"],
        "output",
    )

    assert routes["/uploads/docx"].status_code == 201
    assert upload_count() == 1
    assert uploads_response[0]["upload_id"] == upload_response["upload_id"]
    assert upload_status["available"] is True
    assert upload_response["file_name"] == "article_api_upload_apply.docx"
    assert Path(upload_response["stored_path"]).exists()
    assert create_response["status"] == "queued"
    assert create_response["request"]["upload_id"] == upload_response["upload_id"]
    assert "file_path" not in create_response["request"]
    assert status["summary"]["document_name"] == "article_api_upload_apply.docx"
    assert Path(status["summary"]["output_path"]).name.startswith("article_api_upload_apply_格式修复_V")
    assert result["result"]["document"]["name"] == "article_api_upload_apply.docx"
    assert result["runtime"]["source_upload_id"] == upload_response["upload_id"]
    assert download_response["job_id"] == create_response["job_id"]
    assert download_response["artifact_role"] == "output"
    assert download_response["filename"].startswith("article_api_upload_apply_格式修复_V")


def test_fake_app_feedback_download_excludes_document_artifacts(monkeypatch, tmp_path):
    clear_jobs()
    clear_uploads()
    state_root = tmp_path / "state"
    runtime_root = tmp_path / "runtime"
    monkeypatch.setenv(storage.STATE_ROOT_ENV_VAR, str(state_root))
    monkeypatch.setenv("ARTICLE_API_RUNTIME_ROOT", str(runtime_root))
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    docx_path = runtime_root / "uploads" / "paper.docx"
    docx_path.parent.mkdir(parents=True)
    docx_path.write_bytes(b"docx")
    log_path = runtime_root / "jobs" / "job-1" / "logs" / "worker.log"
    log_path.parent.mkdir(parents=True)
    log_path.write_text("failed\n", encoding="utf-8")

    response = routes["/feedback/download"].endpoint()

    assert response["filename"] == "反馈包.zip"
    assert response["media_type"] == "application/zip"
    with zipfile.ZipFile(response["path"]) as archive:
        names = set(archive.namelist())
        manifest = archive.read("manifest.json").decode("utf-8")
    assert "logs/runtime_root/jobs/job-1/logs/worker.log" in names
    assert not any(name.endswith((".docx", ".pdf", ".png")) for name in names)
    assert str(tmp_path) not in manifest


def test_fake_app_feedback_route_uses_app_module_download_response_builder(monkeypatch):
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)

    monkeypatch.setattr(
        app_module,
        "_build_feedback_download_response",
        lambda: {"path": "patched-feedback.zip", "filename": "patched.zip", "media_type": "application/zip"},
    )

    response = routes["/feedback/download"].endpoint()

    assert response == {"path": "patched-feedback.zip", "filename": "patched.zip", "media_type": "application/zip"}


def test_fake_app_upload_endpoint_rejects_corrupt_docx_without_storing(monkeypatch, tmp_path):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)

    class FakeUpload:
        filename = "broken.docx"

        def __init__(self, payload: bytes):
            self.file = BytesIO(payload)

    with pytest.raises(FakeHTTPException) as exc_info:
        routes["/uploads/docx"].endpoint(
            FakeUpload(b"not a zip archive"),
            runtime_root=str(tmp_path / "runtime"),
        )

    detail = exc_info.value.detail
    assert exc_info.value.status_code == 400
    assert detail["code"] == "invalid_docx"
    assert detail["retryable"] is False
    assert "Word/WPS" in detail["next_action"]
    assert upload_count() == 0
    assert not list((tmp_path / "runtime" / "uploads").glob("*.docx"))


def test_fake_app_upload_normalize_job_keeps_original_name(monkeypatch, tmp_path):
    clear_jobs()
    clear_uploads()
    monkeypatch.setattr(output_naming, "DEFAULT_OUTPUT_DIR", tmp_path / "versioned-outputs")
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = _make_style_conflict_lnu_doc(tmp_path / "article_api_upload_normalize.docx")

    class FakeUpload:
        filename = "article_api_upload_normalize.docx"

        def __init__(self, payload: bytes):
            self.file = BytesIO(payload)

    upload_response = routes["/uploads/docx"].endpoint(
        FakeUpload(Path(source_path).read_bytes()),
        runtime_root=str(tmp_path / "runtime"),
    )
    create_response = routes["/uploads/{upload_id}/jobs/normalize"].endpoint(
        upload_response["upload_id"],
        UploadNormalizeRequest(),
    )
    wait_for_job(create_response["job_id"])
    result = routes["/jobs/{job_id}/result"].endpoint(create_response["job_id"])

    assert create_response["request"]["upload_id"] == upload_response["upload_id"]
    output_path = Path(result["summary"]["output_path"])
    assert output_path.name.startswith("article_api_upload_normalize_结构整理_V")
    assert output_path.parent != output_naming.DEFAULT_OUTPUT_DIR
    assert "intermediate" in output_path.parts
    assert result["result"]["document"]["name"] == "article_api_upload_normalize.docx"
    assert result["runtime"]["source_upload_id"] == upload_response["upload_id"]


def test_fake_app_download_endpoint_maps_missing_artifact_file_to_409(monkeypatch, tmp_docx):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = tmp_docx(make_compliant_doc, filename="article_api_download_missing.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    doc.save(source_path)

    create_response = routes["/jobs/apply"].endpoint(
        ApplyRequest(
            file_path=str(source_path),
            scopes=["headings"],
        )
    )
    wait_for_job(create_response["job_id"])
    status = routes["/jobs/{job_id}"].endpoint(create_response["job_id"])
    output_path = Path(status["summary"]["output_path"])
    output_path.unlink()

    with pytest.raises(FakeHTTPException) as exc_info:
        routes["/jobs/{job_id}/artifacts/{artifact_role}/download"].endpoint(
            create_response["job_id"],
            "output",
        )

    refreshed_status = routes["/jobs/{job_id}"].endpoint(create_response["job_id"])
    refreshed_result = routes["/jobs/{job_id}/result"].endpoint(create_response["job_id"])
    status_without_artifacts = dict(status)
    refreshed_without_artifacts = dict(refreshed_status)
    status_without_artifacts.pop("artifacts")
    refreshed_without_artifacts.pop("artifacts")
    refreshed_output = next(item for item in refreshed_status["artifacts"] if item["role"] == "output")

    assert exc_info.value.status_code == 409
    assert exc_info.value.detail["code"] == "artifact_unavailable"
    assert exc_info.value.detail["user_message"] == "下载文件已经不可用。"
    assert "重新运行修复任务" in exc_info.value.detail["next_action"]
    assert refreshed_without_artifacts == status_without_artifacts
    assert refreshed_output["available"] is False
    assert refreshed_result["status"] == "succeeded"


def test_fake_app_apply_endpoint_creates_missing_output_parent(monkeypatch, tmp_docx, tmp_path):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = tmp_docx(make_compliant_doc, filename="article_api_apply_nested_output.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    doc.save(source_path)
    output_path = tmp_path / "missing" / "nested" / "article_api_apply_nested_output_fixed.docx"

    payload = routes["/apply"].endpoint(
        ApplyRequest(
            file_path=str(source_path),
            output_path=str(output_path),
            scopes=["headings"],
        )
    )

    assert payload["output"]["path"] == str(output_path)
    assert output_path.exists()


def test_fake_app_apply_endpoint_auto_versions_display_name(monkeypatch, tmp_docx, tmp_path):
    clear_jobs()
    clear_uploads()
    monkeypatch.setattr(output_naming, "DEFAULT_OUTPUT_DIR", tmp_path / "versioned-outputs")
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = tmp_docx(make_compliant_doc, filename="runtime_apply_versioned.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    doc.save(source_path)
    existing = output_naming.DEFAULT_OUTPUT_DIR / "article_api_sync_versioned_格式修复_V01.docx"
    existing.parent.mkdir(parents=True, exist_ok=True)
    existing.write_bytes(b"existing")

    payload = routes["/apply"].endpoint(
        ApplyRequest(
            file_path=str(source_path),
            source_display_name="article_api_sync_versioned_格式修复_V01.docx",
            scopes=["headings"],
        )
    )

    assert payload["output"]["path"].endswith("article_api_sync_versioned_格式修复_V02.docx")
    assert "V01_格式修复" not in payload["output"]["path"]


def test_fake_app_upload_verify_job_supports_upload_id_and_stage_input(monkeypatch, tmp_docx, tmp_path):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = tmp_docx(make_compliant_doc, filename="article_api_upload_verify.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    doc.save(source_path)

    class FakeUpload:
        filename = "article_api_upload_verify.docx"

        def __init__(self, payload: bytes):
            self.file = BytesIO(payload)

    upload_response = routes["/uploads/docx"].endpoint(
        FakeUpload(Path(source_path).read_bytes()),
        runtime_root=str(tmp_path / "runtime"),
    )
    create_response = routes["/uploads/{upload_id}/jobs/verify"].endpoint(
        upload_response["upload_id"],
        UploadVerifyRequest(
            scopes=["headings"],
            stage_input=True,
        ),
    )
    wait_for_job(create_response["job_id"])
    result = routes["/jobs/{job_id}/result"].endpoint(create_response["job_id"])

    assert create_response["request"]["upload_id"] == upload_response["upload_id"]
    assert "file_path" not in create_response["request"]
    assert create_response["resolved_request"]["file_path"].startswith(create_response["workspace"]["inputs"])
    assert create_response["resolved_request"]["source_display_name"] == "article_api_upload_verify.docx"
    assert result["result"]["document"]["name"] == "article_api_upload_verify.docx"
    assert result["runtime"]["source_upload_id"] == upload_response["upload_id"]


def test_fake_app_upload_job_missing_upload_id_maps_to_404(monkeypatch):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)

    with pytest.raises(FakeHTTPException) as exc_info:
        routes["/uploads/{upload_id}/jobs/verify"].endpoint(
            "missing-upload",
            UploadVerifyRequest(scopes=["headings"]),
        )

    assert exc_info.value.status_code == 404
    assert job_count() == 0


def test_fake_app_upload_registry_marks_missing_file_unavailable(monkeypatch, tmp_docx, tmp_path):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = tmp_docx(make_compliant_doc, filename="article_api_upload_missing_source.docx")

    class FakeUpload:
        filename = "article_api_upload_missing_source.docx"

        def __init__(self, payload: bytes):
            self.file = BytesIO(payload)

    upload_response = routes["/uploads/docx"].endpoint(
        FakeUpload(Path(source_path).read_bytes()),
        runtime_root=str(tmp_path / "runtime"),
    )
    Path(upload_response["stored_path"]).unlink()

    upload_status = routes["/uploads/{upload_id}"].endpoint(upload_response["upload_id"])
    assert upload_status["available"] is False

    with pytest.raises(FakeHTTPException) as exc_info:
        routes["/uploads/{upload_id}/jobs/apply"].endpoint(
            upload_response["upload_id"],
            UploadApplyRequest(scopes=["headings"]),
        )

    assert exc_info.value.status_code == 409
    assert job_count() == 0


def test_fake_app_upload_cleanup_is_blocked_by_active_job(monkeypatch, tmp_docx, tmp_path):
    import time

    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = tmp_docx(make_compliant_doc, filename="article_api_upload_cleanup_guard.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    doc.save(source_path)

    class FakeUpload:
        filename = "article_api_upload_cleanup_guard.docx"

        def __init__(self, payload: bytes):
            self.file = BytesIO(payload)

    def slow_attempt(operation: str, resolved_request: dict[str, object]) -> dict[str, object]:
        time.sleep(0.2)
        return {
            "ok": True,
            "result": jobs_module.verify_document(
                file_path=str(resolved_request["file_path"]),
                profile_path=resolved_request.get("profile_path"),
                scopes=resolved_request.get("scopes"),
                strict_profile=bool(resolved_request.get("strict_profile")),
            ),
        }

    monkeypatch.setattr(jobs_module, "_execute_job_attempt", slow_attempt)
    upload_response = routes["/uploads/docx"].endpoint(
        FakeUpload(Path(source_path).read_bytes()),
        runtime_root=str(tmp_path / "runtime"),
    )
    create_response = routes["/uploads/{upload_id}/jobs/verify"].endpoint(
        upload_response["upload_id"],
        UploadVerifyRequest(scopes=["headings"]),
    )

    with pytest.raises(FakeHTTPException) as exc_info:
        routes["/uploads/{upload_id}/cleanup"].endpoint(upload_response["upload_id"])

    assert exc_info.value.status_code == 409
    assert create_response["job_id"] in exc_info.value.detail
    wait_for_job(create_response["job_id"])


def test_fake_app_cleanup_endpoint_keeps_job_payload_readable_and_breaks_download(monkeypatch, tmp_docx, tmp_path):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = tmp_docx(make_compliant_doc, filename="article_api_cleanup.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    doc.save(source_path)

    response = routes["/jobs/apply"].endpoint(
        ApplyRequest(
            file_path=str(source_path),
            scopes=["headings"],
            stage_input=True,
            runtime_root=str(tmp_path / "runtime"),
        )
    )
    wait_for_job(response["job_id"])
    before = routes["/jobs/{job_id}/result"].endpoint(response["job_id"])
    cleanup_response = routes["/jobs/{job_id}/cleanup"].endpoint(response["job_id"])
    after = routes["/jobs/{job_id}/result"].endpoint(response["job_id"])

    assert cleanup_response["cleanup"]["state"] == "cleaned"
    assert after["summary"] == before["summary"]
    assert after["result"] == before["result"]

    with pytest.raises(FakeHTTPException) as exc_info:
        routes["/jobs/{job_id}/artifacts/{artifact_role}/download"].endpoint(
            response["job_id"],
            "output",
        )

    assert exc_info.value.status_code == 409
    assert exc_info.value.detail["code"] == "artifact_unavailable"
    assert exc_info.value.detail["user_message"] == "下载文件已经不可用。"


def test_fake_app_retry_endpoint_requeues_upload_backed_job_without_leaking_file_path(monkeypatch, tmp_docx, tmp_path):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = tmp_docx(make_compliant_doc, filename="article_api_retry_upload.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    doc.save(source_path)

    class FakeUpload:
        filename = "article_api_retry_upload.docx"

        def __init__(self, payload: bytes):
            self.file = BytesIO(payload)

    upload_response = routes["/uploads/docx"].endpoint(
        FakeUpload(Path(source_path).read_bytes()),
        runtime_root=str(tmp_path / "runtime"),
    )
    first_job = routes["/uploads/{upload_id}/jobs/verify"].endpoint(
        upload_response["upload_id"],
        UploadVerifyRequest(scopes=["headings"]),
    )
    wait_for_job(first_job["job_id"])

    retried_job = routes["/jobs/{job_id}/retry"].endpoint(first_job["job_id"])
    wait_for_job(retried_job["job_id"])
    retried_status = routes["/jobs/{job_id}"].endpoint(retried_job["job_id"])

    assert routes["/jobs/{job_id}/retry"].status_code == 201
    assert retried_job["request"]["upload_id"] == upload_response["upload_id"]
    assert retried_job["request"]["retry_of_job_id"] == first_job["job_id"]
    assert "file_path" not in retried_job["request"]
    assert retried_status["status"] == "succeeded"
    assert retried_status["runtime"]["retry_of_job_id"] == first_job["job_id"]


def test_fake_app_retry_endpoint_still_works_after_job_cleanup(monkeypatch, tmp_docx, tmp_path):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = tmp_docx(make_compliant_doc, filename="article_api_retry_upload_after_cleanup.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    doc.save(source_path)

    class FakeUpload:
        filename = "article_api_retry_upload_after_cleanup.docx"

        def __init__(self, payload: bytes):
            self.file = BytesIO(payload)

    upload_response = routes["/uploads/docx"].endpoint(
        FakeUpload(Path(source_path).read_bytes()),
        runtime_root=str(tmp_path / "runtime"),
    )
    first_job = routes["/uploads/{upload_id}/jobs/verify"].endpoint(
        upload_response["upload_id"],
        UploadVerifyRequest(scopes=["headings"]),
    )
    wait_for_job(first_job["job_id"])
    cleanup_response = routes["/jobs/{job_id}/cleanup"].endpoint(first_job["job_id"])

    retried_job = routes["/jobs/{job_id}/retry"].endpoint(first_job["job_id"])
    wait_for_job(retried_job["job_id"])
    retried_status = routes["/jobs/{job_id}"].endpoint(retried_job["job_id"])

    assert cleanup_response["cleanup"]["state"] in {"cleaned", "noop"}
    assert Path(upload_response["stored_path"]).exists()
    assert retried_job["request"]["upload_id"] == upload_response["upload_id"]
    assert retried_job["request"]["retry_of_job_id"] == first_job["job_id"]
    assert "file_path" not in retried_job["request"]
    assert retried_status["status"] == "succeeded"
    assert retried_status["runtime"]["retry_of_job_id"] == first_job["job_id"]


def test_fake_app_cleanup_breaks_input_download_but_result_stays_frozen(monkeypatch, tmp_docx, tmp_path):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = tmp_docx(make_compliant_doc, filename="article_api_cleanup_input.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    doc.save(source_path)

    response = routes["/jobs/verify"].endpoint(
        VerifyRequest(
            file_path=str(source_path),
            scopes=["headings"],
            stage_input=True,
            runtime_root=str(tmp_path / "runtime"),
        )
    )
    wait_for_job(response["job_id"])
    before = routes["/jobs/{job_id}/result"].endpoint(response["job_id"])
    routes["/jobs/{job_id}/cleanup"].endpoint(response["job_id"])
    after = routes["/jobs/{job_id}/result"].endpoint(response["job_id"])

    with pytest.raises(FakeHTTPException) as exc_info:
        routes["/jobs/{job_id}/artifacts/{artifact_role}/download"].endpoint(
            response["job_id"],
            "input",
        )

    assert exc_info.value.status_code == 409
    assert exc_info.value.detail["code"] == "artifact_unavailable"
    assert exc_info.value.detail["user_message"] == "下载文件已经不可用。"
    assert after["summary"] == before["summary"]
    assert after["result"] == before["result"]


def test_fake_app_upload_job_rejects_conflicting_runtime_root(monkeypatch, tmp_docx, tmp_path):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = tmp_docx(make_compliant_doc, filename="article_api_upload_runtime_conflict.docx")

    class FakeUpload:
        filename = "article_api_upload_runtime_conflict.docx"

        def __init__(self, payload: bytes):
            self.file = BytesIO(payload)

    upload_response = routes["/uploads/docx"].endpoint(
        FakeUpload(Path(source_path).read_bytes()),
        runtime_root=str(tmp_path / "runtime-a"),
    )

    with pytest.raises(FakeHTTPException) as exc_info:
        routes["/uploads/{upload_id}/jobs/verify"].endpoint(
            upload_response["upload_id"],
            UploadVerifyRequest(
                scopes=["headings"],
                runtime_root=str(tmp_path / "runtime-b"),
            ),
        )

    assert exc_info.value.status_code == 400
    assert job_count() == 0


def test_fake_app_upload_registry_is_atomic_when_storage_write_fails(monkeypatch, tmp_docx, tmp_path):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = tmp_docx(make_compliant_doc, filename="article_api_upload_atomic.docx")
    stored_file = tmp_path / "runtime" / "uploads" / "aaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaa_article_api_upload_atomic.docx"
    stored_file.parent.mkdir(parents=True, exist_ok=True)
    stored_file.write_bytes(Path(source_path).read_bytes())

    class FakeUpload:
        filename = "article_api_upload_atomic.docx"

        def __init__(self, payload: bytes):
            self.file = BytesIO(payload)

    class FakeStoredUpload:
        upload_id = "aaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaa"
        stored_path = str(stored_file)
        file_name = "article_api_upload_atomic.docx"
        workspace_dir = str(stored_file.parent)
        size_bytes = len(Path(source_path).read_bytes())

    monkeypatch.setattr(app_module, "store_uploaded_docx", lambda *_args, **_kwargs: FakeStoredUpload())
    monkeypatch.setattr(app_module.storage, "upsert_upload", lambda *_args, **_kwargs: (_ for _ in ()).throw(RuntimeError("db write failed")))

    with pytest.raises(FakeHTTPException) as exc_info:
        routes["/uploads/docx"].endpoint(
            FakeUpload(Path(source_path).read_bytes()),
            runtime_root=str(tmp_path / "runtime"),
        )

    assert exc_info.value.status_code == 409
    assert upload_count() == 0


def test_fake_app_upload_cleanup_marks_upload_unavailable_and_blocks_future_jobs(monkeypatch, tmp_docx, tmp_path):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = tmp_docx(make_compliant_doc, filename="article_api_upload_cleanup.docx")

    class FakeUpload:
        filename = "article_api_upload_cleanup.docx"

        def __init__(self, payload: bytes):
            self.file = BytesIO(payload)

    upload_response = routes["/uploads/docx"].endpoint(
        FakeUpload(Path(source_path).read_bytes()),
        runtime_root=str(tmp_path / "runtime"),
    )
    cleanup_response = routes["/uploads/{upload_id}/cleanup"].endpoint(upload_response["upload_id"])
    upload_status = routes["/uploads/{upload_id}"].endpoint(upload_response["upload_id"])

    assert cleanup_response["cleanup"]["state"] == "cleaned"
    assert cleanup_response["cleanup"]["attempt_count"] == 1
    assert upload_status["available"] is False

    with pytest.raises(FakeHTTPException) as exc_info:
        routes["/uploads/{upload_id}/jobs/verify"].endpoint(
            upload_response["upload_id"],
            UploadVerifyRequest(scopes=["headings"]),
        )

    assert exc_info.value.status_code == 409
    assert not Path(upload_response["stored_path"]).exists()


def test_fake_app_retention_sweep_cleans_expired_job_and_upload(monkeypatch, tmp_docx, tmp_path):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = tmp_docx(make_compliant_doc, filename="article_api_retention.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    doc.save(source_path)

    class FakeUpload:
        filename = "article_api_retention.docx"

        def __init__(self, payload: bytes):
            self.file = BytesIO(payload)

    upload_response = routes["/uploads/docx"].endpoint(
        FakeUpload(Path(source_path).read_bytes()),
        runtime_root=str(tmp_path / "runtime"),
    )
    create_response = routes["/uploads/{upload_id}/jobs/apply"].endpoint(
        upload_response["upload_id"],
        UploadApplyRequest(scopes=["headings"], stage_input=True),
    )
    wait_for_job(create_response["job_id"])
    monkeypatch.setattr(app_module, "_utcnow", lambda: "2099-01-01T00:00:00Z")

    sweep_response = routes["/ops/retention/sweep"].endpoint(
        RetentionSweepRequest(job_max_age_seconds=1, upload_max_age_seconds=1),
    )
    job_payload = routes["/jobs/{job_id}/result"].endpoint(create_response["job_id"])
    upload_payload = routes["/uploads/{upload_id}"].endpoint(upload_response["upload_id"])

    assert sweep_response["policy"] == "retention"
    assert sweep_response["job_retention"]["cleaned_count"] == 1
    assert sweep_response["upload_retention"]["cleaned_count"] == 1
    assert sweep_response["job_retention"]["items"][0]["job_id"] == create_response["job_id"]
    assert sweep_response["job_retention"]["items"][0]["cleanup"]["policy"] == "retention"
    assert sweep_response["upload_retention"]["items"][0]["upload_id"] == upload_response["upload_id"]
    assert sweep_response["upload_retention"]["items"][0]["cleanup"]["policy"] == "retention"
    assert job_payload["cleanup"]["policy"] == "retention"
    assert upload_payload["cleanup"]["policy"] == "retention"
    assert upload_payload["available"] is False


def test_fake_app_run_default_retention_uses_env_thresholds(monkeypatch, tmp_docx, tmp_path):
    clear_jobs()
    clear_uploads()
    monkeypatch.setenv("ARTICLE_API_JOB_RETENTION_SECONDS", "1")
    monkeypatch.setenv("ARTICLE_API_UPLOAD_RETENTION_SECONDS", "1")
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = tmp_docx(make_compliant_doc, filename="article_api_default_retention.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    doc.save(source_path)

    class FakeUpload:
        filename = "article_api_default_retention.docx"

        def __init__(self, payload: bytes):
            self.file = BytesIO(payload)

    upload_response = routes["/uploads/docx"].endpoint(
        FakeUpload(Path(source_path).read_bytes()),
        runtime_root=str(tmp_path / "runtime"),
    )
    create_response = routes["/uploads/{upload_id}/jobs/apply"].endpoint(
        upload_response["upload_id"],
        UploadApplyRequest(scopes=["headings"], stage_input=True),
    )
    wait_for_job(create_response["job_id"])
    monkeypatch.setattr(app_module, "_utcnow", lambda: "2099-01-01T00:00:00Z")

    sweep_response = routes["/ops/retention/run-defaults"].endpoint(app_module.RetentionDefaultsRunRequest())
    summary_payload = routes["/ops/summary"].endpoint()

    assert sweep_response["job_retention"]["cleaned_count"] == 1
    assert sweep_response["upload_retention"]["cleaned_count"] == 1
    assert summary_payload["retention"]["state"]["last_trigger"] == "manual-defaults"
    assert summary_payload["retention"]["defaults"]["job_max_age_seconds"] == 1.0
    assert summary_payload["retention"]["defaults"]["upload_max_age_seconds"] == 1.0


def test_fake_app_run_default_retention_requires_env_thresholds(monkeypatch):
    clear_jobs()
    clear_uploads()
    monkeypatch.delenv("ARTICLE_API_JOB_RETENTION_SECONDS", raising=False)
    monkeypatch.delenv("ARTICLE_API_UPLOAD_RETENTION_SECONDS", raising=False)
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)

    with pytest.raises(FakeHTTPException) as exc_info:
        routes["/ops/retention/run-defaults"].endpoint(app_module.RetentionDefaultsRunRequest())

    assert exc_info.value.status_code == 400


def test_fake_app_autorun_retention_cleans_expired_job_on_next_write(monkeypatch, tmp_docx, tmp_path):
    clear_jobs()
    clear_uploads()
    monkeypatch.setenv("ARTICLE_API_JOB_RETENTION_SECONDS", "1")
    monkeypatch.setenv("ARTICLE_API_RETENTION_AUTORUN", "true")
    monkeypatch.setenv("ARTICLE_API_RETENTION_AUTORUN_INTERVAL_SECONDS", "1")
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)
    source_path = tmp_docx(make_compliant_doc, filename="article_api_autorun_retention.docx")
    doc = Document(source_path)
    RULE_MUTATORS["H02"](doc)
    doc.save(source_path)

    first_job = routes["/jobs/verify"].endpoint(
        VerifyRequest(
            file_path=str(source_path),
            scopes=["headings"],
            stage_input=True,
            runtime_root=str(tmp_path / "runtime"),
        )
    )
    wait_for_job(first_job["job_id"])

    job_result = routes["/jobs/{job_id}/result"].endpoint(first_job["job_id"])
    monkeypatch.setattr(app_module, "_utcnow", lambda: "2099-01-01T00:00:00Z")
    app_module._RETENTION_STATE["last_monotonic"] = 0.0
    routes["/jobs/verify"].endpoint(
        VerifyRequest(
            file_path=str(source_path),
            scopes=["headings"],
        )
    )
    retained_job = routes["/jobs/{job_id}/result"].endpoint(first_job["job_id"])
    summary_payload = routes["/ops/summary"].endpoint()

    assert job_result["cleanup"] is None
    assert retained_job["cleanup"]["policy"] == "retention"
    assert summary_payload["retention"]["state"]["last_trigger"] == "autorun"


def test_fake_app_retention_sweep_requires_threshold(monkeypatch):
    clear_jobs()
    clear_uploads()
    app = _build_fake_app(monkeypatch)
    routes = _routes_by_path(app)

    with pytest.raises(FakeHTTPException) as exc_info:
        routes["/ops/retention/sweep"].endpoint(RetentionSweepRequest())

    assert exc_info.value.status_code == 400
