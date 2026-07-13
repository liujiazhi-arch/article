from __future__ import annotations

import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document

import fix_docx_io


def test_write_docx_atomically_keeps_docx_openable(tmp_path: Path):
    source_path = tmp_path / "source.docx"
    output_path = tmp_path / "output.docx"

    fix_docx_io.Document = Document

    doc = Document()
    doc.add_paragraph("测试输出写盘")
    doc.save(source_path)

    with zipfile.ZipFile(source_path, "r") as source_zip:
        document_xml = source_zip.read("word/document.xml")

    fix_docx_io.write_docx_atomically(
        str(source_path),
        str(output_path),
        {"word/document.xml": document_xml},
    )

    reopened = Document(output_path)
    assert output_path.exists()
    assert reopened.paragraphs[0].text == "测试输出写盘"


def test_write_docx_atomically_registers_missing_jpg_content_type(tmp_path: Path):
    valid_path = tmp_path / "valid.docx"
    source_path = tmp_path / "source.docx"
    output_path = tmp_path / "output.docx"

    doc = Document()
    doc.add_picture(str(Path(__file__).parents[1] / "scripts/article_api/static/assets/lnu-emblem.jpg"))
    doc.save(valid_path)

    with zipfile.ZipFile(valid_path, "r") as source_zip, zipfile.ZipFile(source_path, "w") as target_zip:
        for item in source_zip.infolist():
            content = source_zip.read(item.filename)
            if item.filename == "[Content_Types].xml":
                root = ET.fromstring(content)
                for default in root.findall(f"{{{fix_docx_io.CONTENT_TYPES_NS}}}Default"):
                    if default.get("Extension", "").lower() == "jpg":
                        root.remove(default)
                content = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target_zip.writestr(item, content)

    with zipfile.ZipFile(source_path, "r") as source_zip:
        document_xml = source_zip.read("word/document.xml")

    fix_docx_io.write_docx_atomically(
        str(source_path),
        str(output_path),
        {"word/document.xml": document_xml},
    )

    assert output_path.exists()
    assert Document(output_path).inline_shapes
