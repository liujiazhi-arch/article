from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
import zipfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK


_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


@dataclass(frozen=True)
class CompatSample:
    id: str
    description: str
    required_parts: tuple[str, ...]
    required_markers: tuple[str, ...]
    builder: Callable[[Path], Path]


def _save_base_docx(path: Path, *, title: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_paragraph(title)
    doc.add_paragraph("1 绪论").style = "Heading 1"
    doc.add_paragraph("1.1 研究背景").style = "Heading 2"
    doc.add_paragraph("这是用于公开回归的脱敏兼容性样本文档。")
    doc.save(path)
    return path


def _rewrite_docx(path: Path, writer: Callable[[zipfile.ZipFile, zipfile.ZipFile], None]) -> None:
    temp_path = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as target:
        writer(source, target)
    temp_path.replace(path)


def _copy_all(source: zipfile.ZipFile, target: zipfile.ZipFile, replacements: dict[str, str] | None = None) -> None:
    replacements = replacements or {}
    for info in source.infolist():
        if info.filename in replacements:
            target.writestr(info, replacements[info.filename])
        else:
            target.writestr(info, source.read(info.filename))


def _add_content_type_override(content_types: str, *, part_name: str, content_type: str) -> str:
    marker = f'PartName="{part_name}"'
    if marker in content_types:
        return content_types
    insert = f'<Override PartName="{part_name}" ContentType="{content_type}"/>'
    return content_types.replace("</Types>", f"{insert}</Types>")


def _add_document_relationship(rels: str, *, rel_id: str, rel_type: str, target: str) -> str:
    if f'Id="{rel_id}"' in rels:
        return rels
    insert = f'<Relationship Id="{rel_id}" Type="{rel_type}" Target="{target}"/>'
    return rels.replace("</Relationships>", f"{insert}</Relationships>")


def _build_wps_style_docx(path: Path) -> Path:
    return _save_base_docx(path, title="WPS 保存风格样本")


def _build_word_toc_field_docx(path: Path) -> Path:
    _save_base_docx(path, title="Word 目录域样本")

    def writer(source: zipfile.ZipFile, target: zipfile.ZipFile) -> None:
        document_xml = source.read("word/document.xml").decode("utf-8")
        field_xml = (
            '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
            '<w:r><w:instrText xml:space="preserve"> TOC \\\\o "1-3" \\\\h \\\\z \\\\u </w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
            '<w:r><w:t>目录项占位</w:t></w:r>'
            '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
        )
        document_xml = document_xml.replace("<w:body>", f"<w:body>{field_xml}", 1)
        _copy_all(source, target, {"word/document.xml": document_xml})

    _rewrite_docx(path, writer)
    return path


def _build_footnote_docx(path: Path) -> Path:
    _save_base_docx(path, title="脚注样本")

    def writer(source: zipfile.ZipFile, target: zipfile.ZipFile) -> None:
        document_xml = source.read("word/document.xml").decode("utf-8")
        document_xml = document_xml.replace(
            "</w:p>",
            '<w:r><w:footnoteReference w:id="2"/></w:r></w:p>',
            1,
        )
        rels = source.read("word/_rels/document.xml.rels").decode("utf-8")
        rels = _add_document_relationship(
            rels,
            rel_id="rIdCompatFootnotes",
            rel_type=f"{_R_NS}/footnotes",
            target="footnotes.xml",
        )
        content_types = source.read("[Content_Types].xml").decode("utf-8")
        content_types = _add_content_type_override(
            content_types,
            part_name="/word/footnotes.xml",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        )
        footnotes_xml = (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<w:footnotes xmlns:w="{_W_NS}">'
            '<w:footnote w:type="separator" w:id="-1"><w:p/></w:footnote>'
            '<w:footnote w:type="continuationSeparator" w:id="0"><w:p/></w:footnote>'
            '<w:footnote w:id="2"><w:p><w:r><w:t>脱敏脚注文本</w:t></w:r></w:p></w:footnote>'
            "</w:footnotes>"
        )
        _copy_all(
            source,
            target,
            {
                "word/document.xml": document_xml,
                "word/_rels/document.xml.rels": rels,
                "[Content_Types].xml": content_types,
            },
        )
        target.writestr("word/footnotes.xml", footnotes_xml)

    _rewrite_docx(path, writer)
    return path


def _build_comments_docx(path: Path) -> Path:
    _save_base_docx(path, title="批注样本")

    def writer(source: zipfile.ZipFile, target: zipfile.ZipFile) -> None:
        document_xml = source.read("word/document.xml").decode("utf-8")
        document_xml = document_xml.replace(
            "</w:p>",
            '<w:commentRangeStart w:id="0"/><w:r><w:t>含批注文本</w:t></w:r>'
            '<w:commentRangeEnd w:id="0"/><w:r><w:commentReference w:id="0"/></w:r></w:p>',
            1,
        )
        rels = source.read("word/_rels/document.xml.rels").decode("utf-8")
        rels = _add_document_relationship(
            rels,
            rel_id="rIdCompatComments",
            rel_type=f"{_R_NS}/comments",
            target="comments.xml",
        )
        content_types = source.read("[Content_Types].xml").decode("utf-8")
        content_types = _add_content_type_override(
            content_types,
            part_name="/word/comments.xml",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
        )
        comments_xml = (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<w:comments xmlns:w="{_W_NS}">'
            '<w:comment w:id="0" w:author="tester" w:date="2026-01-01T00:00:00Z">'
            "<w:p><w:r><w:t>脱敏批注</w:t></w:r></w:p>"
            "</w:comment></w:comments>"
        )
        _copy_all(
            source,
            target,
            {
                "word/document.xml": document_xml,
                "word/_rels/document.xml.rels": rels,
                "[Content_Types].xml": content_types,
            },
        )
        target.writestr("word/comments.xml", comments_xml)

    _rewrite_docx(path, writer)
    return path


def _build_revision_docx(path: Path) -> Path:
    _save_base_docx(path, title="修订样本")

    def writer(source: zipfile.ZipFile, target: zipfile.ZipFile) -> None:
        document_xml = source.read("word/document.xml").decode("utf-8")
        revision_xml = (
            '<w:p><w:ins w:id="1" w:author="tester" w:date="2026-01-01T00:00:00Z">'
            "<w:r><w:t>插入修订文本</w:t></w:r></w:ins>"
            '<w:del w:id="2" w:author="tester" w:date="2026-01-01T00:00:00Z">'
            "<w:r><w:delText>删除修订文本</w:delText></w:r></w:del></w:p>"
        )
        document_xml = document_xml.replace("</w:body>", f"{revision_xml}</w:body>", 1)
        _copy_all(source, target, {"word/document.xml": document_xml})

    _rewrite_docx(path, writer)
    return path


def _build_floating_image_docx(path: Path) -> Path:
    _save_base_docx(path, title="浮动图片样本")

    def writer(source: zipfile.ZipFile, target: zipfile.ZipFile) -> None:
        document_xml = source.read("word/document.xml").decode("utf-8")
        anchor_xml = (
            '<w:p><w:r><w:drawing><wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
            'simplePos="0" relativeHeight="0" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
            '<wp:extent cx="914400" cy="914400"/><wp:docPr id="99" name="脱敏浮动图片"/>'
            "</wp:anchor></w:drawing></w:r></w:p>"
        )
        document_xml = document_xml.replace("</w:body>", f"{anchor_xml}</w:body>", 1)
        _copy_all(source, target, {"word/document.xml": document_xml})

    _rewrite_docx(path, writer)
    return path


def _build_formula_docx(path: Path) -> Path:
    _save_base_docx(path, title="公式样本")

    def writer(source: zipfile.ZipFile, target: zipfile.ZipFile) -> None:
        document_xml = source.read("word/document.xml").decode("utf-8")
        formula_xml = (
            '<w:p><m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
            "<m:r><m:t>E=mc</m:t></m:r></m:oMath></w:p>"
        )
        document_xml = document_xml.replace("</w:body>", f"{formula_xml}</w:body>", 1)
        _copy_all(source, target, {"word/document.xml": document_xml})

    _rewrite_docx(path, writer)
    return path


def _build_nested_table_docx(path: Path) -> Path:
    path = _save_base_docx(path, title="嵌套表格样本")
    doc = Document(path)
    outer = doc.add_table(rows=1, cols=1)
    outer.cell(0, 0).text = "外层表格"
    inner = outer.cell(0, 0).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "内层表格"
    doc.save(path)
    return path


def _build_section_break_docx(path: Path) -> Path:
    path = _save_base_docx(path, title="异常 section break 样本")
    doc = Document(path)
    doc.add_section(WD_SECTION.CONTINUOUS)
    paragraph = doc.add_paragraph("连续分节后的正文")
    paragraph.runs[0].add_break(WD_BREAK.PAGE)
    doc.save(path)
    return path


def _build_mixed_dirty_stress_docx(path: Path) -> Path:
    path = _save_base_docx(path, title="综合脏结构压力样本")
    doc = Document(path)
    outer = doc.add_table(rows=1, cols=1)
    outer.cell(0, 0).text = "综合样本外层表格"
    inner = outer.cell(0, 0).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "综合样本内层表格"
    doc.add_section(WD_SECTION.CONTINUOUS)
    paragraph = doc.add_paragraph("综合样本连续分节后的正文")
    paragraph.runs[0].add_break(WD_BREAK.PAGE)
    doc.save(path)

    def writer(source: zipfile.ZipFile, target: zipfile.ZipFile) -> None:
        document_xml = source.read("word/document.xml").decode("utf-8")
        toc_xml = (
            '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
            '<w:r><w:instrText xml:space="preserve"> TOC \\\\o "1-3" \\\\h \\\\z \\\\u </w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
            '<w:r><w:t>综合目录项占位</w:t></w:r>'
            '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
        )
        revision_xml = (
            '<w:p><w:ins w:id="11" w:author="tester" w:date="2026-01-01T00:00:00Z">'
            "<w:r><w:t>综合插入修订文本</w:t></w:r></w:ins>"
            '<w:del w:id="12" w:author="tester" w:date="2026-01-01T00:00:00Z">'
            "<w:r><w:delText>综合删除修订文本</w:delText></w:r></w:del></w:p>"
        )
        drawing_xml = (
            '<w:p><w:r><w:drawing><wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
            'simplePos="0" relativeHeight="0" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
            '<wp:extent cx="914400" cy="914400"/><wp:docPr id="199" name="综合脱敏浮动图片"/>'
            "</wp:anchor></w:drawing></w:r></w:p>"
        )
        formula_xml = (
            '<w:p><m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
            "<m:r><m:t>a2+b2=c2</m:t></m:r></m:oMath></w:p>"
        )
        document_xml = document_xml.replace("<w:body>", f"<w:body>{toc_xml}", 1)
        document_xml = document_xml.replace(
            "</w:p>",
            '<w:commentRangeStart w:id="0"/><w:r><w:t>综合含批注文本</w:t></w:r>'
            '<w:commentRangeEnd w:id="0"/><w:r><w:commentReference w:id="0"/></w:r>'
            '<w:r><w:footnoteReference w:id="2"/></w:r></w:p>',
            1,
        )
        document_xml = document_xml.replace("</w:body>", f"{revision_xml}{drawing_xml}{formula_xml}</w:body>", 1)

        rels = source.read("word/_rels/document.xml.rels").decode("utf-8")
        rels = _add_document_relationship(
            rels,
            rel_id="rIdMixedFootnotes",
            rel_type=f"{_R_NS}/footnotes",
            target="footnotes.xml",
        )
        rels = _add_document_relationship(
            rels,
            rel_id="rIdMixedComments",
            rel_type=f"{_R_NS}/comments",
            target="comments.xml",
        )
        content_types = source.read("[Content_Types].xml").decode("utf-8")
        content_types = _add_content_type_override(
            content_types,
            part_name="/word/footnotes.xml",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        )
        content_types = _add_content_type_override(
            content_types,
            part_name="/word/comments.xml",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
        )
        footnotes_xml = (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<w:footnotes xmlns:w="{_W_NS}">'
            '<w:footnote w:type="separator" w:id="-1"><w:p/></w:footnote>'
            '<w:footnote w:type="continuationSeparator" w:id="0"><w:p/></w:footnote>'
            '<w:footnote w:id="2"><w:p><w:r><w:t>综合脱敏脚注文本</w:t></w:r></w:p></w:footnote>'
            "</w:footnotes>"
        )
        comments_xml = (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<w:comments xmlns:w="{_W_NS}">'
            '<w:comment w:id="0" w:author="tester" w:date="2026-01-01T00:00:00Z">'
            "<w:p><w:r><w:t>综合脱敏批注</w:t></w:r></w:p>"
            "</w:comment></w:comments>"
        )
        _copy_all(
            source,
            target,
            {
                "word/document.xml": document_xml,
                "word/_rels/document.xml.rels": rels,
                "[Content_Types].xml": content_types,
            },
        )
        target.writestr("word/footnotes.xml", footnotes_xml)
        target.writestr("word/comments.xml", comments_xml)

    _rewrite_docx(path, writer)
    return path


COMPAT_SAMPLES: tuple[CompatSample, ...] = (
    CompatSample("wps_basic", "WPS 保存风格", ("word/document.xml",), ("WPS 保存风格样本",), _build_wps_style_docx),
    CompatSample("word_toc_field", "Word 目录域", ("word/document.xml",), ("fldChar", " TOC "), _build_word_toc_field_docx),
    CompatSample("footnote", "脚注", ("word/footnotes.xml",), ("footnoteReference",), _build_footnote_docx),
    CompatSample("comments", "批注", ("word/comments.xml",), ("commentRangeStart",), _build_comments_docx),
    CompatSample("revision", "修订", ("word/document.xml",), ("<w:ins ", "<w:del "), _build_revision_docx),
    CompatSample("floating_image", "浮动图片", ("word/document.xml",), ("wp:anchor",), _build_floating_image_docx),
    CompatSample("formula", "公式", ("word/document.xml",), ("m:oMath",), _build_formula_docx),
    CompatSample("nested_table", "嵌套表格", ("word/document.xml",), ("内层表格",), _build_nested_table_docx),
    CompatSample("section_break", "异常 section break", ("word/document.xml",), ("w:sectPr",), _build_section_break_docx),
    CompatSample(
        "mixed_dirty_stress",
        "综合脏结构压力样本",
        ("word/document.xml", "word/footnotes.xml", "word/comments.xml"),
        ("fldChar", "footnoteReference", "commentRangeStart", "<w:ins ", "wp:anchor", "m:oMath", "综合样本内层表格", "w:sectPr"),
        _build_mixed_dirty_stress_docx,
    ),
)


def build_compat_sample(sample: CompatSample, output_dir: Path) -> Path:
    return sample.builder(output_dir / f"{sample.id}.docx")
